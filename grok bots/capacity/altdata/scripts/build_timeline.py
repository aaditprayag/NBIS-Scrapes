#!/usr/bin/env python3
"""Builds altdata_timelines.xlsx: Timeline (evidence rows), Vineland_MW (power-bounded estimate),
Imagery_metrics (detail/crop metrics per site), Thermal (Landsat ST), Sites_AOI."""
import csv, glob, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
OUT='/workspace/nebius/capacity/altdata'
S2='Sentinel-2 L2A (Element84 Earth Search STAC); chips in imagery/'
T=[ # site, date, evidence type, observation, source, confidence
# ---------------- VINELAND
('Vineland NJ','2024-02','Zoning','City zoning officer determination: data center is a permitted use under redevelopment plan (per city solicitor, Aug 2026 hearing)','https://www.nj.com/business/2026/08/you-sold-us-out-nj-town-approves-next-phase-of-massive-ai-data-center-as-residents-erupt.html','medium'),
('Vineland NJ','2024-06','Planning','Original site plan (June 2024) approved later in 2024; amended 2025','https://www.burlingtoncountytimes.com/story/money/business/2026/08/20/vineland-planning-board-hearing-vote-dataone-lincoln-avenue-data-center-new-site-plan/91338935007/','medium'),
('Vineland NJ','2024-09','Utility','City/Vineland Municipal Electric commissioned Atlantic City Electric (ACE) study for 300 MW; VMEU "not able to accommodate the Service Request" (peak system load 164 MW)','https://measuredai.substack.com/p/microsoft-new-jersey-ai-data-center (secondary, cites ACE deposit request Mar 2025)','medium'),
('Vineland NJ','2025-02','Utility','ACE offered temporary 50 MW at 138 kV via tap off Lincoln-Union line to customer-owned substation (~$6.2M +/-50%); 300 MW "under study"','https://measuredai.substack.com/p/microsoft-new-jersey-ai-data-center (secondary)','medium'),
('Vineland NJ','2025-02-26','Land','NEP Real Estate (Northeast Precast affiliate) transferred campus lots to DataOne; groundbreaking Feb 2025','https://measuredai.substack.com/p/microsoft-new-jersey-ai-data-center (secondary, county deed records)','medium'),
('Vineland NJ','2025-03-25','Imagery','Active earthworks/sand grading across campus; first steel visible NW corner',S2+' vineland_nj_2025-03-25.png','medium'),
('Vineland NJ','2025-06-03','Imagery','First large white roof (Phase 1 building, NW of campus) visible',S2+' vineland_nj_2025-06-03.png','medium'),
('Vineland NJ','2025-07','Power (temp)','Rented Caterpillar gas generators begin arriving on site (DataOne letter to NJDEP)','https://floodlightnews.org/vineland-nj-data-center-air-pollution/','high'),
('Vineland NJ','2025-08-30','Mgmt/Developer','DataOne CEO concedes build ~2 months behind schedule','https://measuredai.substack.com/p/microsoft-new-jersey-ai-data-center (secondary)','medium'),
('Vineland NJ','2025-09-23','Imagery','Phase 1 roof plus second roof segment; bright-roof in crop ~4.9 ha (incl. pre-existing NE Precast roofs ~3-4 ha)',S2+' vineland_nj_2025-09-23.png','medium'),
('Vineland NJ','2025-10','Power (temp)','16 generators start operating (~31.7 MW nameplate) -> energization of Phase 1 on temporary power','https://floodlightnews.org/vineland-nj-data-center-air-pollution/','high'),
('Vineland NJ','2025-11','Mgmt','Microsoft tranche 1 delivered (Nebius); MSFT payment obligations begin early Nov 2025','site_list.csv / Nebius Q3-25 letter','high'),
('Vineland NJ','2025-11-24','Utility (gas)','DataOne-South Jersey Gas 15-yr firm gas transport agreement (Jul 2026-Jun 2041) sized for up to 125 MW generation; filed with NJBPU Feb 2 2026, still pending Aug 12 2026','https://measuredai.substack.com/p/microsoft-new-jersey-ai-data-center (secondary, NJBPU docket)','medium'),
('Vineland NJ','2025-12-22','Imagery','Rows of grey units appear E/SE of building (consistent with generator rows; not resolvable as individual units at 10 m)',S2+' vineland_nj_2025-12-22.png','low'),
('Vineland NJ','2026-01-13','Utility','City Council authorizes alternative ACE study at Northeast Precast request (request reduced to 49 MW)','https://measuredai.substack.com/p/microsoft-new-jersey-ai-data-center (secondary)','medium'),
('Vineland NJ','2026-02','Power (temp)','+16 generators online (32 total, ~63 MW nameplate); city inspection report "Phase 1 generators are online"; ~50 MW energized by tranche 2 (Feb 2026)','https://floodlightnews.org/vineland-nj-data-center-air-pollution/ ; measuredai (secondary)','high'),
('Vineland NJ','2026-02','Mgmt','Microsoft tranche 2 delivered (Nebius)','site_list.csv / Nebius filings','high'),
('Vineland NJ','2026-03-10','Regulatory','Cumberland County Health Dept Notice of Violation: night noise >50 dBA (generators/construction)','https://measuredai.substack.com/p/microsoft-new-jersey-ai-data-center (secondary)','medium'),
('Vineland NJ','2026-05','Permit','DataOne withdraws air-permit application for 32 Bergen gas engines after NJDEP found deficiencies','https://measuredai.substack.com/p/microsoft-new-jersey-ai-data-center ; https://www.theguardian.com/environment/2026/aug/27/new-jersey-vineland-datacenter-permits','medium'),
('Vineland NJ','2026-05-14','Contract','Nebius-Bloom agreement: ~250 MW guaranteed / 328 MW installed fuel cells (ES6.5)','https://www.enr.com/articles/63695-inside-the-vineland-data-center-power-dispute','high'),
('Vineland NJ','2026-06','Power (temp)','All 62 generators fired up (122.9 MW nameplate); Floodlight thermal: ~25 running late June','https://floodlightnews.org/vineland-nj-data-center-air-pollution/ ; https://www.theguardian.com/environment/2026/aug/27/new-jersey-vineland-datacenter-permits','high'),
('Vineland NJ','2026-06-03','Imagery','Bright roof in crop jumps to ~8.8 ha (from ~6 ha Apr 2026): large Phase 2 roof closed in',S2+' vineland_nj_2026-06-03.png','medium'),
('Vineland NJ','2026-06','Jobs','Nebius "Data Center Technician - Vineland NJ colocation" posted Jun 27 2026 (reposted Aug 20 2026); text references H200 cluster (template)','https://meterwork.com/job/data-center-technician_at_nebius_QHp5B ; https://haystackapp.io/jobs/d272c86f-d7b1-42c5-bdff-69a6f9a604fd','medium'),
('Vineland NJ','2026-04-20..09-19','Shipping','ImportYeti: 18 sea shipments to consignee "Nebius, Vineland NJ": Schroff racks/enclosures Apr-Jul; Bergen Engines components Aug 14-28; Ingeteam power gear Aug 28; Guntner coolers (12 containers) Sep 19','https://www.importyeti.com/company/nebius','medium'),
('Vineland NJ','2026-07-08','Planning','Revised site plan filed: 717,602 sf building (Ph1 129,622 + Ph2 587,980 sf 2-story), 9 Bloom areas (328 MW), LNG tank, chiller bldg; only 5 MW from Vineland Electric','https://www.vinelandcity.org/wp-content/uploads/2026/07/Data-ONe-Public-Hearing-Post-7.24.pdf ; measuredai (secondary)','high'),
('Vineland NJ','2026-07 (late)','Utility/PJM','ACE withdrew the 49 MW Cumberland County load need from PJM transmission-planning process; ACE large-load adjustment = 0 MW 2026-2046 (not independently verified in PJM docs)','https://measuredai.substack.com/p/microsoft-new-jersey-ai-data-center (secondary)','medium'),
('Vineland NJ','2026-07-29','Regulatory','NJDEP inspection finds 62 Cat G3516H gas engines (1,982 kW each) installed/operated without preconstruction permits','https://dep.nj.gov/wp-content/uploads/newsrel/2026/dataone75894-aonocapa.pdf','high'),
('Vineland NJ','2026-07-31','Jobs','Staffing agency hiring event in Vineland for Bloom Energy assembly technicians, start "in August"','https://hntrbrk.com/breaking-news/vineland','medium'),
('Vineland NJ','2026-08-05','Planning','6-hour Planning Board hearing on amended plan (Bloom/LNG) adjourned without vote','https://hntrbrk.com/breaking-news/vineland ; https://whyy.org/articles/vineland-data-center-hearing-power-generation/','high'),
('Vineland NJ','2026-08-06','Permit','City stop-construction order: LNG tank installed with no approvals/permits','https://whyy.org/articles/vineland-data-center-stop-work-orders/ ; https://hntrbrk.com/breaking-news/vineland','high'),
('Vineland NJ','2026-08-10','Permit','City stop-construction order: Bloom Energy units - "no prior approvals, plans, or permits"','https://www.scribd.com/document/1073069744/Stop-Construction-Order-Bloom-Energy-Units ; https://hntrbrk.com/breaking-news/vineland','high'),
('Vineland NJ','2026-08-12','Mgmt','Q2-26 letter: "delivered all capacity tranches to Microsoft... including two tranches over the past month"; call: building finished, "on track"; stop orders not mentioned','https://assets.nebius.com/assets/4462517b-ce83-41f2-96ed-f2ac1bc06a05/SHLQ226%20%281%29.pdf','high'),
('Vineland NJ','2026-08 (early)','Developer','DataOne: at least 3 of 14 data halls operating','https://measuredai.substack.com/p/microsoft-new-jersey-ai-data-center (secondary)','medium'),
('Vineland NJ','2026-08-17','Planning','Planning Board approves amended site plan 8-1 (special meeting); building permits to lift stop orders not publicly confirmed','https://www.burlingtoncountytimes.com/story/money/business/2026/08/20/vineland-planning-board-hearing-vote-dataone-lincoln-avenue-data-center-new-site-plan/91338935007/ ; https://whyy.org/articles/vineland-planning-board-approves-data-center-plan/','high'),
('Vineland NJ','2026-08 (mid/late)','Thermal (3rd party)','Floodlight thermal drone: >=45 of 62 generators running','https://floodlightnews.org/vineland-nj-data-center-air-pollution/','high'),
('Vineland NJ','2026-09-08','Imagery','Full white roof; new dark rectangular arrays E/SE of building vs Jun 2026 (consistent with Bloom pads/units being placed; not confirmable at 10 m)',S2+' vineland_nj_2026-09-08.png ; strips/vineland_nj_strip.png','low'),
('Vineland NJ','2026-09 (~22-25)','Regulatory','NJDEP AO & $1.072M penalty; 45 days from receipt to file and obtain permits or cease installation/operation of generators (=> ~early/mid Nov 2026)','https://dep.nj.gov/wp-content/uploads/newsrel/2026/dataone75894-aonocapa.pdf ; https://www.enr.com/articles/63695-inside-the-vineland-data-center-power-dispute','high'),
('Vineland NJ','2023-2026','Thermal (Landsat)','Landsat 8/9 daytime ST campus-minus-background, Jun-Aug mean: 2023 +1.35C (n=1), 2024 +0.37 (n=2), 2025 +1.42 (n=3), 2026 +1.92 (n=5). Weak upward trend; confounded by roof/sand albedo; 100 m TIRS; not a power meter','thermal/vineland_nj_landsat_st.csv (Planetary Computer landsat-c2-l2)','low'),
# ---------------- MANTSALA
('Mantsala FI','2014-12','History','Original Yandex DC (YIT GC, 10,654 m2) commissioned Dec 2014; office Mar 2015','https://www.yit.fi/projektit/datakeskus','high'),
('Mantsala FI','2024-10-08','Mgmt','Expansion to 75 MW announced; target end-2025/early-2026','https://nebius.com/newsroom/nebius-to-triple-capacity-at-finland-data-center-to-75-mw','high'),
('Mantsala FI','2023-08..2024-09','Imagery','Main DC building (Moreenikatu 6) unchanged; no new structure SW of main hall',S2+' mantsala_fi crops','medium'),
('Mantsala FI','2025-05-11','Imagery','New construction (grey structure) SW of main hall',S2+' mantsala_fi_2025-05-11','medium'),
('Mantsala FI','2025-10-18','Imagery','Long new building SW of main hall with finished roof; persists 2026 (identity as Nebius expansion hall is likely but unconfirmed: could be adjacent utility/heat-recovery plant)',S2+' mantsala_fi_2025-10-18 ; strips/mantsala_fi_strip.png','low'),
('Mantsala FI','2026-03-31','Mgmt','Nebius: expansion to 75 MW "completed earlier this year"','site_list.csv (Nebius PR 2026-03-31)','high'),
('Mantsala FI','2026','Utility','Caverion EUR15M 70 MW substation for Mantsala DC#2; completion autumn 2027 (DC#2 in Kapuli, ~1 km away)','https://www.assemblincaverion.com/news-and-media/news/2026/caverion-finland-strengthens-its-position-in-the-data-centre-market-and-will-deliver-a-eur-15-million-substation-for-ai-cloud-company-nebius/','high'),
# ---------------- KANSAS CITY
('Kansas City MO','2024-11-19','Mgmt','Patmos (1601 McGee St) first tenant Nebius, 5 MW, go-live Q1 2025 (launched Apr 2025)','https://nebius.com/blog/posts/nebius-opens-first-availability-zone-in-us','high'),
('Kansas City MO','2023-2026','Imagery','No measurable change at 10 m (interior retrofit of existing 400k sf printing plant); rooftop gear not resolvable',S2+' kansascity_mo','high'),
('Kansas City MO','2025-09-23','Landlord','Patmos: facility at 35 MW total, 10 MW new colo available now; +20 MW in 2026','https://www.datacenterdynamics.com/en/news/patmos-adds-10mw-of-capacity-at-kansas-city-data-center/','high'),
('Kansas City MO','2026-09-17','Jobs','Nebius "Data Center IT Technician - Kansas City MO" posted (H200 cluster; onsite)','https://careers.nebius.com/?gh_jid=4975997101','medium'),
# ---------------- BIRMINGHAM
('Birmingham AL','2025-09','Land','Nebius acquires 75 acres at 201 Milan Pkwy (ex-Regions Lakeshore ops center)','https://nebius.com/birminghamal','high'),
('Birmingham AL','2026-01','Permit','Demolition permit for existing building','https://bhamnow.com/2026/05/18/40-million-permits-birmingham-data-center/','high'),
('Birmingham AL','2026-01-09','Jobs','DC Operations Manager - Critical Infrastructure (Birmingham AL / Independence MO) posted','https://careers.nebius.com/?gh_jid=4705231101','medium'),
('Birmingham AL','2026-01-16..04-20','Imagery','Existing Regions office building demolished; site mass-graded by Apr 20 2026',S2+' strips/birmingham_al_strip.png','high'),
('Birmingham AL','2026-03-26..04-09','Zoning','ZBA votes on substation/switching station; city attorney memo Apr 9: not "utility substations", no special exception needed; Nebius appeals denials (May 2026)','https://www.al.com/news/birmingham/2026/04/birmingham-walks-back-on-zoning-decisions-for-ai-factorys-power-infrastructure.html','high'),
('Birmingham AL','2026-05-11..15','Permit','Three building permits to Hoar Construction totaling ~$40M','https://bhamnow.com/2026/05/18/40-million-permits-birmingham-data-center/','high'),
('Birmingham AL','2026 (letter)','Utility','Nebius to city: 115 kV feed to Nebius-owned substation in service NLT Dec 31 2026; 230 kV switching station (APCO-owned) Oct 2027; ~32 x 400 kW temp gas gensets for ~3-month commissioning','https://abc3340.com/resources/pdf/9222f76a-3063-46e4-aee7-796b30ed611f-NebiusResponsestoCityLetter_Redacted.pdf','high'),
('Birmingham AL','2026-09-20','Imagery','Slab/steel framing over main pad + second pad to N; no closed white roof yet',S2+' birmingham_al_2026-09-20','medium'),
# ---------------- INDEPENDENCE
('Independence MO','2025-12-05','Land','Nebius closed on 398 acres (Eastgate, from NorthPoint)','https://www.independencemo.gov/sites/default/files/2025-12/EASTGATE%20UPDATE%20Presentation%2012-8.pdf','high'),
('Independence MO','2026-01-20','Utility','City pass-through power contracts: NextEra Jun 2026-May 2027 ($10.7M), Evergy Jun-Nov 2027; Blue Valley (IPP) plant reopening ~Oct 2027, 250 MW 2027 -> up to 1,100 MW 2029','https://www.govtech.com/artificial-intelligence/independence-mo-to-reopen-old-power-plant-for-ai-data-center','high'),
('Independence MO','2026-01-12..04-20','Imagery','Earthworks begin NE of Little Blue Pkwy/MO-78 interchange',S2+' strips/independence_mo_strip.png','medium'),
('Independence MO','2026-04-22 / 05-12','Mgmt','Site handed to GC ~Apr 20-22; groundbreaking May 12 2026; Phase 1 online ~Q3 2027','https://www.kshb.com/news/local-news/missouri/jackson-county/q-a-taking-independence-residents-concerns-to-nebius-vp-of-public-affairs ; https://ded.mo.gov/press-room/nebius-breaks-ground-gigawatt-scale-ai-factory-independence','high'),
('Independence MO','2026-07-17','Local','Residents allege night-work permit violation (ARCO Construction 4pm-4am notice)','https://fox4kc.com/news/independence-residents-allege-permit-violation-at-nebius-data-center-site/','medium'),
('Independence MO','2026-09-02','Imagery','Building pads/structures and bright surfaces on Phase 1 area (consistent with "going vertical by August")',S2+' independence_mo_2026-09-02','medium'),
# ---------------- BUTLER PA
('Butler Twp PA','2025-09','Zoning','Data center overlay adopted for original parcel','https://www.republicanherald.com/2026/06/26/ai-company-starts-development-at-highridge-business-park/','high'),
('Butler Twp PA','2026-05-01','Land','Nebius closes ~358 acres from NorthPoint, $187.5M','https://www.republicanherald.com/2026/06/26/ai-company-starts-development-at-highridge-business-park/','high'),
('Butler Twp PA','2025-10-16..2026-06-03','Imagery','New tree clearing/earthworks N of I-81 at Highridge interchange (identity as Nebius parcel probable, not confirmed against parcel 04-09-0037.001 map)',S2+' strips/butler_pa_strip.png','low'),
('Butler Twp PA','2026-08-31','Zoning','Supervisors 2-1 add SEDCO land to overlay (600-acre campus); final land-development plan still required','https://shensentinel.com/news/butler-supervisors-clear-hurdle-for-data-center-campus-expansion-at-highridge/','high'),
('Butler Twp PA','2027-10 (target)','Mgmt','Nebius site page: 260 MW initial phase operational by Oct 2027','https://nebius.com/schuylkillcountypa','high'),
# ---------------- BETHUNE
('Bethune FR','2025-10-30 / 12-04','Environmental','ICPE "cas par cas" filing Oct 30 2025; decision Dec 4 2025 (SIG Reverse Industry site, ex-Bridgestone, 575 av. George Washington)','https://www.pas-de-calais.gouv.fr/contenu/telechargement/87192/545817/file/PPVE%20SIG%20B%C3%A9thune_%20sommaire%20d%C3%A9taill%C3%A9%20des%20pi%C3%A8ces%20avec%20compl%C3%A9ments.pdf','high'),
('Bethune FR','2025-12-25','Imagery','New bright roofing on plant halls first visible (~3 ha vs ~2 ha through Aug 2025)',S2+' strips/bethune_fr_strip.png','medium'),
('Bethune FR','2026-02-04','Environmental','Porter a Connaissance filed (Azur Datacenter: data hall, cooling, power & backup gensets); works planned H1 2026','https://www.pas-de-calais.gouv.fr/contenu/telechargement/87090/545272/file/R002%201624293FTH%20V01%20-%20Porter%20%C3%A0%20Connaissance%20-%2004%2002%2026%20(1).pdf','high'),
('Bethune FR','2026-02-12','Mgmt','Announced 240 MW; first capacity late summer 2026; ~120 MW by YE2026','https://www.reuters.com/technology/amsterdam-based-ai-firm-nebius-build-240mw-data-centre-near-lille-france-2026-02-12/','high'),
('Bethune FR','2026-03..08','Imagery','Bright roof area grows ~4.3 ha (Mar) -> ~5.3 (May) -> ~6.5 ha (Aug 12)',S2+' bethune_fr detail_metrics.csv','medium'),
('Bethune FR','2026-05-07','Environmental','DREAL inspection report on the PAC with observations requiring response; operator reply Jul 2026; prefecture letter Jul 15 2026 ends preliminary review','https://www.pas-de-calais.gouv.fr/contenu/telechargement/87091/545277/file/20260506_SIG_BETHUNE_RAP_0007000473[1].pdf','high'),
('Bethune FR','2026-08-26..09-09','Environmental','Public e-participation (PPVE) on datacenter modification; prefect to issue arrete de prescriptions complementaires afterwards (not found as of Oct 4 2026); includes >=20 MW combustion & EU ETS (SEQE) documents','https://www.pas-de-calais.gouv.fr/Publications/Consultation-du-public/Participation-du-public-par-voie-electronique/Societe-SIG-REVERSE-INDUSTRY-a-Bethune','high'),
# ---------------- OTHERS
('Paris FR (Equinix PA10)','n/a','Imagery','Not measurable: dense urban colocation; Nebius cages inside existing Equinix building','-','high'),
('Keflavik IS (Verne)','2023-2026','Imagery','Chips/strip produced (strips/keflavik_is_strip.png); colocation in existing Verne campus - not interpreted for Nebius-specific change','-','low'),
('Longcross UK (Ark)','2023-2026','Imagery','Chips/strip produced; Ark LP02 new building (announced Jun 2026) could be tracked; not interpreted','-','low'),
('Lappeenranta FI','2023-2026','Imagery','AOI at city-level coordinate; Pajarila campus NOT located - strip marked _UNLOCATED; no interpretation','-','low'),
('Israel sites','n/a','Imagery','Not attempted (colocation in Mega/Or DCs; small; deprioritized)','-','n/a'),
]
wb=Workbook(); ws=wb.active; ws.title='Timeline'
hdr=['site','date','evidence type','observation','source URL / file','confidence']; ws.append(hdr)
for r in T: ws.append(list(r))
for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='305496')
for col,w in zip('ABCDEF',[18,18,18,95,70,11]): ws.column_dimensions[col].width=w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
# Vineland MW
m=wb.create_sheet('Vineland_MW')
m.append(['as of','gensets installed/operating (source)','genset nameplate MW (x1.982)','grid MW','gross power MW','gensets observed running (thermal)','running-gen output MW (@~80% load)','IT MW est. low (PUE 1.35)','IT MW est. high (PUE 1.2, nameplate-bound)','notes'])
rows=[('2025-10',16,None,5,None,None,None,None,None,'Floodlight/DataOne letter: 16 operating Oct 2025; tranche 1 Nov 2025'),
      ('2026-02',32,None,5,None,None,None,None,None,'32 operating; ~50 MW energized (secondary); tranche 2 Feb 2026'),
      ('2026-06 (end Q2)',62,None,5,None,25,None,None,None,'all 62 fired by June; ~25 running late June (thermal)'),
      ('2026-08/09 (Q3)',62,None,5,None,45,None,None,None,'>=45 running mid/late Aug; Bloom stop-ordered Aug 10, plan approved Aug 17')]
for i,(d,g,_,grid,__,run,___,____,_____,n) in enumerate(rows, start=2):
    m.append([d,g,f'=B{i}*1.982',grid,f'=C{i}+D{i}',run,(f'=F{i}*1.982*0.8' if run else None),
              (f'=(F{i}*1.982*0.8+D{i})/1.35' if run else f'=E{i}*0.8/1.35'),f'=E{i}/1.2',n])
m.append([]); m.append(['Method: power-bounded. Upper bound = all nameplate + grid / PUE 1.2 (no redundancy). Low = running units at ~80% load + grid / PUE 1.35. Tranche sizes are undisclosed, so this does not prove/disprove "all tranches delivered"; it bounds physical IT MW. Contracted ~300 MW (315 MW IT per secondary source).'])
m.append(['Cross-check: DataOne "at least 3 of 14 data halls operating" (early Aug 2026); if ~300 MW/14 halls ~21 MW/hall => ~65 MW, consistent with Q3 low-end.'])
for col in 'ABCDEFGHIJ': m.column_dimensions[col].width=20
m.column_dimensions['J'].width=70
# imagery metrics
im=wb.create_sheet('Imagery_metrics'); first=True
for f in sorted(glob.glob(f'{OUT}/imagery/*/detail_metrics.csv')):
    for r in csv.DictReader(open(f)):
        if first: im.append(list(r.keys())); first=False
        im.append(list(r.values()))
im2=wb.create_sheet('Imagery_AOI_metrics'); first=True
for f in sorted(glob.glob(f'{OUT}/imagery/[!_]*/metrics.json')):
    for r in json.load(open(f))['rows']:
        if not r.get('date'): continue
        keys=['site','period','date','scene','scene_cloud','aoi_cloud','aoi_snow','valid_frac','ndvi_mean','ndbi_mean','bsi_mean','builtbare_ha','brightroof_ha','chip']
        if first: im2.append(keys); first=False
        im2.append([r.get(k) for k in keys])
th=wb.create_sheet('Thermal'); first=True
for f in sorted(glob.glob(f'{OUT}/thermal/*_landsat_st.csv')):
    for r in csv.DictReader(open(f)):
        if first: th.append(list(r.keys())); first=False
        th.append(list(r.values()))
sa=wb.create_sheet('Sites_AOI'); sa.append(['site','lat','lon','half_m','cadence','coordinate basis'])
import importlib.util,sys; sys.path.insert(0,f'{OUT}/scripts'); from sites import SITES
for k,v in SITES.items(): sa.append([k,v['lat'],v['lon'],v['half_m'],v['cadence'],v['basis']])
wb.save(f'{OUT}/altdata_timelines.xlsx'); print('rows',len(T))
