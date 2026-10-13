import openpyxl, csv
Q=['Q3 2026','Q4 2026','Q1 2027','Q2 2027','Q3 2027','Q4 2027']
wb=openpyxl.load_workbook('cap/nebius_capacity.xlsx',data_only=True); ws=wb['Quarterly MW']
model={k:[ws.cell(r,c).value for c in range(15,21)] for k,r in {'low':119,'high':120,'base':121}.items()}
# ACTIVE adjustments vs model (MW IT/stated basis), per site: dict low/base/high -> deltas (proposed - model)
adj={}
def put(name,low,base,high,mlow,mbase,mhigh,basis):
    adj[name]=dict(d={'low':[a-b for a,b in zip(low,mlow)],'base':[a-b for a,b in zip(base,mbase)],'high':[a-b for a,b in zip(high,mhigh)]},
                   p={'low':low,'base':base,'high':high},basis=basis)
z=[0]*6
put('Vineland (shutdown low case)',[57,4.2,57,107,160,230],[66,81.8,118.5,168.5,205,240],[75,106.6,180,230,250,250],
    [57,57,57,107,160,230],[66,81.8,118.5,168.5,205,240],[75,106.6,180,230,250,250],
    'Low Q4 2026 = grid-only 4.2 MW IT if gensets are still off at Dec 31 (3-month shutdown case); base unchanged (prob-weighted Q4 impact ~$54M)')
put('Oklahoma colo (new row)',[5,10,25,50,75,100],[10,25,50,75,100,100],[25,50,100,120,150,150],z,z,z,
    'Q2 2026 letter map: OK = "Sites >100MW"; 20-F: OK colo contracts secured Feb 2026; 8 on-site roles since Mar 2 2026 (site manager Sep 30, liquid-cooling engineer Oct 2)')
put('Minnesota colo (new row)',[5]*6,[10]*6,[15]*6,z,z,z,
    'us-north1 Woodbury public region; likely IronGate MSP01 (15 MW live, fully leased); map dot <100 MW; jobs since Feb 20 2026')
put('Newport (Vantage CWL1) + UK 3rd (new row)',[0,0,5,10,20,30],[0,5,15,25,35,43],[5,15,30,43,43,43],z,z,z,
    'Jun 8 2026 PR: 3 new UK deployments = 65 MW fully ramped in 2027; minus Harlow 22 = 43 MW (Derived)')
put('Masmiyya (timing slip)',[0,0,0,11,22,22],[0,5.5,11,22,22,22],[0,22,22,22,22,22],
    [0,0,11,22,22,22],[0,11,16.5,22,22,22],[0,22,22,22,22,22],'Mega Or Q2 2026 presentation: Yoav construction completion Q4 2026')
put('Beit Shemesh (timing slip)',[0,0,0,20,39,58],[0,0,10,29,48.5,58],[0,39,58,58,58,58],
    [0,0,20,39,58,58],[0,19.5,39,48.5,58,58],[0,39,58,58,58,58],'Mega Or Q2 2026 presentation: Beit Shemesh completion Q2 2027')
put('AIB CLT1 (hall timing)',[0,0,0,0,0,0],[0,0,0,0,25,25],[0,0,0,0,25,50],
    [0,0,0,0,0,0],[0,0,0,0,12.5,25],[0,0,0,0,25,50],'AIB: ESA starts Oct 1 2026; two 25 MW halls ~month 10 / ~month 14 (~Aug and ~Dec 2027); parent: 0/25/50 by YE27')
act={k:[model[k][i]+sum(a['d'][k][i] for a in adj.values()) for i in range(6)] for k in model}
# CONNECTED (Estimated; gross/shell basis where sites built and powered but not yet active)
conn_sites={ # low, base, high per quarter
'Vineland (gross shell, 300)':([125,300,300,300,300,300],[125,300,300,300,300,300],[125,300,300,300,300,300]),
'Independence MO':([0,0,0,0,100,200],[0,100,200,200,200,400],[0,200,200,200,400,600]),
'Butler PA':([0]*6,[0,0,0,0,0,130],[0,0,0,0,130,260]),
'Birmingham AL':([0,0,0,25,50,100],[0,50,50,100,100,150],[0,100,100,150,200,300]),
'Bethune FR':([0,60,60,120,120,120],[0,120,120,120,180,240],[0,120,180,240,240,240]),
'Lappeenranta FI':([0,0,0,0,35,70],[0,0,0,35,70,70],[0,0,35,70,70,140]),
'Mantsala #1+#2':([75]*5+[75],[75,75,75,75,75,110],[75,75,75,110,145,145]),
'Kansas City':([20,20,20,20,20,40],[40]*6,[40]*6),
'Paris PA10 + Keflavik':([15]*6,[15]*6,[15]*6),
'Longcross + Harlow':([16,19,23.4,23.4,41,41],[16,41,41,41,41,41],[16,41,41,41,41,41]),
'Newport + UK 3rd':([0,0,5,10,20,30],[0,15,30,43,43,43],[0,30,43,43,43,43]),
"Israel (Modi'in, Masmiyya, Beit Shemesh)":([8,8,19,30,69,88],[8,30,52,69,88,88],[8,69,88,88,88,88]),
'Madrid + Tallinn':([6,9,29,40,40,40],[9,40,40,40,40,40],[18,40,40,40,40,40]),
'Minnesota':([5]*6,[15]*6,[15]*6),
'Oklahoma':([25,50,75,100,100,100],[50,100,100,150,150,150],[100,150,150,200,200,200]),
'AIB CLT1':([0,0,0,0,0,25],[0,0,0,0,25,50],[0,0,0,25,50,50]),
}
conn={k:[sum(v[j][i] for v in conn_sites.values()) for i in range(6)] for j,k in enumerate(['low','base','high'])}
mg={'low':[None,800,None,None,None,1800],'base':[None,900,None,None,None,1800],'high':[None,1000,None,None,None,2000]}
rows=[]
for i,q in enumerate(Q):
    for k in ['low','base','high']:
        pass
    rows.append(dict(quarter=q,metric='active (energized) MW',low=act['low'][i],base=act['base'][i],high=act['high'][i],
        model_low=model['low'][i],model_base=model['base'][i],model_high=model['high'][i],
        delta_base=round(act['base'][i]-model['base'][i],1),
        changed='YES' if any(abs(act[k][i]-model[k][i])>0.05 for k in model) else 'no',
        tag='Estimated',basis='mixed IT/stated (as nebius_capacity.xlsx identified-site rows) + adjustments in unseen_capacity.md',
        notes='; '.join(f"{n}: {a['d']['base'][i]:+g}" for n,a in adj.items() if abs(a['d']['base'][i])>0.05 or abs(a['d']['low'][i])>0.05 or abs(a['d']['high'][i])>0.05)))
    rows.append(dict(quarter=q,metric='connected MW',low=conn['low'][i],base=conn['base'][i],high=conn['high'][i],
        model_low='',model_base=ws.cell(125,15+i).value,model_high='',
        delta_base=round(conn['base'][i]-(ws.cell(125,15+i).value or 0),1),
        changed='YES (new bottom-up line; model has only mgmt figure + plug)',
        tag='Estimated',basis='gross/shell where built & powered but not active; mgmt guide YE26 800-1,000 MW connected (Reported, Q2 2026 call)',
        notes=f"mgmt guide: {mg['low'][i]}-{mg['high'][i]}" if mg['low'][i] else 'no mgmt figure'))
for r in rows:
    for k in ['low','base','high','model_low','model_base','model_high']:
        if isinstance(r[k],float): r[k]=round(r[k],1)
with open('/workspace/nebius/capacity/mw_path_proposed.csv','w',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
for r in rows: print(r['quarter'],r['metric'][:9],r['low'],r['base'],r['high'],'| model',r['model_low'],r['model_base'],r['model_high'],r['changed'][:3])
print({n:[conn_sites[n][1][1]] for n in conn_sites})
