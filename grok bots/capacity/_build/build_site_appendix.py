# Builds vineland_site_appendix.docx + .md (one page, landscape). All lines sourced from capacity files.
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import sys

BODY = float(sys.argv[1]) if len(sys.argv) > 1 else 9
SRC = float(sys.argv[2]) if len(sys.argv) > 2 else 8
OUT = '/workspace/nebius/capacity/vineland_site_appendix'

TITLE = "Appendix: Vineland, NJ (DataOne / Microsoft) site status, as of Thu Oct 8, 2026, ~9:00 AM ET"
SUB = ("Nebius Capacity & Power. Public sources only. Tags: R = Reported, D = Derived, E = Estimated. "
       "[n] = source list at the bottom.")
DEFS = [("Definitions [3]: ", "Connected = power delivered into fully built, equipped halls; it can sit idle until GPUs arrive (not the ~5 GW contracted, not billing). "
         "Active = power used by installed, operating IT. Billed = revenue-generating IT MW (active lagged ~1Q; Vineland counted by delivered tranche). "),
        ("Connected is not billing: ", "company YE26 ~941 connected [E] (guide 800 MW–1 GW [R]) vs ~258 billed [D] [17].")]

TR_HDR = ["Tranche", "Status", "Delivered / billing start (base)", "IT MW", "Power", "Src"]
TR = [
 ["T1", "Live", "Delivered Nov 2025 [R]", "26.7 E", "Gensets [R]", "1, 6"],
 ["T2", "Live", "Delivered Feb 2026 [R]", "26.7 E", "Gensets [R]", "1, 6"],
 ["T3", "Live", "Delivered Jul–Aug 2026 [R]", "26.7 E", "Gensets [R]", "2, 6"],
 ["T4", "Live", "Delivered Jul–Aug 2026 [R]", "26.7 E", "Gensets [R]", "2, 6"],
 ["T5", "Bloom-gated", "Not delivered [R]; bills Q2'27 [E]", "26.7 E", "Bloom [R]", "2, 9, 17"],
 ["T6", "Bloom-gated", "Not delivered [R]; bills Q2'27 [E]", "26.7 E", "Bloom [R]", "2, 9, 17"],
 ["T7", "Bloom-gated", "Not delivered [R]; bills Q3'27 [E]", "26.7 E", "Bloom [R]", "2, 9, 17"],
 ["T8", "Bloom-gated", "Not delivered [R]; bills Q3'27 [E]", "26.7 E", "Bloom [R]", "2, 9, 17"],
 ["T9", "Bloom-gated", "Not delivered [R]; bills Q3'27 [E]", "26.7 E", "Bloom [R]", "2, 9, 17"],
]
TR_NOTE = ("9 tranches in 2025–26 [R, 1]; T3–T4 = \"two tranches over the past month\" (Aug 12 letter) [R, 2]. Tranche MW not disclosed: 26.7 = 240 IT ÷ 9 equal split [E, 17]. "
           "Gensets: 16 running from Oct 2025, 32 from Feb 2026, all 62 fired by Jun 2026 [R, 6]. "
           "Missed dates give Microsoft tranche-termination / service-credit rights [R, 1]. "
           "Bloom deal: ~328 MW installed / ~250 MW guaranteed, \"operational during 2026\" [R, 9].")

LK_HDR = ["Locked number", "Value", "Tag", "Src"]
LK = [
 ["YE26 connected, gross shell (site split of 941)", "300 MW", "E", "17"],
 ["YE26 active IT, on gensets", "~82 MW", "D", "17"],
 ["Gap (only the gap; not a forecast)", "218 MW", "D", "17"],
 ["T1–T4 genset ceiling: (62 × 1.982 + ~5 grid) ÷ 1.2", "~107 MW IT", "D", "5, 6, 17"],
 ["MSFT rate: $17,392.9M ÷ 5 yrs ÷ 240 IT", "$14.5M/MW-yr", "D", "4, 17"],
 ["Vineland FY27 billed, v0.9e engine ($2,643.8M ÷ $14.5M)", "182.3 MW-yr", "D", "17"],
 ["Case B, Bloom slips 2Q: 2027 rev $1,159.5M → $193.3M", "−$966M @ ~35%", "D; p E", "17"],
]

BL_HDR = ["Bloom check (T5–T9)", "Finding", "Tag", "Src"]
BL = [
 ["Delivery/install/interconnect", "None found as of Oct 8 (no press after Oct 5)", "R absence", "8,10,15"],
 ["Aug 10 stop order (Bloom)", "No public lift; site plan approved 8-1 on Aug 17", "R; R abs.", "7, 8"],
 ["Bloom NJDEP permit", "None found (DataOne claims approvals; unverified)", "R absence", "8, 15"],
 ["Bloom Q2 PR; $1.7B IDF", "Jul 28 / Jul 16: no Vineland, COD or delivery date", "R", "9, 10"],
 ["Sentinel-2 Oct 6 v Sep 8", "+0.56 ha dark area (1.27→1.83); roof vs pad unresolvable at 10 m: NOT Bloom evidence", "D; read E", "16"],
 ["Bloom Q3'26 results", "Oct 27, 2026 (third-party calendar; not on IR)", "E", "11"],
]

NJ_HDR = ["Date (ET)", "NJDEP / permitting item", "Tag", "Src"]
NJ = [
 ["Sep 22", "$1,072,000 penalty order: 62 unpermitted CAT G3516H (1,982 kW)", "R", "5"],
 ["As of Oct 8", "Genset PCP / operating-cert application: none found (DataOne \"will apply\", Sep 23); no hearing request reported", "R absence", "5,8,14,15"],
 ["As of Oct 8", "Aug 10 stop-construction order: not lifted in any public record", "R absence", "7"],
 ["~Oct 12–15", "Hearing-request deadline (20 days after receipt; receipt date n/d)", "D", "14"],
 ["Oct 14", "Planning Board agenda (posted Oct 7): no DataOne item/resolution", "R", "13"],
 ["Oct–Nov", "Nebius Q3'26: late Oct–mid Nov (Oct 22 or Nov 10 per third-party calendars; not company-announced)", "E", "12"],
 ["~Nov 6–9", "45-day deadline to submit genset applications (Nov 9 if mailed) or stop", "D", "5"],
 ["As of Oct 8", "NJ BPU: no DataOne / Vineland docket (VMEU not BPU-regulated)", "R absence", "15"],
]

WHAT = [("What would change our view: ",
  "(1) Oct 27 Bloom Q3 call [E] or the Nebius Q3 letter (may precede the NJDEP deadline): Vineland delivery/COD, 5+ tranches, or a lift of the Aug 10 order cuts Case B odds; delay talk raises them [11, 12]. "
  "(2) ~Nov 6–9 NJDEP filing: an application lowers T1–T4 (~107 IT) shutdown risk; no filing makes the cease order live [5].")]

SOURCES = [
 "[1] Nebius FY25 20-F (Apr 30 2026) https://www.sec.gov/Archives/edgar/data/1513845/000110465926052948/nbis-20251231x20f.htm",
 "[2] Nebius Q2'26 letter (Aug 12 2026) https://assets.nebius.com/assets/4462517b-ce83-41f2-96ed-f2ac1bc06a05/SHLQ226%20%281%29.pdf ; Q2'26 call https://stockanalysis.com/stocks/nbis/transcripts/655490-q2-2026/",
 "[3] Nebius Q2'25 https://www.sec.gov/Archives/edgar/data/1513845/000110465925075028/tm2522866d1_ex99-2.htm ; Q3'25 https://www.sec.gov/Archives/edgar/data/1513845/000110465925109806/tm2530882d1_ex99-2.htm ; Q1'26 https://assets.nebius.com/assets/aa1bc2e6-df83-40cd-a6a2-95e7cda3d16c/Nebius%20SHL_Q1%202026.pdf",
 "[4] Nebius 6-K, Microsoft $17.4B (Sep 8 2025) https://www.sec.gov/Archives/edgar/data/1513845/000110465925088312/tm2525580d1_6k.htm",
 "[5] NJDEP order https://dep.nj.gov/wp-content/uploads/newsrel/2026/dataone75894-aonocapa.pdf ; https://www.nj.gov/governor/news/2026/20260922a.shtml ; ENR Sep 23 https://www.enr.com/articles/63694-nj-fines-data-center-1m-over-123-mw-of-unpermitted-generation",
 "[6] Genset history: https://www.enr.com/articles/63695-inside-the-vineland-data-center-power-dispute ; https://newjerseymonitor.com/2026/10/01/nj-data-center-polluted-air-for-nearly-a-year-before-regulators-intervened/ ; https://floodlightnews.org/new-jersey-dataones-illegal-still-running/",
 "[7] Stop orders: https://hntrbrk.com/breaking-news/vineland ; https://whyy.org/articles/vineland-data-center-stop-work-orders/",
 "[8] WHYY: https://whyy.org/articles/vineland-planning-board-approves-data-center-plan/ ; https://whyy.org/articles/vineland-data-center-fined/",
 "[9] Nebius–Bloom https://nebius.com/newsroom/nebius-and-bloom-energy-partner-to-power-ai-infrastructure-build-out ; IDF/Oaktree https://www.prnewswire.com/news-releases/industrial-development-funding-and-oaktree-announce-1-7-billion-project-investment-in-bloom-energy-fuel-cells-for-nebius-ai-infrastructure-build-out-302827566.html",
 "[10] Bloom Q2'26 https://investor.bloomenergy.com/press-releases/press-release-details/2026/Bloom-Energy-Reports-Record-Second-Quarter-2026-Financial-Results-and-Raises-Full-Year-2026-Guidance/default.aspx",
 "[11] https://www.marketbeat.com/earnings/reports/2026-10-27-bloom-energy-co-stock/",
 "[12] https://www.akrostec.com/indices/A100PAM/events/2026-10-22t00-00-00z-nebius-group-n-v-nbis-financial-results-da3e72aa ; https://earnings.fru.dev/companies/nebius",
 "[13] https://www.vinelandcity.org/Archive/Planning%20Board/Agendas/2026/Agenda%2010-14-26%20.pdf ; Council https://vinelandnj.api.civicclerk.com/v1/Events",
 "[14] https://morristownminute.org/1-07-million-penalty-vineland-data-center/",
 "[15] https://njems.nj.gov/DataMiner ; https://dep.nj.gov/docminer/ ; https://dep.nj.gov/boss/public-notices/ (all blocked Oct 8 ~8:35 AM ET) ; https://publicaccess.bpu.state.nj.us/",
 "[16] Sentinel-2 S2C_18SVJ_20261006 (Earth Search STAC); capacity/vineland_check/oct8/vineland_s2_sep08_vs_oct06.png",
 "[17] Team files: ye26_active_vs_connected.xlsx; nebius_capacity.xlsx v2; final_billed_vs_connected.xlsx; nebius_model.xlsx v0.9e (Judge_Slide A19, Contracts); vineland_watch.md §0b/§1c/§2; vineland_shutdown.xlsx",
]

# ---------------- docx ----------------
doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(BODY)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
pf = st.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width = Inches(11); sec.page_height = Inches(8.5)
for m in ('left_margin', 'right_margin'): setattr(sec, m, Inches(0.4))
sec.top_margin = Inches(0.3); sec.bottom_margin = Inches(0.3)
sec.header_distance = Inches(0.2); sec.footer_distance = Inches(0.2)

def para(runs, size=None, sa=1, sb=0, align=None, box=None):
    box = box if box is not None else doc
    p = box.add_paragraph()
    p.paragraph_format.space_after = Pt(sa); p.paragraph_format.space_before = Pt(sb)
    if align: p.alignment = align
    for txt, bold in runs:
        r = p.add_run(txt); r.bold = bold
        r.font.size = Pt(size or BODY)
    return p

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); s_ = OxmlElement('w:shd')
    s_.set(qn('w:val'), 'clear'); s_.set(qn('w:color'), 'auto'); s_.set(qn('w:fill'), hexcolor); tcPr.append(s_)

def fix_table(t, widths, l=40, r=40):
    tblPr = t._tbl.tblPr
    m = OxmlElement('w:tblCellMar')
    for k, v in (('left', l), ('right', r), ('top', 0), ('bottom', 0)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tblPr.append(m)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    tw = tblPr.find(qn('w:tblW'))
    if tw is None: tw = OxmlElement('w:tblW'); tblPr.append(tw)
    tw.set(qn('w:w'), str(int(sum(widths) * 1440))); tw.set(qn('w:type'), 'dxa')
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths): gc.set(qn('w:w'), str(int(w * 1440)))
    for row in t.rows:
        for c, w in zip(row.cells, widths): c.width = Inches(w)

def table(box, hdr, rows, widths, size=None, bold_first=False):
    t = box.add_table(rows=1 + len(rows), cols=len(hdr)); t.style = doc.styles['Table Grid']
    t.autofit = False
    fix_table(t, widths)
    for i, row in enumerate([hdr] + rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val); r.font.size = Pt(size or BODY)
            if i == 0: r.bold = True; shade(c, 'D9E2F3')
            elif bold_first and j == 0: r.bold = True
    return t

def heading(box, txt, first=False):
    p = box.paragraphs[0] if first else box.add_paragraph()
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(0 if first else 3)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(BODY + 0.5)
    return p

para([(TITLE, True)], size=BODY + 3, sa=0)
para([(SUB, False)], size=BODY - 0.5, sa=2)
para([(DEFS[0][0], True), (DEFS[0][1], False), (DEFS[1][0], True), (DEFS[1][1], False)], sa=2)

# two-column body via borderless outer table (robust in Word + LibreOffice)
CW = 5.0; GUT = 0.2
outer = doc.add_table(rows=1, cols=2); outer.autofit = False
fix_table(outer, [CW + GUT / 2, CW + GUT / 2], l=0, r=int(GUT / 2 * 1440))
L, R = outer.cell(0, 0), outer.cell(0, 1)

heading(L, "1. Microsoft tranches at Vineland: 4 of 9 live, T5–T9 Bloom-gated", first=True)
table(L, TR_HDR, TR, [0.52, 0.8, 1.9, 0.5, 0.68, 0.6], size=BODY - 0.5, bold_first=True)
para([(TR_NOTE, False)], size=BODY - 0.5, sa=1, sb=1, box=L)
heading(L, "2. Locked numbers")
table(L, LK_HDR, LK, [3.05, 0.95, 0.5, 0.5], size=BODY - 0.5)
heading(R, "3. Bloom fuel-cell timing: no evidence as of Oct 8", first=True)
table(R, BL_HDR, BL, [1.45, 2.55, 0.58, 0.42], size=BODY - 0.5)
heading(R, "4. NJDEP / permitting window (PI 75894)")
table(R, NJ_HDR, NJ, [0.68, 3.2, 0.57, 0.55], size=BODY - 0.5)
para([(WHAT[0][0], True), (WHAT[0][1], False)], sa=0, sb=3, box=R)

# collapse empty spacer paragraphs python-docx leaves after nested tables
from docx.shared import Pt as _Pt
from docx.enum.text import WD_LINE_SPACING
for cell in (L, R):
    for pp in cell.paragraphs:
        if not pp.text.strip():
            pp.paragraph_format.space_before = _Pt(0); pp.paragraph_format.space_after = _Pt(0)
            pp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; pp.paragraph_format.line_spacing = _Pt(2)
p = para([("Sources: ", True), (" · ".join(SOURCES).replace("https://", "").replace("http://", ""), False)], size=SRC, sa=0, sb=3)
for r in p.runs: r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
doc.save(OUT + '.docx')

# ---------------- md ----------------
def mdt(hdr, rows):
    out = ["| " + " | ".join(hdr) + " |", "|" + "---|" * len(hdr)]
    out += ["| " + " | ".join(r) + " |" for r in rows]
    return "\n".join(out)
md = [f"# {TITLE}", f"_{SUB}_", "", f"**{DEFS[0][0].strip()}** {DEFS[0][1]}**{DEFS[1][0].strip()}** {DEFS[1][1]}", "",
      "## 1. Microsoft tranches at Vineland: 4 of 9 live, T5–T9 Bloom-gated", mdt(TR_HDR, TR), "", TR_NOTE, "",
      "## 2. Locked numbers", mdt(LK_HDR, LK), "",
      "## 3. Bloom fuel-cell timing: no evidence as of Oct 8", mdt(BL_HDR, BL), "",
      "## 4. NJDEP / permitting window (PI 75894)", mdt(NJ_HDR, NJ), "",
      f"**{WHAT[0][0].strip()}** {WHAT[0][1]}", "",
      "## Sources", *[f"- {s}" for s in SOURCES], "",
      "_Built by `capacity/_build/build_site_appendix.py`; one page landscape (docx → PDF via LibreOffice)._"]
open(OUT + '.md', 'w').write("\n".join(md) + "\n")
print("ok")
