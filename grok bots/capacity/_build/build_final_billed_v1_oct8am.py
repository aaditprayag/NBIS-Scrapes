# Build final_billed_vs_connected.xlsx  (Nebius Capacity & Power, Thu Oct 8 2026)
import pickle, openpyxl, datetime
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as CL
data,basis=pickle.load(open('/tmp/fb.pkl','rb'))
Qs=["Q3'26","Q4'26","Q1'27","Q2'27","Q3'27","Q4'27"]
BLUE=Font(color='0000FF'); BOLD=Font(bold=True); FILL=PatternFill('solid',fgColor='DDEBF7'); RED=Font(color='C00000',bold=True)
Q3='late Oct–mid Nov (Oct 22 or Nov 10 per third-party calendars; not company-announced)'
# YE26 connected L/B/H (E) from ye26_active_vs_connected.xlsx (Oct 5), keyed by capacity site name
conn={'Vineland NJ (DataOne)':(300,300,300),'Independence MO (Eastgate)':(0,100,200),'Butler Twp PA (Highridge)':(0,0,0),
'Birmingham AL (BHM01, Oxmoor)':(0,50,100),'Bethune FR (Azur Datacenter, ex-Bridgestone)':(30,120,120),'Lappeenranta FI (Pajarila / Polarnode)':(0,0,0),
'Mantsala FI #1':(75,75,75),'Mantsala FI #2 (Kapuli)':(0,0,0),'Kansas City MO (Patmos)':(20,40,40),'Paris (Equinix PA10)':(5,5,5),'Keflavik IS (Verne)':(10,10,10),
'London Longcross (Ark LP01)':(16,19,19),'Harlow UK (Kao Data KLON-02/03)':(2.2,22,22),'Newport UK (Vantage CWL1)':(0,10,20),
'UK 3rd new deployment (lead: Green Mountain LON-East, Romford)':(0,5,10),"Modi'in IL (Mega Or)":(8,8,8),'Masmiyya IL (Mega DC)':(5.5,22,22),
'Beit Shemesh IL (Mega DC)':(0,0,0),'Madrid-Getafe ES (Merlin Edged)':(13.5,18,18),'Tallinn EE (Greenergy Huuru)':(11,22,22),
'Minnesota colo (us-north1, Woodbury)':(15,15,15),'Oklahoma colo (unnamed)':(50,100,150),'France 2nd colo (unnamed; possibly eu-west2)':None,
'AIB CLT1 (Moore, Spartanburg County SC; Lockhart Power)':(0,0,0)}
# YE27 connected: only where capacity is fully built/fixed (carry-forward, E); others n/d
ye27={'Vineland NJ (DataOne)':300,'Mantsala FI #1':75,'Paris (Equinix PA10)':5,'Keflavik IS (Verne)':10,'London Longcross (Ark LP01)':19,"Modi'in IL (Mega Or)":8}
meta={ # owned/leased, customer, confidence, reason, source
'Vineland NJ (DataOne)':('leased (build-to-suit)','Microsoft','Med',"4 of 9 tranches delivered (R) but MW/tranche is E (26.7 IT); YE26 billed 106.7 (billed-tranche basis) > energized 81.8; FY27 hinges on Bloom (Case B -$966M @~35%)","20-F; Q2'26 letter/call; nebius_capacity.xlsx Vineland tranche track; vineland_watch.md"),
'Independence MO (Eastgate)':('owned','unassigned (Meta #2 candidate, assumed)','High (billed) / Low (connected)',"No IT before ~Q3'27, so YE26 billed 0 is solid; YE26 connected 100 = utility feed/shell plug only","owned_campus_power_path.md; Independence city FAQs/deck"),
'Butler Twp PA (Highridge)':('owned','unassigned','High',"PPL 230 kV switchyard/laterals late-2027 to 2029; nothing bills through FY27","owned_campus_power_path.md; PPL filings"),
'Birmingham AL (BHM01, Oxmoor)':('owned','unassigned','Med',"115 kV NLT Dec 31 2026, 230 kV Oct 2027 (R); first revenue MW 2027; litigation Mar 2027","owned_campus_power_path.md; Alabama Power filings"),
'Bethune FR (Azur Datacenter, ex-Bridgestone)':('leased','Meta (likely, press)','Low-Med',"Mgmt says 120 MW by YE26 (R) vs model energized 30; ICPE decision pending; 1Q lag puts YE26 billed at 0","nebius_capacity.xlsx; CCO remarks; unseen_capacity.md"),
'Lappeenranta FI (Pajarila / Polarnode)':('owned','unassigned','Med',"First building H1'27; substation spring 2027","site_tracker_wed.xlsx; Fira/Polarnode"),
'Mantsala FI #1':('owned','self-serve (assumed)','High',"75 MW built early 2026 (R); IT-normalized at PUE 1.1 (E)","20-F; nebius_capacity.xlsx"),
'Mantsala FI #2 (Kapuli)':('owned','unassigned','Med',"Substation autumn 2027; nothing bills in FY27 on 1Q lag","site_tracker_wed.xlsx"),
'Kansas City MO (Patmos)':('leased','self-serve (assumed)','Med',"Operational; Nebius share of facility not disclosed (5-35 range)","nebius_capacity.xlsx; Patmos"),
'Paris (Equinix PA10)':('leased','self-serve (assumed)','Med',"Operational; 5 MW is third-party figure","site_tracker_wed.xlsx"),
'Keflavik IS (Verne)':('leased','self-serve (assumed)','High',"10 MW live since Mar 2025 (R)","site_tracker_wed.xlsx"),
'London Longcross (Ark LP01)':('leased','self-serve (assumed)','High',"16->19 MW contract (R); stated basis treated as gross (/1.3)","nebius_capacity.xlsx Basis"),
'Harlow UK (Kao Data KLON-02/03)':('leased','n/d','Med',"22 MW contract (R); ramp timing E","unseen_capacity.md; Kao Data"),
'Newport UK (Vantage CWL1)':('leased','n/d','Low',"Nebius MW not disclosed; UK residual 43 MW (D) split E","site_tracker_wed.md"),
'UK 3rd new deployment (lead: Green Mountain LON-East, Romford)':('leased (if confirmed)','n/d','Low',"Site identity unconfirmed (GM 14 MW to unnamed neocloud)","site_tracker_wed.md"),
"Modi'in IL (Mega Or)":('leased','n/d','High',"8 MW live Oct 2025 (R)","site_tracker_wed.xlsx"),
'Masmiyya IL (Mega DC)':('leased','n/d','Med',"22 MW contracted (R); ramp E","unseen_capacity.md"),
'Beit Shemesh IL (Mega DC)':('leased','n/d','Med',"58 MW; landlord completion Q2'27 (R)","site_tracker_wed.xlsx"),
'Madrid-Getafe ES (Merlin Edged)':('leased','n/d','Med',"18 MW contract (R); ramp E","unseen_capacity.md"),
'Tallinn EE (Greenergy Huuru)':('leased','n/d','Med',"22 MW first phase later 2026 (R)","unseen_capacity.md"),
'Minnesota colo (us-north1, Woodbury)':('leased','n/d','Low',"Tenant MW n/d; IronGate MSP01 inference","site_tracker_wed.md"),
'Oklahoma colo (unnamed)':('leased','n/d','Low',"Map '>100 MW' category only; town/landlord n/d","site_tracker_wed.md"),
'France 2nd colo (unnamed; possibly eu-west2)':('leased','n/d','n/d',"Region live (R) but site/MW not disclosed; not invented","site_tracker_wed.md"),
'AIB CLT1 (Moore, Spartanburg County SC; Lockhart Power)':('leased','n/d','Med',"50 MW critical IT (R, AIB PR Sep 30 2026); halls ~2027","AIB PR Sep 30 2026"),
}
sites=list(meta.keys())
wb=openpyxl.Workbook()
# ---------- README
rd=wb.active; rd.title='README'
lines=['Nebius Capacity & Power — FINAL billed IT MW vs connected MW (YE26 / FY27) — locked for Sunday',
f'Built: Thu Oct 8 2026 ~9:30 AM ET. Public sources only; nobody contacted. Tags: R Reported / D Derived / E Estimated.',
'FOR JUDGES: Connected is not billing. We estimate ~941 MW connected at YE26 (E; mgmt guide 800 MW–1 GW, R) but only ~258 MW of IT billing (D), about 27%. Revenue runs on billed MW.',
'',
'Company totals tie to locked anchors (see Final_Table checks): YE26 billed 257.7 base / 284.2 high; FY27 avg billed 475.0 / 645.4; YE26 energized IT 301.7; YE26 connected 941 base / 561 low.',
'NON-TIE flagged: site YE26 connected HIGH sums to 1,156 vs company 1,195 (39 MW residual = France 2nd n/d + range; not invented).',
'',
'Method (site billed, D): replicates the two locked methods exactly.',
"  YE26 billed (base/high) = capacity v2 row 147 method: other sites' IT-normalized energized at Q3'26 (1-quarter lag) + Vineland at billed tranches (4 delivered x 26.7 IT, E).",
"  FY27 avg billed (base/high) = model v0.9d method (Revenue!F12 'Active IT MW (average)', Judge_Slide B25): IT-normalized energized lagged 1Q, quarter-average of start/end, averaged over Q1-Q4'27 = (E_Q3'26 + 2E_Q4'26 + 2E_Q1'27 + 2E_Q2'27 + E_Q3'27)/8 per site.",
'  IT-normalized = energized / Gross-IT factor where the site figure is treated as gross; else as stated (nebius_capacity.xlsx Basis).',
'  Connected is on gross/shell (facility power) basis; billed is on IT basis. On a like-for-like IT basis, 941 connected ≈ 697–784 IT MW (÷1.20–1.35, D) — still ~2.7–3.0x billing.',
'',
f'Q3 results: {Q3} [E]. Short form: late Oct–mid Nov (est.). It may land before the ~Nov 6–9 NJDEP genset filing deadline.',
'Inputs in blue (static copies of Derived/Estimated values from capacity v2, ye26_active_vs_connected.xlsx, model v0.9d); formulas in black.',
'Tabs: Final_Table, Quarterly_IT, Definitions, Change_vs_Oct5, Sources.']
for i,l in enumerate(lines,1): rd.cell(i,1,l)
rd['A1'].font=Font(bold=True,size=13); rd['A3'].font=RED; rd.column_dimensions['A'].width=180
# ---------- Quarterly_IT
qi=wb.create_sheet('Quarterly_IT')
qh=['Site','Treated as','Gross/IT factor']+[f'Base IT {q}' for q in Qs]+[f'High IT {q}' for q in Qs]+['Tag','Source']
for j,h in enumerate(qh,1): qi.cell(1,j,h).font=BOLD
qrow={}
for i,s in enumerate(sites,2):
    t,f=basis[s]; qrow[s]=i
    qi.cell(i,1,s); qi.cell(i,2,t); c=qi.cell(i,3,f); c.font=BLUE; c.fill=FILL
    for k,key in enumerate(['base','high']):
        for j,v in enumerate(data[s][key]):
            c=qi.cell(i,4+k*6+j, round(v,4) if v is not None else 'n/d'); c.font=BLUE; c.fill=FILL
    qi.cell(i,16,'D (from E site energized paths)'); qi.cell(i,17,"nebius_capacity.xlsx 'Quarterly MW' energized base/high ÷ Basis factor")
vr=len(sites)+3
qi.cell(vr,1,'Vineland billed tranches (YE26)').font=BOLD
qi.cell(vr+1,1,'Tranches delivered by YE26 (T1 Nov-25, T2 Feb-26, T3-T4 Jul/Aug-26)'); c=qi.cell(vr+1,2,4); c.font=BLUE; c.fill=FILL; qi.cell(vr+1,3,'R (20-F; Q2-26 letter/call)')
qi.cell(vr+2,1,'IT MW per tranche (240 IT / 9, equal split)'); c=qi.cell(vr+2,2,'=240/9'); qi.cell(vr+2,3,'E (sizes not disclosed)')
qi.cell(vr+3,1,'Vineland YE26 billed IT MW'); qi.cell(vr+3,2,f'=B{vr+1}*B{vr+2}'); qi.cell(vr+3,3,'D')
VB=f'Quarterly_IT!$B${vr+3}'
qi.column_dimensions['A'].width=60
# ---------- Final_Table
ft=wb.create_sheet('Final_Table',1)
ft['A1']='FINAL: YE26 / FY27 billed IT MW vs connected MW by site (locked for Sunday)'; ft['A1'].font=Font(bold=True,size=13)
ft['A2']=f'Thu Oct 8 2026 ~9:30 AM ET. Billed = revenue-generating IT MW (D). Connected = built & equipped halls with power, gross/shell basis (E). Q3 results: late Oct–mid Nov (est.).'
ft['A3']='FOR JUDGES: Connected is not billing. ~941 MW connected at YE26 (E; guide 800 MW–1 GW, R) vs ~258 MW billing (D), about 27%. Revenue runs on billed MW.'; ft['A3'].font=RED
H=['Site','Owned/leased','Customer','YE26 connected LOW (E)','YE26 connected BASE (E)','YE26 connected HIGH (E)','YE26 energized IT base (ref, D)',
'YE26 billed IT BASE (D)','YE26 billed IT HIGH (D)','FY27 avg billed IT BASE (D)','FY27 avg billed IT HIGH (D)','YE27 connected (E; fixed sites only)',
'YE26 gap: connected base − billed base','Confidence','One-line reason','R/D/E tag','Source']
HR=5
for j,h in enumerate(H,1):
    c=ft.cell(HR,j,h); c.font=BOLD; c.alignment=Alignment(wrap_text=True)
r0=HR+1
for i,s in enumerate(sites):
    r=r0+i; q=qrow[s]; m=meta[s]
    ft.cell(r,1,s); ft.cell(r,2,m[0]); ft.cell(r,3,m[1])
    cv=conn[s]
    for j in range(3):
        c=ft.cell(r,4+j, cv[j] if cv else 'n/d'); c.font=BLUE; c.fill=FILL
    nd = data[s]['base'][0] is None
    if nd:
        for j in range(7,12): ft.cell(r,j,'n/d')
    else:
        ft.cell(r,7,f'=Quarterly_IT!E{q}')
        if s.startswith('Vineland'):
            ft.cell(r,8,f'={VB}'); ft.cell(r,9,f'={VB}')
        else:
            ft.cell(r,8,f'=Quarterly_IT!D{q}'); ft.cell(r,9,f'=Quarterly_IT!J{q}')
        ft.cell(r,10,f'=(Quarterly_IT!D{q}+2*Quarterly_IT!E{q}+2*Quarterly_IT!F{q}+2*Quarterly_IT!G{q}+Quarterly_IT!H{q})/8')
        ft.cell(r,11,f'=(Quarterly_IT!J{q}+2*Quarterly_IT!K{q}+2*Quarterly_IT!L{q}+2*Quarterly_IT!M{q}+Quarterly_IT!N{q})/8')
    c=ft.cell(r,12, ye27.get(s,'n/d'));
    if s in ye27: c.font=BLUE; c.fill=FILL
    ft.cell(r,13, f'=IF(ISNUMBER(E{r}),E{r}-IF(ISNUMBER(H{r}),H{r},0),"n/d")')
    ft.cell(r,14,m[2]); ft.cell(r,15,m[3])
    ft.cell(r,16,'Connected E (site split); billed D on E site inputs' + ('; Vineland tranches delivered R, MW/tranche E' if s.startswith('Vineland') else ''))
    ft.cell(r,17,m[4])
rl=r0+len(sites)-1; T=rl+1
ft.cell(T,1,'COMPANY TOTAL (sites; France 2nd n/d excluded)').font=BOLD
for j in range(4,14):
    L=CL(j); c=ft.cell(T,j,f'=SUM({L}{r0}:{L}{rl})'); c.font=BOLD
ft.cell(T,12,'n/d (company 1,882 base, see below)'); ft.cell(T,14,'see checks'); ft.cell(T,16,'E / D')
A=T+2
ft.cell(A,1,'LOCKED COMPANY ANCHORS and TIE CHECKS').font=BOLD
for j,h in enumerate(['Anchor','Locked value','Tag','Site-sum (formula)','Diff','Tie?','Source'],1): ft.cell(A+1,j,h).font=BOLD
anchors=[('YE26 billed IT BASE',257.7,'D (team lock; workbook 257.66)','H','vineland_watch.md §0c; capacity v2 row 147; model Judge_Slide B34'),
('YE26 billed IT HIGH',284.2,'D (team lock)','I','model Judge_Slide B35'),
('FY27 avg billed IT BASE',474.96,'D (team lock ~475)','J','model v0.9d Judge_Slide B25 (=SUMIF(Q_Yr,2027,MW_Avg)/4)'),
('FY27 avg billed IT HIGH',645.4,'D (team lock ~645)','K','model v0.9d Judge_Slide B26'),
('YE26 energized IT BASE',301.72,'D','G','capacity v2 Quarterly MW row 136'),
('YE26 connected BASE',941,'E','E','capacity v2 Quarterly MW row 142; unseen_capacity.md'),
('YE26 connected LOW',561,'E','D','capacity v2 row 141'),
('YE26 connected HIGH',1195,'E','F','capacity v2 row 143 — residual ~39 MW not attributed (France 2nd n/d + range)')]
for k,(n,v,t,L,src) in enumerate(anchors):
    r=A+2+k; ft.cell(r,1,n); c=ft.cell(r,2,v); c.font=BLUE; c.fill=FILL; ft.cell(r,3,t)
    ft.cell(r,4,f'={L}{T}'); ft.cell(r,5,f'=D{r}-B{r}'); ft.cell(r,6,f'=IF(ABS(E{r})<0.5,"TIE","NON-TIE (flagged)")'); ft.cell(r,7,src)
r=A+2+len(anchors)
extra=[('Mgmt YE26 connected guide (MW)','800–1,000','R',"Q1'26/Q2'26 letters; Q2'26 call"),
('YE26 energized IT LOW / HIGH (company)','189.5 / 414.0','D','capacity v2 rows 135/137'),
('YE27 connected LOW / BASE / HIGH (company)','1,249 / 1,882 / 2,517','E','capacity v2 rows 141-143, Q4 2027'),
('Contracted power (land + power), contrast only','~5 GW','R',"Company target; NOT 'connected'")]
for k,(n,v,t,src) in enumerate(extra):
    ft.cell(r+k,1,n); c=ft.cell(r+k,2,v); c.font=BLUE; c.fill=FILL; ft.cell(r+k,3,t); ft.cell(r+k,7,src)
r=r+len(extra)+1
ft.cell(r,1,'DERIVED RATIOS (D)').font=BOLD
ft.cell(r+1,1,'YE26 billed base ÷ connected base'); ft.cell(r+1,2,f'=H{T}/E{T}'); ft.cell(r+1,2).number_format='0%'
ft.cell(r+2,1,'YE26 connected base − billed base (MW)'); ft.cell(r+2,2,f'=E{T}-H{T}')
ft.cell(r+3,1,'YE26 connected base, IT-equivalent at ÷1.35 / ÷1.20'); ft.cell(r+3,2,f'=E{T}/1.35'); ft.cell(r+3,3,f'=E{T}/1.2')
ft.cell(r+4,1,'FY27 avg billed base ÷ YE27 connected base (1,882)'); ft.cell(r+4,2,f'=J{T}/1882'); ft.cell(r+4,2).number_format='0%'
r=r+6
ft.cell(r,1,'BASIS FLAGS (read before quoting site numbers)').font=RED
flags=["1) Vineland YE26 billed 106.7 (4 delivered tranches × 26.7 IT, E) > YE26 energized base 81.8 (genset-midpoint). The lock bills delivered tranches; if early tranches are smaller (Q2'26 top-customer 24% check), true billed is nearer ~82. Locked pair for the deck stays 300 connected vs ~82 active.",
"2) Vineland FY27 avg billed = 125.5 under the model's lock method (energized lagged 1Q). The judge slide's 'Vineland ~187 of ~475' uses tranche revenue ($1,551.5M + $1,159.5M) ÷ $14.5M, which is a different basis. On a consistent tranche basis the FY27 base total would be ~536 (=475.0 − 125.5 + 187.0, D). The lock stays 475; Modeling should reword the line as '~187 MW-equivalent of revenue', not as a share of 475.",
"3) Connected is gross/shell basis; billed is IT basis (PUE ~1.1–1.35). The gap partly reflects the basis difference. Even IT-equivalent connected (~697–784) is ~2.7–3.0x billed.",
"4) Site-level billed is Derived from Estimated site paths; only company totals are locks. Nebius discloses no per-site connected, active or billed MW."]
for k,fl in enumerate(flags): ft.cell(r+1+k,1,fl)
widths=[44,18,22,11,11,11,12,11,11,12,12,13,13,14,70,30,45]
for j,w in enumerate(widths,1): ft.column_dimensions[CL(j)].width=w
ft.freeze_panes=ft.cell(HR+1,2)
# ---------- Definitions
de=wb.create_sheet('Definitions')
for j,h in enumerate(['Term','Definition (one line)','Tag','Source'],1): de.cell(1,j,h).font=BOLD
defs=[('connected',"Power delivered into fully built and equipped data-center halls (shell + fit-out ready); can sit idle until GPUs arrive. NOT contracted land+power (~5 GW); NOT substation/grid access alone; NOT billing.",'R (definition); E (941 and site splits)',"Q2'25 letter ('fully provisioned… can be activated immediately upon GPU installation'); Q3'25 letter ('power connected to built data centers'); Q1'26 letter"),
('active','Power consumed by installed, operational IT.','R (definition); D (site MW)',"Q1'26 letter; nebius_capacity.xlsx"),
('billed','Revenue-generating IT MW (team metric): active IT lagged ~1 quarter for onboarding/acceptance; Vineland at delivered Microsoft tranches.','D','vineland_watch.md §0c; model v0.9d'),
('contracted (contrast)','Land plus power rights (~5 GW target). Not connected, not billing.','R','Company materials; Filings'),
('YE26 connected guide','800 MW–1 GW (management). Our 941 is an Estimated reconstruction inside the guide.','R (guide) / E (941)',"Q1'26 / Q2'26 letters; Q2'26 call")]
for i,d in enumerate(defs,2):
    for j,v in enumerate(d,1): de.cell(i,j,v)
de.column_dimensions['A'].width=22; de.column_dimensions['B'].width=110; de.column_dimensions['C'].width=28; de.column_dimensions['D'].width=80
# ---------- Change_vs_Oct5
ch=wb.create_sheet('Change_vs_Oct5')
for j,h in enumerate(['Item','Oct 5 value','Oct 8 final','Changed?','Note'],1): ch.cell(1,j,h).font=BOLD
chg=[('YE26 billed base / high','~258 / ~284 (257.7 / 284.2)','257.7 / 284.2','No','Now split by site (new); totals tie'),
('FY27 avg billed base / high','~475 / ~645','475.0 / 645.4','No','Now split by site (new); totals tie'),
('YE26 energized IT base / low / high','301.7 / 189.5 / 414.0','same','No',''),
('YE26 connected low / base / high','561 / 941 / 1,195 (E)','same','No','Site high sum 1,156: 39 MW residual still unattributed'),
('Vineland pair','300 connected vs ~82 active','same','No','218 is only the gap; ~107 IT = T1–T4 genset ceiling'),
('MSFT tranches','4 of 9 live; T5–T9 Bloom-gated','same','No','Oct 8 refresh: no new primary evidence'),
('Case B','≈ −$966M 2027 revenue @ ~35%','same','No',''),
('YE27 connected (company)','1,882 base (E)','same','No','Site-level only for fixed/fully built sites; others n/d'),
('Q3 results date','~Nov 10 (est.)',Q3,'Convention only','Team lock Oct 8: may precede Nov 6–9 NJDEP deadline'),
('NEW flag','—',"Vineland FY27 125.5 (lock method) vs ~187 (judge-slide tranche-revenue basis)",'Flag','Consistent tranche basis would give ~536 FY27 base; lock unchanged')]
for i,d in enumerate(chg,2):
    for j,v in enumerate(d,1): ch.cell(i,j,v)
for j,w in enumerate([34,34,50,16,70],1): ch.column_dimensions[CL(j)].width=w
# ---------- Sources
so=wb.create_sheet('Sources')
srcs=[('/workspace/nebius/capacity/nebius_capacity.xlsx (v2)','Quarterly MW energized base/high per site, Basis factors, rows 135-148 company series, Vineland tranche track'),
('/workspace/nebius/capacity/ye26_active_vs_connected.xlsx (Oct 5)','YE26 connected L/B/H site splits (E) summing to 941'),
('/workspace/nebius/capacity/site_tracker_wed.xlsx/.md','Owned/leased, customer, colo MW gaps, billing lock notes'),
('/workspace/nebius/capacity/owned_campus_power_path.md','Independence / Butler / Birmingham power path; connected = feed/shell not IT'),
('/workspace/nebius/capacity/gpu_mw_capex.xlsx','PUE / IT-gross conversion context'),
('/workspace/nebius/capacity/uncontracted_split.csv','Contracted vs uncontracted flags (context)'),
('/workspace/nebius/model/nebius_model_recalc_preview_v09d.xlsx','Judge_Slide B25/B26/B34/B35 locks; Revenue!F12 quarter-average method'),
('/workspace/nebius/capacity/vineland_watch.md (§0c, §1c Oct 8)','Locked billing path; Vineland status Oct 8'),
("Nebius Q2'25, Q3'25, Q1'26, Q2'26 shareholder letters; FY25 20-F",'Connected definition; YE26 guide 800 MW–1 GW; Microsoft 9 tranches')]
so.cell(1,1,'Source').font=BOLD; so.cell(1,2,'Used for').font=BOLD
for i,(a,b) in enumerate(srcs,2): so.cell(i,1,a); so.cell(i,2,b)
so.column_dimensions['A'].width=70; so.column_dimensions['B'].width=100
out='/workspace/nebius/capacity/_rebuild/final_billed_vs_connected.xlsx'
wb.save(out); print('saved',out,'total row',T,'anchor row',A)
