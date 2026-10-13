# Builds independence_tro_note.docx + .md (one page, landscape letter). Adapted from build_site_appendix.py.
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import sys

BODY = float(sys.argv[1]) if len(sys.argv) > 1 else 9
SRC = float(sys.argv[2]) if len(sys.argv) > 2 else 8
OUT = '/workspace/nebius/capacity/independence_tro_note'

TITLE = "Note: Independence, MO (Eastgate) TRO, hearing Oct 22, 2026. As of Thu Oct 8, 2026, ~11:15 PM ET"
SUB = ("Nebius Capacity & Power. Public sources only; nobody contacted. Tags: R = Reported, D = Derived, E = Estimated. "
       "[n] = source list at the bottom.")
BOTTOM = ("Bottom line: ",
  "Independence carries 0 MW billed at YE26 and only 4.6 base / 9.3 high MW-yr of FY27 billed IT [D, 9], 0.9% / 1.4% of the locked 531.8 / 659.2 [D]. "
  "Even a TRO that is granted and sustained only moves Phase 1 (~Q3'27) timing, so it cannot touch YE26 and caps the FY27 hit at those MW-yr. No lock changes.")

CASE_HDR = ["Item", "Detail", "Tag", "Src"]
CASE = [
 ["Plaintiffs", "Independence GUARD Alliance (Guardians United Against Reckless Development) + Bly Road / Little Blue Valley residents; counsel incl. Steve Jeffery (Jeffery Law Group)", "R", "1, 3"],
 ["Defendants", "City of Independence; Missouri ADC Holdings LLC (Nebius subsidiary, developer); Independence Power Partners LP (Blue Valley generator)", "R", "1, 3"],
 ["Court / filed", "Jackson County Circuit Court; 16-count, 52-page petition filed Aug 2026 with TRO + preliminary-injunction motion. Case number n/d (not in press; Case.net not checked)", "R; n/d", "3, 4, 5"],
 ["Claims", "Zoning, Sunshine Law and city-charter violations; non-enforcement of noise/hours/truck rules. Alleges 2022 rezoning without notice, 8 NDAs on power deal, >$150B tax-abated bond authority. City FAQ: site already zoned, no rezoning needed", "R (allegations)", "3, 4, 6"],
 ["Relief sought", "Pause construction; block further city bond issuance", "R", "4"],
 ["Hearing", "Thu Oct 22, 2026, 9:00 AM CT (10:00 AM ET), Eastern Jackson County Courthouse, Independence: plaintiffs' TRO motion AND defendants' motion to dismiss; evidence + witnesses", "R", "1, 2"],
 ["Bond", "A TRO may require a plaintiff bond; if plaintiffs cannot post it, no TRO issues and the suit continues. GUARD founder: \"good chance\" of TRO, \"also a chance we won't\"", "R", "2"],
 ["Responses", "Nebius: \"cannot comment on pending litigation… committed to complying with all applicable laws\". City: disagrees, no further comment", "R", "1, 2, 4"],
]

LK_HDR = ["Independence in our lock", "Value", "Tag", "Src"]
LK = [
 ["YE26 billed IT (no IT before ~Q3'27)", "0 MW", "D", "9"],
 ["YE26 connected plug (utility feed / shell only, not IT)", "100 MW (0–200)", "E", "9, 10"],
 ["Energized IT path: Q3'27 base / high (gross ÷ 1.35)", "37.0 / 74.1 MW", "D on E", "9"],
 ["FY27 avg billed = E_Q3'27 ÷ 8 (1Q lag; only Q3'27 counts)", "4.6 / 9.3 MW-yr", "D", "9"],
 ["Share of locked FY27 billed 531.8 / 659.2", "0.9% / 1.4%", "D", "9"],
 ["Q4'27 energized IT exit (FY28 matter) base / high", "92.6 / 148.1 MW", "D on E", "9"],
 ["Power path: IPL feed from Oct 2026; Blue Valley Ph1 COD Oct 2027", "200 / ~200–250 MW", "R", "6, 7, 8"],
]

SC_HDR = ["Oct 22 outcome", "Effect on our numbers", "FY27 billed base / high (MW-yr)", "Tag"]
SC = [
 ["TRO denied, or motion to dismiss granted", "None. Phase 1 ~Q3'27 holds", "531.8 / 659.2 (0 / 0)", "D"],
 ["TRO granted but bond not posted", "None (no TRO issues; suit proceeds)", "531.8 / 659.2 (0 / 0)", "D; R mech."],
 ["Short TRO (weeks) in Q4'26", "Shell work pauses; YE26 billed still 0. Phase 1 ~Q3'27 likely holds (slack to Blue Valley Oct 2027 COD)", "531.8 / 659.2 (0 / 0)", "D; timing E"],
 ["TRO/PI slips Phase 1 by ≥1 quarter (Q3'27 → Q4'27+)", "Independence FY27 billed → 0. Worst case for FY27 = full site", "527.2 / 649.9 (−4.6 / −9.3; −0.9% / −1.4%)", "D"],
 ["PI until case decided (multi-quarter)", "FY27 hit capped as above; risk moves to the YE27 exit 92.6 / 148.1 IT MW (FY28) and the YE26 100 MW connected plug [E]", "527.2 / 649.9 (capped)", "D"],
]
SC_NOTE = ("Gap to the 768.75 MW needed at $16M widens only 236.9 → 241.5 base / 109.6 → 118.9 high [D, 9]. "
           "Illustrative revenue at the ~$14.5M/IT MW-yr engine rate: ~$67M / ~$134M of FY27 [E, 9].")

WHAT = ("What would change our view: ",
  "(1) A TRO or PI that runs past ~Q1'27 or blocks IPP's Blue Valley work (IPP is a named defendant), which would push Phase 1 and the Oct 2027 Ph1 COD [6, 8]. "
  "(2) Any Nebius 6-K / Q3 letter naming Independence as a Meta #2 or other contracted site with 1H27 dates, which would raise the MW at stake [10]. "
  "(3) A city bond-issuance freeze, if it reaches project financing. Watch: Oct 22 ruling; Q3 letter (late Oct–mid Nov, E).")

SOURCES = [
 "[1] KSHB 41, \"Hearing set for Oct. 22 on TRO for Nebius data center project\" (Oct 8, 2026, 4:15 AM ET) https://www.kshb.com/news/local-news/hearing-set-for-october-22-on-temporary-restraining-order-for-nebius-data-center-project",
 "[2] FOX4 KC, \"Judge could order Nebius to halt data center construction in Independence\" (Oct 7, 2026, 3:23 PM ET; upd. 6:23 PM ET) https://fox4kc.com/news/judge-could-order-nebius-to-halt-data-center-construction-in-independence/",
 "[3] KSHB 41, \"Bly Road residents file lawsuit…\" (Aug 2026) https://www.kshb.com/news/local-news/missouri/jackson-county/independence-guard-alliance-files-lawsuit-to-halt-construction-of-independence-data-center",
 "[4] KCTV5 (Aug 13, 2026) https://www.kctv5.com/2026/08/13/lawsuit-aims-block-construction-independence-data-center/",
 "[5] suedatacenters.org case page (no case no.; Case.net not machine-accessible) https://suedatacenters.org/lawsuits/independence-missouri-guard-alliance-data-center-lawsuit",
 "[6] City of Independence data-center FAQs https://www.independencemo.gov/data-center-faqs",
 "[7] Nebius x Independence deck (Jan/Feb 2026) https://www.independencemo.gov/sites/default/files/2026-02/Nebius%20x%20Independence%2C%20MO.cleaned.pdf",
 "[8] MO DED groundbreaking (May 12, 2026) https://ded.mo.gov/press-room/nebius-breaks-ground-gigawatt-scale-ai-factory-independence ; KC Star / citizenportal on Blue Valley (via owned_campus_power_path.md)",
 "[9] Team files: final_billed_vs_connected.xlsx (Final_Table row 7, Quarterly_IT row 3; locked Oct 8); nebius_model.xlsx v0.9e Judge_Slide",
 "[10] owned_campus_power_path.md §1; site_verification.md §2; site_tracker_wed.xlsx; unseen_capacity.md; Nebius FY25 20-F (Meta #2 1H27)",
]

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(BODY)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
pf = st.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width = Inches(11); sec.page_height = Inches(8.5)
for m in ('left_margin', 'right_margin'): setattr(sec, m, Inches(0.4))
sec.top_margin = Inches(0.3); sec.bottom_margin = Inches(0.3)
sec.header_distance = Inches(0.2); sec.footer_distance = Inches(0.2)

def para(runs, size=None, sa=1, sb=0, box=None):
    box = box if box is not None else doc
    p = box.add_paragraph()
    p.paragraph_format.space_after = Pt(sa); p.paragraph_format.space_before = Pt(sb)
    for txt, bold in runs:
        r = p.add_run(txt); r.bold = bold; r.font.size = Pt(size or BODY)
    return p

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); s_ = OxmlElement('w:shd')
    s_.set(qn('w:val'), 'clear'); s_.set(qn('w:color'), 'auto'); s_.set(qn('w:fill'), hexcolor); tcPr.append(s_)

def fix_table(t, widths, l=40, r=40):
    tblPr = t._tbl.tblPr
    m = OxmlElement('w:tblCellMar')
    for k, v in (('left', l), ('right', r), ('top', 0), ('bottom', 0)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tblPr.append(m)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    tw = tblPr.find(qn('w:tblW'))
    if tw is None: tw = OxmlElement('w:tblW'); tblPr.append(tw)
    tw.set(qn('w:w'), str(int(sum(widths) * 1440))); tw.set(qn('w:type'), 'dxa')
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths): gc.set(qn('w:w'), str(int(w * 1440)))
    for row in t.rows:
        for c, w in zip(row.cells, widths): c.width = Inches(w)

def table(box, hdr, rows, widths, size=None, bold_first=False):
    t = box.add_table(rows=1 + len(rows), cols=len(hdr)); t.style = doc.styles['Table Grid']
    t.autofit = False; fix_table(t, widths)
    for i, row in enumerate([hdr] + rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val); r.font.size = Pt(size or BODY)
            if i == 0: r.bold = True; shade(c, 'D9E2F3')
            elif bold_first and j == 0: r.bold = True
    return t

def heading(box, txt, first=False):
    p = box.paragraphs[0] if first else box.add_paragraph()
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(0 if first else 3)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(BODY + 0.5)
    return p

para([(TITLE, True)], size=BODY + 3, sa=0)
para([(SUB, False)], size=BODY - 0.5, sa=2)
para([(BOTTOM[0], True), (BOTTOM[1], False)], sa=2)

CW = 5.0; GUT = 0.2
outer = doc.add_table(rows=1, cols=2); outer.autofit = False
fix_table(outer, [CW + GUT / 2, CW + GUT / 2], l=0, r=int(GUT / 2 * 1440))
L, R = outer.cell(0, 0), outer.cell(0, 1)

heading(L, "1. The case and the TRO (as reported)", first=True)
table(L, CASE_HDR, CASE, [0.8, 3.37, 0.43, 0.4], size=BODY - 0.5, bold_first=True)
heading(R, "2. Independence in the locked numbers", first=True)
table(R, LK_HDR, LK, [3.2, 1.0, 0.45, 0.35], size=BODY - 0.5)
heading(R, "3. What Oct 22 outcomes would (and would not) change")
table(R, SC_HDR, SC, [1.35, 2.0, 1.25, 0.4], size=BODY - 0.5, bold_first=True)
para([(SC_NOTE, False)], size=BODY - 0.5, sa=1, sb=1, box=R)
para([(WHAT[0], True), (WHAT[1], False)], sa=0, sb=3, box=L)

for cell in (L, R):
    for pp in cell.paragraphs:
        if not pp.text.strip():
            pp.paragraph_format.space_before = Pt(0); pp.paragraph_format.space_after = Pt(0)
            pp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; pp.paragraph_format.line_spacing = Pt(2)
p = para([("Sources: ", True), (" · ".join(SOURCES).replace("https://", ""), False)], size=SRC, sa=0, sb=3)
for r in p.runs: r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
doc.save(OUT + '.docx')

def mdt(hdr, rows):
    out = ["| " + " | ".join(hdr) + " |", "|" + "---|" * len(hdr)]
    return "\n".join(out + ["| " + " | ".join(r) + " |" for r in rows])
md = [f"# {TITLE}", f"_{SUB}_", "", f"**{BOTTOM[0].strip()}** {BOTTOM[1]}", "",
      "## 1. The case and the TRO (as reported)", mdt(CASE_HDR, CASE), "",
      "## 2. Independence in the locked numbers", mdt(LK_HDR, LK), "",
      "## 3. What Oct 22 outcomes would (and would not) change", mdt(SC_HDR, SC), "", SC_NOTE, "",
      f"**{WHAT[0].strip()}** {WHAT[1]}", "",
      "## Sources", *[f"- {s}" for s in SOURCES], "",
      "_Court-news refresh Oct 8 ~11:15 PM ET: nothing new after Oct 7 beyond KSHB Oct 8 [1] (already logged); FOX4 Oct 7 [2] adds bond/evidentiary detail._",
      "_Built by `capacity/_build/build_independence_note.py`; one page landscape letter (docx → PDF via LibreOffice)._"]
open(OUT + '.md', 'w').write("\n".join(md) + "\n")
print("ok")
