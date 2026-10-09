"""Task 5: insider selling cadence (Form 4 executed + Form 144 notices). Data from scripts/insiders.json (parse_insiders.py)."""
import json,os,sys
from collections import defaultdict
sys.path.insert(0,os.path.dirname(__file__))
from openpyxl.styles import Font,Alignment
D=json.load(open(os.path.join(os.path.dirname(__file__),'insiders.json')))
HFIA='https://www.sec.gov/about/divisions-offices/division-corporation-finance/holding-foreign-insiders-accountable-act-frequently-asked-questions'
def qtr(d):
    if '/' in d: m,dd,y=d.split('/')
    else: y,m,dd=d.split('-')
    return f"Q{(int(m)-1)//3+1}-{y[2:]}"
QS=[f"Q{q}-{y}" for y in ('24','25','26') for q in (1,2,3,4)][:11]  # Q1-24..Q3-26
def canon(n):
    n=n.upper()
    for k,v in [('VOLOZH','Arkady Volozh (CEO/Dir)'),('KOROLENKO','Andrey Korolenko (CIO-Infra)'),('BOYNTON','John Boynton IV (Dir/Officer; incl. trust)'),('BUNINA','Elena Bunina (Dir)'),('RYAN','Charles Ryan (Dir)'),('NAVE','Ophir Nave (Officer/Dir)'),('RADINSKY','Kira Radinsky (Dir)'),('ALONSO','Maria Alonso Sanchez (CFO)'),('BORODITSKY','Marc Boroditsky (CRO)'),('SHTAN','Danila Shtan (CTO)'),('TAL','Boaz Tal (GC)')]:
        if k in n: return v
    return n.title()
def f4_sales():
    rows=D['f4']
    amended={(r['owner'],r['date']) for r in rows if r['form']=='4A'}
    out=[r for r in rows if r['code']=='S' and (r['form']=='4A' or (r['owner'],r['date']) not in amended)]
    return out
EVENTS=[('2024-10-21','Trading resumes on Nasdaq'),('2024-10-31','Q3-24 results'),('2024-12-02','$700M private placement'),('2025-02-20','Q4-24 results'),
('2025-05-20','Q1-25 results'),('2025-06-02','$1.0B converts'),('2025-08-07','Q2-25 results'),('2025-09-08','Microsoft deal announced'),('2025-09-15','$1.15B equity + $3.16B converts'),
('2025-11-11','Q3-25 results / Meta deal'),('2026-02-12','Q4-25 results'),('2026-03-11','NVIDIA $2B'),('2026-03-18','HFIA Act: Section 16 applies'),('2026-03-20','$4.34B converts'),
('2026-05-13','Q1-26 results'),('2026-05-15','ATM (May-Jun) $2.85B'),('2026-07-17','$775M secured facility'),('2026-08-12','Q2-26 results'),('2026-08-24','$5.75B converts')]
def aggregates():
    s=f4_sales(); A=defaultdict(lambda:defaultdict(lambda:[0.0,0.0])); P=defaultdict(lambda:[0.0,0.0,0.0])
    for r in s:
        k=canon(r['owner']); q=qtr(r['date']); A[k][q][0]+=r['shares']; A[k][q][1]+=r['value'] or 0
        P[q][0 if r['kind']=='10b5-1 plan' else 1 if r['kind'].startswith('RSU') else 2]+=r['value'] or 0
    N=defaultdict(lambda:defaultdict(lambda:[0.0,0.0])); NP=defaultdict(lambda:[0.0,0.0])
    for r in D['f144']:
        k=canon(r['seller']); q=qtr(r['approx_date'] or r['filed']); N[k][q][0]+=r['shares']; N[k][q][1]+=r['aggval']
        NP[q][0 if r['plan10b51'] else 1]+=r['aggval']
    return A,P,N,NP
def write(wb,sheet,header,widths,url):
    A,P,N,NP=aggregates()
    ws=sheet('Insiders')
    ws['A1']='Insider selling cadence. Form 4 = EXECUTED sales (only exist from Mar-31-2026: Section 16(a) applied to foreign-private-issuer insiders from Mar-18-2026 under the HFIA Act). Form 144 = NOTICE of proposed sale (aggregate market value at notice; not proof of execution); available from Mar-2025.'; ws['A1'].font=Font(bold=True)
    ws['A2']='Source: SEC EDGAR submissions CIK0001513845 + each filing raw XML (URLs in detail tables). HFIA Act effective date: '+HFIA+' . 10% holders of FPIs are still exempt from Form 4. Form 4/A rows replace the original for the same owner+date.'
    r=4; ws.cell(row=r,column=1,value='A) FORM 4 EXECUTED SALES (code S) by insider by quarter, $M').font=Font(bold=True); r+=1
    header(ws,r,['Insider']+QS+['Total $M','Total shares']); r+=1
    for k in sorted(A,key=lambda k:-sum(v[1] for v in A[k].values())):
        ws.cell(row=r,column=1,value=k)
        for j,q in enumerate(QS,2): ws.cell(row=r,column=j,value=round(A[k][q][1]/1e6,2) if q in A[k] else None)
        ws.cell(row=r,column=len(QS)+2,value=round(sum(v[1] for v in A[k].values())/1e6,2)); ws.cell(row=r,column=len(QS)+3,value=int(sum(v[0] for v in A[k].values()))); r+=1
    for lab,i in [('of which 10b5-1 plan',0),('of which RSU tax sell-to-cover',1),('of which discretionary/other',2)]:
        ws.cell(row=r,column=1,value=lab)
        for j,q in enumerate(QS,2): ws.cell(row=r,column=j,value=round(P[q][i]/1e6,2) if q in P else None)
        r+=1
    r+=1; ws.cell(row=r,column=1,value='B) FORM 144 NOTICES of proposed sales by seller by quarter (approx. sale date), aggregate market value $M').font=Font(bold=True); r+=1
    header(ws,r,['Seller']+QS+['Total $M','Total shares']); r+=1
    for k in sorted(N,key=lambda k:-sum(v[1] for v in N[k].values())):
        ws.cell(row=r,column=1,value=k)
        for j,q in enumerate(QS,2): ws.cell(row=r,column=j,value=round(N[k][q][1]/1e6,2) if q in N[k] else None)
        ws.cell(row=r,column=len(QS)+2,value=round(sum(v[1] for v in N[k].values())/1e6,2)); ws.cell(row=r,column=len(QS)+3,value=int(sum(v[0] for v in N[k].values()))); r+=1
    for lab,i in [('of which flagged 10b5-1 (plan adoption date or remark)',0),('of which not flagged (sell-to-cover / vesting)',1)]:
        ws.cell(row=r,column=1,value=lab)
        for j,q in enumerate(QS,2): ws.cell(row=r,column=j,value=round(NP[q][i]/1e6,2) if q in NP else None)
        r+=1
    r+=1; ws.cell(row=r,column=1,value='C) KEY EVENTS OVERLAY (dates from 6-K exhibits; see Raise_to_Deploy and Guidance sheets for URLs)').font=Font(bold=True); r+=1
    header(ws,r,['Date','Quarter','Event']); r+=1
    for d,e in EVENTS: ws.cell(row=r,column=1,value=d); ws.cell(row=r,column=2,value=qtr(d)); ws.cell(row=r,column=3,value=e); r+=1
    r+=1; ws.cell(row=r,column=1,value='D) FORM 144 DETAIL').font=Font(bold=True); r+=1
    cols=['Filed','Seller','Relationship','Shares','Agg. market value $','Approx sale date','10b5-1?','Plan adoption date','Nature of acquisition','Remarks','Source URL']
    header(ws,r,cols); r+=1
    for x in D['f144']:
        for j,v in enumerate([x['filed'],x['seller'],x['rel'],int(x['shares']),x['aggval'],x['approx_date'],'Y' if x['plan10b51'] else 'N',x['plan_adopt'],x['acquired'],x['remarks'],x['url']],1): ws.cell(row=r,column=j,value=v)
        r+=1
    r+=1; ws.cell(row=r,column=1,value='E) FORM 4 TRANSACTION DETAIL (all codes; S=open-market sale, M=option/RSU exercise, A=grant)').font=Font(bold=True); r+=1
    cols=['Filed','Form','Owner','Title','Txn date','Code','Shares','Price $','Value $','Shares after','Classification','10b5-1 checkbox','Footnote (trunc.)','Source URL']
    header(ws,r,cols); r+=1
    for x in D['f4']:
        for j,v in enumerate([x['filed'],x['form'],x['owner'],x['title'],x['date'],x['code'],x['shares'],x['price'],x['value'],x['after'],x['kind'] if x['code']=='S' else '',x['plan_flag'],x['notes'][:250],x['url']],1): ws.cell(row=r,column=j,value=v)
        r+=1
    widths(ws,[34]+[11]*len(QS)+[12,12,40,40])
