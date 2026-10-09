import json,os,sys
sys.path.insert(0,os.path.dirname(__file__))
from language import FAM,POS_PRICE,NEG_PRICE,EXCL,Q
from openpyxl.styles import Font,Alignment
R=json.load(open(os.path.join(os.path.dirname(__file__),'language_results.json')))
def write(wb,sheet,header,widths,url):
    ws=sheet('Language')
    ws['A1']='Language text-mining: phrase-family counts per 10,000 words (alphabetic tokens; forward-looking-statement boilerplate removed). Calls = earnings-call transcripts; letters = shareholder letters (ex 99.2).'; ws['A1'].font=Font(bold=True)
    ws['A2']='Q1-24: no call, no letter (Yandex divestment period; trading halted to Oct-21-2024). Q2-24: no call. Q2-24..Q4-24 "letter" row uses the earnings PRESS RELEASE (no shareholder letter was published), so it is flagged and not comparable to later letters. Supply phrases are matched first and masked so "demand exceeds supply" is not double-counted as demand.'
    r=4
    for kind,label in [('call','EARNINGS CALL TRANSCRIPTS'),('letter','SHAREHOLDER LETTERS (Q2-24..Q4-24 = press release*)')]:
        ws.cell(row=r,column=1,value=label).font=Font(bold=True); r+=1
        header(ws,r,['Metric']+Q); r+=1
        docs={d['q']:d for d in R if d['kind'].strip('*')==kind or (kind=='letter' and d['kind']=='release*')}
        for lab,key in [('Words (alpha tokens)','words'),('Supply-constrained / sold out (per 10k)','supply_constrained_per10k'),('Demand / pipeline / customer interest (per 10k)','demand_pipeline_per10k'),('Pricing / price pressure / competition (per 10k)','pricing_competition_per10k'),('  of which price UP / premium phrases (per 10k)','pos_price_per10k'),('  of which price DOWN / pressure phrases (per 10k)','neg_price_per10k'),('Supply-constrained raw count','supply_constrained'),('Demand raw count','demand_pipeline'),('Pricing raw count','pricing_competition'),('Document date','date'),('Source URL','url')]:
            ws.cell(row=r,column=1,value=lab)
            for j,q in enumerate(Q,2):
                v=docs[q][key] if q in docs else ('not available' if key in ('url','date','words') else None)
                if q in docs and key=='url' and docs[q]['kind']=='release*': v='[press release] '+v
                ws.cell(row=r,column=j,value=v)
            r+=1
        r+=1
    ws.cell(row=r,column=1,value='PHRASE DICTIONARY (Python regex, case-insensitive; full code scripts/language.py)').font=Font(bold=True); r+=1
    for fam,pats in list(FAM.items())+[('price_UP subset',POS_PRICE),('price_DOWN subset',NEG_PRICE),('EXCLUDED (financing/share-price usages masked before counting)',EXCL)]:
        ws.cell(row=r,column=1,value=fam); ws.cell(row=r,column=2,value=' | '.join(pats)).alignment=Alignment(wrap_text=True); r+=1
    ws.cell(row=r+1,column=1,value='Raw mined text with source URL headers: /workspace/nebius/filings/altdata/mined_text/ and transcripts/. Per-document phrase hit lists: scripts/language_results.json.')
    widths(ws,[48]+[13]*10)
