from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
BLUE=Font(color='0000FF'); BOLD=Font(bold=True); HDR=Font(bold=True,color='FFFFFF'); FILL=PatternFill('solid',fgColor='1F3864')
wb=Workbook()
# ---------------- Inputs
I=wb.active; I.title='Inputs'
I.append(['Input','Value','Unit','Tag','Source'])
inputs=[
 ('A_out','Class A outstanding 6/30/26',238400165,'sh','R','Q2 PR outstanding shares; 6-K 8/12/26 acc 0001104659-26-094568'),
 ('B_out','Class B outstanding 6/30/26',33455053,'sh','R','Q2 PR'),
 ('Warr','NVIDIA pre-funded warrant (unexercised)',21065936,'sh','R','Q2FS Note 14; 6-K 8/12/26 acc 0001104659-26-094844'),
 ('IfC_old','If-converted shares, Q2 (anti-dilutive, weighted avg)',65416804,'sh','R','Q2FS Note 1 EPS'),
 ('P29x','Jun-2029 2.00% notes exchanged (orig. principal)',400,'$M','R','6-K 8/24/26 Ex.99.1 acc 0001104659-26-100347'),
 ('P31x','Jun-2031 3.00% notes exchanged (orig. principal)',400,'$M','R','same'),
 ('CR_old','Conversion rate Jun-2029/Jun-2031 notes',19.4363,'sh per $1,000','R','Jun-2025 offering 6-K; Q2FS Note 12'),
 ('XP29','Notes Exchange Price, 2029 notes',4403.18,'$ per $1,000','R','Form of Exchange Agreement, Ex.4.5 to 6-K 8/24/26'),
 ('XP31','Notes Exchange Price, 2031 notes',4413.18,'$ per $1,000','R','same'),
 ('C29','Coupon 2029 notes',0.02,'p.a.','R','notes terms'),
 ('C31','Coupon 2031 notes',0.03,'p.a.','R','notes terms'),
 ('Days','Accrued-interest days 6/5/26 to settlement (30/360)',76,'days','E','ASSUMPTION: 30/360, T+2 settlement 8/21/26; 79 days if 8/24 (adds ~744 sh)'),
 ('SP','Share Price = last reported sale price on pricing date 8/19/26',223.90,'$','R','Pricing PR, 6-K 8/20/26 acc 0001104659-26-098924'),
 ('X_co','Company-stated exchange shares ("approximately")',15800000,'sh','R','6-K 8/20 and 8/24/26: "approximately 15.8 million"'),
 ('P30','New 0.50% Feb-2030 notes, orig. principal',3450,'$M','R','6-K 8/24/26 (greenshoe fully exercised)'),
 ('CR30','Conversion rate 2030 notes',3.1902,'sh per $1,000','R','6-K 8/20/26 ($313.46)'),
 ('P34','New 4.50% Feb-2034 notes, orig. principal',2300,'$M','R','6-K 8/24/26'),
 ('CR34','Conversion rate 2034 notes',3.0802,'sh per $1,000','R','6-K 8/20/26 ($324.65)'),
 ('Post105','Share changes disclosed after 10/5/26',0,'sh','D','EDGAR re-check 10/8/26: no 6-K/424B/F-3/S-8/13G change; see edgar_recheck_2026-10-08.md'),
 ('RSU','Unvested RSUs 6/30/26 (memo, excluded)',6234091,'sh','R','Q2FS Note 13'),
 ('OPT','Share options outstanding 6/30/26 (memo, excluded)',6740600,'sh','R','Q2FS Note 13; WAEP $88.65'),
 ('A_iss','Class A issued 6/30/26',288585891,'sh','R','Q2FS Note 14'),
 ('Gross_new','New notes gross proceeds',5750,'$M','R','6-K 8/24/26'),
 ('Net_new','New notes net proceeds (company estimate, full greenshoe)',5680,'$M','R','6-K 8/20/26: "approximately $5.68 billion if... fully exercise" (exercised in full per 8/24)'),
 ('Acc30','2030 notes accretion at maturity',1.10,'x orig.','R','fact base / 6-K 8/20/26'),
 ('Acc34','2034 notes accretion at maturity',1.25,'x orig.','R','same'),
 ('Carry_cv','Convertible notes carrying value 6/30/26 (all series)',8499.0,'$M','R','Q2 PR balance sheet; Q2FS Note 12'),
 ('Carry29','Jun-2029 notes carrying 6/30/26 (487.5 principal)',504.1,'$M','R','Q2FS Note 12'),
 ('Carry31','Jun-2031 notes carrying 6/30/26 (487.5 principal)',498.4,'$M','R','Q2FS Note 12'),
 ('Rem_old','Jun-2029 / Jun-2031 principal outstanding at 6/30 (each)',487.5,'$M','D','500 - 12.5 converted in 2025 (FY25 20-F)'),
 ('Debt_cur','Debt, current 6/30/26',46.7,'$M','R','Q2 PR balance sheet'),
 ('Cash','Cash & equivalents 6/30/26 (excl. restricted 1,056.0)',8042.1,'$M','R','Q2 PR balance sheet'),
 ('Fac','Senior secured facility size (drawn amount ND)',775,'$M','R','6-K/A 7/17/26 acc 0001104659-26-084606'),
 ('Px','Reference share price (10/7/26 close)',237.15,'$','R','Investing.com 10/7/26 (third-party quote) - moneyness check only'),
 ('Pool_pct','AGM Item 13 authority',0.20,'of issued A+B','R','6-K 8/26/26 AGM results'),
]
name_row={}
for k,lab,v,u,t,s in inputs:
    I.append([lab,v,u,t,s]); r=I.max_row; name_row[k]=r; I.cell(r,2).font=BLUE
def R(k): return f"Inputs!$B${name_row[k]}"
for c in I[1]: c.font=HDR; c.fill=FILL
for col,w in zip('ABCDE',[58,16,16,6,90]): I.column_dimensions[col].width=w
for r in range(2,I.max_row+1): I.cell(r,2).number_format='#,##0.####'

# ---------------- Bridge
B=wb.create_sheet('FD_Bridge')
B.append(['#','Step','Shares','Tag','Source / math'])
rows=[
 ('1','Class A outstanding 6/30/26',f'={R("A_out")}','R','Q2 PR'),
 ('2','Class B outstanding 6/30/26',f'={R("B_out")}','R','Q2 PR'),
 ('3','NVIDIA pre-funded warrant (unexercised; $0.0001 strike)',f'={R("Warr")}','R','Q2FS Note 14'),
 ('4','If-converted, all convert series at 6/30 (Q2 weighted avg)',f'={R("IfC_old")}','R','Q2FS Note 1'),
 ('FD0','FD at 6/30/26 (team PT base) = rows 1-4','=SUM(C2:C5)','D','Should equal 358,337,958'),
 ('5','Less: conversion shares of $400M Jun-2029 notes exchanged',f'=-{R("P29x")}*1000*{R("CR_old")}','D','400,000 bonds x 19.4363'),
 ('6','Less: conversion shares of $400M Jun-2031 notes exchanged',f'=-{R("P31x")}*1000*{R("CR_old")}','D','400,000 bonds x 19.4363'),
 ('7','Plus: Aug exchange shares, 2029 notes',f'=ROUNDDOWN({R("P29x")}*1000*({R("XP29")}+1000*{R("C29")}*{R("Days")}/360)/{R("SP")},0)','D (day count E)','Ex.4.5: (Exchange Price x amount + accrued int.) / Share Price, rounded down'),
 ('8','Plus: Aug exchange shares, 2031 notes',f'=ROUNDDOWN({R("P31x")}*1000*({R("XP31")}+1000*{R("C31")}*{R("Days")}/360)/{R("SP")},0)','D (day count E)','same'),
 ('9','Plus: if-converted, new $3,450M 0.50% Feb-2030 notes',f'={R("P30")}*1000*{R("CR30")}','D (rate R)','3,450,000 bonds x 3.1902'),
 ('10','Plus: if-converted, new $2,300M 4.50% Feb-2034 notes',f'={R("P34")}*1000*{R("CR34")}','D (rate R)','2,300,000 bonds x 3.0802'),
 ('11','Post-10/5/26 changes disclosed',f'={R("Post105")}','D','None found on EDGAR / newsroom through 10/8 AM'),
 ('FD1','FD PRO FORMA after Aug-26 transactions = FD0 + rows 5-11','=C6+SUM(C7:C13)','D','Target ~376,648,956'),
]
for r in rows: B.append(list(r))
fd0=6; fd1=B.max_row
B.append([]); 
B.append(['','Change vs 6/30 FD (shares)',f'=C{fd1}-C{fd0}','D',''])
B.append(['','Change vs 6/30 FD (%)',f'=C{fd1}/C{fd0}-1','D',''])
B.append(['','Check: FD0 - 358,337,958 (should be 0)',f'=C{fd0}-358337958','D',''])
B.append(['','Check: FD1 - 376,648,956 (should be 0)',f'=C{fd1}-376648956','D',''])
B.append(['','Derived exchange shares total (rows 7+8)','=C9+C10','D','vs company "approximately 15.8M"'])
B.append(['','Inducement vs straight conversion (rows 5..8)','=C7+C8+C9+C10','D','extra shares over conversion value'])
B.append([])
B.append(['Memo','Alt: FD using company "~15.8M" instead of derived count',f'=C{fd1}-(C9+C10)+{R("X_co")}','R/D','Use if a reader wants only Reported exchange count'])
B.append(['Memo','Sensitivity: 79 accrual days (8/24 settlement) adds',f'=ROUNDDOWN({R("P29x")}*1000*({R("XP29")}+1000*{R("C29")}*79/360)/{R("SP")},0)+ROUNDDOWN({R("P31x")}*1000*({R("XP31")}+1000*{R("C31")}*79/360)/{R("SP")},0)-(C9+C10)','E','immaterial'])
B.append(['Memo','Class A outstanding pro forma (Q3 baseline, before ATM/RSU deliveries)',f'={R("A_out")}+C9+C10','D','Anything above at 9/30 => ATM/other issuance'])
B.append(['Memo','Point-in-time cross-check: old if-converted from principal x rates (pre-exchange)','=975000*19.4363+3162500*7.2072+2587500*5.4579+1750000*5.546','D','vs Q2 WA 65,416,804 (~0.2% diff)'])
B.append(['Memo','RSUs outstanding 6/30 (EXCLUDED from FD)',f'={R("RSU")}','R','Footnote: team FD excludes RSUs/options'])
B.append(['Memo','Options outstanding 6/30 (EXCLUDED; WAEP $88.65)',f'={R("OPT")}','R','Treasury-stock method not applied'])
B.append(['Memo','FD incl. all RSUs + options (gross upper bound, no TSM)',f'=C{fd1}+{R("RSU")}+{R("OPT")}','D','upper bound only'])
B.append(['Memo','Alt PT convention: new 2030/2034 notes OTM at ref. price -> treat as debt, drop their shares',f'=C{fd1}-C11-C12','D','pair with keeping new notes in debt (see PT_Consistency)'])
for c in B[1]: c.font=HDR; c.fill=FILL
for r in (fd0,fd1):
    for c in B[r]: c.font=BOLD
for col,w in zip('ABCDE',[6,74,18,16,70]): B.column_dimensions[col].width=w
for r in range(2,B.max_row+1): B.cell(r,3).number_format='#,##0;(#,##0)'
B.cell(fd1+3,3).number_format='0.00%'

# ---------------- PT consistency
P=wb.create_sheet('PT_Consistency')
P.append(['Item','$M','Tag','Treatment / math'])
prow=[
 ('Convertible notes carrying 6/30 (all old series)',f'={R("Carry_cv")}','R','Remove from net debt if all old series are if-converted (they are in the 65.4M)'),
 ('of which exchanged $800M (pro-rata carrying of Jun-29/Jun-31)',f'=({R("P29x")}/{R("Rem_old")})*{R("Carry29")}+({R("P31x")}/{R("Rem_old")})*{R("Carry31")}','D','Retired 8/24 for shares; no cash. Already gone post-Aug; do NOT also remove again'),
 ('Old converts remaining carrying post-exchange','=B2-B3','D','Remove from debt (if-converted)'),
 ('New notes original principal (face) - remove from debt when if-converted',f'={R("Gross_new")}','R','$3,450M + $2,300M; Q3 carrying ND'),
 ('New notes accreted principal at maturity (memo)',f'={R("P30")}*{R("Acc30")}+{R("P34")}*{R("Acc34")}','D','3,450x1.10 + 2,300x1.25; the debt claim if NOT converted'),
 ('Cash to KEEP: new notes net proceeds',f'={R("Net_new")}','R (co. estimate)','Add to 6/30 cash 8,042.1 for pro forma (ignores Q3 burn)'),
 ('Exchange cash flow','=0','R','Exchange was notes-for-shares; accrued interest settled in shares'),
 ('Capped calls / hedges','=0','R','None on any series (fact base / offering docs)'),
 ('Debt NOT removed: $775M secured facility (drawn ND)',f'={R("Fac")}','R','Keep as debt; full draw is an assumption [E]'),
 ('Debt NOT removed: current debt 6/30',f'={R("Debt_cur")}','R','Keep'),
 ('Pro forma cash (6/30 cash + new-note net proceeds)',f'={R("Cash")}+B7','D','Excl. restricted 1,056.0; excl. Q3 OCF/capex'),
]
for p in prow: P.append(list(p))
P.append([])
P.append(['Moneyness at reference price $237.15 (10/7 close) [R third-party]','Conv. price $','ITM?',''])
for nm,cp in [('Jun-2029 / Jun-2031 (remaining $87.5M each)',51.45),('Sep-2030 / Sep-2032',138.75),('Mar-2031',183.22),('Mar-2033',180.31),('NEW Feb-2030',313.46),('NEW Feb-2034',324.65)]:
    P.append([nm,cp,f'=IF({R("Px")}>B{P.max_row+1},"ITM","OTM")','Conv. prices R'])
P.append([])
P.append(['Rule','','','Shares and debt must match: if a series is if-converted, drop its debt; if treated as debt (e.g. OTM new notes), drop its shares. Team default = all if-converted (FD 376.6M) => remove ALL convert debt and keep new-note cash.'])
for c in P[1]: c.font=HDR; c.fill=FILL
for col,w in zip('ABCD',[70,14,16,95]): P.column_dimensions[col].width=w
for r in range(2,P.max_row+1): P.cell(r,2).number_format='#,##0.0'

# ---------------- Headroom
H=wb.create_sheet('Headroom')
H.append(['Item','(i) Exchange shares from TREASURY','(ii) Exchange shares NEWLY ISSUED','Tag / math'])
xs="FD_Bridge!C9+FD_Bridge!C10"
H.append(['Class A issued (post-exchange)',f'={R("A_iss")}',f'={R("A_iss")}+{xs}','R / D'])
H.append(['Treasury (6/30 = issued - outstanding; less exchange shares if from treasury)',f'={R("A_iss")}-{R("A_out")}-({xs})',f'={R("A_iss")}-{R("A_out")}','D'])
H.append(['AGM 20% pool = 20% x (Class A issued + Class B)',f'={R("Pool_pct")}*(B2+{R("B_out")})',f'={R("Pool_pct")}*(C2+{R("B_out")})','D'])
H.append(['Gross no-vote headroom','=B3+B4','=C3+C4','D'])
H.append(['Less: unexercised NVIDIA warrant',f'=-{R("Warr")}',f'=-{R("Warr")}','R'])
H.append(['NET no-vote headroom','=B5+B6','=C5+C6','D'])
H.append(['Memo: team 6/30 snapshot net headroom',93527979,93527979,'D (prior note)'])
H.append([])
H.append(['Reasoning','Exchange settled 8/24, one day BEFORE AGM authority (8/25). Source of shares (treasury vs new issue) ND. Q2 ATM shares came from treasury [R] -> treasury case (i) is the conservative ~78M. Case (ii) ~97M if newly issued under prior authority (raises issued capital and the 20% pool). Further reduced share-for-share by any Q3 ATM sales (ND) and RSU/option deliveries from treasury (ND, small).','',''])
for c in H[1]: c.font=HDR; c.fill=FILL
for col,w in zip('ABCD',[70,30,30,40]): H.column_dimensions[col].width=w
for r in range(2,9):
    for c in (2,3): H.cell(r,c).number_format='#,##0;(#,##0)'
H.cell(10,2).alignment=Alignment(wrap_text=True)
wb.save('/workspace/nebius/filings/fd_share_bridge.xlsx'); print('saved')
