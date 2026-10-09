import sys; sys.path.insert(0,'.')
from watchlist_rows import ROWS, COLS
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
# GUARD 10/8/26: the LOCKED q3_watchlist.xlsx/.md/.docx were finalized by in-place edits
# (Change_Log rows 3-6: billed-only MW test, path-keyed $22.7M/$24.51M break, Bitdeer basis note)
# that this script does NOT reproduce. A rerun writes to _rebuild/ only, never over the locked files.
import os
out='/workspace/nebius/filings/_rebuild/'
os.makedirs(out, exist_ok=True)
def esc(x): return str(x).replace('|','\\|').replace('\n',' ')
hdr = """# Q3 2026 watch-list: disclosures that CONFIRM or BREAK the NBIS short

**As of:** Mon Oct 5, 2026 (ET) · **Owner:** Filings & Capital · **Public information only**
**Tags:** **[R]** Reported · **[D]** Derived · **[E]** Estimated / team judgement · **ND** = not disclosed. **Every threshold is a team judgement [E]** unless marked [D]. Team locks (billing MW 257.7/284.2 YE26, 475/645 FY27 avg; long unlock ≥$22–24M/MW mid-term ACV with disclosed MW; comps band $9.7–16.7M/MW; MSFT 4 of 9 on gensets) are used as given.

**Q3 results date:** **not announced by Nebius as of 10/5/26** (nebius.com newsroom/financials checked). Third-party estimate **Tue Nov 10, 2026** [E]; one calendar shows Oct 22. Watch for the "announces date of third quarter 2026 results" PR (for Q2 it came out in July, ~3–4 weeks ahead).

**Where Q3 disclosures will land** (based on the Q2 pattern): 6-K with Ex.99.1 PR (financials, outstanding shares) + Ex.99.2 shareholder letter (ARR, deal economics, capacity, funding), and a separate 6-K with interim FS (RPO, deferred-revenue release, share capital/ATM, warrant, debt notes, concentration). Plus the earnings call (capex guide is call-only).

## Watch-list

"""
lines=["| "+" | ".join(COLS)+" |","|"+"---|"*len(COLS)]
for i,r in enumerate(ROWS,1):
    lines.append("| "+" | ".join([str(i)]+[esc(x) for x in r])+" |")
ftr = """

## Top 5 to read first on results day

1. **Sept ARR vs $4.6B / $5.5B** and whether the **$7–9B YE26 ARR** guide holds [E thresholds].
2. **Any active/billing MW figure.** Compare with the **257.7 / 284.2** YE26 locks. Silence again counts as confirm-lean.
3. **RPO at 9/30 vs $37.5B** (≤$39.5B confirm / ≥$42.5B break) and whether ">$40B" is still quoted.
4. **Class A count vs ~254.2M pro forma** and **ATM remaining vs 12,270,507**. Any gap means equity was used.
5. **Microsoft tranche language.** Count of tranches on permanent power, SLA credits, and the NJDEP genset deadline (~Nov 6) landing days before results.

## Source files
- `post630_capital_update.md` (share count / FD / headroom math) · `backlog_rpo_reconciliation.md` · `deferred_rpo_note.md` · `capacity/site_tracker_wed.md` · `capacity/vineland_watch.md` · `nebius_fact_base.md`
- Companion workbook: `q3_watchlist.xlsx` (one row per metric, same columns)
"""
open(out+'q3_watchlist.md','w').write(hdr+"\n".join(lines)+ftr)

wb=Workbook(); ws=wb.active; ws.title='Q3_Watchlist'
ws.append(COLS)
for i,r in enumerate(ROWS,1): ws.append([i]+list(r))
thin=Side(style='thin',color='BBBBBB')
for c in ws[1]:
    c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F3864'); c.alignment=Alignment(wrap_text=True,vertical='center')
widths=[4,26,34,48,40,44,44,40]
for j,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(j)].width=w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=Border(top=thin,bottom=thin,left=thin,right=thin)
    row[5].fill=PatternFill('solid',fgColor='FDE9E7'); row[6].fill=PatternFill('solid',fgColor='E8F3E8')
ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
n=wb.create_sheet('Notes')
for l in ["NBIS Q3'26 watch-list. As of Mon Oct 5, 2026 (ET). Public information only.",
          "Tags: [R] Reported, [D] Derived, [E] Estimated / team judgement, ND not disclosed. All thresholds are team judgement [E] unless marked [D].",
          "Q3 results date: NOT announced by Nebius as of 10/5/26; third-party estimate Nov 10, 2026.",
          "Key bases: FD 358,337,958 @6/30 [D] -> ~376.6M pro forma post-Aug [D]; Class A pro forma ~254,169,553 [D]; RPO $37,490.6M [R]; ARR $3.0B [R]; YE26 ARR guide $7-9B [R].",
          "Companion notes: post630_capital_update.md, backlog_rpo_reconciliation.md, q3_watchlist.md"]:
    n.append([l])
n.column_dimensions['A'].width=160
wb.save(out+'q3_watchlist.xlsx')
print('ok', len(ROWS))
