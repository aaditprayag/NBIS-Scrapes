import re, sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.section import WD_ORIENT
def add_runs(par, text):
    # handle **bold** and `code`
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
    for p in parts:
        if not p: continue
        if p.startswith('**') and p.endswith('**'): r=par.add_run(p[2:-2]); r.bold=True
        elif p.startswith('`') and p.endswith('`'): r=par.add_run(p[1:-1]); r.font.name='Consolas'
        else: par.add_run(p.replace('\\|','|'))
def split_row(l):
    l=l.strip().strip('|')
    cells=re.split(r'(?<!\\)\|', l)
    return [c.strip().replace('\\|','|') for c in cells]
def convert(md, out, landscape=False):
    doc=Document()
    st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10)
    if landscape:
        s=doc.sections[0]; s.orientation=WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = s.page_height, s.page_width
    for s in doc.sections: s.left_margin=s.right_margin=Inches(0.6)
    lines=open(md).read().split('\n'); i=0
    while i<len(lines):
        l=lines[i]
        if l.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                if not re.match(r'^\|\s*:?-{2,}', lines[i]): rows.append(split_row(lines[i]))
                i+=1
            ncol=max(len(r) for r in rows)
            t=doc.add_table(rows=len(rows), cols=ncol); t.style='Table Grid'
            for ri,r in enumerate(rows):
                for ci in range(ncol):
                    cell=t.cell(ri,ci); cell.text=''
                    p=cell.paragraphs[0]; add_runs(p, r[ci] if ci<len(r) else '')
                    for run in p.runs:
                        run.font.size=Pt(8 if ncol>5 else 9)
                        if ri==0: run.bold=True
            doc.add_paragraph()
            continue
        m=re.match(r'^(#{1,4})\s+(.*)', l)
        if m: doc.add_heading(m.group(2).replace('**',''), level=min(len(m.group(1)),3))
        elif re.match(r'^\s*[-*]\s+', l): add_runs(doc.add_paragraph(style='List Bullet'), re.sub(r'^\s*[-*]\s+','',l))
        elif re.match(r'^\s*\d+\.\s+', l): add_runs(doc.add_paragraph(style='List Number'), re.sub(r'^\s*\d+\.\s+','',l))
        elif l.strip()=='---': pass
        elif l.strip(): add_runs(doc.add_paragraph(), l.rstrip('  '))
        i+=1
    doc.save(out)
if __name__=='__main__':
  base='/workspace/nebius/filings/'
  for f,land in [('post630_capital_update',False),('backlog_rpo_reconciliation',True),('q3_watchlist',True)]:
    convert(base+f+'.md', base+f+'.docx', land); print('wrote',f+'.docx')
