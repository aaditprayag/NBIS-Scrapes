from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.section import WD_ORIENT
d=Document(); st=d.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(8.5)
s=d.sections[0]; s.orientation=WD_ORIENT.LANDSCAPE; s.page_width,s.page_height=s.page_height,s.page_width
for m in ('left_margin','right_margin'): setattr(s,m,Inches(0.45))
s.top_margin=s.bottom_margin=Inches(0.4)
def P(t,b=False,sz=None,i=False):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(t); r.bold=b; r.italic=i
    if sz: r.font.size=Pt(sz)
    return p
def T(rows,widths=None,sz=7.5):
    t=d.add_table(rows=len(rows),cols=len(rows[0])); t.style='Table Grid'
    for ri,row in enumerate(rows):
        for ci,v in enumerate(row):
            c=t.cell(ri,ci); c.text=''; r=c.paragraphs[0].add_run(str(v)); r.font.size=Pt(sz); r.bold=(ri==0) or str(v).startswith('**')
            if str(v).startswith('**'): r.text=str(v).strip('*')
            if widths: c.width=Inches(widths[ci])
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
P('Exhibit: Nebius capital structure and dilution capacity (NBIS short)',True,12)
P('Filings & Capital · locked figures as of Fri Oct 9, 2026 (ET) · public information only · [R] reported · [D] derived · [E] estimate · ND not disclosed',sz=7.5,i=True)
P('Takeaway: every long case needs 2.7–4.2× the share authority Nebius has without a new shareholder vote. Funding the build at disclosed contract prices means a fresh vote, a large equity raise, or more converts.',True,9.5)
P('1. Share count used for price targets',True,9)
T([['Tier','FD shares','Applies when','Tag / source'],
   ['Debt tier (all team PTs)','**358,558,306','PT below $313.46: Feb-30 / Feb-34 notes treated as debt','D · fd_share_bridge (6/30 FD 358,337,958 − exchanged conversion shares + 15.77M exchange shares)'],
   ['Feb-30 converts in the money','369,564,496','PT $313.46–$324.65','D · + 11,006,190 if-converted (3,450,000 × 3.1902)'],
   ['Both new notes in the money','376,648,956','PT above $324.65','D · + 7,084,460 if-converted (2,300,000 × 3.0802)']],[2.0,1.2,3.0,4.0])
P('2. Authority available without a new vote',True,9)
T([['Component','Shares','Tag','Source'],
   ['AGM 20% pool (items 13–14 passed 8/25/26)','64,408,189','D','20% × (Class A issued 288,585,891 + Class B 33,455,053); AGM results 6-K 8/26/26'],
   ['Treasury after the Aug-26 exchange (conservative: exchange shares from treasury)','34,416,338','D','50,185,726 at 6/30 [R] − 15,769,388 exchange shares [D]; settlement source ND'],
   ['Less: NVIDIA pre-funded warrant (unexercised)','(21,065,936)','R','Q2FS Note 14; NVIDIA 13G 7/20/26'],
   ['**Net authority available','**77,758,591 ≈ 77.8M','D','Upper case ~96.7M only if exchange shares were newly issued (ND); Q3 interim treasury count resolves it']],[3.4,1.3,0.5,5.0])
P('3. Authority needed vs available (shares of authority = new shares + convert reserve; model v0.9e full solves, Judge basis, debt tier)',True,9)
T([['Case','PT','Authority needed','× of 77.8M','New shares (memo)','Tag / source'],
   ['Case (a) funding gap','—','~215M','2.8×','—','D · q3_watchlist H16; ratio vs 77.8M [D]'],
   ['A4 long case (every uncontracted MW at $24M, base path)','$223 (below $237.15 spot)','229.7M','3.0×','184.1M','D · model v0.9e Judge basis; $24M input E'],
   ['D3 long case (every uncontracted MW at $24M, 659.2 high path): stress case, not a path','$281','225.7M','2.9×','189.5M','D · model v0.9e Judge basis; $24M input E'],
   ['High-path mixed routes (part at $16.7M band top, rest above Bitdeer ceiling)','spot to D3','~226–273M','2.9–3.5×','—','D · model v0.9e; band top P (Lambda), Bitdeer D/basis P'],
   ['Bitdeer IT-equivalent $24.6M/MW','$239.51 base / $301.47 high','212.5M / 208.7M','~2.7×','170.0M / 174.9M','D · ≤$22.4M D, basis P; IT read D'],
   ['Lambda generous gross reading $21.7M/MW','$156.78 base / $204.21 high','325.7M / 322.0M','4.1–4.2×','~261–272M','D · Lambda $16.67M P; gross reading D']],[3.0,1.6,1.3,0.75,1.15,1.8])
P('4. Convertible notes outstanding after the Aug-26 exchange and new issue',True,9)
T([['Series','Principal ($M)','Conv. price','Point-in-time if-converted shares','Tag'],
   ['2.00% Jun-2029 / 3.00% Jun-2031 (after $400M each exchanged)','87.5 / 87.5','$51.45','1,700,676 each','R principal / D shares'],
   ['1.00% Sep-2030 / 2.75% Sep-2032','1,581.25 / 1,581.25','$138.75','11,396,385 each','R / D'],
   ['1.25% Mar-2031 / 2.625% Mar-2033','2,587.5 / 1,750.0','$183.22 / $180.31','14,122,316 / 9,705,500','R / D'],
   ['0.50% Feb-2030 (new, Aug-26)','3,450.0','$313.46','11,006,190','R / D'],
   ['4.50% Feb-2034 (new, Aug-26)','2,300.0','$324.65','7,084,460','R / D'],
   ['**Total','**13,425.0','','**68,112,589','D. Older notes $7,675.0M face; claim at maturity $9,051.9M [D]']],[3.6,1.4,1.3,2.2,1.7])
P('Other debt: $775M senior secured facility signed 7/10/26 (SOFR + 2.50%, matures 10/31/30); drawn amount ND [R, 6-K/A 7/17/26]. ATM: 12,270,507 shares of capacity left at 6/30 [R]; Q3 sales ND. FD excludes RSUs/options. The FD tiers use the reported Q2 weighted-average if-converted count (65,416,804) less exchanged notes, so they differ slightly from the point-in-time total above.',sz=7.5)
P('Sources: Q2\'26 PR and interim FS (6-K 8/12/26, acc 0001104659-26-094568 / -094844); AGM results 6-K 8/26/26 (tm2624077d1_ex99-1); convert pricing 6-K 8/20/26 and closing + Ex.4.5 6-K 8/24/26 (acc 0001104659-26-100347); NVIDIA 13G 7/20/26; EDGAR re-checked 10/9/26 6:25 PM ET: no new issuer filing, no Q3 date release. Workpapers: fd_share_bridge.xlsx, appendix_filings_capital_final.docx.',sz=7,i=True)
d.save('capital_structure_exhibit.docx'); print('saved')
