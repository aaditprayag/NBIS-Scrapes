from content import *
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
# summary table rows
SUM=[(i+1,L['t'].split(':')[0],L['p'],L['e']) for i,L in enumerate(LEGS)]
md=[f"# {TITLE}",f"*{META_LINE}*","",TAGS,"","## The hurdle (what's priced in)"]+[f"- {s}" for s in SETUP]
md+=["","## Ranked legs (expected $/share = impact × probability; overlapping, not additive)","","| # | Leg | Probability | Expected impact |","|---|---|---|---|"]+[f"| {a} | {b} | {c} | {d} |" for a,b,c,d in SUM]
for i,L in enumerate(LEGS):
    md+=["",f"## {i+1}. {L['t']}",f"**Claim.** {L['claim']}","","**Evidence.**"]+[f"- {x}" for x in L['ev']]+["",f"**Quantification.** {L['q']}",f"**Probability / expected.** {L['p']} → {L['e']}",f"**Bull counterpoint.** {L['bull']}",f"**What would change our mind / watch.** {L['watch']}"]
md+=["","## Five claims to test","","| # | Claim | Owner | Evidence needed | Timeline / kill criteria |","|---|---|---|---|---|"]+[f"| {i+1} | {c['c']} | {c['o']} | {c['ev']} | {c['tl']} |" for i,c in enumerate(CLAIMS)]
md+=["","## Gaps / caveats"]+[f"- {g}" for g in GAPS]
md+=["","Files: `01_priced_in.xlsx/.md` (corrected Oct 4), `03_depreciation.xlsx` (reverse DCF, cohort model), `charts/03c_reverse_dcf_yield.png`."]
open('/workspace/nebius/skeptic/02_bear_outline.md','w').write("\n".join(md)+"\n")
d=Document(); st=d.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10)
d.add_heading(TITLE,0); p=d.add_paragraph(META_LINE); p.runs[0].italic=True
d.add_paragraph(TAGS).runs[0].font.size=Pt(8.5)
d.add_heading("The hurdle (what's priced in)",1)
for s in SETUP: d.add_paragraph(s,style='List Bullet')
d.add_heading("Ranked legs",1)
d.add_paragraph("Expected $/share = impact × probability (Red Team judgment). Legs 1, 2 and 5 share one equation (yield × life), so the figures are not additive.")
t=d.add_table(rows=1,cols=4); t.style='Light Grid Accent 1'
for c,h in zip(t.rows[0].cells,["#","Leg","Probability","Expected impact"]): c.text=h
for r in SUM:
    cs=t.add_row().cells
    for c,v in zip(cs,r): c.text=str(v)
def lab(par,label,text):
    r=par.add_run(label+" "); r.bold=True; par.add_run(text)
for i,L in enumerate(LEGS):
    d.add_heading(f"{i+1}. {L['t']}",1)
    lab(d.add_paragraph(),"Claim.",L['claim'])
    d.add_paragraph().add_run("Evidence.").bold=True
    for x in L['ev']: d.add_paragraph(x,style='List Bullet')
    lab(d.add_paragraph(),"Quantification.",L['q'])
    lab(d.add_paragraph(),"Probability / expected.",f"{L['p']} → {L['e']}")
    lab(d.add_paragraph(),"Bull counterpoint.",L['bull'])
    lab(d.add_paragraph(),"Watch.",L['watch'])
d.add_heading("Five claims to test",1)
t=d.add_table(rows=1,cols=5); t.style='Light Grid Accent 1'
for c,h in zip(t.rows[0].cells,["#","Claim","Owner","Evidence needed","Timeline / kill criteria"]): c.text=h
for i,c in enumerate(CLAIMS):
    cs=t.add_row().cells
    for cell,v in zip(cs,[str(i+1),c['c'],c['o'],c['ev'],c['tl']]):
        cell.text=v
        for pp in cell.paragraphs:
            for rr in pp.runs: rr.font.size=Pt(8.5)
d.add_heading("Gaps / caveats",1)
for g in GAPS: d.add_paragraph(g,style='List Bullet')
d.add_picture('/workspace/nebius/skeptic/charts/03c_reverse_dcf_yield.png',width=Pt(470))
d.save('/workspace/nebius/skeptic/02_bear_outline.docx'); print('ok')
