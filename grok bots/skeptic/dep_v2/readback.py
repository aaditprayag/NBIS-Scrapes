import json,subprocess,shutil,os,sys
from openpyxl import load_workbook
def recalc(src="/workspace/nebius/skeptic/03_depreciation.xlsx"):
    shutil.rmtree("/tmp/lo",ignore_errors=True)
    subprocess.run(["soffice","--headless","--calc","--convert-to","xlsx","--outdir","/tmp/lo",src],capture_output=True,timeout=180)
    return load_workbook("/tmp/lo/"+os.path.basename(src),data_only=True)
if __name__=="__main__":
    wb=recalc()
    for sh in sys.argv[1:]:
        ws=wb[sh]; print("==",sh)
        for r in ws.iter_rows(values_only=True):
            if any(v is not None for v in r): print([round(v,3) if isinstance(v,float) else v for v in r][:22])
