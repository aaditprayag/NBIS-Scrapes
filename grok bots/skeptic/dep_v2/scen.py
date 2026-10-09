import json,subprocess,shutil,os
from openpyxl import load_workbook
C=json.load(open("cells.json"))
SRC="/workspace/nebius/skeptic/03_depreciation.xlsx"
def run(changes,extra=None):
    wb=load_workbook(SRC)
    for k,v in changes.items():
        sh,cell=C[k].split("!"); wb[sh][cell.replace("$","")].value=v
    if extra:
        for sh,cell,v in extra: wb[sh][cell].value=v
    wb.save("/tmp/scen.xlsx"); shutil.rmtree("/tmp/los",ignore_errors=True)
    subprocess.run(["soffice","--headless","--calc","--convert-to","xlsx","--outdir","/tmp/los","/tmp/scen.xlsx"],capture_output=True,timeout=180)
    w=load_workbook("/tmp/los/scen.xlsx",data_only=True)
    g=lambda ref: w[ref.split("!")[0]][ref.split("!")[1].replace("$","")].value
    an=w["Annual"]; rd=w["Reverse_DCF"]; M=w["Model"]; t=C['rd_tab']
    out=dict(rv27=g(C['rv27']),gap27=g(C['gap27']),gapps=g(C['gapps']),ms27=g(C['ms27']),ms28=g(C['ms28']),mspv=g(C['mspv']),
      y1=M[f"B{C['r_y1']}"].value,y2=M[f"B{C['r_y2']}"].value,ymw=M[f"B{C['r_ymw']}"].value,rmw=M[f"B{C['r_rmw']}"].value,
      gapq=M[f"B{C['r_gapq']}"].value,peak=M[f"B{C['r_peak']}"].value,
      bu=[an.cell(C['A_bu'],2+j).value for j in range(5)],
      ebm={L:[an.cell(C['A_ebit'][L]+3,2+j).value for j in range(5)] for L in ("L1","L2","L3")},
      pts=[an.cell(C['A_pts'],2+j).value for j in range(5)],
      fcf=[an.cell(C['A_fcf'],2+j).value for j in range(5)],fcfx=[an.cell(C['A_fcfx'],2+j).value for j in range(5)],liq=[an.cell(C['A_liq'],2+j).value for j in range(5)],
      rd=[[rd.cell(t+1+k,c).value for c in range(1,10)] for k in range(3)])
    return out
S={"base":({},None),"mw_low":({"mwsc":1},None),"mw_high":({"mwsc":3},None),"lag2":({"lagb":2},None),
   "ms_letter":({"mssc":1},None),"ms_bear":({"mssc":3},None),
   "ms_bear_T5_9_cancel":({"mssc":3},[("Microsoft",f"F{r}",1) for r in range(7,12)]),
   "ms_sla5":({"sla":0.05},None),"unseen25":({"unseen":25},None),
   "all_leased":({"ownmode":2},None),"all_owned":({"ownmode":3},None),"v1basis":({"ybasis":2},None),"no_prepay":({"pptog":0},None),
   "low_lag2_bear":({"mwsc":1,"lagb":2,"mssc":3},None)}
if __name__=="__main__":
    R={k:run(*v) for k,v in S.items()}
    json.dump(R,open("scen.json","w"),indent=1,default=str)
    for k,o in R.items(): print(k,'FY27',round(o['rv27']),'gap',round(o['gap27']),'$/sh',round(o['gapps'],1),'ms27',round(o['ms27']),'y2',round(o['y2'],3),'gapq',o['gapq'],'impl5',o['rd'][1][1], 'gap$5',round(o['rd'][1][6]))
