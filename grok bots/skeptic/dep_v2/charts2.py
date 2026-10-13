import json,matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
R=json.load(open("scen.json")); b=R["base"]; yrs=[2026,2027,2028,2029,2030]
w=load_workbook("/tmp/lo/03_depreciation.xlsx",data_only=True); an=w["Annual"]
C=json.load(open("cells.json"))
eb=[an.cell(C['A_eb'],2+j).value for j in range(5)]
ebit={L:[an.cell(C['A_ebit'][L],2+j).value for j in range(5)] for L in ("L1","L2","L3")}
O="/workspace/nebius/skeptic/charts/"
f,ax=plt.subplots(figsize=(8,4.5)); ax.plot(yrs,[v/1e3 for v in eb],'k-o',label="Adj. EBITDA (cons. to FY28)")
for L,n_,c in (("L1","4-yr","C00000"),("L2","5-yr (policy)","1F3864"),("L3","6-yr","7F7F7F")): ax.plot(yrs,[v/1e3 for v in ebit[L]],'-o',color="#"+c,label=f"EBIT, {n_} GPU life")
ax.set_ylabel("$B"); ax.set_title("NBIS: EBITDA vs EBIT by GPU useful life (v2)"); ax.legend(fontsize=8); ax.grid(alpha=.3); plt.tight_layout(); plt.savefig(O+"v2_ebit_by_life.png",dpi=130); plt.close()
rd=w["Reverse_DCF"]; grid=[rd.cell(2,2+j).value for j in range(13)]
f,ax=plt.subplots(figsize=(8,4.5))
for k,(n_,c) in enumerate((("4-yr","#C00000"),("5-yr","#1F3864"),("6-yr","#7F7F7F"))):
    ax.plot([g*100 for g in grid],[rd.cell(C['evrow'][("L1","L2","L3")[k]],2+j).value/1e3 for j in range(13)],color=c,label=f"DCF EV, {n_} life")
ax.axhline(74.727,color='g',ls='--',label="Market EV $74.7B"); ax.axvline(b['y2']*100,color='orange',ls=':',label=f"FY28 yield on billing capital {b['y2']*100:.0f}%")
ax.axvline(b['y1']*100,color='purple',ls=':',label=f"v1 in-service basis {b['y1']*100:.0f}%")
ax.set_xlabel("FY29+ EBITDA yield on gross billing GPU capital (%)"); ax.set_ylabel("EV $B"); ax.set_title("Reverse DCF: yield the market EV requires"); ax.legend(fontsize=7); ax.grid(alpha=.3); plt.tight_layout(); plt.savefig(O+"v2_reverse_dcf.png",dpi=130); plt.close()
sc=[("mw_low","Capacity low"),("low_lag2_bear","Low + 2q lag + MSFT bear"),("lag2","Base, 2q lag"),("ms_bear_T5_9_cancel","Base, MSFT T5-9 cancelled"),("ms_bear","Base, MSFT bear"),("base","Base (1q lag)"),("ms_letter","Base, MSFT 'all delivered'"),("unseen25","Base + 25 MW/qtr unseen"),("mw_high","Capacity high")]
f,ax=plt.subplots(figsize=(8,4.5)); v=[R[k]['rv27']/1e3 for k,_ in sc]
ax.barh([l for _,l in sc],v,color=["#C00000" if x<12.3 else "#2E7D32" for x in v]); ax.axvline(12.3,color='k',ls='--',label="FY27 consensus $12.3B")
for i,x in enumerate(v): ax.text(x+0.1,i,f"${x:.1f}B",va='center',fontsize=8)
ax.set_xlabel("FY27 revenue, $B (billing-MW path)"); ax.set_title("Leg 3: FY27 revenue from Capacity's quarterly billing MW"); ax.legend(fontsize=8); plt.tight_layout(); plt.savefig(O+"v2_fy27_revenue.png",dpi=130); plt.close()
wb=load_workbook("/workspace/nebius/skeptic/03_depreciation.xlsx"); ch=wb["Charts"]
for i,p in enumerate(["v2_ebit_by_life","v2_reverse_dcf","v2_fy27_revenue"]):
    im=Image(O+p+".png"); im.width,im.height=640,360; ch.add_image(im,f"A{1+i*20}")
wb.save("/workspace/nebius/skeptic/03_depreciation.xlsx"); print("ok")
