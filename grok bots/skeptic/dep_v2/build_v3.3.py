import sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as CL
BLUE=Font(color="0000FF"); B=Font(bold=True); RED=Font(color="C00000",bold=True)
HDR=PatternFill("solid",fgColor="1F3864"); INF=PatternFill("solid",fgColor="DDEBF7"); PH=PatternFill("solid",fgColor="FFF2CC")
FB="/workspace/nebius/filings/nebius_fact_base.md"; Q2FS="https://www.sec.gov/Archives/edgar/data/1513845/000110465926094844/nbis-20260812xex99d2.htm"
F20="https://www.sec.gov/Archives/edgar/data/1513845/000110465926052948/nbis-20251231x20f.htm"; MS="https://www.marketscreener.com/quote/stock/NEBIUS-GROUP-8037501/finances/"
CAPX="/workspace/nebius/capacity/gpu_mw_capex.xlsx + .md (team Capacity & Power)"; SITES="/workspace/nebius/capacity/site_list.xlsx (team Capacity & Power)"
PRICE="/workspace/nebius/market/price_history_raw.csv (team GPU Market)"; Q2L="https://www.sec.gov/Archives/edgar/data/1513845/000110465926094568/tm2622968d1_ex99-2.htm"
GPX="/workspace/nebius/market/gpu_price_history.xlsx (A100_Decay; team GPU Market)"
PARENT="Team inputs relayed 2026-10-04 18:20 ET (Capacity physical evidence / Filings)"
MS6K="https://www.sec.gov/Archives/edgar/data/1513845/000110465925088312/tm2525580d1_6k.htm"
wb=Workbook(); A=wb.active; A.title="Assumptions"; A.append(["Input","Value","Unit","Tag","Source / note"]); C={}
def inp(k,label,v,unit,tag,src,formula=False,ph=False):
    A.append([label,v,unit,tag,src]); r=A.max_row; C[k]=f"Assumptions!$B${r}"
    c=A.cell(r,2)
    if not formula: c.font=BLUE; c.fill=PH if ph else INF
def sec(t): A.append([t]); A.cell(A.max_row,1).font=B
def a(k): return C[k]
sec("Lives")
inp("L1","Life scenario A (yrs)",4,"yrs","Reported (pre-2026 policy)",F20)
inp("L2","Life scenario B (yrs) - Nebius policy from 1/1/26",5,"yrs","Reported",Q2FS+" Note 1")
inp("L3","Life scenario C (yrs)",6,"yrs","Reported (peer policy: CRWV/GOOGL/ORCL)","v1 Peers sheet")
inp("shl","Shell & power (owned) / colo fit-out life",18,"yrs","Estimated (user assumption)","User instruction 2026-10-04")
inp("ngl","Existing 6/30 non-server PP&E blended life",12,"yrs","Estimated","As v1")
sec("Capex per IT MW by generation ($M) - Capacity's base case")
for g,gn,sh,colo in [("hop","Hopper H100/H200",25.9,26.9),("b200","B200 (HGX)",32.4,33.4),("gb200","GB200 NVL72",29.2,30.2),("b300","B300 (HGX)",40.1,41.1),("gb300","GB300 NVL72",40.2,41.3),("vr","Vera Rubin NVL72",30.3,31.3)]:
    inp(g+"_g",f"{gn}: GPU/server + network/storage per IT MW",gn and sh,"$M/IT MW","Estimated (Capacity base)",CAPX+" CapexPerMW base split")
    shell={"hop":10,"b200":10,"gb200":12.5,"b300":10,"gb300":12.5,"vr":12.5}[g]
    inp(g+"_s",f"{gn}: DC shell + power, owned, per IT MW",shell,"$M/IT MW","Estimated (Capacity base)",CAPX)
    inp(g+"_c",f"{gn}: colo capex (no shell) per IT MW",colo,"$M/IT MW","Estimated (Capacity base)",CAPX+" 'Colo capex Base'")
    inp(g+"_o",f"{gn}: owned all-in per IT MW",f"={a(g+'_g').split('!')[1]}+{a(g+'_s').split('!')[1]}","$M/IT MW","Derived","GPU+net + shell (Capacity: B200 42, GB200 42, GB300 53)",formula=True)
inp("pue","Gross (facility) MW per IT MW",1.25,"x","Derived (Capacity)","Vineland 300 MW gross ~ 230-250 MW IT => 1.2-1.3; Capacity PUE base 1.20-1.30")
inp("acvb","Nebius ACV/MW basis = gross MW? (1 = gross, 0 = IT)",1,"flag","Estimated","Microsoft $11.6M/gross MW-yr ~ the '2026 base ~$12M' in Q2 letter => ACV quoted on gross/contracted MW (Capacity). Not stated by Nebius")
sec("Opex per MW")
inp("pwr","Power cost",139,"$/MWh","Estimated (Capacity range 93-185, midpoint)",PARENT+"; Bloom is a power purchase (opex), not capex")
inp("lf","Average load factor on gross MW",0.9,"%","Estimated","Calibrates to Capacity's ~$1.1M per gross MW-yr at midpoint: 139 x 8,760 x 0.9 = $1.10M")
inp("pwr_g","Power cost per gross MW-yr",f"={a('pwr').split('!')[1]}*8760*{a('lf').split('!')[1]}/1000000","$M","Derived","",formula=True)
inp("rent","Colo rent per IT MW-yr (leased sites only)",2.25,"$M","Estimated (Capacity range 2.0-2.5)",CAPX+" peers: colo rent $165-205/kW-mo")
inp("segm","Q2'26 AI cloud segment adj. EBITDA margin (calibration)",0.497,"%","Reported",FB)
inp("calacv","Calibration ACV (installed book, per gross MW)",12,"$M","Reported (approx.)",Q2L+" p.3-4")
inp("oth","Other cash opex % of revenue (staff, network, support, SG&A)",f"=1-{a('segm').split('!')[1]}-{a('pwr_g').split('!')[1]}/{a('calacv').split('!')[1]}-{a('rent').split('!')[1]}/({a('calacv').split('!')[1]}*{a('pue').split('!')[1]})","%","Derived (calibration)","Assumes the Q2 segment margin = installed book on leased sites: 1 - 49.7% - power/rev - rent/rev",formula=True)
sec("Cohort pricing & terms")
for k,l,v,t,s in [("acv1","Installed book ACV per gross MW-yr",12,"Reported (approx.)",Q2L),("acv2","Q2 new deals ACV - low",20,"Reported",Q2L),("acv3","Q2 new deals ACV - high",25,"Reported",Q2L),("acv4","Q3 short-term deals ACV",40,"Reported ('>$40M')",Q2L),("acv5","Microsoft ACV per gross MW-yr",f"=17392.9/5/300","Derived","$17,392.9M / 5 yrs / 300 MW gross = $11.6M (= ~$1,324/MWh)"),("acv6","Microsoft ACV if contract covers 200 MW",f"=17392.9/5/200","Derived","Capacity: ~$17.4M at 200 MW")]:
    inp(k,l,v,"$M",t,s,formula=isinstance(v,str))
inp("term_lt","Term - installed book / Q2 deals",5,"yrs","Estimated","Strategic contracts (MSFT, Meta) are 5-yr (R); Q2-deal terms not disclosed")
inp("term_st","Term - short-term deals",1,"yrs","Estimated","'short-term' not defined by Nebius")
inp("term_ms","Term - Microsoft (per tranche)",5,"yrs","Reported","60 months per tranche ("+MS6K+")")
inp("sla","Microsoft SLA / service-credit haircut on billed revenue",0.0,"%","Scenario input","Take-or-pay 'irrespective of actual utilization' - haircut is service credits, not demand; 0% base, test 2-5%")
inp("ppct","Prepayment, % of related capex, for a prepaid deal",0.55,"%","Reported (company: 50-60%)",FB)
inp("pinc","Share of deals carrying prepayments",0.70,"%","Reported (company)",FB)
inp("msup","Microsoft upfront share of TCV",f"=6958.1/17392.9","%","Derived","~$6,958.1M upfront / $17,392.9M TCV (20-F Note 1)",formula=True)
sec("Residual value after contract term (re-rent) - GPU Market A100 evidence (v3.1)")
inp("util","Realized utilization on post-term re-rent",0.60,"%","Estimated (user bear case)","User instruction")
inp("ressel","Residual curve: 1 = hyperscaler A100 list (-1.8%/yr), 2 = spot median ~flat, 3 = neocloud composite (+12%/yr), 4 = A100 80GB median, 5 = v3 placeholder (cross-section)",1,"flag","Scenario input","Default = most conservative LIST series; all are list prices except 2 (spot)")
inp("g_hyp","A100 hyperscaler median list, CAGR Nov-20 to Oct-26",-0.018,"%/yr","Reported (GPU Market)",GPX+": $4.096 -> $3.673")
inp("g_spot","A100 spot/marketplace median CAGR (treated as ~flat)",0.0,"%/yr","Reported (GPU Market; file CAGR +2.2%, $0.93 -> $1.055)",GPX)
inp("g_neo","A100 neocloud composite median CAGR Jul-22 to Oct-26",0.12,"%/yr","Reported (GPU Market)",GPX+": $1.234 -> $2.00")
inp("g_80","A100 80GB median CAGR (relay: $1.87 -> $2.50 over ~6 yrs)",f"=(2.5/1.87)^(1/6)-1","%/yr","Derived (GPU Market relay; file's 80GB neocloud median is $2.04 Apr-23 -> $2.50, +6%/yr)",GPX,formula=True)
inp("rescap","Cap on residual index (x launch price)",1.0,"x","Estimated","Do not assume re-rent above the cohort's own launch price")
ph=[1.00,0.82,0.64,0.53,0.46,0.39]
for i,v in enumerate(ph):
    inp(f"resph{i+1}",f"v3 placeholder index age {i+1} (superseded)",v,"x","Superseded placeholder","A100 vs then-flagship cross-section (v3)")
for i in range(6):
    k=i
    inp(f"res{i+1}",f"Residual price index at asset age {i+1} yrs (x the cohort's own launch price)",
        f"=MIN({a('rescap').split('!')[1]},CHOOSE({a('ressel').split('!')[1]},(1+{a('g_hyp').split('!')[1]})^{k},(1+{a('g_spot').split('!')[1]})^{k},(1+{a('g_neo').split('!')[1]})^{k},(1+{a('g_80').split('!')[1]})^{k},{a(f'resph{i+1}').split('!')[1]}))","x","Derived (GPU Market A100 evidence)",
        "LIST prices; spot median roughly flat. A100 (Nov-20 launch) still rents at ~list after 6 yrs => 5-6 yr economic life supported",formula=True)
inp("acvref","Re-rent reference price per gross MW-yr (new-capacity market rate)",f"={a('acv2').split('!')[1]}","$M","Derived","Q2-deal low ($20M); re-rent = MIN(cohort ACV, ref) x index x utilization",formula=True)
sec("Fleet build mix (Sites sheet)")
inp("ownmode","Owned share of capex: 1 = Capacity tracker (owned sites' share of energized MW added in FY), 2 = all leased (0%), 3 = all owned (100%)",1,"flag","Scenario input","Nebius is mostly leased today (started leases $1,625M; $12.05B signed not started). Unknown-basis sites run both via modes 2/3")
inp("gpu26","GPU+network share of FY26 capex and not-yet-in-use",0.79,"%","Estimated (mgmt via secondary)","~79% GPUs / 20% DC / <1% power (Q4'25 call, secondary). 2027+ share derived from mix and generation")
inp("vrq","First capex quarter buying Vera Rubin (index; 0 = Q3'26)",4,"idx","Estimated","Meta #2 Vera Rubin from early 2027 (R); Q3'27 capex onward priced as VR")
sec("Lags (Modeling pushback)")
inp("lage","Capex -> energized / in service",2,"quarters","Estimated","Capacity gives no explicit lag (gpu_mw_capex.md caveat 5: GPUs bought ~ at energization); v1 value kept")
inp("cxbasis","Capex basis: 1 = consensus capex, 2 = path-consistent (Capacity energized-MW adds x capex per IT MW) - BASE CASE = 2 from v3.2",2,"flag","Scenario input","v3.1. Consensus capex implies ~157-200 IT MW bought/qtr, faster than Capacity's base path")
inp("pglag","Path basis: GPU/IT spend this many quarters BEFORE billing",1,"quarters","Estimated (team)","Test 2")
inp("shlag","Path basis: shell/power spend this many quarters before energization (owned)",2,"quarters","Estimated (user)","")
inp("lease_gb200","Path basis: leased-site capex per IT MW (IT only, no shell) - GB200",30,"$M/IT MW","Estimated (Capacity, medium confidence; corrected 2026-10-04 19:02 ET)",CAPX)
inp("lease_gb300","Path basis: leased-site capex per IT MW - GB300",41,"$M/IT MW","Estimated (Capacity, medium confidence)",CAPX)
inp("lease_vr","Path basis: leased-site capex per IT MW - Vera Rubin (midpoint of $31-33M)",32,"$M/IT MW","Estimated (Capacity, medium confidence)",CAPX)
inp("lagb","Connected/energized -> billing (revenue)",1,"quarters","Estimated (Filings: mgmt's 'few months' lag runs from connected power to revenue; Capacity uses 1-2 qtrs)","/workspace/nebius/filings/nebius_filings.xlsx; run 1 (default) and 2")
inp("mwsc","Capacity MW scenario: 1 = low, 2 = base, 3 = high",2,"flag","Scenario input","/workspace/nebius/capacity/nebius_capacity.xlsx 'Quarterly MW' energized rows")
inp("unseen","Unseen-site MW added (gross), e.g. Oklahoma / Minnesota colos - per quarter from Q4'26",0,"MW/qtr","Scenario input","Capacity: MN, OK, France 2nd, UK 3rd MW not disclosed (plug to reach guide ~496 MW at YE26 base)")
inp("acvinc","ACV on incremental (post-Q2'26) MW, per gross MW-yr",22.5,"$M","Reported range midpoint ($20-25M Q2 deals)",Q2L)
inp("odsh","On-demand + Token Factory share of revenue",0.04,"%","Reported (team GPU Market revenue mix: take-or-pay 45%, reserved ~50%, on-demand+TF ~4%)","/workspace/nebius/market (GPU Market)")
inp("hike","Oct 1 2026 on-demand list-price hike",0.19,"%","Reported (17-21%; H100 3.85->4.50, B300 7.85->9.50)","nebius.com/prices via GPU Market")
sec("Price cycle (v3.1 bear input)")
inp("pc_cut","Blackwell $/GPU-hr cut in the 12 months after Rubin launch",0.275,"%","Scenario input (25-30%)","Team bear input; 5-yr life unchanged underneath")
inp("pc_l","Rubin launch quarter (Rev_MW index; 0 = Q3'26, 2 = Q1'27)",2,"idx","Estimated","GPU Market: Nebius Rubin GA ~Dec'26-Mar'27; cut phases in linearly over 4 quarters")
inp("ressh","Reserved share of revenue",0.50,"%","Reported (GPU Market mix)","take-or-pay 45% (not exposed), reserved ~50%, on-demand + Token Factory ~4%")
inp("pc_sc","Reserved-book renewal scenario (agreed with Modeling): 1 = bull (Modeling's original schedule, -$75M FY27), 2 = BASE 50% renews within FY27, 3 = bear 100% renews within FY27",2,"flag","Scenario input","Scenarios from Nebius pricing-page terms (reserved/commit terms), NOT a filing; Modeling's earlier schedule was a model assumption")
inp("pc_ren1","Bull: reserved share repricing within FY27 implied by Modeling's original schedule (back-solved, set below)",0.048,"%","Derived (calibrated to Modeling's -$75M FY27 on the bottom-up basis)","Linear back-solve: (75 - on-demand component) / reserved component at 100%; formula written after Rev_MW is built",formula=True)
inp("pc_ren2","Base: share of reserved book renewing within FY27",0.5,"%","Scenario input (page terms)","Agreed base")
inp("pc_ren3","Bear: share of reserved book renewing within FY27",1.0,"%","Scenario input (page terms)","")
inp("pc_ren","Renewal share used",f"=CHOOSE({a('pc_sc').split('!')[1]},{a('pc_ren1').split('!')[1]},{a('pc_ren2').split('!')[1]},{a('pc_ren3').split('!')[1]})","%","Derived","",formula=True)
sec("Asset base at 30-Jun-2026 (Reported)")
for k,l,v,s in [("srv","Server & network gross in service",5197.7,Q2FS+" Note 7"),("age","Avg age of in-service server base (quarters)",4,"Estimated as v1"),("othpp","Other in-service PP&E gross",1115.4,Q2FS),("niu","Assets not yet in use",7836.4,Q2FS),("niu3","Share of not-yet-in-use energized Q3'26 (rest Q4'26)",0.5,"Estimated")]:
    inp(k,l,v,"$M" if k not in("age","niu3") else "", "Reported" if k in("srv","othpp","niu") else "Estimated",s)
sec("Capex / EBITDA (consensus; extrapolation)")
for k,l,v,t,s in [("cx26","FY2026 capex",23207,"Estimated (consensus)",MS),("h1cx","1H26 capex",8130.3,"Reported",FB),("cx27","FY2027 capex",29135,"Estimated (consensus)",MS),("cx28","FY2028 capex",38885,"Estimated (consensus)",MS),("g29","FY29 capex growth",0,"Estimated","flat"),("g30","FY30 capex growth",0,"Estimated","flat"),("eb28","FY28 adj. EBITDA cons",13978,"Estimated (consensus)",MS),("rv28","FY28 revenue cons",22983,"Estimated (consensus)",MS)]:
    inp(k,l,v,"$M" if "g" not in k[:1] else "%",t,s)
for i,w in enumerate([0.22,0.24,0.26,0.28]): inp(f"w{i}",f"FY28 weight Q{i+1}",w,"%","Estimated","")
inp("ybasis","FY29+ EBITDA yield basis: 1 = FY28 yield on BILLING GPU capital (v2 default), 2 = v1 in-service basis",1,"flag","Scenario input","Modeling pushback: ramp-period in-service yield understates steady state")
inp("yov","FY29+ yield override (blank = per basis)",None,"%","Scenario input","")
sec("1H26 actuals")
for k,l,v in [("h1rv","1H26 revenue",981.3),("h1eb","1H26 adj. EBITDA",365.7),("h1da","1H26 D&A",471.7),("h1ben","1H26 D&A cut from 4->5 yr",86.1)]: inp(k,l,v,"$M","Reported",FB)
sec("Prepayments (working-capital funding)")
inp("pptog","Prepayment toggle (1 = on)",1,"flag","Scenario input","v1 had none")
inp("dr0","Deferred revenue at 6/30/26",5975.2,"$M","Reported",FB)
inp("drn","Amortization period for deferred revenue",20,"quarters","Estimated","5-yr contracts")
sec("Financing")
tr=[("2.00% 2029",87.5,0.02),("3.00% 2031",87.5,0.03),("1.00% 9/2030",1581.25,0.01),("2.75% 9/2032",1581.25,0.0275),("1.25% 3/2031",2587.5,0.0125),("2.625% 3/2033",1750,0.02625),("0.50% 2/2030",3450,0.005),("4.50% 2/2034",2300,0.045)]
rows=[]
for n,p,c in tr:
    A.append([f"Convert {n}: principal / coupon",p,"$M",c,"Reported ("+FB+")"]); r=A.max_row; rows.append(r)
    for cc in (2,4): A.cell(r,cc).font=BLUE; A.cell(r,cc).fill=INF
inp("sofr","Term SOFR assumption",0.04,"%","Estimated","verify")
inp("loan","Secured loan",775,"$M","Reported",FB)
inp("int_q","Cash interest per quarter","=("+"+".join(f"B{r}*D{r}" for r in rows)+f"+{a('loan').split('!')[1]}*({a('sofr').split('!')[1]}+0.025))/4","$M","Derived","",formula=True)
inp("rate","Cost of incremental funding",0.07,"%","Estimated","between NBIS 5.1% and CRWV 8.4%")
inp("liqsel","Starting liquidity basis: 1 = v3.1 cash basis ($14,565M), 2 = v1 pro forma cash",1,"flag","Scenario input","")
inp("liq1","Starting liquidity: 6/30 unrestricted cash + Aug converts + secured loan",14565,"$M","Derived (Filings, v3.1)","8,040 + 5,750 + up to 775. v3 used $17.2B - a raise-vs-capex table, not cash, double counting ~$4.4B of prepayments already received")
inp("liq2","v1 pro forma cash",14497.1,"$M","Estimated","01_priced_in")
inp("liq","Starting liquidity used",f"=IF({a('liqsel').split('!')[1]}=1,{a('liq1').split('!')[1]},{a('liq2').split('!')[1]})","$M","Derived","",formula=True)
sec("Valuation")
for k,l,v,t,s in [("wacc","WACC",0.11,"Estimated","assumption"),("tg","Terminal growth",0.03,"Estimated","assumption"),("tax","Terminal tax rate",0.258,"Reported (NL statutory)",""),("ev","Market EV (base, ex-leases)",74727,"Derived","01_priced_in v3: 308.69M economic basic x $242.81 + 14,270.7 debt - 14,497.1 cash (Filings share-count reversal 2026-10-04 18:27 ET)"),("nc","Pro forma net cash",226.4,"Derived","01_priced_in"),("sh","Economic basic shares (incl. NVIDIA warrant)",308.69,"Derived","01_priced_in v3 / Filings CapStructure"),("px","Price 10/2/26",242.81,"Reported","StockAnalysis"),("ener","Energy & power purchase commitments (undiscounted)",5324.1,"Reported",Q2FS+": 81.7+541.8+542.0+540.5+535.8+3,082.3"),("lnc","Leases signed not yet commenced (undiscounted)",12054.1,"Reported",Q2FS+" (prepaid $612.4M)"),("olease","Operating lease liabilities on BS",1625.3,"Reported","01_priced_in")]:
    inp(k,l,v,"$M" if k not in ("wacc","tg","tax") else "%",t,s)
inp("evx","EV incl. off-BS commitments (debt-like variant)",f"={a('ev').split('!')[1]}+{a('ener').split('!')[1]}+{a('lnc').split('!')[1]}+{a('olease').split('!')[1]}","$M","Derived","CAUTION: EBITDA is already after rent and power, so pair with EBITDAR for consistency; undiscounted",formula=True)
for c in A[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=HDR
for k,w in zip("ABCDE",[70,12,10,26,120]): A.column_dimensions[k].width=w

# ---------------- Sites
S=wb.create_sheet("Sites")
S.append(["#","Site","Country","Tracker type","Ownership tag (v2)","YE26 MW low","YE26 MW base","YE26 MW high","MW basis","Owned base MW","Source"])
site=[(1,"Mantsala DC #1","Finland","Owned greenfield","Owned",75,75,75,"not disclosed (PUE ~1.1)"),(2,"Mantsala DC #2","Finland","Owned (implied)","Owned",0,0,0,"n/d"),(3,"Lappeenranta","Finland","Owned per Nebius; developer says Nebius 'operating Polarnode's campus'","Unknown (conflict)",0,0,0,"n/d"),(4,"Paris Equinix PA10","France","Colocation","Leased",5,5,5,"n/d"),(5,"Bethune","France","Colocation per 20-F","Leased",0,120,120,"n/d"),(6,"France 2nd colo","France","Colocation","Leased",0,0,0,"n/d"),(7,"Keflavik Verne","Iceland","Colocation","Leased",10,10,10,"n/d"),(8,"London Longcross (Ark)","UK","Colocation","Leased",16,19,19,"gross (Capacity inference)"),(9,"Harlow Kao","UK","Colocation","Leased",0,0,4.4,"n/d"),(10,"UK 3rd site","UK","not disclosed (likely colo - inference)","Unknown",0,0,0,"n/d"),(11,"Modi'in Mega Or","Israel","Colocation","Leased",8,8,8,"IT (Capacity)"),(12,"Masmiyya","Israel","Colocation","Leased",0,22,22,"n/d"),(13,"Beit Shemesh","Israel","Colocation","Leased",0,29,58,"n/d"),(14,"Madrid-Getafe","Spain","Colocation","Leased",9,18,18,"n/d"),(15,"Tallinn Huuru","Estonia","Colocation","Leased",0,22,22,"n/d"),(16,"Kansas City Patmos","USA","Colocation","Leased",5,20,40,"n/d"),(17,"Vineland DataOne","USA","Build-to-suit, Nebius tenant","Leased (gross MW)",200,200,300,"GROSS; physical cap 122.9 MW gross (62 x 1.98 MW rented gensets) until Bloom live"),(18,"Independence MO","USA","Owned greenfield","Owned",0,0,0,"n/d"),(19,"Birmingham BHM01","USA","Owned brownfield","Owned",0,50,100,"'electrical capacity' (press)"),(20,"Pennsylvania","USA","Owned greenfield","Owned",0,0,0,"n/d"),(21,"Oklahoma colo","USA","Colocation","Leased",0,0,0,"n/d"),(22,"Minnesota colo","USA","Colocation","Leased",0,0,0,"n/d"),(23,"AIB southeastern US","USA","Colocation/leased 12-yr","Leased",0,0,0,"IT (50 MW critical IT)")]
for s in site:
    r=S.max_row+1
    S.append(list(s)+[f'=IF(LEFT(E{r},5)="Owned",G{r},0)',SITES])
    for c in (6,7,8): S.cell(r,c).font=BLUE
S.append(["","TOTAL","","","",f"=SUM(F2:F24)",f"=SUM(G2:G24)",f"=SUM(H2:H24)","",f"=SUM(J2:J24)"])
S.append([]); S.append(["","Owned share of YE26 base MW","","","","","","","",f"=J25/G25"])
S.append(["","Leased (incl. Vineland) share","","","","","","","",f'=SUMPRODUCT((LEFT(E2:E24,6)="Leased")*G2:G24)/G25'])
S.append(["","Unknown: Lappeenranta, UK 3rd (0 MW in YE26 base) - run owned and leased (Cohorts both; fleet owned-share 0%/100% in md)"])
for c in S[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=HDR
S.column_dimensions['B'].width=26; S.column_dimensions['D'].width=40; S.column_dimensions['E'].width=20; S.column_dimensions['I'].width=50

# ======================= PART 2: MW, Microsoft, FY27 revenue =======================
from capdata import CAP
Q=["Q2 2026"]+[f"Q{q} {y}" for y in range(2026,2031) for q in range(1,5)][2:]   # MW sheet: Q2'26..Q4'30 (19)
MW=wb.create_sheet("MW"); MW.append(["Capacity energized MW (identified sites) - "+"/workspace/nebius/capacity/nebius_capacity.xlsx 'Quarterly MW' (pasted values, blue)",""]+Q)
mc=[CL(3+i) for i in range(len(Q))]
def mrow(label,vals,src="",blue=True):
    MW.append([label,src]+vals); r=MW.max_row
    for i in range(len(vals)):
        c=MW.cell(r,3+i); c.number_format='#,##0.0'
        if blue: c.font=BLUE; c.fill=INF
    return r
R={}
for sc in ("low","base","high"):
    R["tot"+sc]=mrow(f"Energized MW, identified sites - {sc} (mixed basis)",CAP['tot'][sc],"Capacity Quarterly MW (sum of site rows; ties to COMPANY rows)")
    R["it"+sc]=mrow(f"  of which IT-basis sites (Vineland, AIB) - {sc}",CAP['it'][sc],"Capacity Basis column = IT")
    R["vin"+sc]=mrow(f"  of which Vineland (IT) - {sc}",CAP['vin'][sc],"Capacity")
    R["own"+sc]=mrow(f"  of which owned sites (Independence, Butler PA, Birmingham, Lappeenranta, Mantsala) - {sc}",CAP['own'][sc],"Capacity Sites 'Ownership type'")
MW.append([]); 
sel=lambda k: "CHOOSE({m},{a},{b},{c})".format(m=a('mwsc'),a="{c}"+str(R[k+'low']),b="{c}"+str(R[k+'base']),c="{c}"+str(R[k+'high']))
def frow(label,f,fmt='#,##0.0',bold=False):
    MW.append([label,""]+[f.format(c=c,i=i,p=(mc[i-1] if i>0 else None)) for i,c in enumerate(mc)]); r=MW.max_row
    for i in range(len(mc)): MW.cell(r,3+i).number_format=fmt
    if bold: MW.cell(r,1).font=B
    return r
r_sel=frow("Selected energized MW (mixed)","="+sel("tot"))
r_it=frow("Selected IT-basis MW","="+sel("it"))
r_vin=frow("Selected Vineland IT MW","="+sel("vin"))
r_ownmw=frow("Selected owned-site MW","="+sel("own").replace("{c}"+str(R['ownlow']),"{c}"+str(R['ownlow'])))
r_g=frow("Energized GROSS MW (unknown basis treated as gross; IT x PUE)","={c}%d-{c}%d+{c}%d*%s"%(r_sel,r_it,r_it,a('pue')),bold=True)
r_un=frow("+ Unseen-site gross MW (cumulative from Q4'26)","=IF({i}>=2,({i}-1)*%s,0)"%a('unseen'))
r_gt=frow("Energized gross MW incl. unseen","={c}%d+{c}%d"%(r_g,r_un),bold=True)
r_ite=frow("IT-equivalent energized MW incl. unseen (gross / PUE)","={c}%d/%s"%(r_gt,a('pue')))
r_add=frow("IT MW energized in quarter (additions)","=IF({i}=0,0,{c}%d-{p}%d)"%(r_ite,r_ite))
C['mw_add']=r_add
r_vg=frow("  Vineland gross MW","={c}%d*%s"%(r_vin,a('pue')))
r_bill=frow("BILLING gross MW (energized shifted by billing lag; lag 0/1/2 supported)","=IF(%s=0,{c}%d,IF(%s=1,IF({i}>=1,INDEX($C$%d:$U$%d,{i}),{c}%d),IF({i}>=2,INDEX($C$%d:$U$%d,{i}-1),$C$%d)))"%(a('lagb'),r_gt,a('lagb'),r_gt,r_gt,r_gt,r_gt,r_gt,r_gt),bold=True)
r_billnv=frow("  Billing gross MW ex-Vineland","=IF(%s=0,{c}%d-{c}%d,IF(%s=1,IF({i}>=1,INDEX($C$%d:$U$%d,{i})-INDEX($C$%d:$U$%d,{i}),{c}%d-{c}%d),IF({i}>=2,INDEX($C$%d:$U$%d,{i}-1)-INDEX($C$%d:$U$%d,{i}-1),$C$%d-$C$%d)))"%(a('lagb'),r_gt,r_vg,a('lagb'),r_gt,r_gt,r_vg,r_vg,r_gt,r_vg,r_gt,r_gt,r_vg,r_vg,r_gt,r_vg))
r_ownsh=frow("Owned share of cumulative energized MW (selected)","=IF({c}%d>0,{c}%d/{c}%d,0)"%(r_sel,r_ownmw,r_sel),'0%')
MW.append(["Note: FY owned share of energized ADDITIONS (feeds capex mix) is on the Model sheet. Q2'26 column = starting point (energized at 6/30/26 per Capacity)."])
MW.column_dimensions['A'].width=78; MW.column_dimensions['B'].width=30
for c in MW[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=HDR
C['mw_bill']=r_bill; C['mw_billnv']=r_billnv; C['mw_gt']=r_gt; C['mw_sel']=r_sel; C['mw_own']=r_ownmw

# ---- Microsoft
MSQ=[f"Q{q} {y}" for y in range(2025,2032) for q in range(1,5)][3:]  # Q4'25..Q4'31 (25)
MSs=wb.create_sheet("Microsoft"); MSs.append(["Microsoft (Vineland) tranche model - take-or-pay 'irrespective of actual utilization'; haircut = SLA/service credits + per-tranche termination (NOT demand)"])
MSs.append(["Tranche","Scheduled (contract/letter)","Capacity base billing start","Capacity bear billing start","Selected start","Terminate? (1 = cancelled)","Quarterly revenue $M (TCV/9/20)","Source / tag"])
sched=["Q4 2025","Q1 2026","Q3 2026","Q3 2026","Q3 2026","Q3 2026","Q3 2026","Q3 2026","Q3 2026"]
base=["Q4 2025","Q1 2026","Q3 2026","Q3 2026","Q2 2027","Q2 2027","Q3 2027","Q3 2027","Q3 2027"]
bear=["Q4 2025","Q1 2026","Q3 2026","Q3 2026","Q3 2027","Q3 2027","Q4 2027","Q4 2027","Q1 2028"]
inp("mssc","Microsoft scenario: 1 = letter ('all delivered', T3-T9 billing Q3'26), 2 = Capacity base, 3 = Capacity bear",2,"flag","Scenario input","Capacity 'Vineland tranche track': T1-T4 within ~107 MW IT genset ceiling (equal-size assumption, ~27 MW IT each); T5-T9 need Bloom (stop-ordered Aug 6/10; NJDEP 45-day permit deadline ~early/mid Nov)")
inp("mstcv","Microsoft TCV",17392.9,"$M","Reported","20-F Note 1 / "+FB)
for i in range(9):
    r=MSs.max_row+1
    MSs.append([f"T{i+1}",sched[i],base[i],bear[i],f"=CHOOSE({a('mssc')},B{r},C{r},D{r})",0,f"={a('mstcv')}/9/20*(1-{a('sla')})*(1-F{r})","Equal split (Estimated); dates: 20-F (T1 Nov-25, T2 Feb-26), Q2 letter (two tranches Jul/Aug-26), Capacity estimates T5-T9"])
    for c in (2,3,4,6): MSs.cell(r,c).font=BLUE; MSs.cell(r,c).fill=INF
t0=3; t9=11
MSs.append([]); MSs.append(["Quarter"]+MSQ); rq=MSs.max_row
MSs.append(["Revenue $M"]+[f'=SUMPRODUCT(($G${t0}:$G${t9})*(MATCH($E${t0}:$E${t9},$B${rq}:$Z${rq},0)<=COLUMN()-1)*(MATCH($E${t0}:$E${t9},$B${rq}:$Z${rq},0)>COLUMN()-21))' for _ in MSQ]); rr=MSs.max_row
MSs.append(["Billing tranches"]+[f'=SUMPRODUCT((1-$F${t0}:$F${t9})*(MATCH($E${t0}:$E${t9},$B${rq}:$Z${rq},0)<=COLUMN()-1)*(MATCH($E${t0}:$E${t9},$B${rq}:$Z${rq},0)>COLUMN()-21))' for _ in MSQ])
for c in MSs[rr][1:]: c.number_format='#,##0'
MSs.append([])
MSs.append(["Q2'26 check: equal-split T1+T2 = $193M/qtr vs top customer 24% x $582.3M = ~$140M (Capacity) - either T1/T2 smaller than 1/9 or customer C is not Microsoft"])
MSs.append(["FY2026 Microsoft revenue",f"=SUMPRODUCT(($B${rr}:$Z${rr})*(RIGHT($B${rq}:$Z${rq},4)=\"2026\"))"]); C['ms26']=f"Microsoft!$B${MSs.max_row}"
MSs.append(["FY2027 Microsoft revenue",f"=SUMPRODUCT(($B${rr}:$Z${rr})*(RIGHT($B${rq}:$Z${rq},4)=\"2027\"))"]); C['ms27']=f"Microsoft!$B${MSs.max_row}"
MSs.append(["FY2028 Microsoft revenue",f"=SUMPRODUCT(($B${rr}:$Z${rr})*(RIGHT($B${rq}:$Z${rq},4)=\"2028\"))"]); C['ms28']=f"Microsoft!$B${MSs.max_row}"
MSs.append(["Total billed over life (TCV after SLA/terminations)",f"=SUM(G{t0}:G{t9})*20"]); C['mslife']=f"Microsoft!$B${MSs.max_row}"
MSs.append(["PV at Q3'26 of Microsoft revenue Q3'26-Q4'31 at WACC (Microsoft through 2031 only partly visible in the 2030 horizon)",f"=SUMPRODUCT($B${rr}:$Z${rr}*(COLUMN($B${rr}:$Z${rr})>=5)*(1+{a('wacc')})^(-(COLUMN($B${rr}:$Z${rr})-5)/4))"]); C['mspv']=f"Microsoft!$B${MSs.max_row}"
MSs.append(["Genset ceiling run-rate (4 tranches) $M/yr",f"=4*{a('mstcv')}/9/5"])
MSs.column_dimensions['A'].width=50
for k in "BCDE": MSs.column_dimensions[k].width=14
C['ms_rq']=rq; C['ms_rr']=rr

# ---- FY27 revenue bottom-up (leg 3)
RV=wb.create_sheet("Rev_MW"); RQ=Q[1:]  # Q3'26..Q4'30 (18)
RV.append(["Bottom-up revenue from billing-MW path ($M) - leg 3 re-run",""]+RQ)
rc=[CL(3+i) for i in range(len(RQ))]
def rrow(label,f,fmt='#,##0',bold=False):
    RV.append([label,""]+[f.format(c=c,i=i,mwc=CL(4+i)) for i,c in enumerate(rc)]); r=RV.max_row
    for i in range(len(rc)): RV.cell(r,3+i).number_format=fmt
    if bold: RV.cell(r,1).font=B
    return r
r1=rrow("Billing gross MW ex-Vineland","=MW!{mwc}%d"%r_billnv)
r0=rrow("  Installed-book MW (Q2'26 billing level, at $12M)","=MIN({c}%d,MW!$C$%d)"%(r1,r_billnv))
r2=rrow("  Incremental MW (at Q2-deal ACV)","={c}%d-{c}%d"%(r1,r0))
r3=rrow("Non-Microsoft revenue","=({c}%d*%s+{c}%d*%s)/4*(1+%s*%s)"%(r0,a('acv1'),r2,a('acvinc'),a('odsh'),a('hike')))
r4=rrow("Microsoft revenue (Microsoft sheet)","=INDEX(Microsoft!$B$%d:$Z$%d,MATCH({c}$1,Microsoft!$B$%d:$Z$%d,0))"%(rr,rr,rq,rq))
r5=rrow("TOTAL bottom-up revenue","={c}%d+{c}%d"%(r3,r4),bold=True)
cons=[904.1,1452,2175,2646,3304,4052]
r_bw=rrow("Blackwell (incremental-MW) revenue, non-Microsoft","={c}%d*%s/4*(1+%s*%s)"%(r2,a('acvinc'),a('odsh'),a('hike')))
r_rp=rrow("Price-cycle phase-in (0-100% over 4 qtrs from Rubin launch)","=MIN(1,MAX(0,({i}-%s+1)/4))"%a('pc_l'),fmt='0%')
r_pc=rrow("PRICE-CYCLE revenue hit = cut x phase x Blackwell rev x (on-demand + reserved x renewal share)","=-%s*{c}%d*{c}%d*(%s+%s*%s)"%(a('pc_cut'),r_rp,r_bw,a('odsh'),a('ressh'),a('pc_ren')),bold=True)
C['rv_pc']=r_pc
RV.append(["Consensus quarterly revenue (MarketScreener)",""]+cons+[None]*12); r6=RV.max_row
for i in range(6): RV.cell(r6,3+i).font=BLUE
RV.append([]); 
RV.append(["FY2027 bottom-up revenue",f"=SUMPRODUCT(($C${r5}:$T${r5})*(RIGHT($C$1:$T$1,4)=\"2027\"))"]); rb=RV.max_row; C['rv27']=f"Rev_MW!$B${rb}"
RV.append(["FY2027 consensus revenue",12300]); RV.cell(RV.max_row,2).font=BLUE
RV.append(["Gap vs consensus",f"=B{rb}-B{rb+1}"]); C['gap27']=f"Rev_MW!$B${rb+2}"
RV.append(["Gap %",f"=B{rb+2}/B{rb+1}"])
RV.append(["$/share at constant EV/FY27 revenue (308.69M sh)",f"=B{rb+2}*({a('ev')}/B{rb+1})/{a('sh')}"]); C['gapps']=f"Rev_MW!$B${rb+4}"
RV.append(["FY2026 H2 bottom-up vs consensus H2 (2,356)",f"=SUMPRODUCT(($C${r5}:$T${r5})*(RIGHT($C$1:$T$1,4)=\"2026\"))"])
RV.append(["FY2027 price-cycle revenue hit (= EBITDA hit; costs fixed; 5-yr life unchanged)",f"=SUMPRODUCT(($C${r_pc}:$T${r_pc})*(RIGHT($C$1:$T$1,4)=\"2027\"))"]); C['pc27']=f"Rev_MW!$B${RV.max_row}"
RV.append(["  as % of FY27 bottom-up revenue",f"=B{RV.max_row}/{C['rv27'].split('!')[1]}"]); C['pc27p']=f"Rev_MW!$B${RV.max_row}"
RV.append(["  scaled to FY27 consensus revenue ($12.3B)",f"=B{RV.max_row}*12300"]); C['pc27c']=f"Rev_MW!$B${RV.max_row}"
RV.append(["  as % of FY27 consensus EBITDA ($6,728M)",f"=B{RV.max_row}/6728"]); C['pc27e']=f"Rev_MW!$B${RV.max_row}"
RV.append(["FY27 hit split: on-demand component (independent of renewal)",f"=-{a('pc_cut')}*SUMPRODUCT(($C${r_rp}:$T${r_rp})*($C${r_bw}:$T${r_bw})*(RIGHT($C$1:$T$1,4)=\"2027\"))*{a('odsh')}"]); rod=RV.max_row
RV.append(["FY27 hit split: reserved component per 100% renewal",f"=-{a('pc_cut')}*SUMPRODUCT(($C${r_rp}:$T${r_rp})*($C${r_bw}:$T${r_bw})*(RIGHT($C$1:$T$1,4)=\"2027\"))*{a('ressh')}"]); rrs=RV.max_row
for lab,k in (("BULL (Modeling original schedule)","pc_ren1"),("BASE (50% renews within FY27)","pc_ren2"),("BEAR (100% renews within FY27)","pc_ren3")):
    RV.append([f"FY27 revenue/EBITDA hit - {lab}",f"=B{rod}+B{rrs}*{a(k)}"]); C['pcs_'+k]=f"Rev_MW!$B${RV.max_row}"
for lab,k in (("BULL","pc_ren1"),("BASE 50%","pc_ren2"),("BEAR 100%","pc_ren3")):
    RV.append([f"  same, scaled to FY27 consensus revenue ($12.3B) - {lab}",f"=({a('pcs_'+k).split('!')[1]})/{C['rv27'].split('!')[1]}*12300"]); C['pcc_'+k]=f"Rev_MW!$B${RV.max_row}"
A[a('pc_ren1').split('!')[1].replace('$','')].value=f"=(75+Rev_MW!B{rod})/(-Rev_MW!B{rrs})"
RV.append(["Note: renewal scenarios come from Nebius pricing-page terms, not a filing; Modeling's earlier schedule was a model assumption"])
RV.append(["FY2028 price-cycle revenue hit",f"=SUMPRODUCT(($C${r_pc}:$T${r_pc})*(RIGHT($C$1:$T$1,4)=\"2028\"))"]); C['pc28']=f"Rev_MW!$B${RV.max_row}"
RV.append(["Calibration: Q2'26 implied = MW!C ex-Vin (Q1'26 energized ~134 MW) x $12M/4 + MSFT T1+T2 $193M = ~$595M vs $582.3M reported"])
RV.append(["Oct 1 hike effect = on-demand share 4% x 19% = ~+0.8% of revenue (applied to non-MSFT revenue)"])
RV.column_dimensions['A'].width=58
C['rv_tot']=r5

# ======================= PART 3: Fleet Model =======================
QM=Q[1:]; n=len(QM); mcol=[CL(3+i) for i in range(n)]; YR=[int(q[-4:]) for q in QM]
M=wb.create_sheet("Model"); M.append(["Fleet quarterly model ($M)",""]+QM)
def row(label,cells,fmt='#,##0',bold=False,note=""):
    M.append([label,note]+cells); r=M.max_row
    for i in range(n): M.cell(r,3+i).number_format=fmt
    if bold: M.cell(r,1).font=B
    return r
def mf(label,fn,**k): return row(label,[fn(i,mcol[i]) for i in range(n)],**k)
def ref(r,i): return f"{mcol[i]}{r}"
x=lambda k: a(k)
r_i=row("Quarter index (0 = Q3'26)",list(range(n)),fmt='0'); 
r_y=row("Fiscal year",YR,fmt='0')
# capex
def cxf(i,c):
    y=YR[i]
    if y==2026: return f"=({x('cx26')}-{x('h1cx')})/2"
    if y==2027: return f"={x('cx27')}/4"
    if y==2028: return f"={x('cx28')}*{x('w%d'%(i-6))}"
    if y==2029: return f"={x('cx28')}*(1+{x('g29')})/4"
    return f"={x('cx28')}*(1+{x('g29')})*(1+{x('g30')})/4"
r_cxc=mf("Capex - consensus basis (FY26-28; FY29-30 extrapolated)",cxf)
# owned share of FY energized additions (MW sheet columns: C=Q2'26, F=Q1'27..; Q4 of year y)
q4col={2026:"E",2027:"I",2028:"M",2029:"Q",2030:"U"}; prev={2026:"C",2027:"E",2028:"I",2029:"M",2030:"Q"}
def ownf(i,c):
    y=YR[i]; o=C['mw_own']; t=C['mw_sel']
    add=f"(MW!{q4col[y]}{o}-MW!{prev[y]}{o})/(MW!{q4col[y]}{t}-MW!{prev[y]}{t})"
    cum=f"MW!{q4col[y]}{o}/MW!{q4col[y]}{t}"
    return f"=CHOOSE({x('ownmode')},IFERROR(IF(MW!{q4col[y]}{t}-MW!{prev[y]}{t}>0,{add},{cum}),0),0,1)"
r_own=mf("Owned share of capex (FY energized additions, Capacity tracker)",ownf,fmt='0%')
gen=lambda i,s: f"IF({mcol[i]}{r_i}>={x('vrq')},{x('vr_'+s)},{x('gb300_'+s)})"
r_gsh=mf("GPU+network share of capex (consensus basis)",lambda i,c: f"={x('gpu26')}" if YR[i]==2026 else f"={gen(i,'g')}/({gen(i,'g')}+{c}{r_own}*{gen(i,'s')}+(1-{c}{r_own})*({gen(i,'c')}-{gen(i,'g')}))",fmt='0%')
ad=C['mw_add']
def addat(i,c,lagexpr):
    m=f"({c}{r_i}+1+{lagexpr})"
    return f"IF({m}+1<=19,INDEX(MW!$C${ad}:$U${ad},{m}+1),AVERAGE(MW!$R${ad}:$U${ad}))"
r_ag=mf("Path: IT MW whose GPUs are bought this qtr (energized at t + lagb - pglag)",lambda i,c: "="+addat(i,c,f"{x('lagb')}-{x('pglag')}"),fmt='#,##0.0')
r_as=mf("Path: IT MW whose shell is built this qtr (energized at t + shlag)",lambda i,c: "="+addat(i,c,x('shlag')),fmt='#,##0.0')
r_pg=mf("Path: GPU/IT capex (owned: GPU+net by gen; leased: leased $/IT MW by gen - GB300 before VR switch, VR after)",lambda i,c: f"={c}{r_ag}*({c}{r_own}*{gen(i,'g')}+(1-{c}{r_own})*IF({c}{r_i}>={x('vrq')},{x('lease_vr')},{x('lease_gb300')}))")
r_ps=mf("Path: shell/power capex (owned share only)",lambda i,c: f"={c}{r_as}*{c}{r_own}*{gen(i,'s')}")
r_cxp=mf("Capex - path-consistent basis",lambda i,c: f"={c}{r_pg}+{c}{r_ps}")
r_cx=mf("CAPEX USED (basis toggle)",lambda i,c: f"=IF({x('cxbasis')}=1,{c}{r_cxc},{c}{r_cxp})",bold=True)
r_itmw=mf("Capex-implied IT MW bought (consensus capex / all-in $ per IT MW)",lambda i,c: f"={c}{r_cxc}/({gen(i,'g')}+{c}{r_own}*{gen(i,'s')}+(1-{c}{r_own})*({gen(i,'c')}-{gen(i,'g')}))",fmt='#,##0')
r_gcx=mf("GPU+network capex",lambda i,c: f"=IF({x('cxbasis')}=1,{c}{r_cx}*{c}{r_gsh},{c}{r_pg})")
r_ncx=mf("Shell/power/fit-out capex",lambda i,c: f"={c}{r_cx}-{c}{r_gcx}")
M.append(["Effective capex->in-service lag for GPUs (consensus: lage; path: lagb - pglag)",f"=IF({x('cxbasis')}=1,{x('lage')},MAX(0,{x('lagb')}-{x('pglag')}))"]); r_ise=M.max_row
def gadd(i,c):
    if i==0: return f"={x('niu')}*{x('gpu26')}*{x('niu3')}"
    if i==1: return f"={x('niu')}*{x('gpu26')}*(1-{x('niu3')})"
    return f"=IF({c}{r_i}>=$B${r_ise},INDEX($C${r_gcx}:$T${r_gcx},{c}{r_i}-$B${r_ise}+1),0)"
r_gin=mf("GPU+network placed in service (capex lagged by capex->energized; 6/30 not-in-use first)",gadd)
r_nin=mf("Non-GPU placed in service",lambda i,c: (f"={x('niu')}*(1-{x('gpu26')})*"+(x('niu3') if i==0 else f"(1-{x('niu3')})")) if i<2 else f"=IF({c}{r_i}>={x('lage')},INDEX($C${r_ncx}:$T${r_ncx},{c}{r_i}-{x('lage')}+1),0)")
r_gcum=mf("Gross GPU+network in service (end of qtr)",lambda i,c: f"={x('srv')}+SUM($C${r_gin}:{c}{r_gin})")
r_gbill=mf("Gross GPU+network BILLING (in service lagged by energized->billing)",lambda i,c: f"=IF({c}{r_i}>={x('lagb')},INDEX($C${r_gcum}:$T${r_gcum},{c}{r_i}-{x('lagb')}+1),{x('srv')})")
r_ncum=mf("Gross non-GPU in service",lambda i,c: f"={x('othpp')}+SUM($C${r_nin}:{c}{r_nin})")
# D&A by life
r_da={}
for L in ("L1","L2","L3"):
    q=f"(4*{x(L)})"
    r_da[L]=mf(f"GPU+network D&A at life {L}",lambda i,c,q=q: f"=IF({c}{r_i}<{q}-{x('age')},{x('srv')}/{q},0)+SUMPRODUCT($C${r_gin}:{c}{r_gin},((({c}{r_i}-$C${r_i}:{c}{r_i})=0)*0.5+((({c}{r_i}-$C${r_i}:{c}{r_i})>0)*(({c}{r_i}-$C${r_i}:{c}{r_i})<{q}))+((({c}{r_i}-$C${r_i}:{c}{r_i})={q})*0.5))/{q})")
r_dan=mf("Non-GPU D&A (existing at ngl, new at shell life)",lambda i,c: f"={x('othpp')}/(4*{x('ngl')})+SUMPRODUCT($C${r_nin}:{c}{r_nin},((({c}{r_i}-$C${r_i}:{c}{r_i})=0)*0.5+(({c}{r_i}-$C${r_i}:{c}{r_i})>0))/(4*{x('shl')}))")
# EBITDA
ebq=[355,671,1277,1581,1883,2372]; rvq=[904.1,1452,2175,2646,3304,4052]
r_rcons=row("Revenue - consensus (Q3'26-Q4'27 quarterly, FY28 weights)",[rvq[i] if i<6 else f"={x('rv28')}*{x('w%d'%(i-6))}" if YR[i]==2028 else None for i in range(n)])
for i in range(6): M.cell(r_rcons,3+i).font=BLUE
r_econs=row("Adj. EBITDA - consensus",[ebq[i] if i<6 else f"={x('eb28')}*{x('w%d'%(i-6))}" if YR[i]==2028 else None for i in range(n)])
for i in range(6): M.cell(r_econs,3+i).font=BLUE
# yields FY28
fy28=[mcol[i] for i in range(n) if YR[i]==2028]; f28=lambda r: f"AVERAGE({fy28[0]}{r}:{fy28[-1]}{r})"
M.append([]); 
M.append(["FY28 EBITDA yield on avg gross GPU IN SERVICE (v1 basis)",f"={x('eb28')}/{f28(r_gcum)}"]); r_y1=M.max_row
M.append(["FY28 EBITDA yield on avg gross GPU BILLING capital (v2 default)",f"={x('eb28')}/{f28(r_gbill)}"]); r_y2=M.max_row
M.append(["FY28 EBITDA per avg billing gross MW (Capacity path, selected scenario) $M/MW",f"={x('eb28')}/AVERAGE(MW!J{C['mw_bill']}:MW!M{C['mw_bill']})"]); r_ymw=M.max_row
M.append(["FY28 consensus revenue per avg billing gross MW $M/MW",f"={x('rv28')}/AVERAGE(MW!J{C['mw_bill']}:MW!M{C['mw_bill']})"]); r_rmw=M.max_row
M.append(["FY29+ yield used",f"=IF(ISNUMBER({x('yov')}),{x('yov')},IF({x('ybasis')}=1,B{r_y2},B{r_y1}))"]); r_yu=M.max_row
M.append(["FY28 EBITDA margin (FY29+ revenue = EBITDA / margin)",f"={x('eb28')}/{x('rv28')}"]); r_mg=M.max_row
for rr_ in (r_y1,r_y2,r_yu,r_mg): M.cell(rr_,2).number_format='0.0%'
for rr_ in (r_ymw,r_rmw): M.cell(rr_,2).number_format='0.0'
M.append([])
basis=lambda i,c: f"IF({x('ybasis')}=1,AVERAGE(INDEX($C${r_gbill}:$T${r_gbill},{c}{r_i}),INDEX($C${r_gbill}:$T${r_gbill},{c}{r_i}+1)),AVERAGE(INDEX($C${r_gcum}:$T${r_gcum},{c}{r_i}),INDEX($C${r_gcum}:$T${r_gcum},{c}{r_i}+1)))"
r_eb=mf("Adj. EBITDA used",lambda i,c: f"={c}{r_econs}" if YR[i]<=2028 else f"=$B${r_yu}/4*{basis(i,c)}",bold=True)
r_rv=mf("Revenue used",lambda i,c: f"={c}{r_rcons}" if YR[i]<=2028 else f"={c}{r_eb}/$B${r_mg}")
r_ebit={}; r_ebm={}
for L in ("L1","L2","L3"):
    r_ebit[L]=mf(f"EBIT at life {L}",lambda i,c,L=L: f"={c}{r_eb}-{c}{r_da[L]}-{c}{r_dan}",bold=True)
for L in ("L1","L2","L3"):
    r_ebm[L]=mf(f"EBIT margin at life {L}",lambda i,c,L=L: f"={c}{r_ebit[L]}/{c}{r_rv}",fmt='0.0%')
# prepayments & FCF
r_pp=mf("New customer prepayments (pinc x ppct x capex)",lambda i,c: f"={x('pptog')}*{x('pinc')}*{x('ppct')}*{c}{r_cx}")
r_am=mf("Deferred revenue recognized (non-cash revenue in EBITDA)",lambda i,c: f"={x('pptog')}*({x('dr0')}/{x('drn')}+SUMPRODUCT($C${r_pp}:{c}{r_pp},(({c}{r_i}-$C${r_i}:{c}{r_i})<{x('drn')})*1)/{x('drn')})")
r_int=mf("Cash interest (converts + loan) + interest on funding gap",lambda i,c: f"={x('int_q')}" if i==0 else f"={x('int_q')}+MAX(0,-{mcol[i-1]}{{LIQ}})*{x('rate')}/4")
r_fcf=mf("FCF (EBITDA - capex - interest + prepayments - deferred revenue recognized)",lambda i,c: f"={c}{r_eb}-{c}{r_cx}-{c}{r_int}+{c}{r_pp}-{c}{r_am}",bold=True)
r_fcfx=mf("FCF ex-prepayment effects",lambda i,c: f"={c}{r_eb}-{c}{r_cx}-{c}{r_int}")
r_liq=mf("Liquidity (starting liquidity + cumulative FCF; negative = funding need)",lambda i,c: f"={x('liq')}+SUM($C${r_fcf}:{c}{r_fcf})",bold=True)
for i in range(1,n):
    cell=M.cell(r_int,3+i); cell.value=cell.value.replace("{LIQ}",str(r_liq))
M.append([]); M.append(["First quarter with negative liquidity",f'=IFERROR(INDEX($C$1:$T$1,MATCH(TRUE,INDEX($C${r_liq}:$T${r_liq}<0,0),0)),"none by Q4 2030")']); r_gapq=M.max_row
M.append(["Peak funding need ($M)",f"=-MIN(0,MIN($C${r_liq}:$T${r_liq}))"]); r_peak=M.max_row
M.append(["Peak funding need in new shares at $242.81 (M)",f"=B{r_peak}/{x('px')}"]); r_peaksh=M.max_row
M.column_dimensions['A'].width=72
for c in M[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=HDR
for k,v in dict(r_cx=r_cx,r_eb=r_eb,r_rv=r_rv,r_fcf=r_fcf,r_fcfx=r_fcfx,r_liq=r_liq,r_gcum=r_gcum,r_gbill=r_gbill,r_ncum=r_ncum,r_dan=r_dan,r_y1=r_y1,r_y2=r_y2,r_yu=r_yu,r_ymw=r_ymw,r_rmw=r_rmw,r_gapq=r_gapq,r_peak=r_peak,r_peaksh=r_peaksh,r_itmw=r_itmw,r_own=r_own,r_cxc=r_cxc,r_cxp=r_cxp,r_pp=r_pp,r_am=r_am,r_int=r_int).items(): C[k]=v
C['r_da']=r_da; C['r_ebit']=r_ebit; C['r_ebm']=r_ebm

# ---- Annual
AN=wb.create_sheet("Annual"); yrs=[2026,2027,2028,2029,2030]; AN.append(["Annual roll-up ($M; FY26 = 1H actual + 2H model)"]+yrs)
def an(label,r,add=None,fmt='#,##0'):
    AN.append([label]+[f"=SUMPRODUCT((Model!$C$%d:$T$%d=%d)*Model!$C$%d:$T$%d)"%(r_y,r_y,y,r,r)+(f"+{add}" if (add and y==2026) else "") for y in yrs]); rr_=AN.max_row
    for j in range(5): AN.cell(rr_,2+j).number_format=fmt
    return rr_
A_rv=an("Revenue",r_rv,x('h1rv')); A_eb=an("Adj. EBITDA",r_eb,x('h1eb')); A_cx=an("Capex",r_cx,x('h1cx'))
A_da={L:an(f"Total D&A at life {L} (FY26 H1 actual at 5y; 4y adds back {86.1})",r_da[L],None) for L in ("L1","L2","L3")}
for L,rr_ in A_da.items():
    for j in range(5): AN.cell(rr_,2+j).value=AN.cell(rr_,2+j).value+f"+SUMPRODUCT((Model!$C${r_y}:$T${r_y}={yrs[j]})*Model!$C${r_dan}:$T${r_dan})"+(f"+{x('h1da')}"+(f"+{x('h1ben')}" if L=="L1" else "") if j==0 else "")
A_ebit={L:AN.max_row+1+k for k,L in enumerate(("L1","L2","L3"))}
for L in ("L1","L2","L3"): AN.append([f"EBIT at life {L}"]+[f"={CL(2+j)}{A_eb}-{CL(2+j)}{A_da[L]}" for j in range(5)])
for L in ("L1","L2","L3"): AN.append([f"EBIT margin at life {L}"]+[f"={CL(2+j)}{A_ebit[L]}/{CL(2+j)}{A_rv}" for j in range(5)]); [setattr(AN.cell(AN.max_row,2+j),'number_format','0.0%') for j in range(5)]
AN.append(["EBIT-margin points added by 4->5 yr life change (governance leg)"]+[f"=({CL(2+j)}{A_da['L1']}-{CL(2+j)}{A_da['L2']})/{CL(2+j)}{A_rv}" for j in range(5)]); A_pts=AN.max_row
for j in range(5): AN.cell(A_pts,2+j).number_format='0.0%'
A_fcf=an("FCF (H2 only for 2026)",r_fcf); A_fcfx=an("FCF ex-prepayments (H2 only for 2026)",r_fcfx)
AN.append(["Liquidity at year end"]+[f"=INDEX(Model!$C${r_liq}:$T${r_liq},MATCH({y},Model!$C${r_y}:$T${r_y},1))" for y in yrs]); A_liq=AN.max_row
AN.append(["Bottom-up revenue from billing MW (Rev_MW)"]+[f"=SUMPRODUCT((RIGHT(Rev_MW!$C$1:$T$1,4)=\"{y}\")*Rev_MW!$C${C['rv_tot']}:$T${C['rv_tot']})"+(f"+{x('h1rv')}" if y==2026 else "") for y in yrs]); A_bu=AN.max_row
AN.column_dimensions['A'].width=66
for r_ in range(2,AN.max_row+1):
    for j in range(5):
        if AN.cell(r_,2+j).number_format=='General': AN.cell(r_,2+j).number_format='#,##0'
C.update(A_rv=A_rv,A_eb=A_eb,A_ebit=A_ebit,A_pts=A_pts,A_fcf=A_fcf,A_fcfx=A_fcfx,A_liq=A_liq,A_bu=A_bu,A_da=A_da)

# ---- Horizons (6m = Q1'27, 1y = Q3'27, 3y = Q3'29 vs Q3'26 start)
H=wb.create_sheet("Horizons"); H.append(["Metric (quarter at horizon)","6-month (Q1 2027)","1-year (Q3 2027)","3-year (Q3 2029)"])
hc=["E","G","O"]
for lab,r_ in [("Revenue",r_rv),("Adj. EBITDA",r_eb),("EBIT 4y",r_ebit['L1']),("EBIT 5y",r_ebit['L2']),("EBIT 6y",r_ebit['L3']),("FCF",r_fcf),("Liquidity",r_liq),("Gross GPU billing capital",r_gbill)]:
    H.append([lab]+[f"=Model!{c}{r_} " .strip() for c in hc])
H.append(["Billing gross MW (Capacity path)"]+[f"=MW!{chr(ord(c)+1)}{C['mw_bill']}" for c in hc])
H.column_dimensions['A'].width=36

# ======================= PART 4: Reverse DCF, Cohorts =======================
r_avb=mf("Avg billing GPU capital in quarter (for yield)",lambda i,c: f"=AVERAGE({c}{r_gbill},{(mcol[i-1] if i else c)}{r_gbill})")
r_avs=mf("Avg in-service GPU capital in quarter",lambda i,c: f"=AVERAGE({c}{r_gcum},{(mcol[i-1] if i else c)}{r_gcum})")
RD=wb.create_sheet("Reverse_DCF")
grid=[round(0.10+0.05*k,2) for k in range(13)]   # 10%..70%
RD.append(["Unlevered DCF EV ($M) by FY29+ EBITDA yield on billing GPU capital (columns) and GPU life (rows). Pre-2029 = consensus. Discounted from Q3'26 at WACC; terminal at Q4'30 run-rate with maintenance capex = capital/life"]+[""]*1)
RD.append(["Life \\ yield"]+grid); 
for j in range(len(grid)): RD.cell(2,2+j).number_format='0%'; RD.cell(2,2+j).font=BLUE
w=x('wacc'); tg=x('tg'); tx=x('tax'); rng=lambda r,a_="C",b_="T": f"Model!${a_}${r}:${b_}${r}"
disc=f"(1+{w})^(-(Model!$C${r_i}:$T${r_i}+0.5)/4)"
pre=f"(Model!$C${r_y}:$T${r_y}<=2028)"; post=f"(Model!$C${r_y}:$T${r_y}>=2029)"
evrow={}
for L in ("L1","L2","L3"):
    da=f"(Model!$C${r_da[L]}:$T${r_da[L]}+Model!$C${r_dan}:$T${r_dan})"
    cells=[]
    for j in range(len(grid)):
        y=f"{CL(2+j)}$2"
        eb=f"(({pre})*Model!$C${r_eb}:$T${r_eb}+({post})*{y}/4*Model!$C${r_avb}:$T${r_avb})"
        ebit=f"({eb}-{da})"
        fcf=f"({eb}-Model!$C${r_cx}:$T${r_cx}-{tx}*({ebit}>0)*{ebit})"
        capT=f"(Model!$T${r_gbill}+Model!$T${r_ncum})"
        maint=f"(Model!$T${r_gbill}/{x(L)}+Model!$T${r_ncum}/{x('shl')})"
        ebT=f"({y}*Model!$T${r_gbill})"
        fT=f"({ebT}-{tx}*MAX(0,{ebT}-{maint})-{maint}-{tg}*{capT})"
        cells.append(f"=SUMPRODUCT({fcf}*{disc})+{fT}*(1+{tg})/({w}-{tg})*(1+{w})^(-18/4)")
    RD.append([f"EV at life {L}"]+cells); evrow[L]=RD.max_row
RD.append([])
RD.append(["Market EV (ex-leases)",f"={x('ev')}"]); r_mev=RD.max_row
RD.append(["Market EV incl. off-BS commitments (debt-like variant)",f"={x('evx')}"]); r_mevx=RD.max_row
RD.append([]); RD.append(["Life","Implied FY29+ yield (market EV)","Implied yield (EV incl. off-BS)","FY28 yield on billing capital (v2 default)","FY28 yield in service (v1 basis)","EV at v2 default yield","Value gap $/share at v2 yield (308.69M sh)","EV at v1-basis yield","Value gap $/share at v1 yield"])
hdr_r=RD.max_row
def interp(L,target):
    e=f"$B${evrow[L]}:$N${evrow[L]}"; g="$B$2:$N$2"
    k=f"MATCH({target},{e},1)"
    return f"=IFERROR(INDEX({g},{k})+({target}-INDEX({e},{k}))*(INDEX({g},{k}+1)-INDEX({g},{k}))/(INDEX({e},{k}+1)-INDEX({e},{k})),IF({target}<$B${evrow[L]},\"<10%\",\">70%\"))"
def evat(L,yc):
    e=f"$B${evrow[L]}:$N${evrow[L]}"; g="$B$2:$N$2"
    k=f"MIN(12,MAX(1,MATCH({yc},{g},1)))"
    return f"(INDEX({e},{k})+({yc}-INDEX({g},{k}))*(INDEX({e},{k}+1)-INDEX({e},{k}))/(INDEX({g},{k}+1)-INDEX({g},{k})))"
for L in ("L1","L2","L3"):
    r=RD.max_row+1
    RD.append([f"=Assumptions!B{int(x(L).split('$')[-1])}",interp(L,f"$B${r_mev}"),interp(L,f"$B${r_mevx}"),f"=Model!B{r_y2}",f"=Model!B{r_y1}",f"={evat(L,f'D{r}')}",f"=(F{r}-$B${r_mev})/{x('sh')}",f"={evat(L,f'E{r}')}",f"=(H{r}-$B${r_mev})/{x('sh')}"])
    for cc in (2,3,4,5): RD.cell(r,cc).number_format='0.0%'
    for cc in (6,8): RD.cell(r,cc).number_format='#,##0'
    for cc in (7,9): RD.cell(r,cc).number_format='$#,##0'
C['rd_tab']=hdr_r; C['evrow']=evrow
RD.append([]); RD.append(["Note: off-BS variant compares an EBITDA-after-rent-and-power DCF with an EV that also counts lease/energy commitments - directionally conservative (double-counts), shown as a bound only."])
RD.column_dimensions['A'].width=30
for k in "BCDEFGHI": RD.column_dimensions[k].width=17

# ---- Cohorts
CO=wb.create_sheet("Cohorts")
hdr=["Case","Cohort","ACV $M/gross MW-yr","Owned (1)/Leased (0)","GPU life (yrs)","Contract term (yrs)","Capex $M/gross MW","Contract EBITDA margin"]+[f"Project CF yr{k}" for k in range(7)]+["Project IRR","Project payback (yrs)","Project NPV @WACC","Breakeven EBITDA margin (NPV=0)","Accounting breakeven margin (D&A/ACV)","Prepayment $M/MW (prepaid deal)"]+[f"Equity CF yr{k}" for k in range(7)]+["Equity IRR (prepaid deal)","Equity payback (yrs)"]+[f"cumP{k}" for k in range(7)]+[f"cumE{k}" for k in range(7)]
CO.append(hdr)
res1=int(x('res1').split('$')[-1])
cohorts=[("Installed book ($12M, Hopper-era capex)","acv1","hop","term_lt",False),("Q2 deals low ($20M)","acv2","gb300","term_lt",False),("Q2 deals high ($25M)","acv3","gb300","term_lt",False),("Short-term (>$40M, 1-yr)","acv4","gb300","term_st",False),("Microsoft ($11.6M = 300 MW gross basis, leased)","acv5","gb300","term_ms",True),("Microsoft ($17.4M = 200 MW basis, leased)","acv6","gb300","term_ms",True)]
case=0; co_rows=[]
for nm,acv,g,term,ms in cohorts:
    for own in ((0,) if ms else (1,0)):
        for L in ("L1","L2","L3"):
            r=CO.max_row+1; case+=1
            capex=f"IF(D{r}=1,{x(g+'_g')}+{x(g+'_s')},{x(g+'_c')})/{x('pue')}"
            row_=[case,nm,f"={x(acv)}*(1-{x('sla') if ms else 0})",own,f"={x(L)}",f"={x(term)}",f"={capex}"]
            opc=f"({x('pwr_g')}+(1-D{r})*{x('rent')}/{x('pue')})"
            row_.append(f"=1-{x('oth')}-{opc}/C{r}")
            cfs=[f"=-G{r}"]
            for k in range(1,7):
                rev=f"IF({k}>E{r},0,IF({k}<=F{r},C{r},MIN(C{r},{x('acvref')})*INDEX(Assumptions!$B${res1}:$B${res1+5},{k})*{x('util')}))"
                op=f"IF({k}>E{r},0,{rev}*{x('oth')}+{x('pwr_g')}*IF({k}<=F{r},1,{x('util')})+(1-D{r})*{x('rent')}/{x('pue')})"
                sal=f"IF({k}=E{r},D{r}*{x(g+'_s')}/{x('pue')}*MAX(0,1-E{r}/{x('shl')}),0)"
                cfs.append(f"={rev}-{op}+{sal}")
            row_+=cfs
            P=f"I{r}:O{r}"
            row_.append(f"=IFERROR(IRR({P}),-1)")
            cumP=[f"{CL(31+k)}{r}" for k in range(7)]; cumE=[f"{CL(38+k)}{r}" for k in range(7)]
            pb=lambda cr,cf0: f'=IFERROR(MATCH(TRUE,INDEX({cr[0]}:{cr[-1]}>=0,0),0)-2+(-INDEX({cr[0]}:{cr[-1]},MATCH(TRUE,INDEX({cr[0]}:{cr[-1]}>=0,0),0)-1))/INDEX({cf0},MATCH(TRUE,INDEX({cr[0]}:{cr[-1]}>=0,0),0)),"never")'
            row_.append(pb(cumP,P))
            row_.append(f"=NPV({x('wacc')},J{r}:O{r})+I{r}")
            pvrev="+".join(f"IF({k}>E{r},0,IF({k}<=F{r},C{r},MIN(C{r},{x('acvref')})*INDEX(Assumptions!$B${res1}:$B${res1+5},{k})*{x('util')}))/(1+{x('wacc')})^{k}" for k in range(1,7))
            pvsal=f"D{r}*{x(g+'_s')}/{x('pue')}*MAX(0,1-E{r}/{x('shl')})/(1+{x('wacc')})^E{r}"
            row_.append(f"=(G{r}-{pvsal})/({pvrev})")
            row_.append(f"=(IF(D{r}=1,{x(g+'_g')}/E{r}+{x(g+'_s')}/{x('shl')},{x(g+'_g')}/E{r}+({x(g+'_c')}-{x(g+'_g')})/{x('shl')})/{x('pue')})/C{r}")
            row_.append(f"={x('msup')}*C{r}*F{r}" if ms else f"={x('ppct')}*G{r}")
            eq=[f"=I{r}+U{r}"]+[f"={CL(9+k)}{r}-IF({k}<=F{r},$U{r}/F{r},0)" for k in range(1,7)]
            row_+=eq
            E_=f"V{r}:AB{r}"
            row_.append(f"=IFERROR(IRR({E_}),IF(V{r}>=0,\"no equity at risk\",-1))")
            row_.append(pb(cumE,E_))
            row_+= [f"=SUM($I{r}:{CL(9+k)}{r})" for k in range(7)]+[f"=SUM($V{r}:{CL(22+k)}{r})" for k in range(7)]
            CO.append(row_); co_rows.append(r)
            for cc in list(range(9,16))+list(range(22,29))+[7,18,21]: CO.cell(r,cc).number_format='#,##0.0'
            for cc in (8,16,19,20,29): CO.cell(r,cc).number_format='0.0%'
            for cc in (17,30): CO.cell(r,cc).number_format='0.0'
CO.append([]); CO.append(["Notes: per gross MW. Capex from Capacity base (owned = GPU/net + shell; leased = colo capex incl. fit-out, no shell). Opex = power (Capacity $93-185/MWh midpoint) + colo rent (leased) + other cash opex calibrated to Q2 segment margin. After term: re-rent at MIN(ACV, Q2-low ref) x PROVISIONAL residual index x 60% utilization. Owned shell salvage = NBV at end of GPU life. Equity = prepaid deal (prepayment = 55% of capex; Microsoft = 40% of TCV upfront), repaid via lower cash revenue over the term; blended book = 70% prepaid. Short-term deal capex is a full new build - optimistic if it re-uses installed GPUs."])
for c in CO[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=HDR
CO.column_dimensions['B'].width=34
C['co_rows']=co_rows

# ---- Charts placeholder sheet & save
wb.create_sheet("Charts")
wb.save("/workspace/nebius/skeptic/03_depreciation.xlsx")
json.dump({k:v for k,v in C.items()},open("cells.json","w"),default=str)
print("saved")
