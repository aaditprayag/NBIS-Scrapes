#!/usr/bin/env python3
# Final appendix-ready peer KPI pack (Nebius Peer KPIs, built 2026-10-09 ET for the Sun Oct 11 finals).
# Nebius vs CoreWeave, IREN, Oracle (OCI), Nscale, Lambda, Crusoe: $/MW-yr, contracted vs live MW, capex/MW.
# NO new research: every value is READ from the locked deliverables (peers/comparison.xlsx, benchmark/peer_benchmark.xlsx,
# price_or_mw.xlsx, sites/sites.xlsx + site appendices, model v0.9e via nebius_locks.py). The script asserts every lock and
# the cell-for-cell agreement of overlapping figures, and STOPS (SystemExit) on any drift. Writes only peers/final_appendix/.
# Run: python3 build_peer_kpi_appendix.py   (recalcs the xlsx with model/recalc.py and converts the docx with soffice)
import os, re, sys, shutil, subprocess, statistics, textwrap
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
sys.path.insert(0, "/workspace/nebius/peers")
import nebius_locks as NL

W = "/workspace/nebius"; P = f"{W}/peers"; D = f"{P}/final_appendix"
XLSX, DOCX, PDF = f"{D}/peer_kpi_appendix.xlsx", f"{D}/peer_kpi_appendix.docx", f"{D}/peer_kpi_appendix.pdf"
PNG1, PNG2, PNG3 = (f"{D}/{n}.png" for n in ("chart_usd_per_mw_yr", "chart_contracted_vs_live_mw", "chart_capex_per_mw"))
AS_OF = "Fri Oct 9, 2026 (ET); frozen for Sun Oct 11 finals"
PROBLEMS = []
def need(cond, msg):
    if not cond: PROBLEMS.append(msg)
def stop_if_problems(stage):
    if PROBLEMS:
        print(f"LOCK / CONSISTENCY DRIFT at {stage} - build stopped:"); [print("  -", p) for p in PROBLEMS]; raise SystemExit(2)
def close(a, b, tol=1e-6): return isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) <= tol

# ---------------------------------------------------------------- load sources (read-only)
CMP = openpyxl.load_workbook(f"{P}/comparison.xlsx", data_only=True)
BMK = openpyxl.load_workbook(f"{P}/benchmark/peer_benchmark.xlsx", data_only=True)
POM = openpyxl.load_workbook(f"{P}/price_or_mw.xlsx", data_only=True)
SIT = openpyxl.load_workbook(f"{P}/sites/sites.xlsx", data_only=True)["Sites"]
MOD = openpyxl.load_workbook(f"{W}/model/nebius_model.xlsx", data_only=True)
CAPQ = openpyxl.load_workbook(f"{W}/capacity/nebius_capacity.xlsx", data_only=True)["Quarterly MW"]
ML = NL.model_locks()
# v2 sources (read-only): GPU Market Print_Table + final band slide; Capacity Meta #1 cells; Oracle I45E
PT = openpyxl.load_workbook(f"{W}/market/contract_economics.xlsx", data_only=True)["Print_Table"]
PT8 = [r for r in PT.iter_rows(min_row=6, max_row=13, values_only=True)]
SLIDE = f"{W}/market/charts/contract_band_slide_final.png"
CAPT = openpyxl.load_workbook(f"{W}/capacity/nebius_capacity.xlsx", data_only=True)["Capacity tranches"]
ORI = openpyxl.load_workbook(f"{P}/oracle/oracle_kpis.xlsx", data_only=True)["Inputs"]
SITE_MD = {k: open(f"{P}/sites/appendix_{k}.md", encoding="utf-8").read() for k in ("coreweave_denton", "iren_childress", "crusoe_abilene", "oracle_port_washington")}

cmpC = CMP["Comparison"]; COLS = {"Nebius": 4, "CoreWeave": 7, "IREN": 10, "Oracle OCI": 13, "Nscale": 16, "Lambda": 19, "Crusoe": 22}
KROW = {cmpC.cell(r, 1).value: r for r in range(3, cmpC.max_row + 1) if cmpC.cell(r, 1).value}
def K(kid, peer):  # (value, tag, as_of) from comparison.xlsx Comparison
    r = KROW[kid]; c = COLS[peer]; return cmpC.cell(r, c).value, cmpC.cell(r, c + 1).value, cmpC.cell(r, c + 2).value
def Kmed(kid): return cmpC.cell(KROW[kid], 25).value
cmpK = CMP["Contracts_$MWyr"]; CON = {r[0]: r for r in cmpK.iter_rows(min_row=4, values_only=True) if r[0]}
ni = {r[0]: r for r in CMP["Nebius_Inputs"].iter_rows(min_row=3, values_only=True) if r[0]}
cd = CMP["Chart_Data"]; CDR = {cd.cell(r, 1).value: r for r in range(1, cd.max_row + 1) if cd.cell(r, 1).value}
bK = {r[0]: r for r in BMK["Contracts"].iter_rows(min_row=3, values_only=True) if r[0]}
bP = {r[0]: r for r in BMK["KPI_Prints"].iter_rows(min_row=3, values_only=True) if r[0]}
bL = {r[0]: r for r in BMK["Nebius_Locks"].iter_rows(min_row=3, values_only=True) if r[0]}
bB = {r[1]: r for r in BMK["Comp_Band"].iter_rows(min_row=3, values_only=True) if r[1]}
bS = {r[0]: r for r in BMK["Sources"].iter_rows(min_row=2, values_only=True) if r[0]}
bench = BMK["Benchmark"]; BROW = {bench.cell(r, 1).value: r for r in range(3, 12) if bench.cell(r, 1).value}
def bmw(name): return str(bench.cell(BROW[name], 7).value)
pomN, pomI = POM["Nebius_Check"], POM["Implied_Price"]
sites = {(r[0], str(r[1]).split(" (")[0]): r for r in SIT.iter_rows(min_row=2, values_only=True) if r[0] and r[1]}

# ---------------------------------------------------------------- 1) LOCKS (model v0.9e) - stop on drift
m = lambda s, c: MOD[s][c].value
LOCKS = [  # key, label, expected (display), model value, tol, tag, model cell
 ("MSFT", "Nebius-Microsoft $M per IT MW-yr [E, house]", 14.49, m("Contracts", "E12"), 0.005, "E (house)", "model Contracts!E12"),
 ("HURDLE", "Ex-MSFT hurdle (peer comparator), $M per billed IT MW-yr", 27.63, m("Drivers", "B9"), 0.005, "D", "model Drivers!B9"),
 ("BLENDED", "$23.13M = Street FY27 ÷ our 531.8 billed MW, all customers", 23.13, m("Drivers", "B10"), 0.005, "D", "model Drivers!B10"),
 ("CLEAR_BASE", "Clear-spot BASE, $M per billed IT MW-yr", 24.51, m("Drivers", "B6"), 0.005, "D", "model Drivers!B6"),
 ("CLEAR_HIGH", "Clear-spot HIGH", 22.68, m("Debate_Grid", "A18"), 0.005, "D", "model Debate_Grid!A18"),
 ("TP_BASE", "Turn-positive BASE", 17.05, m("Debate_Grid", "A13"), 0.005, "D", "model Debate_Grid!A13"),
 ("TP_HIGH", "Turn-positive HIGH", 16.70, m("Debate_Grid", "A10"), 0.005, "D", "model Debate_Grid!A10"),
 ("EN_Q1", "Energized MW 31-Mar-26", 181.5, CAPQ["M130"].value, 1e-9, "E", "capacity Quarterly MW!M130"),
 ("EN_Q2", "Energized MW 30-Jun-26", 193, CAPQ["N130"].value, 1e-9, "E", "capacity Quarterly MW!N130"),
 ("YE26_BASE", "Billed IT MW YE26 base", 257.7, ML["YE26_BASE"]["value"], 0.05, "D", "model " + ML["YE26_BASE"]["cell"]),
 ("YE26_HIGH", "Billed IT MW YE26 high", 284.2, ML["YE26_HIGH"]["value"], 0.05, "D", "model " + ML["YE26_HIGH"]["cell"]),
 ("FY27_BASE", "FY27 avg billed IT MW base", 531.84, ML["FY27_BASE"]["value"], 0.005, "D", "model " + ML["FY27_BASE"]["cell"]),
 ("FY27_HIGH", "FY27 avg billed IT MW high", 659.16, ML["FY27_HIGH"]["value"], 0.005, "D", "model " + ML["FY27_HIGH"]["cell"]),
]
LK = {}
for key, lab, exp, val, tol, tag, cell in LOCKS:
    need(isinstance(val, (int, float)) and abs(val - exp) <= tol, f"{key}: model {cell} = {val} vs locked {exp}")
    LK[key] = dict(label=lab, exp=exp, val=val, tag=tag, cell=cell)
need(ML["VERSION"]["value"] == "v0.9e", "model version is not v0.9e")
need(close(ML["HURDLE_XMSFT"]["value"], LK["HURDLE"]["val"]) and close(ML["BLENDED"]["value"], LK["BLENDED"]["val"]), "nebius_locks vs model cells")
# K13 and the comp band
k13n = K("K13", "Nebius")[0]; k13med = Kmed("K13")
need(close(round(k13n, 1), 77.9, 1e-9), f"K13 Nebius {k13n} != 77.9"); need(close(round(k13med, 2), 33.18, 1e-9), f"K13 median {k13med} != 33.18")
BAND_IDS = ["P1", "NEW1", "P2", "O1", "NEW2", "O2"]
def con_best(cid):  # IT basis where available (col W), else stated MW (col V) - the comparison.xlsx band convention
    r = CON[cid]; return r[22] if isinstance(r[22], (int, float)) else r[21]
band = [con_best(c) for c in BAND_IDS]
need(close(min(band), 9.7) and close(round(max(band), 1), 16.7, 1e-9), f"band {min(band)}-{max(band)} != 9.7-16.7")
need(close(bB["Band low (MIN)"][2], min(band)) and close(bB["Band top (MAX)"][2], max(band)), "benchmark Comp_Band != comparison band")
# Bitdeer: only disclosed-MW contract print >= $20M/MW (benchmark Contracts; comparison Contracts current rows have none >= 20)
bd = bK["BD-A201"][10]; need(close(round(bd, 1), 22.4, 1e-9), f"Bitdeer A201 {bd} != 22.4")
ge20 = [k for k, r in bK.items() if isinstance(r[10], (int, float)) and r[10] >= 20 and isinstance(r[7], (int, float))]
need(ge20 == ["BD-A201"], f"disclosed-MW contracts >= $20M in benchmark: {ge20}")
ge20c = [k for k, r in CON.items() if not str(r[1]).startswith("HISTORY") and any(isinstance(r[i], (int, float)) and r[i] >= 20 for i in (21, 22))]
need(ge20c == [], f"comparison Contracts rows >= $20M: {ge20c}")
# v2: Nebius-Meta #1 = $11.52M [E, MW is a team estimate of ~50 IT MW]: comparison N2 == benchmark N2 == GPU Market Print_Table #3; one of eight band prints; ends 9.7/16.7
META1 = CON["N2"][22]
need(close(META1, 2880.7 / 5 / 50, 1e-9) and CON["N2"][11] == 50 and str(CON["N2"][12]).startswith("E"), f"comparison N2 {META1} / MW {CON['N2'][11]}")
need([r[0] for r in PT8] == list(range(1, 9)), "Print_Table A is not eight band prints")
_pm = [r for r in PT8 if "Nebius #1" in str(r[1])]
need(len(_pm) == 1 and close(_pm[0][8], round(META1, 2), 1e-9) and _pm[0][9] == "E" and str(_pm[0][6]) == "~50", f"Print_Table Meta #1 {_pm}")
need(close(min(r[8] for r in PT8), 9.7, 1e-9) and close(max(r[8] for r in PT8), 16.67, 1e-9), "Print_Table eight-print ends != 9.7 / 16.67")
need(os.path.exists(SLIDE), "final band slide missing")
need(str(CAPT["A3"].value).startswith("Meta #1") and CAPT["G3"].value == "n/d" and CAPT["H3"].value == "n/d", "Capacity Meta #1 MW cells changed")
META_MW_MODEL = MOD["Contracts"]["D23"].value + MOD["Contracts"]["D24"].value
need(close(META_MW_MODEL, 48.0117, 0.01), f"model Contracts!D23+D24 = {META_MW_MODEL}")
need(str(ORI["A48"].value) == "I45E" and close(ORI["P48"].value, 2050) and "Aug-26" in str(ORI["P2"].value), "oracle_kpis Inputs!P48 (I45E Aug-26) != 2,050")
stop_if_problems("locks")

# ---------------------------------------------------------------- 2) cell-for-cell agreement of overlapping figures
# 2a Nebius locks: model == comparison Nebius_Inputs == peer_benchmark Nebius_Locks == price_or_mw
pairs = [
 ("MSFT", CON["N1"][22], "comparison Contracts_$MWyr W(N1)"), ("MSFT", bK["N1-H"][10], "benchmark Contracts K(N1-H)"), ("MSFT", pomN["B22"].value, "price_or_mw Nebius_Check!B22"),
 ("HURDLE", ni["D_STREET_FY27_NONMSFT"][2], "comparison Nebius_Inputs D_STREET_FY27_NONMSFT"), ("HURDLE", bL["H1"][2], "benchmark Nebius_Locks H1"),
 ("HURDLE", bL["HURDLE"][2], "benchmark Nebius_Locks HURDLE"), ("HURDLE", pomN["B37"].value, "price_or_mw Nebius_Check!B37"),
 ("BLENDED", ni["D_STREET_FY27_BASE"][2], "comparison Nebius_Inputs D_STREET_FY27_BASE"), ("BLENDED", bL["H2"][2], "benchmark Nebius_Locks H2"), ("BLENDED", pomN["B6"].value, "price_or_mw Nebius_Check!B6"),
 ("CLEAR_BASE", bL["CLEAR_BASE"][2], "benchmark CLEAR_BASE"), ("CLEAR_HIGH", bL["CLEAR_HIGH"][2], "benchmark CLEAR_HIGH"),
 ("TP_BASE", bL["TP_BASE"][2], "benchmark TP_BASE"), ("TP_HIGH", bL["TP_HIGH"][2], "benchmark TP_HIGH"),
 ("EN_Q1", ni["N16"][2], "comparison N16"), ("EN_Q2", ni["N17"][2], "comparison N17"), ("EN_Q1", bL["EN_Q1"][2], "benchmark EN_Q1"), ("EN_Q2", bL["EN_Q2"][2], "benchmark EN_Q2"),
 ("YE26_BASE", ni["N_BILL_YE26_BASE"][2], "comparison"), ("YE26_HIGH", ni["N_BILL_YE26_HIGH"][2], "comparison"),
 ("FY27_BASE", ni["N_BILL_FY27_BASE"][2], "comparison"), ("FY27_HIGH", ni["N_BILL_FY27_HIGH"][2], "comparison"),
 ("YE26_BASE", bL["YE26_BASE"][2], "benchmark"), ("YE26_HIGH", bL["YE26_HIGH"][2], "benchmark"),
 ("FY27_BASE", bL["FY27_BASE"][2], "benchmark"), ("FY27_HIGH", bL["FY27_HIGH"][2], "benchmark"),
]
for key, v, where in pairs: need(close(v, LK[key]["val"], 1e-6), f"{key}: {where} = {v} vs model {LK[key]['val']}")
need(close(pomN["B37"].value, 27.6278, 5e-5), f"price_or_mw B37 = {pomN['B37'].value} != 27.6278")
# 2b named contracts: comparison Contracts_$MWyr (best basis) == benchmark Contracts col K
CONTRACTS = [  # cid (comparison), bench id, company, label, tag, MW basis note
 ("N1", "N1-H", "Nebius", "Nebius-Microsoft (house)", "E (house)", "IT, est. 240"),
 ("N2", "N2", "Nebius", "Nebius-Meta #1 (team MW est.)", "E", "IT, team est. ~50"),
 ("P1", "P1", "IREN", "IREN-Microsoft", "D (R inputs)", "IT 200"),
 ("NEW1", "NEW1", "IREN", "IREN-NVIDIA", "D (R inputs)", "stated ~60"),
 ("P2a", "P2a", "Nscale", "Nscale-Microsoft Narvik", "D", "IT 106 (GPU-derived)"),
 ("P2", "P2", "Nscale", "Nscale-Microsoft (all SOWs)", "P (term)", "IT 596 (GPU-derived)"),
 ("O2", "O2", "Oracle OCI", "Oracle-OpenAI (stated 4.5 GW)", "P", "stated 4,500"),
 ("O2", "O2-IT", "Oracle OCI", "Oracle-OpenAI (IT read)", "P", "IT read 3,750"),
 ("O1", "O1", "Nscale", "Nscale-Anthropic", "P (term)", "IT 460"),
 ("NEW2", "NEW2", "Lambda", "Lambda-Anthropic", "P", "stated ~350"),
]
CV = {}
for cid, bid, co, lab, tag, basis in CONTRACTS:
    v = CON[cid][21] if bid == "O2" else con_best(cid)
    need(close(v, bK[bid][10], 1e-9), f"contract {lab}: comparison {v} vs benchmark {bK[bid][10]}"); CV[bid] = v
# 2c KPI prints (benchmark KPI_Prints) - text cross-check in Benchmark col C
PR = {k: (bP[k][3], bP[k][4], bP[k][5]) for k in ("CW-K1", "IR-K1", "IR-K2", "CR-K1", "NB-K1", "NB-K2", "NB-K3")}
# 2d MW: comparison Chart_Data (active / contracted) and K04, benchmark Benchmark col G text, sites.xlsx / site appendices
act = {c: cd.cell(CDR["Active / energized MW"], j).value for j, c in enumerate(COLS, 2)}
con = {c: cd.cell(CDR["Contracted MW"], j).value for j, c in enumerate(COLS, 2)}
MW = {  # company: (contracted, c_tag, live, l_tag, basis, pack note) - values as in comparison Chart_Data unless noted
 "Nebius":     (con["Nebius"], "R (lower bound)", act["Nebius"], "E", "IT/energized", "contracted >3.5 GW is a lower bound (5 GW YE26 target) [R]"),
 "CoreWeave":  (con["CoreWeave"], "R", act["CoreWeave"], "R", "IT-like active", "~4,200 contracted Aug-11 [R]; Denton ~260 IT live [D]"),
 "IREN":       (con["IREN"], "D", act["IREN"], "R", "IT (AI)", "Childress ~75 IT active Oct [D, site]"),
 "Oracle OCI": (con["Oracle OCI"], "R (lower bound)", act["Oracle OCI"], "E", "cum. delivered (likely gross)", "'>10 GW secured'; Port Washington 0 live [R]"),
 "Nscale":     (con["Nscale"], "R", act["Nscale"], "D (upper bound)", "IT", "Monarch 460 IT contracted / 0 live [R]"),
 "Lambda":     (None, "n/d", act["Lambda"], "P", "unstated", "Beacon Point ~350 [P]; KC ~12 IT live [E]"),
 "Crusoe":     (con["Crusoe"], "R", act["Crusoe"], "R", "GROSS", "Abilene 618 IT live of 1.2 GW gross [D on R]"),
}
need(close(con["Oracle OCI"], 10000) and close(act["Oracle OCI"], ORI["P48"].value), f"comparison Chart_Data Oracle MW {act['Oracle OCI']} != oracle_kpis I45E Aug-26 2,050 (v2)")
need(close(pomI["D16"].value, 2050) and "2,050" in bmw("Oracle (OCI)"), "Oracle ~2,050 cumulative delivered (Aug-26) not in price_or_mw r16 / benchmark")
need(con["Lambda"] is None, "Lambda contracted MW now in comparison")
txt = {"Nebius": None, "CoreWeave": ["3,700", "1,500", "4,200"], "IREN": ["260", "~40"], "Nscale": ["1,370", "81"], "Lambda": ["47", "~350"], "Crusoe": ["6 GW+", "~1 GW", "618"]}
for co, toks in txt.items():
    if toks:
        for t in toks: need(t in bmw(co), f"benchmark MW text for {co} lacks '{t}'")
need("3.5 GW" in bmw("Nebius") or "3,500" in bmw("Nebius"), "benchmark Nebius contracted text")
for co, kid_ok in [("Nebius", True), ("CoreWeave", True), ("Nscale", True), ("Crusoe", True)]:
    k04 = K("K04", co)[0]; c_, l_ = MW[co][0], MW[co][2]
    need(close(k04, c_ / l_, 1e-3), f"K04 {co}: comparison {k04} vs {c_}/{l_}")
need(K("K04", "IREN")[0] == "n/d", "comparison K04 IREN no longer n/d")
need(close(Kmed("K04"), statistics.median([MW[c][0] / MW[c][2] for c in ("CoreWeave", "Nscale", "Crusoe")]), 1e-6), "K04 median")
SITE = {"CoreWeave": sites[("CoreWeave", "Denton TX")][6], "IREN": sites[("IREN", "Childress TX")][6], "Crusoe": sites[("Crusoe", "Abilene TX")][6],
        "Oracle OCI": sites[("Oracle (OCI)", "Port Washington WI")][6], "Lambda": sites[("Lambda", "Kansas City MO")][6], "Nscale": sites[("Nscale", "Monarch Compute Campus, Mason County WV")][6]}
need(SITE == {"CoreWeave": 260, "IREN": 75, "Crusoe": 618, "Oracle OCI": 0, "Lambda": 12, "Nscale": 0}, f"sites.xlsx active base changed: {SITE}")
for k, t in [("coreweave_denton", "| Active | ~260"), ("iren_childress", "| Active | 75"), ("crusoe_abilene", "| Active | 618"), ("oracle_port_washington", "| Active | 0")]:
    need(t in SITE_MD[k], f"site appendix {k} lacks '{t}'")
# 2e revenue per MW and capex (comparison Comparison) + forward (price_or_mw Implied_Price)
FWD = {"Nebius": (pomI["F5"].value, "D", "Street FY27 / 531.8 billed (all customers)"), "CoreWeave": (pomI["F8"].value, "E", "Street FY27 / base avg MW 2,325"),
       "IREN": (pomI["F12"].value, "D (R/R)", "YE26 ARR $4B / 300 IT MW"), "Oracle OCI": (pomI["F16"].value, "D/E", "OCI annualized / ~2,050 cum. MW")}
need(close(FWD["Nebius"][0], LK["BLENDED"]["val"], 1e-6) and close(pomI["F18"].value, LK["HURDLE"]["val"], 1e-6), "price_or_mw Implied_Price F5/F18 vs model")
need(close(FWD["CoreWeave"][0], 26260 / 2325, 1e-6) and close(FWD["IREN"][0], 4000 / 300, 1e-6) and close(FWD["Oracle OCI"][0], pomI["B16"].value / 2050, 1e-6), "price_or_mw forward rows")
need(close(Kmed("K01"), statistics.median([K("K01", c)[0] for c in ("CoreWeave", "IREN", "Nscale", "Lambda", "Crusoe")]), 1e-9), "K01 median")
need(close(Kmed("K14"), statistics.median([K("K14", c)[0] for c in ("CoreWeave", "Nscale")]), 1e-9), "K14 median (Crusoe excluded)")
need(close(k13med, statistics.median([K("K13", c)[0] for c in ("CoreWeave", "Oracle OCI")]), 1e-9), "K13 median")
# benchmark Benchmark $ text agrees with the contract values shown here
for co, toks in {"IREN": ["9.7", "11.3"], "Oracle (OCI)": ["13.3", "16.0"], "Nscale": ["16.2", "13.4", "11.7"], "Lambda": ["16.7"], "CoreWeave": ["$40M"], "Crusoe": ["$23.3M"]}.items():
    for t in toks: need(t in str(bench.cell(BROW[co], 3).value), f"benchmark $ text {co} lacks {t}")
need("$14.49M per IT MW [E, house]" in str(bench.cell(BROW["Nebius"], 3).value), "benchmark Nebius MSFT text")
need(close(bK["N2"][10], META1, 1e-9) and bK["N2"][7] == 50 and bK["N2"][9] == "E", f"benchmark N2 {bK['N2'][10]} != comparison {META1}")
need("$11.52M [E, MW is a team estimate of ~50 IT MW]" in str(bench.cell(BROW["Nebius"], 3).value), "benchmark Nebius $ text lacks the Meta #1 wording")
stop_if_problems("cross-checks")
N_CHECKS = len(pairs) + len(CONTRACTS) + 40

# ---------------------------------------------------------------- 3) table content
TG = lambda t: {"Derived": "D", "Estimated": "E", "Reported": "R"}.get(str(t), str(t)).replace("E (inputs P/R)", "E").replace("D (inputs P)", "D (P inputs)")
def kp(kid, co):
    v, t, a = K(kid, co)
    return (v, TG(t), a) if isinstance(v, (int, float)) else (None, "n/d", None)
PEERS = ["Nebius", "CoreWeave", "IREN", "Oracle OCI", "Nscale", "Lambda", "Crusoe"]
NAMED = {  # text, low, high, tag, footnotes
 "Nebius": (f"{CV['N1-H']:.2f} Microsoft [E, house]; Meta #1 ${CV['N2']:.2f}M [E, MW is a team estimate of ~50 IT MW]", CV["N2"], CV["N1-H"], "E", "1,2,11"),
 "CoreWeave": ("n/d (no MW on any named deal)", None, None, "n/d", "1"),
 "IREN": (f"{CV['P1']:.1f}–{CV['NEW1']:.1f} (Microsoft IT; NVIDIA stated)", CV["P1"], CV["NEW1"], "D", "1"),
 "Oracle OCI": (f"{CV['O2']:.1f} stated / {CV['O2-IT']:.1f} IT read (OpenAI)", CV["O2"], CV["O2-IT"], "P", "1"),
 "Nscale": (f"{CV['P2a']:.1f}–{CV['O1']:.1f} (Narvik D; MSFT {CV['P2']:.1f} P; Anthropic P)", CV["P2a"], CV["O1"], "D / P", "1"),
 "Lambda": (f"{CV['NEW2']:.1f} Anthropic (band top)", CV["NEW2"], CV["NEW2"], "P", "1"),
 "Crusoe": ("n/d (no value with MW)", None, None, "n/d", "1"),
}
PRINTS = {
 "Nebius": (f"{PR['NB-K1'][0]:.0f}–{PR['NB-K1'][1]:.0f} (Q2'26 deals); >{PR['NB-K2'][0]:.0f} mgmt comment¹; base ~{PR['NB-K3'][0]:.0f}", "R", "3"),
 "CoreWeave": (f"~{PR['CW-K1'][0]:.0f} (3–6-mo deals)", "R", "3"), "IREN": (f">{PR['IR-K1'][0]:.0f} (3-yr); ~{PR['IR-K2'][0]:.0f} unsigned", "R", "3"),
 "Oracle OCI": ("n/d", "n/d", ""), "Nscale": ("n/d", "n/d", ""), "Lambda": ("n/d", "n/d", ""),
 "Crusoe": (f"~{PR['CR-K1'][0]:.1f} per gross MW, whole term (not /yr)", "D", "3"),
}
ROWS = []
for co in PEERS:
    k01, k13, k14 = kp("K01", co), kp("K13", co), kp("K14", co)
    f = FWD.get(co)
    c_, ct, l_, lt, basis, note = MW[co]
    ROWS.append(dict(co=co, named=NAMED[co], prt=PRINTS[co], k01=k01, fwd=f, con=(c_, ct), live=(l_, lt), basis=basis, mwnote=note, k13=k13, k14=k14))
# Python-side expected ratio / medians (the xlsx carries live formulas; checked after recalc)
EXP_RATIO = {r["co"]: (r["con"][0] / r["live"][0] if r["con"][0] and r["live"][0] else None) for r in ROWS}

# ---------------------------------------------------------------- 4) workbook
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Peer_KPI_Table"
NAVY, GREY = "1F4E8C", "808080"; HF = PatternFill("solid", fgColor=NAVY); NEB = PatternFill("solid", fgColor="E8EEF7"); MED = PatternFill("solid", fgColor="F2F2F2")
WR = Alignment(wrap_text=True, vertical="top"); thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
ws["A1"] = "Peer KPI appendix - Nebius vs CoreWeave, IREN, Oracle (OCI), Nscale, Lambda, Crusoe: $ per MW-yr, contracted vs live MW, capex per MW"; ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = (f"As of {AS_OF}. Values READ from the locked team files (no new research); tags: R reported, D derived, P press, E estimate, n/d not disclosed. "
            "Superscripts = footnotes below; source keys -> Sources tab. Nebius Q3'26 results ~Nov 10 (est.), not announced.")
HD = ["Company", "Named-contract $M per MW-yr", "Low", "High", "Tag", "Fn", "Company $/MW print ($M)", "Tag", "Rev / active MW K01 ($M/MW-yr)", "Tag", "As of",
      "Fwd rev / billed MW ($M)", "Tag", "Fwd basis", "Contracted MW", "Tag", "Live MW", "Tag", "MW basis", "Contracted / live (x)", "Tag",
      "Capex / MW K13 ($M)", "Tag", "Net PP&E / active MW K14 ($M)", "Tag", "MW notes (key site)", "Source keys"]
HR = 4
for j, h in enumerate(HD, 1):
    c = ws.cell(HR, j, h); c.font = Font(bold=True, color="FFFFFF", size=9); c.fill = HF; c.alignment = WR; c.border = BOX
SK = {"Nebius": "NB-MSFT|NB-META|MKT-PT|MKT-SLIDE|NB-CAPMETA|NB-Q2|NB-CALL|NB-MODEL|NB-MW|CMP|BMK|POM", "CoreWeave": "CW-DECK|CW-Q2|CW-SITE|CMP|BMK|POM|SITES", "IREN": "IR-MSFT|IR-NV|IR-FY26|IR-SITE|CMP|BMK|POM|SITES",
      "Oracle OCI": "OR-OAI|OR-CALL|OR-SITE|CMP|BMK|POM|SITES|ORCL", "Nscale": "NS-S1|NS-TERM|NS-NARV|NS-SITE|CMP|BMK", "Lambda": "LM-ANTH|LM-MW|LM-SITE|CMP|BMK", "Crusoe": "CR-F|CR-DEALS|CR-SITE|CMP|BMK|SITES"}
R0 = HR + 1; RIDX = {}
for i, r in enumerate(ROWS):
    rr = R0 + i; RIDX[r["co"]] = rr
    vals = [r["co"], r["named"][0], r["named"][1], r["named"][2], r["named"][3], r["named"][4], r["prt"][0], r["prt"][1],
            r["k01"][0], r["k01"][1], r["k01"][2], r["fwd"][0] if r["fwd"] else None, r["fwd"][1] if r["fwd"] else "n/d", r["fwd"][2] if r["fwd"] else None,
            r["con"][0], r["con"][1], r["live"][0], r["live"][1], r["basis"],
            f'=IF(AND(ISNUMBER(O{rr}),ISNUMBER(Q{rr})),O{rr}/Q{rr},"n/d")', "D" if EXP_RATIO[r["co"]] else "n/d",
            r["k13"][0], r["k13"][1], r["k14"][0], r["k14"][1], r["mwnote"], SK[r["co"]]]
    for j, v in enumerate(vals, 1):
        c = ws.cell(rr, j, "n/d" if v is None and j in (3, 4, 9, 12, 15, 17, 22, 24) else v); c.alignment = WR; c.border = BOX
        if isinstance(v, float) or (isinstance(v, str) and v.startswith("=")): c.number_format = "0.00" if j in (3, 4, 9, 12) else ("#,##0" if j in (15, 17) else "0.0")
        if j in (15, 17) and isinstance(v, (int, float)): c.number_format = "#,##0"
        if r["co"] == "Nebius": c.fill = NEB
    ws.cell(rr, 1).font = Font(bold=True)
MR = R0 + len(ROWS); RIDX["MED"] = MR
rx = lambda co, col: f"{col}{RIDX[co]}"
med = {"I": f"=MEDIAN({rx('CoreWeave','I')},{rx('IREN','I')},{rx('Nscale','I')},{rx('Lambda','I')},{rx('Crusoe','I')})",
       "T": f"=MEDIAN({rx('CoreWeave','T')},{rx('Nscale','T')},{rx('Crusoe','T')})",
       "V": f"=MEDIAN({rx('CoreWeave','V')},{rx('Oracle OCI','V')})", "X": f"=MEDIAN({rx('CoreWeave','X')},{rx('Nscale','X')})"}
ws.cell(MR, 1, "Peer median (comparison.xlsx sets)").font = Font(bold=True)
for col, f in med.items(): c = ws[f"{col}{MR}"]; c.value = f; c.number_format = "0.00"
for col in "JUWY": ws[f"{col}{MR}"] = "D"
ws[f"Z{MR}"] = "Same sets as comparison.xlsx col Y: K01 ex-Nebius; K04 CRWV/Nscale/Crusoe (IREN ratio computed here, not in K04); K13 CRWV/Oracle; K14 CRWV/Nscale (Crusoe shell+power excluded)."
for j in range(1, len(HD) + 1): ws.cell(MR, j).fill = MED; ws.cell(MR, j).border = BOX; ws.cell(MR, j).alignment = WR
BR = MR + 1
bdrow = bK["BD-A201"]
for j, v in enumerate(["Bitdeer AI (reference; not one of the six)", f"<= {bd:.1f} A201 Johor (ceiling)", None, bd, "D (basis P)", "10", "n/d", "n/d", "n/d", "n/d", None, "n/d", "n/d", None,
                       bdrow[7], "R (floor: >70% of ~21.7)", 0, "R", "unlabeled", "n/m", "n/d", "n/d", "n/d", "n/d", "n/d", "Energized 0 at A201/A102/A103 (Q1'27-Q1'28) [R]", "BD-A201|BD-MKT|BMK"], 1):
    c = ws.cell(BR, j, v); c.alignment = WR; c.border = BOX; c.font = Font(color="C00000", italic=True)
ws.cell(BR, 4).number_format = "0.00"
# ---- Nebius locks block + live checks (model value vs pack-display value)
LR = BR + 2
ws.cell(LR, 1, "NEBIUS LOCKS (model v0.9e; required revenue per billed IT MW unless stated - NOT contract prints)").font = Font(bold=True, color=NAVY)
for j, h in enumerate(["Key", "Lock", "Locked (display)", "Model value", "Tag", "Model cell", "Diff (formula)", "Check (formula)"], 1):
    c = ws.cell(LR + 1, j, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = HF
for i, (key, lab, exp, val, tol, tag, cell) in enumerate(LOCKS):
    r_ = LR + 2 + i
    for j, v in enumerate([key, lab, exp, val, tag, cell, f"=D{r_}-C{r_}", f'=IF(ABS(G{r_})<={max(tol, 0.005)},"OK","DRIFT")'], 1): ws.cell(r_, j, v)
    ws.cell(r_, 4).number_format = "0.0000"; ws.cell(r_, 7).number_format = "0.0000"
r_ = LR + 2 + len(LOCKS)
extra = [("K13_NEB", "Nebius K13 capex / MW vs peer median 33.18 (comparison.xlsx K13)", 77.9, k13n, "E", "comparison Comparison K13", 0.05),
         ("K13_MED", "K13 peer median (CRWV, Oracle)", 33.18, f"=V{MR}", "D", "this sheet V (= comparison Y)", 0.005),
         ("BAND_LO", "Comp band low (IREN-Microsoft)", 9.7, min(band), "D", "comparison Contracts_$MWyr P1", 0.005),
         ("BAND_HI", "Comp band top (Lambda-Anthropic, P)", 16.7, max(band), "P", "comparison Contracts_$MWyr NEW2", 0.05),
         ("BITDEER", "Bitdeer A201 ceiling - only disclosed-MW print >= $20M", 22.4, bd, "D (basis P)", "benchmark Contracts BD-A201", 0.05),
         ("META1", "Nebius-Meta #1 [E, MW is a team estimate of ~50 IT MW] (= Print_Table #3)", 11.52, META1, "E", "comparison Contracts_$MWyr N2 = benchmark N2", 0.005),
         ("BAND8_LO", "GPU Market eight band prints - low", 9.7, min(r[8] for r in PT8), "R / D", "market Print_Table A", 0.005),
         ("BAND8_HI", "GPU Market eight band prints - top", 16.7, max(r[8] for r in PT8), "R (secondary) / D", "market Print_Table A", 0.05),
         ("ORCL_LIVE", "Oracle live MW (cum. delivered since Jun-25, Aug-26)", 2050, act["Oracle OCI"], "E", "comparison Chart_Data (v2) = oracle_kpis Inputs!P48", 0.5)]
for i, (key, lab, exp, val, tag, cell, tol) in enumerate(extra):
    rr = r_ + i
    for j, v in enumerate([key, lab, exp, val, tag, cell, f"=D{rr}-C{rr}", f'=IF(ABS(G{rr})<={tol},"OK","DRIFT")'], 1): ws.cell(rr, j, v)
    ws.cell(rr, 4).number_format = "0.0000"; ws.cell(rr, 7).number_format = "0.0000"
CHK_FIRST, CHK_LAST = LR + 2, r_ + len(extra) - 1
# ---- footnotes
FN = [
 f"1 Named contracts = TCV ÷ term ÷ MW, IT MW where disclosed, else stated MW (comparison.xlsx Contracts_$MWyr cols V/W = peer_benchmark.xlsx Contracts col K, asserted equal). Comp band $9.7–16.7M = IREN-Microsoft 9.7 to Lambda-Anthropic 16.7 (P, Medium); same ends on the memo set (P1/NEW1/P2/O1/NEW2/O2, Nebius excluded) and on GPU Market's eight band prints (incl. Nebius-Meta #1 11.52 and Nebius-Microsoft 14.49).",
 f"2 Nebius-Microsoft $14.49M per IT MW [E, house] = model Contracts!E12 = 17,392.9 / (5 x 240 IT MW est.); neither company states the contract MW (300 MW = Vineland site size). Agrees with the Drivers 14.5 input.",
 "3 Company prints are not single contracts; MW basis undefined; short tenors and GPU-inclusive pricing. Nebius >$40M = management comment on Q3-to-date short-term pricing (Q2'26 letter KPI chart; CEO on the Aug 12 2026 call), not reported Q3. Crusoe ~$23.3M is whole-term TCV per GROSS MW, not per year.",
 "4 K01 = latest-quarter segment revenue x4 ÷ period-end active MW (comparison.xlsx; bases differ: Crusoe FY26E revenue / ~1 GW gross; Lambda Q2'25 press). Fwd = price_or_mw.xlsx Implied_Price (FY27 or run-rate ÷ avg billed MW).",
 f"5 Nebius hurdles (model v0.9e) are REQUIRED revenue per billed IT MW: ex-MSFT ${LK['HURDLE']['val']:.2f}M (Drivers!B9, peer comparator); $23.13M = Street FY27 ÷ our 531.8 billed MW, all customers [D] (Drivers!B10, memo); clear-spot ${LK['CLEAR_BASE']['val']:.2f}M / ${LK['CLEAR_HIGH']['val']:.2f}M (Drivers!B6 / Debate_Grid!A18); turn-positive ${LK['TP_BASE']['val']:.2f}M / ${LK['TP_HIGH']['val']:.2f}M (Debate_Grid!A13 / A10).",
 f"6 MW bases are mixed (IT vs gross). Nebius live = energized 181.5 -> 193 MW (Q1 -> Q2'26, Capacity base case, E); contracted >3.5 GW is a lower bound (R). Billed IT MW (model v0.9e): YE26 {LK['YE26_BASE']['val']:.1f} / {LK['YE26_HIGH']['val']:.1f}; FY27 avg {LK['FY27_BASE']['val']:.2f} / {LK['FY27_HIGH']['val']:.2f}.",
 "7 Oracle live ~2,050 = cumulative AI MW delivered since Jun-25 at Aug-26 (E; FQ1 FY27 call + peers/oracle Inputs!P48 I45E); identical in comparison.xlsx Chart_Data (v2, was 1,200 May-26), peer_benchmark and price_or_mw r16. Contracted '>10 GW secured' (call) is a lower bound -> ratio is a lower bound.",
 "8 IREN contracted ~260 IT (Microsoft 200 + NVIDIA ~60, D) vs ~40 MW AI operating at Jun-30 (R); ratio computed in this pack (comparison K04 has IREN n/d) and kept out of the median. Childress site ~75 IT active in Oct (sites appendix, D).",
 "9 K13 = trailing-4Q cash capex ÷ trailing-4Q new active MW; K14 = net PP&E ÷ active MW (comparison.xlsx). Nebius capex funds not-yet-live MW (Worse vs median). Crusoe K14 = Abilene JV shell+power only, excluded from the median.",
 f"10 Bitdeer A201 <= ${bd:.1f}M is a CEILING (contracted MW is a floor: >70% of ~21.7 MW; 15.7 if 100%); MW basis P. It is the only contract with disclosed MW at >= $20M/MW (peer_benchmark Contracts; comparison has none).",
 f"11 Nebius-Meta #1 ${META1:.2f}M [E, MW is a team estimate of ~50 IT MW] = $2,880.7M / 5 / ~50 (comparison.xlsx N2 = peer_benchmark N2). Sources: GPU Market market/contract_economics.xlsx Print_Table #3 and final band slide (contract_band_slide_final.png); Nebius filings and Capacity (Capacity tranches G3:H3 = n/d) state no MW; model Contracts!D23:D24 = {META_MW_MODEL:.1f} IT MW. One of the eight band prints; band ends unchanged at $9.7M / $16.7M.",
]
FR = CHK_LAST + 2
ws.cell(FR, 1, "Footnotes").font = Font(bold=True, color=NAVY)
for i, t in enumerate(FN): ws.cell(FR + 1 + i, 1, t).alignment = Alignment(wrap_text=False)
widths = [16, 34, 7, 7, 9, 5, 30, 6, 10, 9, 11, 9, 7, 24, 10, 13, 9, 11, 14, 9, 6, 9, 9, 10, 7, 40, 34]
for j, w in enumerate(widths, 1): ws.column_dimensions[CL(j)].width = w
ws.row_dimensions[HR].height = 42; ws.freeze_panes = "B5"
# ---- Sources tab
wsS = wb.create_sheet("Sources")
for j, h in enumerate(["Key", "Document", "URL / path", "Accessed / note"], 1): c = wsS.cell(1, j, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = HF
LOCAL = [("CMP", "peers/comparison.xlsx (Comparison K01/K04/K13/K14 + medians; Contracts_$MWyr; Chart_Data MW; Nebius_Inputs)", f"local:{P}/comparison.xlsx"),
         ("BMK", "peers/benchmark/peer_benchmark.xlsx (Benchmark, Contracts, KPI_Prints, Nebius_Locks, Comp_Band, Sources)", f"local:{P}/benchmark/peer_benchmark.xlsx"),
         ("POM", "peers/price_or_mw.xlsx (Implied_Price forward rev per billed MW; Nebius_Check B22/B37/B6)", f"local:{P}/price_or_mw.xlsx"),
         ("SITES", "peers/sites/sites.xlsx + appendix_{coreweave_denton,iren_childress,crusoe_abilene,oracle_port_washington}.md (active base MW)", f"local:{P}/sites/"),
         ("MODEL", "model/nebius_model.xlsx v0.9e (read-only) via peers/nebius_locks.py: Contracts!E12, Drivers!B6/B9/B10/B34/C34, Judge_Slide!B34/B35, Debate_Grid!A10/A13/A18", f"local:{W}/model/nebius_model.xlsx"),
         ("CAP", "capacity/nebius_capacity.xlsx Quarterly MW!M130/N130 (energized 181.5 / 193, base case)", f"local:{W}/capacity/nebius_capacity.xlsx"),
         ("ORCL", "peers/oracle/oracle_kpis.xlsx Inputs!P48 I45E (cumulative AI MW delivered since 1-Jun-25 = 2,050 at FQ1 FY27 Aug-26, E; was 1,200 at O48 May-26)", f"local:{P}/oracle/oracle_kpis.xlsx")]
rS = 2
for k, d_, u in LOCAL: [wsS.cell(rS, j, v) for j, v in enumerate([k, d_, u, "read 2026-10-09 ET (frozen)"], 1)]; rS += 1
used = set("|".join(SK.values()).split("|")) | {"BD-A201", "BD-MKT", "NB-LOCK"}
for k, r in bS.items():
    if k in used: [wsS.cell(rS, j, v) for j, v in enumerate(r[:4], 1)]; rS += 1
for col, w in zip("ABCD", [10, 110, 80, 40]): wsS.column_dimensions[col].width = w
missing = sorted(k for k in used if k not in bS and k not in {x[0] for x in LOCAL}); need(not missing, f"source keys without a Sources row: {missing}")
stop_if_problems("workbook")
wb.save(XLSX)
# recalc (LibreOffice) -> must be 0 errors; then verify formula results vs Python
rc = subprocess.run(["python3", f"{W}/model/recalc.py", XLSX], capture_output=True, text=True, timeout=300)
print(rc.stdout.strip().splitlines()[-1] if rc.stdout.strip() else rc.stderr[-300:])
need(rc.returncode == 0 and os.path.exists("/tmp/rc/peer_kpi_appendix.xlsx"), f"recalc failed: {rc.stderr[-300:]}"); stop_if_problems("recalc")
shutil.copy("/tmp/rc/peer_kpi_appendix.xlsx", XLSX)
wv = openpyxl.load_workbook(XLSX, data_only=True); wf = openpyxl.load_workbook(XLSX)
errs = [(s.title, c.coordinate) for s in wv for row in s.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("#") and c.value[:2] in ("#D", "#N", "#V", "#R", "#!")]
nform = sum(1 for s in wf for row in s.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
need(not errs, f"recalc errors: {errs}")
t = wv["Peer_KPI_Table"]
for co, rr in RIDX.items():
    if co == "MED": continue
    v = t[f"T{rr}"].value; e = EXP_RATIO[co]
    need((v == "n/d" and e is None) or close(v, e, 1e-9), f"ratio {co}: {v} vs {e}")
need(close(t[f"I{MR}"].value, Kmed("K01"), 1e-9) and close(t[f"T{MR}"].value, Kmed("K04"), 1e-9) and close(t[f"V{MR}"].value, Kmed("K13"), 1e-9) and close(t[f"X{MR}"].value, Kmed("K14"), 1e-9), "xlsx medians != comparison medians")
chk = [t.cell(r, 8).value for r in range(CHK_FIRST, CHK_LAST + 1)]
need(all(c == "OK" for c in chk), f"lock check cells: {chk}")
stop_if_problems("post-recalc")
print(f"xlsx ok: {nform} formulas, 0 errors, {len(chk)} lock checks OK")

# ---------------------------------------------------------------- 5) charts
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from matplotlib.lines import Line2D
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8})
CLR = {"Nebius": "#1f4e9c", "CoreWeave": "#7f7f7f", "IREN": "#2ca02c", "Oracle OCI": "#d62728", "Nscale": "#9467bd", "Lambda": "#8c564b", "Crusoe": "#ff7f0e", "Bitdeer": "#c00000"}
def clip_check(fig, name):
    fig.canvas.draw(); rd = fig.canvas.get_renderer(); fb = fig.bbox; bad = []
    for ax in fig.axes:
        ab = ax.get_window_extent(rd)
        for tx in ax.texts:
            if not tx.get_text().strip(): continue
            bb = tx.get_window_extent(rd)
            if bb.x0 < ab.x0 - 1 or bb.x1 > ab.x1 + 1 or bb.y0 < ab.y0 - 1 or bb.y1 > ab.y1 + 1: bad.append(tx.get_text()[:30])
    need(not bad, f"{name}: data labels outside axes: {bad}")
def esc(s): return s.replace("$", r"\$")
LBOX = dict(facecolor="white", edgecolor="none", pad=0.6, alpha=0.85)
# Chart 1: $ per MW-yr
bars = [("IREN-Microsoft (IT)", CV["P1"], "IREN", "D", ""), ("IREN-NVIDIA (stated MW)", CV["NEW1"], "IREN", "D", ""),
        ("Nebius-Meta #1 (team est. ~50 IT MW)", CV["N2"], "Nebius", "E", ""),
        ("Nscale-Microsoft Narvik (IT)", CV["P2a"], "Nscale", "D", ""), ("Oracle-OpenAI (stated 4.5 GW)", CV["O2"], "Oracle OCI", "P", ""),
        ("Nscale-Microsoft all SOWs (IT)", CV["P2"], "Nscale", "P", ""), ("Nebius-Microsoft (house, IT est.)", CV["N1-H"], "Nebius", "E", ""),
        ("Oracle-OpenAI (IT read)", CV["O2-IT"], "Oracle OCI", "P", ""), ("Nscale-Anthropic (IT)", CV["O1"], "Nscale", "P", ""),
        ("Lambda-Anthropic (stated MW)", CV["NEW2"], "Lambda", "P", ""), ("Bitdeer A201 (ceiling; not one of the six)", bd, "Bitdeer", "D", "≤")]
prints = [("IREN 3-yr deals (company print, floor)", PR["IR-K1"][0], "IREN", "R", ">"), ("Nebius Q2'26 deals (company print)", PR["NB-K1"][1], "Nebius", "R", "20–"),
          ("IREN unsigned discussions (print)", PR["IR-K2"][0], "IREN", "R", "~"), ("CoreWeave 3–6-mo deals (print)", PR["CW-K1"][0], "CoreWeave", "R", "~"),
          ("Nebius Q3-to-date short-term (mgmt comment)", PR["NB-K2"][0], "Nebius", "R", ">")]
items = bars + [None] + prints
fig, ax = plt.subplots(figsize=(10.2, 4.6)); y = list(range(len(items)))[::-1]; ylab = []
ax.axvspan(min(band), max(band), color="#fff2cc", zorder=0)
for yy, it in zip(y, items):
    if it is None: ylab.append(""); continue
    lab, v, co, tg, pre = it; isprint = it in prints
    ax.barh(yy, v, 0.62, color=CLR[co], alpha=0.45 if isprint else 0.95, hatch="////" if isprint else None, edgecolor="white" if not isprint else CLR[co], zorder=2)
    vt = f"{v:.2f}" if co == "Nebius" and not isprint else (f"{v:.0f}" if isprint else f"{v:.1f}")
    ax.text(v + 0.4, yy, f"{pre}{vt} [{tg}]", va="center", fontsize=7.2, zorder=6, bbox=LBOX)
    ylab.append(lab)
ax.set_yticks(y); ax.set_yticklabels(ylab, fontsize=7.4)
HL = [(LK["HURDLE"]["val"], "#5c0000", "-", 1.8, f"Ex-MSFT hurdle {LK['HURDLE']['val']:.2f} (Drivers!B9)"),
      (LK["CLEAR_BASE"]["val"], "#1f4e9c", "--", 1.3, f"Clear-spot base {LK['CLEAR_BASE']['val']:.2f} / high {LK['CLEAR_HIGH']['val']:.2f}"),
      (LK["CLEAR_HIGH"]["val"], "#1f4e9c", "--", 1.3, None),
      (LK["BLENDED"]["val"], "#7f7f7f", ":", 1.5, f"Memo: {LK['BLENDED']['val']:.2f} = Street FY27 ÷ our 531.8 billed MW, all customers [D]"),
      (LK["TP_BASE"]["val"], "#2ca02c", "-.", 1.3, f"Turn-positive base {LK['TP_BASE']['val']:.2f} / high {LK['TP_HIGH']['val']:.2f}"),
      (LK["TP_HIGH"]["val"], "#2ca02c", "-.", 1.3, None)]
hs = [Patch(color="#fff2cc", label=f"Comp band {min(band):.1f}–{max(band):.1f} (same ends with GPU Market's 8 prints)")]
for x, c, ls, lw, lab in HL:
    ax.axvline(x, color=c, ls=ls, lw=lw, zorder=3)
    if lab: hs.append(Line2D([0], [0], color=c, ls=ls, lw=lw, label=lab))
hs.append(Patch(facecolor="white", edgecolor="#555555", hatch="////", label="Company print (not a contract; MW basis undefined)"))
ax.set_xlim(0, 52); ax.set_xlabel("$M per MW-yr (contracts: TCV ÷ term ÷ MW; Nebius lines = required revenue per billed IT MW)")
ax.axhline(y[len(bars)], color="#bbbbbb", lw=0.6)
ax.set_title("$ per MW-yr: disclosed contracts and company prints vs the Nebius hurdles (model v0.9e)", fontweight="bold", fontsize=9.5, loc="left")
ax.legend(handles=hs, loc="upper center", bbox_to_anchor=(0.5, 0.115), bbox_transform=fig.transFigure, ncol=3, fontsize=6.6, frameon=False)
for s in ("top", "right"): ax.spines[s].set_visible(False)
fig.text(0.005, 0.0, textwrap.fill(esc("Sources: comparison.xlsx Contracts_$MWyr = peer_benchmark.xlsx Contracts/KPI_Prints; model v0.9e. Nebius-Microsoft $14.49M per IT MW [E, house] = model Contracts!E12; Nebius-Meta #1 $11.52M [E, MW is a team estimate of ~50 IT MW]. "
         "Bitdeer ≤22.4 is the only disclosed-MW print ≥ $20M. Nebius >40 = mgmt comment on Q3-to-date short-term pricing, not reported Q3."), 200), fontsize=6, color="dimgray")
fig.tight_layout(rect=(0, 0.13, 1, 1)); clip_check(fig, "chart1"); fig.savefig(PNG1, dpi=200, bbox_inches="tight"); plt.close(fig)
# Chart 2: contracted vs live MW (log)
fig, ax = plt.subplots(figsize=(5.4, 3.6)); xs = range(len(PEERS)); w = 0.38
for i, r in enumerate(ROWS):
    c_, l_ = r["con"][0], r["live"][0]; col = CLR[r["co"]]
    if c_: ax.bar(i - w / 2, c_, w, color=col, alpha=0.95, zorder=2); ax.text(i - w / 2, c_ * 1.12, (">" if "lower" in r["con"][1] else "") + f"{c_:,.0f}", ha="center", fontsize=6.2, rotation=90, va="bottom", bbox=LBOX, zorder=6)
    else: ax.text(i - w / 2, 12, "n/d", ha="center", fontsize=6.2, rotation=90, va="bottom")
    if l_: ax.bar(i + w / 2, l_, w, color=col, alpha=0.4, hatch="....", edgecolor=col, zorder=2); ax.text(i + w / 2, l_ * 1.12, ("≤" if "upper" in r["live"][1] else "") + f"{l_:,.0f}", ha="center", fontsize=6.2, rotation=90, va="bottom", bbox=LBOX, zorder=6)
    e = EXP_RATIO[r["co"]]
    lab = "n/d" if e is None else ((">" if "lower" in r["con"][1] else "") + f"{e:.1f}x" + ("*" if r["co"] == "IREN" else ""))
    ax.text(i, 1.6e5, lab, ha="center", fontsize=6.8, fontweight="bold", color=col)
ax.scatter([0.5] * 4, [LK["YE26_BASE"]["val"], LK["YE26_HIGH"]["val"], LK["FY27_BASE"]["val"], LK["FY27_HIGH"]["val"]], marker="_", s=60, color="#000000", zorder=5, linewidths=1.4)
ax.set_yscale("log"); ax.set_ylim(10, 4e5); ax.set_xticks(list(xs)); ax.set_xticklabels([c.replace(" OCI", "") for c in PEERS], fontsize=7)
ax.set_ylabel("MW (log; bases differ: IT vs gross)", fontsize=7)
ax.set_title("Contracted vs live MW (latest) and contracted / live (x)", fontweight="bold", fontsize=8.5, loc="left")
ax.legend(handles=[Patch(color="#555555", label="Contracted"), Patch(facecolor="#dddddd", edgecolor="#555555", hatch="....", label="Live (energized / active)"),
                   Line2D([0], [0], marker="_", color="black", ls="", markersize=9, label="Nebius billed IT MW: YE26 257.7/284.2, FY27 531.8/659.2")],
          loc="upper center", bbox_to_anchor=(0.5, -0.1), ncol=2, fontsize=6, frameon=False)
for s in ("top", "right"): ax.spines[s].set_visible(False)
fig.text(0.005, -0.04, esc("Nebius 3,500 is a lower bound; live = energized 181.5 -> 193 (E). Oracle >10,000 'secured' (lower bound) / ~2,050 cum. delivered (E). Crusoe GROSS. *IREN ratio computed here (260 D / 40 R)."), fontsize=5.4, color="dimgray", wrap=True)
fig.tight_layout(); clip_check(fig, "chart2"); fig.savefig(PNG2, dpi=200, bbox_inches="tight"); plt.close(fig)
# Chart 3: capex / MW vs peer median
fig, ax = plt.subplots(figsize=(5.4, 3.6))
for i, r in enumerate(ROWS):
    col = CLR[r["co"]]
    for off, (v, tg), hatch, alpha in [(-w / 2, r["k13"][:2], None, 0.95), (w / 2, r["k14"][:2], "////", 0.45)]:
        if v is not None:
            ax.bar(i + off, v, w, color=col, alpha=alpha, hatch=hatch, edgecolor=col if hatch else "white", zorder=2)
            ax.text(i + off, v + 1.5, f"{v:.1f}", ha="center", fontsize=6.2, rotation=90, va="bottom", bbox=LBOX, zorder=6)
        else: ax.text(i + off, 2, "n/d", ha="center", fontsize=5.8, rotation=90, va="bottom", color="gray")
ax.axhline(Kmed("K13"), color="black", ls="--", lw=1.1, zorder=3); ax.axhline(Kmed("K14"), color="black", ls=":", lw=1.1, zorder=3)
ax.set_ylim(0, 100); ax.set_xticks(list(xs)); ax.set_xticklabels([c.replace(" OCI", "") for c in PEERS], fontsize=7); ax.set_ylabel("$M per MW", fontsize=7)
ax.set_title(f"Capex per MW: Nebius K13 {k13n:.1f} vs peer median {Kmed('K13'):.1f}", fontweight="bold", fontsize=8.5, loc="left")
ax.legend(handles=[Patch(color="#555555", label="K13 T4Q cash capex / new MW"), Patch(facecolor="#dddddd", edgecolor="#555555", hatch="////", label="K14 net PP&E / active MW"),
                   Line2D([0], [0], color="black", ls="--", label=f"K13 peer median {Kmed('K13'):.1f} (CRWV, Oracle)"),
                   Line2D([0], [0], color="black", ls=":", label=f"K14 peer median {Kmed('K14'):.1f} (CRWV, Nscale)")],
          loc="upper center", bbox_to_anchor=(0.5, -0.1), ncol=2, fontsize=6, frameon=False)
for s in ("top", "right"): ax.spines[s].set_visible(False)
fig.text(0.005, -0.04, esc("comparison.xlsx K13/K14 (tags: Nebius E; CRWV D; Oracle E; Nscale D; Crusoe D). Nebius capex funds not-yet-live MW. Crusoe 12.5 = Abilene shell+power only (excluded)."), fontsize=5.4, color="dimgray", wrap=True)
fig.tight_layout(); clip_check(fig, "chart3"); fig.savefig(PNG3, dpi=200, bbox_inches="tight"); plt.close(fig)
stop_if_problems("charts")

# ---------------------------------------------------------------- 6) docx -> pdf (1 table page + 1 chart page)
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
doc = Document(); sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = Inches(11), Inches(8.5)
sec.left_margin = sec.right_margin = Inches(0.35); sec.top_margin = Inches(0.3); sec.bottom_margin = Inches(0.3)
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(6.5); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_before = Pt(0); st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.0
NV = RGBColor(0x1F, 0x4E, 0x8C)
def para(text, size=6.5, bold=False, color=None, after=0, italic=False):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after); r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    return p
def shade(cell, hx):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hx); tcPr.append(s)
def margins(tbl, m=25):
    tblPr = tbl._tbl.tblPr; mar = OxmlElement("w:tblCellMar")
    for side in ("top", "bottom", "left", "right"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(m if side in ("left", "right") else 10)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)
para("Peer KPI appendix: Nebius vs CoreWeave, IREN, Oracle (OCI), Nscale, Lambda and Crusoe", 11, True, NV)
para(f"$ per MW-yr, contracted vs live MW, capex per MW. As of {AS_OF}. Read from the locked team files (comparison.xlsx, peer_benchmark.xlsx, price_or_mw.xlsx, site appendices, model v0.9e); "
     "every value tagged [R reported | D derived | P press | E estimate | n/d]; superscripts = footnotes. Not a valuation; Nebius Q3'26 results ~Nov 10 (est.).", 6.5, False, None, 3)
def f2(v, d=1): return "n/d" if v is None else f"{v:,.{d}f}"
def tg(t): return "" if t in ("n/d", None) else f" [{t}]"
HDD = ["Company", "Named-contract $M per MW-yr¹", "Company $/MW print ($M)³", "Rev / active MW (K01)⁴", "Fwd rev / billed MW⁴", "Contracted MW⁶", "Live MW⁶", "Contracted / live", "Capex / MW (K13)⁹", "Net PP&E / MW (K14)⁹", "MW notes (key site)"]
WD = [0.85, 1.75, 1.35, 0.72, 0.95, 0.78, 0.72, 0.62, 0.62, 0.66, 1.28]
tbl = doc.add_table(rows=1, cols=len(HDD)); tbl.style = "Table Grid"; tbl.autofit = False; margins(tbl)
for j, h in enumerate(HDD):
    cl = tbl.rows[0].cells[j]; cl.width = Inches(WD[j]); shade(cl, "1F4E8C"); r = cl.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(6.4); r.font.color.rgb = RGBColor(255, 255, 255)
fnsup = {"Nebius": {"named": "²,¹¹", "live": "", "con": ""}, "Oracle OCI": {"live": "⁷", "con": "⁷"}, "IREN": {"con": "⁸"}, "Crusoe": {"prt": "³"}}
for r in ROWS:
    co = r["co"]; s = fnsup.get(co, {})
    ratio = EXP_RATIO[co]; rtxt = "n/d" if ratio is None else ((">" if "lower" in r["con"][1] else "") + f"{ratio:.1f}x [D]" + ("⁸" if co == "IREN" else ""))
    fwdtxt = "n/d" if not r["fwd"] else f"{r['fwd'][0]:.2f} [{r['fwd'][1]}] ({r['fwd'][2]})"
    if co == "Nebius": fwdtxt = f"{LK['BLENDED']['val']:.2f} all customers [D]; ex-MSFT hurdle {LK['HURDLE']['val']:.2f} [D]⁵"
    cont = "n/d" if r["con"][0] is None else ((">" if "lower" in r["con"][1] else "") + f"{r['con'][0]:,.0f}{tg(r['con'][1].split(' ')[0])}")
    live = "n/d" if r["live"][0] is None else (("≤" if "upper" in r["live"][1] else ("~" if co in ("Oracle OCI", "IREN", "Crusoe") else "")) + f"{r['live'][0]:,.0f}{tg(r['live'][1].split(' ')[0])}")
    if co == "Nebius": live = f"181.5→193 [E]; billed YE26 {LK['YE26_BASE']['val']:.1f}/{LK['YE26_HIGH']['val']:.1f}, FY27 {LK['FY27_BASE']['val']:.2f}/{LK['FY27_HIGH']['val']:.2f} [D]"
    vals = [co, r["named"][0] + ("" if co == "Nebius" else tg(r["named"][3])) + s.get("named", ""), r["prt"][0] + tg(r["prt"][1]), f2(r["k01"][0], 2) + tg(r["k01"][1]) + (f" ({r['k01'][2]})" if r["k01"][2] else ""),
            fwdtxt, cont + s.get("con", ""), live + s.get("live", ""), rtxt, f2(r["k13"][0]) + tg(r["k13"][1]), f2(r["k14"][0]) + tg(r["k14"][1]), f"{r['basis']}; {r['mwnote']}"]
    cells = tbl.add_row().cells
    for j, v in enumerate(vals):
        cells[j].width = Inches(WD[j]); p = cells[j].paragraphs[0]; rr = p.add_run(v.replace("Q2'26 deals); >40 mgmt comment¹", "Q2'26 deals); >40 mgmt comment")); rr.font.size = Pt(6.3); rr.bold = (j == 0)
        if j == 0: rr.font.color.rgb = NV
        if co == "Nebius": shade(cells[j], "E8EEF7")
medv = ["Peer median", "band 9.7–16.7 (Nebius excl.)¹", "", f"{Kmed('K01'):.2f} [D]", "", "", "", f"{Kmed('K04'):.1f}x [D]", f"{Kmed('K13'):.1f} [D]", f"{Kmed('K14'):.1f} [D]", "comparison.xlsx sets; IREN ratio and Crusoe K14 excluded"]
cells = tbl.add_row().cells
for j, v in enumerate(medv):
    cells[j].width = Inches(WD[j]); rr = cells[j].paragraphs[0].add_run(v); rr.font.size = Pt(6.3); rr.bold = True; shade(cells[j], "F2F2F2")
bdv = ["Bitdeer AI (ref.)", f"≤{bd:.1f} A201 Johor, ceiling [D, basis P]¹⁰", "n/d", "n/d", "n/d", f"≥{bdrow[7]:.1f} of ~21.7 [R]", "0 [R]", "n/m", "n/d", "n/d", "not one of the six; energized Q1'27–Q1'28"]
cells = tbl.add_row().cells
for j, v in enumerate(bdv):
    cells[j].width = Inches(WD[j]); rr = cells[j].paragraphs[0].add_run(v); rr.font.size = Pt(6.3); rr.italic = True; rr.font.color.rgb = RGBColor(0xC0, 0, 0)
para("", 3)
p = para(f"NEBIUS LOCKS (model v0.9e, asserted by the build): Microsoft ${LK['MSFT']['val']:.2f}M per IT MW [E, house] (Contracts!E12) only · ex-MSFT hurdle ${LK['HURDLE']['val']:.2f}M [D] · "
         f"$23.13M = Street FY27 ÷ our 531.8 billed MW, all customers [D] · clear-spot ${LK['CLEAR_BASE']['val']:.2f}M / ${LK['CLEAR_HIGH']['val']:.2f}M · turn-positive ${LK['TP_BASE']['val']:.2f}M / ${LK['TP_HIGH']['val']:.2f}M · "
         f"energized 181.5→193 MW [E] · billed YE26 {LK['YE26_BASE']['val']:.1f}/{LK['YE26_HIGH']['val']:.1f}, FY27 {LK['FY27_BASE']['val']:.2f}/{LK['FY27_HIGH']['val']:.2f} [D] · K13 capex/MW {k13n:.1f} vs peer median {Kmed('K13'):.2f} · "
         f"comp band ${min(band):.1f}–{max(band):.1f}M · Bitdeer ≤${bd:.1f}M is the only disclosed-MW print ≥ $20M.", 6.6, True, NV, 2)
para("HEADLINES: " + " | ".join([
    f"Nebius: ex-MSFT hurdle ${LK['HURDLE']['val']:.2f}M is ~1.7x the band top; live 193 MW vs >3.5 GW contracted (18.1x); capex/MW {k13n:.1f} vs median {Kmed('K13'):.1f}.",
    "CoreWeave: no MW on named deals; ~$40M print is 3–6-mo pricing; 2.5x contracted/live; lowest capex/MW (20.0).",
    f"IREN: named {CV['P1']:.1f}–{CV['NEW1']:.1f} (band floor); >$20M only on 3-yr deals; 260 vs ~40 MW.",
    f"Oracle: OpenAI {CV['O2']:.1f}–{CV['O2-IT']:.1f} [P]; >10 GW secured vs ~2,050 delivered; capex/MW 46.4.",
    f"Nscale: {CV['P2a']:.1f}–{CV['O1']:.1f}; ~17x contracted/live (≤81 live).",
    f"Lambda: Anthropic {CV['NEW2']:.1f} sets the band top [P]; company MW n/d.",
    "Crusoe: no per-deal $/MW; ~$23.3M is whole-term per gross MW; 6x."]), 6.3, False, None, 2)
for t_ in FN:
    n, rest = t_.split(" ", 1); p = doc.add_paragraph(); r1 = p.add_run(n + " "); r1.bold = True; r1.font.size = Pt(5.6); r2 = p.add_run(rest); r2.font.size = Pt(5.6)
para("Sources: comparison.xlsx; benchmark/peer_benchmark.xlsx (+ its Sources keys: SEC 8-K/6-K/S-1, company PRs, WSJ/Reuters/Bloomberg); price_or_mw.xlsx; sites/sites.xlsx + site appendices; model/nebius_model.xlsx v0.9e; capacity/nebius_capacity.xlsx. Full list: peer_kpi_appendix.xlsx Sources tab.", 5.6, False, RGBColor(0x60, 0x60, 0x60))
pb = doc.add_paragraph(); pb.add_run().add_break(WD_BREAK.PAGE)
para("Charts: $ per MW-yr vs Nebius hurdles; contracted vs live MW; capex per MW vs peer median", 9, True, NV, 2)
doc.add_picture(PNG1, width=Inches(8.6)); doc.paragraphs[-1].alignment = 1
t2 = doc.add_table(rows=1, cols=2); t2.autofit = False
for j, pth in enumerate((PNG2, PNG3)):
    cl = t2.rows[0].cells[j]; cl.width = Inches(5.1); cl.paragraphs[0].add_run().add_picture(pth, width=Inches(4.5)); cl.paragraphs[0].alignment = 1
para("All values from the table on page 1 (same files, same tags). Charts are context for the appendix table, not new estimates.", 5.6, False, RGBColor(0x60, 0x60, 0x60))
doc.save(DOCX)
out = "/tmp/pka_pdf"; shutil.rmtree(out, ignore_errors=True); os.makedirs(out)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", out, DOCX], check=True, capture_output=True, timeout=240)
shutil.copy(f"{out}/peer_kpi_appendix.pdf", PDF)
pages = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", PDF], capture_output=True, text=True).stdout).group(1))
need(pages <= 2, f"PDF has {pages} pages (max 2)"); stop_if_problems("pdf")
txt = subprocess.run(["pdftotext", "-layout", PDF, "-"], capture_output=True, text=True).stdout
bad = [t for t in ("11.60", "17.08", "17.1 ", "12.9", "12.93") if t in txt]; need(not bad, f"retired values in PDF: {bad}"); stop_if_problems("pdf text")
print(f"PDF {pages} pages; charts {PNG1}, {PNG2}, {PNG3}; checks passed")
