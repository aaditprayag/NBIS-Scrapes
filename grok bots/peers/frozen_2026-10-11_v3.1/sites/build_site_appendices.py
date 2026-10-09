#!/usr/bin/env python3
# Builds one-page landscape site appendices (md + docx + pdf + MW chart PNG) for the peer key sites, in the style and with the
# definitions of capacity/vineland_site_appendix (capacity/_build/build_site_appendix.py). Content: peers/sites/site_data.py.
# Each docx is converted with LibreOffice and must be exactly 1 page; the font steps down until it fits. Combined PDF via pdfunite.
# Nebius Peer KPIs, Thu Oct 8 2026 ET. Run: python3 /workspace/nebius/peers/sites/build_site_appendices.py
import sys, os, subprocess, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from site_data import SITES, DEFS
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt

D = "/workspace/nebius/peers/sites"; CH = f"{D}/charts"; os.makedirs(CH, exist_ok=True)

def chart(key, c):
    lab = ["Announced\n(IT)", "Grid /\nutility", "Connected\n(halls, IT)", "Active\n(IT)", "Billed\n(IT)"]
    lo, base, hi = c["act"]
    vals = [c["ann"], c["grid"], c["conn"], base, c["bill"]]
    col = ["#bfbfbf", "#ffffff", "#9dc3e6", "#2e75b6", "#1f3864"]
    fig, ax = plt.subplots(figsize=(4.6, 1.75), dpi=200)
    bars = ax.bar(range(5), vals, 0.62, color=col, edgecolor="#404040", linewidth=0.6)
    bars[1].set_hatch("////")
    ax.errorbar([3], [base], yerr=[[base - lo], [hi - base]], fmt="none", ecolor="black", capsize=3, lw=0.8)
    top = max(vals + [hi]) or 100
    for i, v in enumerate(vals):
        t = c["grid_lbl"] if i == 1 else (f"{v:,.0f}" + (f" ({lo}–{hi})" if i == 3 and hi != lo else ""))
        ax.text(i, v + top * 0.03, f"{t} [{c['tags'][i]}]", ha="center", va="bottom", fontsize=5.4)
    ax.set_xticks(range(5)); ax.set_xticklabels(lab, fontsize=5.6)
    ax.set_ylim(0, top * 1.42); ax.tick_params(axis="y", labelsize=5.4); ax.set_ylabel("MW", fontsize=5.6); ax.margins(x=0.06)
    ax.spines[["top", "right"]].set_visible(False)
    ax.set_title("MW by stage (hatched = grid/utility, not IT; whisker = active low–high)", fontsize=5.8, loc="left")
    fig.tight_layout(pad=0.3)
    out = f"{CH}/{key}_mw.png"; fig.savefig(out, dpi=200); plt.close(fig); return out

def build(key, S, BODY, SRC, with_chart=True):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(BODY)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    pf = st.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width = Inches(11); sec.page_height = Inches(8.5)
    sec.left_margin = sec.right_margin = Inches(0.35); sec.top_margin = Inches(0.28); sec.bottom_margin = Inches(0.25)
    sec.header_distance = Inches(0.15); sec.footer_distance = Inches(0.15)

    def para(runs, size=None, sa=1, sb=0, box=None, color=None):
        box = box if box is not None else doc
        p = box.add_paragraph(); p.paragraph_format.space_after = Pt(sa); p.paragraph_format.space_before = Pt(sb)
        for txt, bold in runs:
            r = p.add_run(txt); r.bold = bold; r.font.size = Pt(size or BODY)
            if color: r.font.color.rgb = color
        return p
    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr(); s_ = OxmlElement("w:shd")
        s_.set(qn("w:val"), "clear"); s_.set(qn("w:color"), "auto"); s_.set(qn("w:fill"), hexcolor); tcPr.append(s_)
    def fix_table(t, widths, l=35, r=35):
        tblPr = t._tbl.tblPr; m = OxmlElement("w:tblCellMar")
        for k, v in (("left", l), ("right", r), ("top", 0), ("bottom", 0)):
            e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
        tblPr.append(m)
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        tw = tblPr.find(qn("w:tblW"))
        if tw is None: tw = OxmlElement("w:tblW"); tblPr.append(tw)
        tw.set(qn("w:w"), str(int(sum(widths) * 1440))); tw.set(qn("w:type"), "dxa")
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(int(w * 1440)))
        for row in t.rows:
            for c, w in zip(row.cells, widths): c.width = Inches(w)
    def table(box, hdr, rows, widths, size=None, bold_first=True):
        t = box.add_table(rows=1 + len(rows), cols=len(hdr)); t.style = doc.styles["Table Grid"]; t.autofit = False
        fix_table(t, widths)
        for i, row in enumerate([hdr] + rows):
            for j, val in enumerate(row):
                c = t.cell(i, j); p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
                r = p.add_run(val); r.font.size = Pt(size or BODY)
                if i == 0: r.bold = True; shade(c, "D9E2F3")
                elif bold_first and j == 0: r.bold = True
        return t
    def heading(box, txt, first=False):
        p = box.paragraphs[0] if first else box.add_paragraph()
        p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(0 if first else 2.5)
        r = p.add_run(txt); r.bold = True; r.font.size = Pt(BODY + 0.5)

    para([(S["title"], True)], size=BODY + 2.5, sa=0)
    para([(S["sub"], False)], size=BODY - 0.5, sa=1.5)
    para([(DEFS, False)], size=BODY - 0.5, sa=2)
    CW = 5.1; GUT = 0.1
    outer = doc.add_table(rows=1, cols=2); outer.autofit = False
    fix_table(outer, [CW + GUT, CW + GUT], l=0, r=int(GUT * 1440))
    L, R = outer.cell(0, 0), outer.cell(0, 1)
    sz = BODY - 0.5
    heading(L, "1. Site facts", first=True)
    table(L, ["Item", "Fact", "Tag", "Src"], S["facts"], [1.0, 3.35, 0.4, 0.35], size=sz)
    heading(L, "2. MW: announced / connected / active / billed")
    table(L, ["Stage", "MW", "Basis", "Tag", "Src"], S["mw"], [1.25, 1.6, 1.45, 0.45, 0.35], size=sz)
    para([(S["mw_note"], False)], size=sz, sa=1, sb=1, box=L)
    if with_chart:
        pc = L.add_paragraph(); pc.paragraph_format.space_before = Pt(1)
        pc.add_run().add_picture(f"{CH}/{key}_mw.png", width=Inches(4.5))
    heading(R, "3. Power and interconnection", first=True)
    table(R, ["Item", "Status", "Tag", "Src"], S["power"], [1.05, 3.3, 0.42, 0.33], size=sz)
    heading(R, "4. Permits (TCEQ / state / local)")
    table(R, ["Permit", "Status / dates", "Tag", "Src"], S["permits"], [1.25, 3.1, 0.42, 0.33], size=sz)
    heading(R, "5. Imagery, construction evidence and contract economics")
    table(R, ["Check", "Finding", "Tag", "Src"], S["evidence"], [1.1, 3.25, 0.42, 0.33], size=sz)
    para([("What would change our view: ", True), (S["what"], False)], size=sz, sa=0, sb=2.5, box=R)
    para([("What changed since Oct 5: ", True), (S["changed"], False)], size=sz, sa=0, sb=1.5, box=R)
    for cell in (L, R):
        for pp in cell.paragraphs:
            if not pp.text.strip() and not pp._p.xpath(".//w:drawing"):
                pp.paragraph_format.space_before = Pt(0); pp.paragraph_format.space_after = Pt(0)
                pp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; pp.paragraph_format.line_spacing = Pt(2)
    src = " · ".join(S["sources"]).replace("https://", "").replace("http://", "")
    p = para([("Sources: ", True), (src, False)], size=SRC, sa=0, sb=2)
    for r in p.runs: r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    out = f"{D}/appendix_{key}.docx"; doc.save(out); return out

def pdf_pages(docx):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", D, docx], check=True, capture_output=True, timeout=180)
    pdf = docx[:-5] + ".pdf"
    n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout).group(1))
    return pdf, n

def md(key, S):
    def mdt(hdr, rows):
        return "\n".join(["| " + " | ".join(hdr) + " |", "|" + "---|" * len(hdr)] + ["| " + " | ".join(r) + " |" for r in rows])
    out = [f"# {S['title']}", f"_{S['sub']}_", "", DEFS, "",
           "## 1. Site facts", mdt(["Item", "Fact", "Tag", "Src"], S["facts"]), "",
           "## 2. MW: announced / connected / active / billed", mdt(["Stage", "MW", "Basis", "Tag", "Src"], S["mw"]), "", S["mw_note"], "",
           f"![MW by stage](charts/{key}_mw.png)", "",
           "## 3. Power and interconnection", mdt(["Item", "Status", "Tag", "Src"], S["power"]), "",
           "## 4. Permits (TCEQ / state / local)", mdt(["Permit", "Status / dates", "Tag", "Src"], S["permits"]), "",
           "## 5. Imagery, construction evidence and contract economics", mdt(["Check", "Finding", "Tag", "Src"], S["evidence"]), "",
           f"**What would change our view:** {S['what']}", "", f"**What changed since Oct 5:** {S['changed']}", "",
           "## Sources", *[f"- {s}" for s in S["sources"]], "",
           "_Built by `peers/sites/build_site_appendices.py` (content `site_data.py`); one page landscape (docx → PDF via LibreOffice)._"]
    open(f"{D}/appendix_{key}.md", "w").write("\n".join(out) + "\n")

if __name__ == "__main__":
    pdfs, report = [], []
    for key, S in SITES.items():
        chart(key, S["chart"]); md(key, S)
        done = False
        for with_chart in (True, False):
            for body in (8.0, 7.6, 7.3, 7.0, 6.8):
                d = build(key, S, body, max(5.6, body - 1.6), with_chart)
                pdf, n = pdf_pages(d)
                if n == 1:
                    report.append((key, body, with_chart, n)); pdfs.append(pdf); done = True; break
            if done: break
        assert done, f"{key} does not fit on one page"
    subprocess.run(["pdfunite", *pdfs, f"{D}/peer_site_appendices.pdf"], check=True)
    for r in report: print("fit", r)
    print("combined", f"{D}/peer_site_appendices.pdf")
