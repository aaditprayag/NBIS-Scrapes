import openpyxl,sys
f=sys.argv[1]
wb=openpyxl.load_workbook(f,data_only=True)
err=0
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('#') : err+=1; print('ERR',ws.title,c.coordinate,c.value)
            if isinstance(c.value,str) and c.value.startswith('Err:'): err+=1; print('ERR',ws.title,c.coordinate,c.value)
print('errors',err)
ws=wb['KPIs']; n=ws.max_column
for r in ws.iter_rows(min_row=3,max_row=ws.max_row):
    if r[0].value and (str(r[0].value)[0] in 'KM' or r[0].value=='ID') and len(str(r[0].value))<5:
        vals=[(f"{c.value:.2f}" if isinstance(c.value,(int,float)) else str(c.value)) for c in r[4:n-1]]
        print(r[0].value, (r[1].value or '')[:34].ljust(34), ' '.join(x[:7].rjust(7) for x in vals))
