# Shared peer-KPI workbook engine (same layout / row IDs / definitions as peers/coreweave/build_crwv_kpis.py).
# Used by peers/iren/build_iren_kpis.py and peers/oracle/build_orcl_kpis.py. Author: Grok Bot for Aadit Prayag, 2026-10-04/05 ET.
# Template contract: Inputs I01-I27 keep CRWV semantics; I28-I32 are new standard inputs; I40+ are peer-specific memo inputs.
# KPIs K01-K40 keep CRWV semantics; K41-K45 are new standard KPIs (n/d for CRWV unless noted). Tags: R/D/E/None(n/d).
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

BLUE = Font(color="0000CC"); ORANGE = Font(color="C65911"); GREY = Font(color="808080", italic=True); BLACK = Font(color="000000")
BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="D9E1F2"); SEC_F = PatternFill("solid", fgColor="F2F2F2")
TAGFONT = {"R": BLUE, "D": BLACK, "E": ORANGE, None: GREY}
TAGNAME = {"R": "Reported", "D": "Derived", "E": "Estimated", None: "n/d"}
RANK = {"R": 0, "D": 1, "E": 2}
wrap = Alignment(wrap_text=True, vertical="top")
C0 = 5

def weakest(tags):
    if any(t is None for t in tags): return None
    return max(tags, key=lambda t: RANK[t]) if tags else "D"

class Peer:
    def __init__(self, peer, out, qs, qend, fq, days, sources):
        self.peer, self.out, self.qs, self.qend, self.fq, self.S = peer, out, qs, qend, fq, sources
        self.N = len(qs); self.days = days
        self.inputs = []; self.derived = []; self.params = []; self.src_rows = []
        self.ROW = {}; self.TAG = {}  # id -> list of tags per col (None = n/d)
        self.kpi_ovr = {}; self.extra_kpis = []
        self.debt = None
    def qc(self, i): return L(C0 + i)
    # ---- inputs ----
    def add(self, iid, item, unit, vals, tags, srcs, sect, note=""):
        N = self.N
        assert len(vals) == N, (iid, len(vals))
        if isinstance(tags, str): tags = [tags] * N
        tags = [t if v is not None else None for t, v in zip(tags, vals)]
        if isinstance(srcs, str): srcs = [srcs] * N
        if isinstance(sect, str): sect = [sect] * N
        assert len(tags) == N and len(srcs) == N and len(sect) == N, iid
        for v, t in zip(vals, tags):
            if v is not None: assert t in ("R", "D", "E"), (iid, t)
        self.inputs.append(dict(id=iid, item=item, unit=unit, vals=vals, tags=tags, srcs=srcs, sect=sect, note=note))
        self.TAG[iid] = tags
    def param(self, pid, name, unit, val, tag, sk, sect):
        self.params.append((pid, name, unit, val, tag, sk, sect)); self.TAG[pid] = [tag] * self.N
    def dtag(self, iid, j):
        if j < 0 or j >= self.N or iid not in self.TAG: return None
        return self.TAG[iid][j]
    def derive(self, iid, item, unit, expr, deps, note, idx=None, force=None):
        """expr(i)-> excel expr (no '='), deps(i)-> list of (id, j); tag = weakest(dep tags) with R->D; force(i)-> override tag"""
        self.derived.append(dict(id=iid, item=item, unit=unit, expr=expr, deps=deps, note=note, idx=idx, force=force))
        tags = []
        for i in range(self.N):
            if idx is not None and i not in idx: tags.append(None); continue
            dt = [self.dtag(a, j) for a, j in deps(i)]
            t = weakest(dt)
            if t is not None and force: t = force(i) or t
            if t == "R": t = "D"
            tags.append(t)
        self.TAG[iid] = tags
    def passthru(self, iid, item, unit, src, note):
        """derived row equal to another row (keeps tag)"""
        self.derived.append(dict(id=iid, item=item, unit=unit, expr=lambda i: f"__REF__{src}__{i}", deps=lambda i: [(src, i)], note=note, idx=None, force=None, keep=True))
        self.TAG[iid] = list(self.TAG[src])

# ---------------- standard KPI library ----------------
def std_kpis():
    K = []
    def k(kid, name, unit, did, expr, deps, note, fmt="0.00", passthru=False, sect=None):
        K.append(dict(kid=kid, name=name, unit=unit, did=did, expr=expr, deps=deps, note=note, fmt=fmt, passthru=passthru, sect=sect))
    A = "A. Revenue per MW ($M per MW-year, annualized: quarterly revenue x 4 / MW)"
    k("K01","Revenue per active MW - period-end MW","$M/MW-yr","D1", lambda I,i: f"{I('I01',i)}*4/{I('I09',i)}", lambda i: [("I01",i),("I09",i)], "Nebius KPIs 'Revenue per energized/active MW' (= rev x4 / MW).", sect=A)
    k("K02","Revenue per active MW - average MW","$M/MW-yr","D1", lambda I,i: f"{I('I01',i)}*4/{I('I09A',i)}", lambda i: [("I01",i),("I09A",i)], "Same numerator on (open+close)/2 MW. Preferred trend line.")
    k("K03","Revenue per contracted MW","$M/MW-yr","D1", lambda I,i: f"{I('I01',i)}*4/{I('I10',i)}", lambda i: [("I01",i),("I10",i)], "Contracted = power under contract incl. not yet built.")
    k("K04","Contracted / active MW (build-out multiple)","x","D1", lambda I,i: f"{I('I10',i)}/{I('I09',i)}", lambda i: [("I10",i),("I09",i)], "")
    k("K05","Adj. EBITDA per active MW (average MW)","$M/MW-yr","D1", lambda I,i: f"{I('I02',i)}*4/{I('I09A',i)}", lambda i: [("I02",i),("I09A",i)], "Nebius KPIs 'Adj. EBITDA per energized MW'.")
    B = "B. Revenue per GPU and implied utilization"
    k("K06","Revenue per GPU ($K per GPU-year, avg GPUs)","$K/GPU-yr","D2", lambda I,i: f"{I('I01',i)}*4000/{I('I11A',i)}", lambda i: [("I01",i),("I11A",i)], "Nebius KPIs 'Revenue per GPU' = rev x 4000 / GPU_Avg.", "0.0", sect=B)
    k("K07","Implied blended realized price ($/GPU-hr)","$/GPU-hr","D2", lambda I,i: f"{I('I01',i)}*1000000/({I('I11A',i)}*{I('I26',i)}*24)", lambda i: [("I01",i),("I11A",i),("I26",i)], "Revenue / (avg GPUs x hours).")
    k("K08","Fleet-weighted on-demand list price ($/GPU-hr)","$/GPU-hr","D3", lambda I,i: f"{I('I12W',i)}", lambda i: [("I12W",i)], "Nebius KPIs 'Fleet-weighted on-demand list price'.", passthru=True)
    k("K09","Implied utilization (realized $/GPU-hr / weighted list)","%","D3", lambda I,i: f"{I('I01',i)}*1000000/({I('I11A',i)}*{I('I26',i)}*24)/{I('I12W',i)}*100", lambda i: [("I01",i),("I11A",i),("I26",i),("I12W",i)], "Mixes realized-price discount with utilization (Nebius caveat). NOT physical utilization - see K41.", "0.0")
    C = "C. Capex per MW ($M per MW)"
    k("K10","Cash capex per new active MW","$M/MW","D4", lambda I,i: f"{I('I04',i)}/{I('I09N',i)}", lambda i: [("I04",i),("I09N",i)], "Nebius KPIs 'Capex per new active MW' = Capex / MW_New. Timing-distorted single quarters.", "0.0", sect=C)
    k("K11","Cash capex per new active MW (incl. estimated MW)","$M/MW","D4", lambda I,i: f"{I('I04',i)}/{I('I09NE',i)}", lambda i: [("I04",i),("I09NE",i)], "", "0.0")
    k("K12","In-service additions per new active MW","$M/MW","D4", lambda I,i: f"{I('I25D',i)}/{I('I09N',i)}", lambda i: [("I25D",i),("I09N",i)], "= delta(gross PP&E - CIP) / delta active MW.", "0.0")
    k("K13","Trailing-4Q cash capex per trailing-4Q new active MW","$M/MW","D4", lambda I,i: f"SUM(Inputs!{L(C0+i-3)}{{r_I04}}:{L(C0+i)}{{r_I04}})/{I('I09T4',i)}", lambda i: [("I04",i-3),("I04",i-2),("I04",i-1),("I04",i),("I09T4",i)], "Smooths timing.", "0.0")
    k("K14","Net PP&E per active MW (EoP)","$M/MW","D4", lambda I,i: f"{I('I07',i)}/{I('I09',i)}", lambda i: [("I07",i),("I09",i)], "Includes CIP for not-yet-live capacity -> overstates.", "0.0")
    k("K15","Gross in-service PP&E per active MW (EoP)","$M/MW","D4", lambda I,i: f"{I('I05S',i)}/{I('I09',i)}", lambda i: [("I05S",i),("I09",i)], "", "0.0")
    D = "D. D&A useful lives and depreciation yield"
    k("K16","Useful life - GPU / server equipment","years","D5", lambda I,i: f"{I('I28',i)}", lambda i: [("I28",i)], "Disclosed useful-life policy.", "0", passthru=True, sect=D)
    k("K17","Depreciation yield (D&A x4 / avg in-service gross PP&E)","%","D5", lambda I,i: f"{I('I03',i)}*4/(({I('I05S',i-1)}+{I('I05S',i)})/2)*100", lambda i: [("I03",i),("I05S",i-1),("I05S",i)], "Mirrors Nebius Capex_DA 'Depreciation yield - reported basis'.", "0.0")
    k("K18","Implied average depreciable life (1 / yield)","years","D5", lambda I,i: f"100/{{K17}}", lambda i: [("I03",i),("I05S",i-1),("I05S",i)], "", "0.0")
    k("K19","D&A as % of revenue","%","D5", lambda I,i: f"{I('I03',i)}/{I('I01',i)}*100", lambda i: [("I03",i),("I01",i)], "Context row.", "0.0")
    E = "E. Backlog / RPO"
    k("K20","RPO (ASC 606, EoP)","$B","D6", lambda I,i: f"{I('I14',i)}", lambda i: [("I14",i)], "Same ASC 606 basis as Nebius.", "0.0", passthru=True, sect=E)
    k("K21","Revenue backlog (company metric / incl. contracted non-RPO, EoP)","$B","D6", lambda I,i: f"{I('I15',i)}", lambda i: [("I15",i)], "", "0.0", passthru=True)
    k("K22","RPO % recognized within 24 months","%","D6", lambda I,i: f"{I('I16',i)}", lambda i: [("I16",i)], "", "0", passthru=True)
    k("K23","RPO / annualized quarterly revenue","years","D6", lambda I,i: f"{I('I14',i)}*1000/({I('I01',i)}*4)", lambda i: [("I14",i),("I01",i)], "Coverage of current run-rate.", "0.0")
    F = "F. Customer concentration (% of quarter revenue; customers <10% not disclosed)"
    k("K24","Largest customer % of revenue","%","D7", lambda I,i: f"{I('I17',i)}", lambda i: [("I17",i)], "", "0", passthru=True, sect=F)
    k("K25","2nd-largest customer %","%","D7", lambda I,i: f"{I('I18',i)}", lambda i: [("I18",i)], "", "0", passthru=True)
    k("K26","Top-N >=10% customers combined %","%","D7", lambda I,i: f"SUM({I('I17',i)},{I('I18',i)},{I('I19',i)})", lambda i: [("I17",i)], "", "0", passthru=True)
    k("K27","Number of customers >=10% of revenue","#","D7", lambda I,i: f"{I('I20',i)}", lambda i: [("I20",i)], "", "0", passthru=True)
    G = "G. Funding cost (% p.a.)"
    K.append(dict(kid="DEBT", sect=G))
    k("K36","Implied cash cost: interest x4 / avg principal","%","D8", lambda I,i: f"{I('I21',i)}*4/{I('I23A',i)}*100", lambda i: [("I21",i),("I23A",i)], "")
    k("K37","Company-reported weighted average interest rate","%","D8", lambda I,i: f"{I('I24',i)}", lambda i: [("I24",i)], "", "0.0", passthru=True)
    k("K38","Total debt principal (EoP)","$B","D8", lambda I,i: f"{I('I23',i)}/1000", lambda i: [("I23",i)], "Excl. leases.", "0.0", passthru=True)
    k("K39","Debt principal / RPO","x","D8", lambda I,i: f"{I('I23',i)}/1000/{I('I14',i)}", lambda i: [("I23",i),("I14",i)], "", "0.00")
    k("K40","Capitalized interest % of contractual interest","%","D8", lambda I,i: f"{I('I22',i)}/{I('I21',i)}*100", lambda i: [("I22",i),("I21",i)], "", "0")
    H = "H. Template extensions added in phase 2 (K41+; n/d for CoreWeave unless stated)"
    k("K41","Company-reported GPU utilization (physical)","%","D3", lambda I,i: f"{I('I29',i)}", lambda i: [("I29",i)], "Company-stated utilization of deployed GPUs (not price-based like K09).", "0.0", passthru=True, sect=H)
    k("K42","RPO % recognized within 12 months","%","D6", lambda I,i: f"{I('I30',i)}", lambda i: [("I30",i)], "", "0", passthru=True)
    k("K43","Incremental annualized revenue per incremental MW (YoY)","$M/MW-yr","D1", lambda I,i: f"({I('I01',i)}-{I('I01',i-4)})*4/{I('I09T4',i)}", lambda i: [("I01",i),("I01",i-4),("I09T4",i)], "= (rev_q - rev_q-4) x 4 / MW added over the same 4 quarters. Marginal economics of new capacity; immune to legacy-base mix.", "0.0")
    k("K44","Annualized revenue / avg gross GPU-compute hardware in service","x","D4", lambda I,i: f"{I('I01',i)}*4/(({I('I31',i-1)}+{I('I31',i)})/2)", lambda i: [("I01",i),("I31",i-1),("I31",i)], "Asset-turn on the compute fleet (hardware only; excludes buildings/power).", "0.00")
    k("K45","Segment gross margin (revenue - cost of revenue excl. D&A) / revenue","%","D1", lambda I,i: f"({I('I01',i)}-{I('I32',i)})/{I('I01',i)}*100", lambda i: [("I01",i),("I32",i)], "Cash gross margin of the KPI perimeter, before depreciation.", "0.0")
    return K

DEFS = [
 ("D1","Revenue per MW","Quarterly revenue x 4 / MW, $M per MW-year. Denominators: (a) active MW at period end; (b) average of opening/closing active MW; (c) contracted MW. EBITDA/MW same construction. K43 = YoY incremental revenue x4 / MW added in the same 4 quarters. K45 = segment cash gross margin.",
  "/workspace/nebius/model/nebius_model.xlsx KPIs rows 'Revenue per BILLING / contracted / energized MW' (=AIRev*4/MW); /workspace/nebius/parts/comps_rev_per_mw.csv; capacity/gpu_mw_capex.md s.4 MW basis."),
 ("D2","Revenue per GPU","Quarterly revenue x 4000 / average GPUs ($K per GPU-year); implied $/GPU-hr = revenue / (avg GPUs x days x 24).",
  "nebius_model.xlsx KPIs 'Revenue per GPU ($K per GPU-yr)' and 'Implied blended $/GPU-hr'."),
 ("D3","Implied utilization","Implied blended $/GPU-hr / fleet-weighted on-demand list price. K41 = company-reported physical GPU utilization (new row).",
  "nebius_model.xlsx KPIs 'Fleet-weighted on-demand list price' / 'Implied utilization (blended $/hr / weighted list)'; prices /workspace/nebius/parts/price_history_quarterly.csv & price_history_raw.csv."),
 ("D4","Capex per MW","Cash purchases of PP&E / new active MW in the quarter ($M per MW); in-service additions per new MW; trailing-4Q; stock checks net PP&E / MW and in-service gross PP&E / MW. K44 = annualized revenue / average gross compute hardware.",
  "nebius_model.xlsx KPIs 'Capex per new active MW'; Hist 'Capex (purchases of PP&E + intangibles, cash)'; capacity/gpu_mw_capex.xlsx 'Nebius checks' and 'Peers'."),
 ("D5","D&A useful lives","Disclosed useful-life table and any change in estimate; depreciation yield = D&A x 4 / average in-service gross PP&E; implied life = 1/yield.",
  "/workspace/nebius/filings/nebius_filings.xlsx D&A_Policy; nebius_model.xlsx Capex_DA 'Depreciation yield - REPORTED basis'."),
 ("D6","Backlog / RPO","RPO = ASC 606 remaining performance obligations ($B, EoP); % within 24 months (K22) and 12 months (K42); backlog incl. non-RPO contracted value shown separately; RPO / (quarterly revenue x 4).",
  "nebius_model.xlsx Hist 'Remaining performance obligations'; nebius_filings.xlsx Quarterly 'RPO'."),
 ("D7","Customer concentration","Each customer >=10% of quarter revenue, top-N combined, count of >=10% customers.",
  "nebius_filings.xlsx Quarterly 'Customer concentration (>10% of revenue, quarter)'."),
 ("D8","Funding cost","Weighted effective cost of debt = sum(principal x effective rate) / sum(principal with a disclosed rate), EoP, excl. leases; split by facility type; cross-checks: interest x4 / avg principal; company-reported rate.",
  "/workspace/nebius/parts/comps_sources.csv 'Wtd effective cost of debt (%)'; shared/findings_log.xlsx cost-of-debt row (Nebius 5.1%)."),
]

def header(ws, row, cols, widths=None):
    for j, h in enumerate(cols, 1):
        c = ws.cell(row=row, column=j, value=h); c.font = BOLD; c.fill = HDR; c.alignment = wrap
    if widths:
        for j, w in enumerate(widths, 1): ws.column_dimensions[L(j)].width = w

def build(P, readme, defs_app, notes, mapping, kpi_ovr, comparison_ids, kpi_order_extra=None, debt=None, memo_kpis=None):
    """P: Peer. defs_app: {Dn: peer application text}. notes: list of (type, note, impact, fix). mapping: list of rows for Fiscal_Map.
    kpi_ovr: {kid: dict(note=..., idx=[...], expr=..., deps=..., name=...)}; debt: dict(SNAP, SNAPQ, SNAPSRC, FAC, CATS, RATE_NOTE, POST)."""
    N = P.N; NC = C0 + N  # notes column
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "README"
    for i, t in enumerate(readme, 1):
        ws.cell(row=i, column=1, value=t).alignment = wrap
    ws.column_dimensions["A"].width = 160
    # ---- Inputs ----
    wsI = wb.create_sheet("Inputs")
    wsI.cell(row=1, column=1, value=f"Inputs - {P.peer}. Reported (blue), derived (black), estimated (orange), n/d (grey). Per-cell source detail on Sources tab. Row 2 = fiscal quarter / period end mapped to each calendar column.").font = BOLD
    for i in range(N): wsI.cell(row=2, column=C0 + i, value=P.fq[i]).font = GREY
    header(wsI, 3, ["ID", "Item", "Unit", "Tag mix"] + P.qs + ["Source summary / notes"], [7, 52, 10, 12] + [10] * N + [90])
    r = 4
    for inp in P.inputs:
        P.ROW[inp["id"]] = r
        for j, v in enumerate([inp["id"], inp["item"], inp["unit"]], 1): wsI.cell(row=r, column=j, value=v)
        wsI.cell(row=r, column=4, value="/".join(sorted(set(TAGNAME[t] for t, v in zip(inp["tags"], inp["vals"]) if v is not None))) or "n/d")
        for i, (v, t) in enumerate(zip(inp["vals"], inp["tags"])):
            c = wsI.cell(row=r, column=C0 + i)
            if v is None: c.value = "n/d"; c.font = GREY
            else:
                c.value = v; c.font = TAGFONT[t]
                doc, url = P.S.get(inp["srcs"][i], ("", ""))
                P.src_rows.append((inp["id"], inp["item"], P.qs[i], v, TAGNAME[t], doc, inp["sect"][i], url))
                if t in ("D", "E") and inp["sect"][i]: c.comment = Comment(inp["sect"][i][:1500], "GrokBot")
        srcset = []
        for sk in inp["srcs"]:
            if sk and sk in P.S and P.S[sk][0] not in srcset: srcset.append(P.S[sk][0])
        wsI.cell(row=r, column=NC, value=("; ".join(srcset) + (" | NOTE: " + inp["note"] if inp["note"] else ""))).alignment = wrap
        r += 1
    r += 1
    wsI.cell(row=r, column=1, value="Parameters (estimation assumptions)").font = BOLD; r += 1
    for pid, name, unit, val, tag, sk, sect in P.params:
        P.ROW[pid] = r
        for j, v in enumerate([pid, name, unit, TAGNAME[tag]], 1): wsI.cell(row=r, column=j, value=v)
        c = wsI.cell(row=r, column=C0, value=val); c.font = TAGFONT[tag]
        wsI.cell(row=r, column=NC, value=P.S[sk][0] + " - " + sect).alignment = wrap
        P.src_rows.append((pid, name, "param", val, TAGNAME[tag], P.S[sk][0], sect, P.S[sk][1]))
        r += 1
    r += 1
    wsI.cell(row=r, column=1, value="Derived & estimated input rows (formulas)").font = BOLD; r += 1
    for d in P.derived: P.ROW[d["id"]] = r; r += 1  # pre-assign rows (forward refs allowed)
    def ref(iid, j):
        if iid.startswith("P"): return f"$E${P.ROW[iid]}"
        return f"{L(C0 + j)}{P.ROW[iid]}"
    for d in P.derived:
        rr = P.ROW[d["id"]]
        for j, v in enumerate([d["id"], d["item"], d["unit"]], 1): wsI.cell(row=rr, column=j, value=v)
        tags = P.TAG[d["id"]]
        wsI.cell(row=rr, column=4, value="/".join(sorted(set(TAGNAME[t] for t in tags if t))) or "n/d")
        for i in range(N):
            t = tags[i]
            c = wsI.cell(row=rr, column=C0 + i)
            if t is None: c.value = "n/d"; c.font = GREY; continue
            if d.get("keep"):
                src = d["deps"](i)[0][0]; c.value = f"={ref(src, i)}"
            else:
                e = d["expr"](i)
                e = e.replace("{", "").replace("}", "") if False else e
                c.value = "=" + fill_refs(e, ref)
            c.font = TAGFONT[t]
        wsI.cell(row=rr, column=NC, value=d["note"]).alignment = wrap
    # ---- Debt ----
    DROW = {}
    if debt:
        wsD = wb.create_sheet("Debt_Facilities")
        SNAP = debt["SNAP"]; NS = len(SNAP)
        wsD.cell(row=1, column=1, value=f"Debt by facility - {P.peer}. Principal ($M) and effective interest rate (%) as printed in each period-end debt table. Excludes leases. Weighted rates exclude facilities without a disclosed rate from the denominator.").font = BOLD
        wsD.cell(row=2, column=1, value=debt["RATE_NOTE"]).alignment = wrap
        cols = ["Facility", "Category"]
        for s in SNAP: cols += [f"{s} principal", f"{s} eff. rate %"]
        header(wsD, 4, cols, [50, 22] + [11] * (2 * NS))
        dr = 5
        for name, cat, vals in debt["FAC"]:
            wsD.cell(row=dr, column=1, value=name); wsD.cell(row=dr, column=2, value=cat)
            for k, v in enumerate(vals):
                if v is not None:
                    a = wsD.cell(row=dr, column=3 + 2 * k, value=v[0]); a.font = BLUE
                    if v[1] is not None:
                        b = wsD.cell(row=dr, column=4 + 2 * k, value=v[1]); b.font = BLUE
                    sk = debt["SNAPSRC"][k]
                    P.src_rows.append(("DEBT", name + " (principal / eff. rate)", SNAP[k], f"{v[0]} / {v[1] if v[1] is not None else 'n/d'}%", "Reported", P.S[sk][0], debt.get("SECT", "Debt note"), P.S[sk][1]))
            dr += 1
        last = dr - 1; dr += 1
        wsD.cell(row=dr, column=1, value="Total principal").font = BOLD
        for k in range(NS):
            pc = L(3 + 2 * k); wsD.cell(row=dr, column=3 + 2 * k, value=f"=SUM({pc}5:{pc}{last})")
        DROW["tot"] = dr; dr += 1
        wsD.cell(row=dr, column=1, value="Check vs Inputs I23 (total debt principal)")
        for k in range(NS):
            wsD.cell(row=dr, column=3 + 2 * k, value=f"={L(3+2*k)}{DROW['tot']}-Inputs!{L(C0+debt['SNAPQ'][k])}{P.ROW['I23']}")
        dr += 1
        wsD.cell(row=dr, column=1, value="WEIGHTED EFFECTIVE COST OF DEBT (principal-weighted, facilities with disclosed rate)").font = BOLD
        for k in range(NS):
            pc, rc = L(3 + 2 * k), L(4 + 2 * k)
            wsD.cell(row=dr, column=4 + 2 * k, value=f"=IFERROR(SUMPRODUCT({pc}5:{pc}{last},{rc}5:{rc}{last})/SUMPRODUCT(({rc}5:{rc}{last}<>\"\")*{pc}5:{pc}{last}),\"n/d\")")
        DROW["wavg"] = dr; dr += 1
        for cat in debt["CATS"]:
            wsD.cell(row=dr, column=1, value=f"  {cat}: principal / weighted eff. rate")
            for k in range(NS):
                pc, rc = L(3 + 2 * k), L(4 + 2 * k)
                wsD.cell(row=dr, column=3 + 2 * k, value=f"=SUMIF($B$5:$B${last},\"{cat}\",{pc}5:{pc}{last})")
                wsD.cell(row=dr, column=4 + 2 * k, value=f"=IFERROR(SUMPRODUCT(($B$5:$B${last}=\"{cat}\")*{pc}5:{pc}{last}*{rc}5:{rc}{last})/SUMPRODUCT(($B$5:$B${last}=\"{cat}\")*({rc}5:{rc}{last}<>\"\")*{pc}5:{pc}{last}),\"n/d\")")
            DROW[cat] = dr; dr += 1
        wsD.cell(row=dr, column=1, value=f"  Share of principal: {debt['CATS'][0]} %")
        for k in range(NS):
            wsD.cell(row=dr, column=4 + 2 * k, value=f"=IFERROR({L(3+2*k)}{DROW[debt['CATS'][0]]}/{L(3+2*k)}{DROW['tot']}*100,\"n/d\")")
        DROW["ddtlsh"] = dr; dr += 2
        for t in debt.get("POST", []):
            wsD.cell(row=dr, column=1, value=t).alignment = wrap; dr += 1
    # ---- KPIs ----
    wsK = wb.create_sheet("KPIs", 1)
    wsK.cell(row=1, column=1, value=f"KPIs - {P.peer} - quarterly, Nebius definitions (Definitions tab). Value row = live formula; row below = tag (weakest input). Row 2 = fiscal quarter mapped to each calendar column.").font = BOLD
    for i in range(N): wsK.cell(row=2, column=C0 + i, value=P.fq[i]).font = GREY
    header(wsK, 3, ["ID", "KPI", "Unit", "Def"] + P.qs + ["Method / caveat (" + P.peer + ")"], [7, 56, 12, 6] + [10] * N + [90])
    KR = {}; kr = [4]
    def section(title):
        wsK.cell(row=kr[0], column=1, value=title).font = BOLD
        for j in range(1, NC + 1): wsK.cell(row=kr[0], column=j).fill = SEC_F
        kr[0] += 1
    Ifn = lambda iid, i: f"Inputs!{ref(iid, i)}" if iid in P.ROW else "NA()"
    def put_kpi(kid, name, unit, did, cells, note, fmt):
        row = kr[0]; KR[kid] = row
        for j, v in enumerate([kid, name, unit, did], 1): wsK.cell(row=row, column=j, value=v)
        for i in range(N):
            f, t = cells[i]
            c = wsK.cell(row=row, column=C0 + i, value=(f if f is not None else "n/d")); c.number_format = fmt
            c.font = TAGFONT.get(t, GREY) if f is not None else GREY
            wsK.cell(row=row + 1, column=C0 + i, value=(TAGNAME.get(t) if f is not None else "n/d")).font = GREY
        wsK.cell(row=row + 1, column=2, value="   tag").font = GREY
        wsK.cell(row=row, column=NC, value=note).alignment = wrap
        kr[0] += 2
    for kd in std_kpis():
        if kd.get("sect"): section(kd["sect"])
        if kd["kid"] == "DEBT":
            if not debt: continue
            SNAPQ = debt["SNAPQ"]
            def dref(key, i, rate=True):
                if i not in SNAPQ: return None
                k = SNAPQ.index(i); return f"Debt_Facilities!{L(4+2*k) if rate else L(3+2*k)}{DROW[key]}"
            o = kpi_ovr.get("K28", {})
            put_kpi("K28", "Weighted effective cost of debt (principal-weighted, EoP) - all facilities", "%", "D8",
                    [((f"=IFERROR({dref('wavg',i)}*1,\"n/d\")", "D") if dref('wavg', i) else (None, None)) for i in range(N)], o.get("note", ""), "0.00")
            for cat, kid in zip(debt["CATS"], ["K29", "K30", "K31", "K32", "K33", "K34"]):
                o = kpi_ovr.get(kid, {})
                put_kpi(kid, f"  of which {cat} - weighted eff. rate", "%", "D8",
                        [((f"={dref(cat,i)}", "D") if dref(cat, i) else (None, None)) for i in range(N)], o.get("note", "n/d where facility type not outstanding."), "0.00")
            o = kpi_ovr.get("K35", {})
            put_kpi("K35", f"{debt['CATS'][0]} share of principal", "%", "D8",
                    [((f"={dref('ddtlsh',i)}", "D") if dref('ddtlsh', i) else (None, None)) for i in range(N)], o.get("note", ""), "0")
            continue
        o = dict(kd); o.update(kpi_ovr.get(kd["kid"], {}))
        cells = []
        for i in range(N):
            if o.get("idx") is not None and i not in o["idx"]: cells.append((None, None)); continue
            deps = o["deps"](i)
            if any(j < 0 or j >= N for _, j in deps): cells.append((None, None)); continue
            tags = [P.dtag(a, j) for a, j in deps]
            t = weakest(tags)
            if t is None: cells.append((None, None)); continue
            if not o.get("passthru") and t == "R": t = "D"
            if o.get("tagf"): t = o["tagf"](i) or t
            e = o["expr"](Ifn, i)
            e = e.replace("{r_I04}", str(P.ROW.get("I04", 0)))
            e = _re.sub(r"\{r(I\w+?)\}", lambda m: str(P.ROW[m.group(1)]), e)
            if "{K17}" in e: e = e.replace("{K17}", f"{L(C0+i)}{KR['K17']}")
            cells.append((f"=IFERROR({e},\"n/d\")", t))
        put_kpi(o["kid"], o.get("name", kd["name"]), o["unit"], o["did"], cells, o.get("note", kd["note"]), o.get("fmt", kd["fmt"]))
    if memo_kpis:
        section("M. Peer-specific memo KPIs (not in template; see method column)")
        for m in memo_kpis:
            cells = []
            for i in range(N):
                deps = m["deps"](i)
                if (m.get("idx") is not None and i not in m["idx"]) or any(j < 0 or j >= N for _, j in deps): cells.append((None, None)); continue
                t = weakest([P.dtag(a, j) for a, j in deps])
                if t is None: cells.append((None, None)); continue
                if not m.get("passthru") and t == "R": t = "D"
                e = _re.sub(r"\{r(I\w+?)\}", lambda mm: str(P.ROW[mm.group(1)]), m['expr'](Ifn, i))
                cells.append((f"=IFERROR({e},\"n/d\")", t))
            put_kpi(m["kid"], m["name"], m["unit"], m.get("did", ""), cells, m["note"], m.get("fmt", "0.0"))
    wsK.freeze_panes = "E4"
    # ---- Sources ----
    wsS = wb.create_sheet("Sources")
    header(wsS, 1, ["Input ID", "Item", "Period", "Value", "Tag", "Document", "Section / page / formula", "URL"], [9, 48, 9, 14, 11, 52, 80, 90])
    for j, row in enumerate(P.src_rows, 2):
        for k, v in enumerate(row, 1): wsS.cell(row=j, column=k, value=v)
    j = len(P.src_rows) + 3
    wsS.cell(row=j, column=1, value="Document register").font = BOLD; j += 1
    for k, (doc, url) in P.S.items():
        wsS.cell(row=j, column=1, value=k); wsS.cell(row=j, column=6, value=doc); wsS.cell(row=j, column=8, value=url); j += 1
    # ---- Definitions ----
    wsF = wb.create_sheet("Definitions")
    header(wsF, 1, ["Def", "KPI group", "Definition used (formula, units)", "Mirrored from (Nebius workspace)", f"{P.peer} application / basis notes"], [6, 24, 70, 70, 90])
    for jj, d in enumerate(DEFS, 2):
        for k, v in enumerate(list(d) + [defs_app.get(d[0], "")], 1): wsF.cell(row=jj, column=k, value=v).alignment = wrap
    # ---- Fiscal map ----
    wsM = wb.create_sheet("Fiscal_Map")
    header(wsM, 1, ["Template column", "Calendar quarter-end used", "Company fiscal quarter", "Fiscal period end", "Offset vs calendar", "Notes"], [16, 18, 22, 16, 18, 90])
    for jj, row in enumerate(mapping, 2):
        for k, v in enumerate(row, 1): wsM.cell(row=jj, column=k, value=v).alignment = wrap
    # ---- Notes ----
    wsN = wb.create_sheet("Notes_Gaps")
    header(wsN, 1, ["#", "Type", "Note / gap", "Impact", "Possible fix"], [4, 14, 90, 50, 50])
    for jj, n in enumerate(notes, 2):
        wsN.cell(row=jj, column=1, value=jj - 1)
        for k, v in enumerate(n, 2): wsN.cell(row=jj, column=k, value=v).alignment = wrap
    # ---- Comparison stub (peer column) ----
    wsC = wb.create_sheet("Comparison")
    wsC.cell(row=1, column=1, value="Per-peer stub. The multi-peer Comparison tab (Nebius, CoreWeave, IREN, Oracle OCI) lives in /workspace/nebius/peers/comparison.xlsx (built by peers/build_comparison.py from each peer workbook's KPIs tab).").font = BOLD
    header(wsC, 3, ["ID", "KPI", "Unit"] + [f"{P.peer} {q}" for q in P.qs[-3:]], [7, 56, 12, 16, 16, 16])
    cr = 4
    for kid in comparison_ids:
        if kid not in KR: continue
        rr = KR[kid]
        wsC.cell(row=cr, column=1, value=kid); wsC.cell(row=cr, column=2, value=f"=KPIs!B{rr}"); wsC.cell(row=cr, column=3, value=f"=KPIs!C{rr}")
        for k, i in enumerate(range(N - 3, N)):
            c = wsC.cell(row=cr, column=4 + k, value=f"=KPIs!{L(C0+i)}{rr}"); c.number_format = "0.00"
        cr += 1
    for _ws in wb:
        for _row in _ws.iter_rows():
            for _c in _row:
                if isinstance(_c.value, str) and _c.value.startswith("= "): _c.value = "calc: " + _c.value[2:]
    wb.save(P.out)
    print("saved", P.out, "inputs", len(P.inputs), "derived", len(P.derived), "src rows", len(P.src_rows))
    return KR

import re as _re
def fill_refs(e, ref):
    """replace tokens like [I01@3] with cell refs"""
    return _re.sub(r"\[([A-Z]\d+[A-Z0-9]*)@(-?\d+)\]", lambda m: ref(m.group(1), int(m.group(2))), e)
