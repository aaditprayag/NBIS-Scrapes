#!/usr/bin/env python3
"""Build /workspace/nebius/peers/payback_dilution.xlsx  (Nebius Peer KPIs bot, 2026-10-04).
Tabs: README, Payback_ACV, Basis_Sensitivity, Dilution_Data, Dilution_Fit, Nebius_Path, Sources.
All derived numbers are live formulas. Tags: R=Reported, P=Press/secondary, D=Derived (from R/P), E=Estimated, n-d=not disclosed.
Reads (never edits): capacity/nebius_capacity.xlsx (Quarterly MW), filings/nebius_kpi_history.xlsx (recalc copy),
peers/coreweave/crwv_kpis.xlsx, peers/iren/iren_kpis.xlsx, parts/contract_economics.xlsx, capacity/gpu_mw_capex.xlsx (recalc copy).
Run:  python3 build_payback_dilution.py && python3 /workspace/nebius/model/recalc.py payback_dilution.xlsx && cp /tmp/rc/payback_dilution.xlsx .
Then: python3 build_payback_dilution.py --charts   (reads recalculated values, writes PNGs to peers/charts/)
"""
import sys, os, math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as CL

ROOT = "/workspace/nebius"
OUT = f"{ROOT}/peers/payback_dilution.xlsx"
BOLD = Font(bold=True); BLUE = Font(color="0000CC"); HDR = PatternFill("solid", fgColor="DDEBF7")
KEY = PatternFill("solid", fgColor="FFF2CC"); WRAP = Alignment(wrap_text=True, vertical="top")

# ---------------- source data pulled from team files (values copied; source cell cited) ----------------
cap = openpyxl.load_workbook(f"{ROOT}/capacity/nebius_capacity.xlsx", data_only=True)["Quarterly MW"]
def caprow(label):
    for r in range(1, cap.max_row + 1):
        if cap.cell(r, 1).value == "COMPANY" and cap.cell(r, 2).value == label: return r
    raise KeyError(label)
QCOLS = {"Q1'26": "M", "Q2'26": "N", "Q3'26": "O", "Q4'26": "P", "Q1'27": "Q", "Q2'27": "R", "Q3'27": "S", "Q4'27": "T"}
MWCASE = {k: caprow(f"Energized MW {k.lower()} - IT-normalized") for k in ("Low", "Base", "High")}
MWMIX = {k: caprow(f"Identified sites - energized {k.lower()}") for k in ("Low", "Base", "High")}
mw_path = {k: [cap[f"{c}{r}"].value for c in QCOLS.values()] for k, r in MWCASE.items()}
mw_mix = {k: (cap[f"P{r}"].value, cap[f"T{r}"].value) for k, r in MWMIX.items()}
_pl = caprow("Total incl. plug (ties to mgmt where reported)")
mgmt_guided = {"YE26": cap[f"P{_pl}"].value, "YE27": cap[f"T{_pl}"].value, "row": _pl}
crwv = openpyxl.load_workbook(f"{ROOT}/peers/coreweave/crwv_kpis.xlsx", data_only=True)["Inputs"]
cr_q = ["Q1'24", "Q2'24", "Q3'24", "Q4'24", "Q1'25", "Q2'25", "Q3'25", "Q4'25", "Q1'26", "Q2'26"]
cr_cols = "FGHIJKLMNO"  # Q1'24..Q2'26 ; E = Q4'23
cr_mw = [crwv[f"{c}36"].value for c in "EFGHIJKLMNO"]  # I09E (incl. 2024 interpolation)
cr_rev = [crwv[f"{c}4"].value for c in cr_cols]
cr_ebitda_q226, cr_isg_q126, cr_isg_q226 = crwv["O5"].value, crwv["N43"].value, crwv["O43"].value

iren = openpyxl.load_workbook(f"{ROOT}/peers/iren/iren_kpis.xlsx", data_only=True)["Inputs"]
ir_q = ["Q2'24", "Q3'24", "Q4'24", "Q1'25", "Q2'25", "Q3'25"]
ir_mw = [iren[f"{c}53"].value for c in "FGHIJKL"]      # I09E Q1'24..Q3'25 (estimated from GPUs/713)
ir_rev = [iren[f"{c}4"].value for c in "GHIJKL"]
ir_rev_q226, ir_mw_q226 = iren["O4"].value, iren["O12"].value

nkh = openpyxl.load_workbook("/tmp/rc/nkh.xlsx", data_only=True)["KPI"]
nb_q = ["Q3'24", "Q4'24", "Q1'25", "Q2'25", "Q3'25", "Q4'25", "Q1'26", "Q2'26"]
nb_mw = [nkh.cell(28, c).value for c in range(4, 13)]   # Q2'24..Q2'26 (open for Q3'24 = Q2'24)
nb_rev = [nkh.cell(6, c).value for c in range(5, 13)]   # Q3'24..Q2'26

gmc = openpyxl.load_workbook("/tmp/rc/gmc.xlsx", data_only=True)["CapexPerMW"]
def _gmrow(gen, scen):
    for r in range(1, gmc.max_row + 1):
        if str(gmc.cell(r, 1).value or "").startswith(gen) and gmc.cell(r, 2).value == scen: return r
    raise KeyError((gen, scen))
def gm(gen, scen, col): return gmc.cell(_gmrow(gen, scen), col).value
# cols: C(3)=GPUs/IT MW, G(7)=all-in $k/GPU (server+network), H(8)=GPU/server $M/MW, I(9)=network $M/MW, K(11)=owned total $M/IT MW
GB300_LOW_COMPUTE = round(gm("GB300", "Low", 8) + gm("GB300", "Low", 9), 2); GB300_HIGH_COMPUTE = round(gm("GB300", "High", 8) + gm("GB300", "High", 9), 2)
GB300_OWNED_BASE = round(gm("GB300", "Base", 11), 2)

wb = openpyxl.Workbook()

def hdr(ws, r, vals, widths=None):
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v); c.font = BOLD; c.fill = HDR; c.alignment = WRAP
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[CL(i)].width = w

# =========================== Sources ===========================
SRC = [
 ("S1", "Nebius Q2'26 shareholder letter (6-K Ex.99.2, Aug 12 2026): cover 'ACV per MW is stepping up(1): $12M 2026 base; >$20M Q2 deals; >$40M Q3 short-term'; 'Payback period(2)(3) 1 year 10 months for deals closed in Q2'; fn(2) 'revenue recognition basis, excluding any prepayment'; fn(3) 'based on forecast costs and contracted future capacity, including capacity not yet built'; body 'yield of $20-25 million per megawatt'; 'expected payback period for the associated capex and related operating costs ... 1 year and 10 months'. NO IT/gross MW basis stated.", "https://www.sec.gov/Archives/edgar/data/1513845/000110465926094568/tm2622968d1_ex99-2.htm ; local /workspace/nebius/src/q2_26.txt"),
 ("S2", "Nebius Q2'26 earnings call (Aug 12 2026): CEO 'mid-term contracts of one to three years ... yield of $20 million-$25 million per megawatt, with upfront payments that cover 50%-60% of the associated CapEx'; CFO 'more than $20 million per megawatt deals and less than two-year payback'. No MW basis.", "/workspace/nebius/filings/altdata/transcripts/transcript_q2-2026.txt"),
 ("S3", "IREN-Microsoft PR (Nov 3 2025): 5-yr, ~$9.7bn TCV, 20% prepayment, GB300, 200MW critical IT load (Horizon 1-4); Dell GPUs + ancillary equipment ~$5.8bn.", "https://www.sec.gov/Archives/edgar/data/1878848/000114036125040072/ef20058139_ex99-1.htm"),
 ("S4", "IREN investor presentation (8-K, Nov 3 2025): '$1.94bn ARR', '85% Estimated Project EBITDA Margin' (after power, R&M, other direct operating costs incl. DC operating costs; before overhead/financing), '$5.8bn estimated GPU capex, including ancillaries' (deployment services, servers, InfiniBand, cabling, software, licensing). 'Microsoft contract scales IREN's footprint to 100k GPUs'.", "https://iren.gcs-web.com/static-files/87a13320-156c-4f88-8589-dee2d10faa81"),
 ("S5", "IREN FY26 results PR (Aug 27 2026): 'Recent 3-year contracts >$20m revenue per MW (IT), representing a ~2 year payback(3)'; 'Active discussions at ~$25m per MW (IT)'; prepayments 45-55% of GPU capex; fn3 'Payback period estimated based on estimated capital expenditure for GPUs and ancillaries divided by contracted revenue less estimated direct costs.'", "https://www.sec.gov/Archives/edgar/data/1878848/000187884826000051/irenreportsfy26results.htm ; local comps_raw/iren_fy26pr__irenreportsfy26results.htm.txt"),
 ("S6", "IREN Horizon 1 delivered (Aug 13 2026): four 50MW (IT load) liquid-cooled deployments; IREN targets 480MW (gross) 2026 / 1.2GW (gross) 2027 - IREN labels both bases.", "https://www.sec.gov/Archives/edgar/data/1878848/000114036126032638/ef20080141_ex99-1.htm"),
 ("S7", "IREN Horizon DC build cost ~$300-350M per 50MW IT = $6-7M/IT MW (Capacity team Peers tab S26).", "/workspace/nebius/capacity/gpu_mw_capex.xlsx (Peers)"),
 ("S8", "CoreWeave revenue, adj. EBITDA, in-service gross PP&E, active MW (8-K PRs, 10-Q/10-K) as compiled in peer KPI workbook (Inputs I01, I02, I05S, I09/I09E).", "/workspace/nebius/peers/coreweave/crwv_kpis.xlsx ; https://www.sec.gov/Archives/edgar/data/1769628/000176962826000366/crwv-20260630.htm"),
 ("S9", "Nscale S-1 (Sep 18 2026): GPU counts by site/customer; Monarch WV 460 IT MW contracted (Anthropic); 'We expect average cash payback to occur within the initial contract term' (no figure); adj. EBITDA negative.", "https://www.sec.gov/Archives/edgar/data/2110365/000119312526395475/ck0002110365-20260918.htm ; local comps_raw/nscale_s1.htm.txt"),
 ("S10", "Contract economics (TCV, term, MW, GPU counts, MW basis) for Nebius-Microsoft/Meta, Nscale-Microsoft/Anthropic, Oracle-OpenAI with per-deal sources.", "/workspace/nebius/parts/contract_economics.xlsx (Deals, Normalized)"),
 ("S11", "GPUs per IT MW and capex per IT MW by generation (third-party price build-up; Estimated): GB300 server+network $81.9k/GPU, 491 GPUs/IT MW -> $40.3M/IT MW compute (low 35.5, high 52.5); VR $94.4k/GPU, 320/IT MW -> $30.3M; owned DC shell+power $12.5M/IT MW base.", "/workspace/nebius/capacity/gpu_mw_capex.xlsx (GPUsPerMW, CapexPerMW)"),
 ("S12", "Oracle Financial Analyst Day (Oct 16 2025): AI infrastructure adjusted gross margin 30-40% over contract life incl. start-up; example 6-yr $60B contract with ~$6.4B cost per year; costs ~35% site/data center/power, ~65% compute/network/storage (Magouyrk).", "https://www.tradingview.com/news/reuters.com,2025:newsml_L6N3VX11P:0-oracle-expects-cloud-sales-of-166-billion-by-2030-as-business-expands/ ; https://longbridge.com/en/news/261873306"),
 ("S13", "Crusoe/Blue Owl Abilene JV ~$15B for ~1.2 GW (DC shell + power, no GPUs) = $12.5M/MW; Oracle/OpenAI Abilene ~450k GB200 on 1.2 GW (375 GPUs/MW -> gross-like).", "/workspace/nebius/capacity/gpu_mw_capex.xlsx (Peers S31/S32)"),
 ("S14", "Modeling team payback inputs: PB_Opex $4.5M/IT MW-yr (rent $2.0-2.5M + power + ops; range 4-5), PB_CapGB leased GB300 $41.3M/IT MW, PB_Target 22 months.", "/workspace/nebius/model/build_part5.py (lines ~479-485) ; model/nebius_model.xlsx Payback_Floor"),
 ("S15", "Nebius AI cloud revenue and energized MW history (Q1'24-Q2'26), rev per average energized MW.", "/workspace/nebius/filings/nebius_kpi_history.xlsx KPI rows 6, 28, 58"),
 ("S16", "Nebius energized MW path by quarter, identified sites low/base/high (rows 119/121/120) and total incl. plug tied to mgmt (row 125: YE26 800, YE27 1,800 = connected guide). Basis mixed IT/gross per Capacity README.", "/workspace/nebius/capacity/nebius_capacity.xlsx (Quarterly MW) ; model/build_model.py block 2 (YE26 185/305/424, YE27 764/1,069/1,374)"),
 ("S17", "Nebius Q2'26 AI cloud adjusted EBITDA $285.7M on revenue $574.9M (segment).", "Nebius Q2'26 letter (S1) / peers/comparison.xlsx Nebius_Inputs"),
 ("S18", "Team findings: Filings 2026-10-04 21:39 ET (basis undisclosed; lean gross at ~$41M capex/IT MW); Capacity 18:08 ET (YE25 in-service servers+network $18.4M per active MW).", "/workspace/nebius/shared/findings_log.xlsx (Findings)"),
 ("S19", "Oracle-OpenAI ~$300B / ~5 yrs / 4.5 GW (WSJ; not confirmed by Oracle by name) - press.", "/workspace/nebius/parts/contract_economics.xlsx (O2)"),
]
src = wb.active; src.title = "Sources"
hdr(src, 1, ["ID", "What it supports (quoted where possible)", "URL / local path"], [6, 120, 90])
for i, (a, b, c) in enumerate(SRC, 2):
    src.cell(i, 1, a); src.cell(i, 2, b).alignment = WRAP; src.cell(i, 3, c).alignment = WRAP

# =========================== Payback_ACV ===========================
pa = wb.create_sheet("Payback_ACV", 0)
pa["A1"] = "Payback / ACV test - peer deals vs Nebius. Payback (months) = capex per MW / (ACV per MW x cash margin) x 12; prepayments EXCLUDED (Nebius fn 2, IREN fn 3 basis). $M; MW."
pa["A1"].font = BOLD
pa["A2"] = ("Within one deal, payback is identical on IT and gross MW (both numerator and denominator scale by the same factor) - columns Z vs AA prove it. "
            "The IT/gross question only bites when a per-MW benchmark on one basis is applied to per-MW revenue on another: see Basis_Sensitivity.")
pa["A2"].alignment = WRAP
hdr(pa, 4, ["Shared assumption", "", "Value", "Tag", "Source / note"])
ASS = [
 ("Opex per IT MW-yr for LEASED capacity (rent + power + ops)", 4.5, "E", "S14 Modeling PB_Opex (range 4-5)"),
 ("Default gross/IT factor (PUE-type) where basis must be converted", 1.3, "E", "S11 base PUE 1.2-1.3; sensitivity 1.2-1.4 in Basis_Sensitivity"),
 ("GB300 NVL72 compute (server+network+storage) $k per GPU", round(gm("GB300", "Base", 7), 2), "E", "S11 CapexPerMW GB300 base"),
 ("GB300 GPUs per IT MW", round(gm("GB300", "Base", 3), 2), "E", "S11"),
 ("Vera Rubin NVL72 compute $k per GPU", round(gm("Vera Rubin", "Base", 7), 2), "E", "S11 VR base"),
 ("Vera Rubin GPUs per IT MW", round(gm("Vera Rubin", "Base", 3), 2), "E", "S11"),
 ("GB200 NVL72 compute $k per GPU", round(gm("GB200", "Base", 7), 2), "E", "S11 GB200 base"),
 ("GB200 GPUs per IT MW", round(gm("GB200", "Base", 3), 2), "E", "S11"),
 ("B300 HGX compute $k per GPU", round(gm("Blackwell Ultra B300", "Base", 7), 2), "E", "S11 B300 base"),
 ("B300 GPUs per IT MW", round(gm("Blackwell Ultra B300", "Base", 3), 2), "E", "S11"),
 ("Owned DC shell + power $M per IT MW", 12.5, "E", "S11 base (range 8-17.6); Crusoe Abilene $12.5M/MW (S13)"),
 ("IREN Horizon DC build $M per IT MW", 6.5, "D", "S7 ($300-350M / 50 MW IT)"),
 ("IREN Microsoft project EBITDA margin", 0.85, "R", "S4"),
 ("Nebius AI cloud adj. EBITDA margin Q2'26 (segment, incl. overhead)", "=285.7/574.9", "D", "S17"),
 ("Oracle AI-infra CASH margin, analyst-day example: 1 - (35% x $6.4B)/$10B", "=1-0.35*6.4/10", "D", "S12 (assumes the 65% compute cost is non-cash depreciation; E element)"),
]
A = {}
for i, (lab, v, t, s) in enumerate(ASS, 5):
    pa.cell(i, 1, lab); c = pa.cell(i, 3, v); c.font = BLUE; pa.cell(i, 4, t); pa.cell(i, 5, s)
    A[i - 4] = f"$C${i}"
OPEX, PUE, GB3K, GB3D, VRK, VRD, GB2K, GB2D, B3K, B3D, DCS, IRDC, IRM, NBM, ORM = [A[i] for i in range(1, 16)]
for i in (14, 18, 19): pa.cell(4 + i, 3).number_format = "0.0%"

H0 = 22
cols = ["ID", "Deal", "Provider", "Chip", "ACV $M/yr", "tag", "Stated MW", "Stated basis (IT/gross)", "MW tag",
        "Gross/IT factor", "tag", "IT MW", "Gross MW", "ACV $M per IT MW-yr", "ACV $M per gross MW-yr",
        "Compute capex $M (GPUs+network+storage)", "tag", "Compute $M per IT MW", "DC shell+power $M per IT MW (0 = leased, rent in opex)", "tag",
        "Total capex $M per IT MW", "Total capex $M per gross MW", "Cash margin (EBITDA/rev)", "tag",
        "Payback compute-only (months)", "Payback total capex, IT basis (months)", "Payback total capex, gross basis (months)",
        "Reported payback (months)", "Payback tag", "Sources", "Notes"]
hdr(pa, H0, cols, [9, 40, 10, 12, 10, 5, 9, 9, 6, 8, 5, 9, 9, 10, 10, 12, 5, 10, 12, 5, 10, 10, 9, 5, 10, 10, 10, 9, 6, 14, 70])
pa.row_dimensions[H0].height = 75
OPXM = lambda r: f"=1-{OPEX}/N{r}"
D = [  # id, deal, prov, chip, acv, acv_t, mw, basis, mw_t, fac, fac_t, capex, capex_t, dc, dc_t, margin, m_t, rep, src, note
 ("P1", "IREN - Microsoft (Childress Horizon 1-4), 5 yr", "IREN", "GB300", "=9700/5", "D", 200, "IT", "R", f"={PUE}", "E",
  5800, "R", f"={IRDC}", "D", f"={IRM}", "R", "n-d", "S3,S4,S7",
  "Capex = Dell GPU+ancillaries $5.8B (R). ACV $1.94B = IREN ARR (R). Gross MW not stated for Horizon; factor is the sensitivity default. 85% = IREN's own project EBITDA margin (owned power)."),
 ("P1b", "IREN - recent 3-yr contracts (FY26 PR), per 1 IT MW", "IREN", "GB300-class", 20, "R", 1, "IT", "R", f"={PUE}", "E",
  "=5800/200", "E", f"={IRDC}", "D", f"={IRM}", "E", 24, "S5,S3,S4",
  "'>$20m revenue per MW (IT), representing a ~2 year payback' (R; IREN def = GPU+ancillaries capex / (revenue - direct costs)). Capex = IREN Microsoft $29M/IT MW applied (E); margin = Microsoft project margin applied (E). Model gives ~20.5 mo vs reported ~24 -> IREN's own implied capex ~$34M/IT MW."),
 ("C1", "CoreWeave fleet Q2'26 (all contracts; no deal-level MW disclosed)", "CoreWeave", "Mixed", f"={cr_rev[-1]}*4", "D", "=(1000+1500)/2", "IT", "E", f"={PUE}", "E",
  f"=({cr_isg_q126}+{cr_isg_q226})/2", "D", 0, "R", f"={cr_ebitda_q226}/{cr_rev[-1]}", "D", "n-d", "S8",
  "Avg active MW (CRWV 'active power' ~ IT-like, E). Capex = avg in-service gross PP&E (incl. leasehold improvements; DCs leased so shell = 0, rent inside adj. EBITDA). Fleet-level: includes older Hopper at lower $/MW and capacity still ramping - an upper bound on deal payback. Deal MW for Meta/OpenAI n-d."),
 ("N1", "Nscale - Microsoft, all SOWs (TX, Narvik, Loughton, Sines), 5.5 yr", "Nscale", "GB300 + VR", "=43800/5.5", "D", 596.47, "IT", "D", f"={PUE}", "E",
  f"=(191600*{GB3K}+66000*{VRK})/1000", "E", f"={DCS}", "E", f"={IRM}", "E", "n-d", "S9,S10,S11",
  "IT MW derived from 191.6k GB300/491 + 66k VR/320 (S10). Capex n-d -> GPU count x unit cost (E). Margin n-d (Nscale adj. EBITDA negative) -> IREN owned-power project margin applied (E). S-1: payback 'within the initial contract term' only."),
 ("N1a", "  of which Nscale/Aker - Microsoft Narvik, 5 yr", "Nscale", "GB300", "=6200/5", "D", 105.91, "IT", "D", f"={PUE}", "E",
  f"=52000*{GB3K}/1000", "E", f"={DCS}", "E", f"={IRM}", "E", "n-d", "S10,S11", "52,000 GB300 (R) / 491 per IT MW (E)."),
 ("N2", "Nscale - Anthropic (Monarch WV), Vera Rubin, 6 yr", "Nscale", "Vera Rubin", "=44600/6", "D", 460, "IT", "R", f"={PUE}", "E",
  f"=460*{VRD}*{VRK}/1000", "E", f"={DCS}", "E", f"={IRM}", "E", "n-d", "S9,S10,S11",
  "460 IT MW contracted (S-1 site table, R). GPUs n-d (redacted) -> density x unit cost (E). Monarch gross/IT 8/6.5 GW = 1.23 (S11 Peers)."),
 ("O1", "Oracle - OpenAI (Stargate incl. Abilene), ~5 yr", "Oracle", "GB200 mix", "=300000/5", "P", 4500, "gross", "E", 1.2, "E",
  f"=L{{r}}*{GB2D}*{GB2K}/1000", "E", 0, "E", f"={ORM}", "D", "n-d", "S19,S12,S13",
  "4.5 GW unlabeled; read as gross (Abilene 450k GB200 / 1.2 GW = 375 GPUs/MW < 449/gross MW) -> factor 1.2 (E). Oracle leases shells from developers (Crusoe/Blue Owl JV) so shell = 0 and rent sits in the 35% site/DC/power cost bucket (S12)."),
 ("NB1", "Nebius - Microsoft (Vineland), 5 yr committed", "Nebius", "GB300", "=17392.9/5", "D", 300, "gross", "E", f"={PUE}", "E",
  f"=100000*{GB3K}/1000", "E", 0, "R", None, "E", "n-d", "S10,S11,S14",
  "'up to 300 MW' unlabeled; Capacity reads gross. GPU-derived IT MW = 100k/491 = 203.7 (ACV $17.1M/IT MW). Capex: >100k GB300 (P) x unit cost (E). Vineland = build-to-suit lease (rent in opex). Margin = 1 - opex/ACV per IT MW (E). Prepayment 40% excluded."),
 ("NB2", "Nebius - Meta #1 (Nov-2025), 5 yr", "Nebius", "n-d (B300?)", "=2880.7/5", "D", 50, "IT", "E", f"={PUE}", "E",
  f"=L{{r}}*{B3D}*{B3K}/1000", "E", 0, "E", None, "E", "n-d", "S10,S11",
  "MW not disclosed: ~50 MW is a LOW-confidence estimate (S10). Chip unverified. Reference only."),
 ("NB3", "Nebius - Meta #2 (Mar-2026), dedicated Vera Rubin, 5 yr", "Nebius", "Vera Rubin", "=12000/5", "D", "n-d", "n-d", "n-d", "n-d", "n-d",
  "n-d", "n-d", "n-d", "n-d", "n-d", "n-d", "n-d", "S10", "MW and GPUs not disclosed -> per-MW metrics and payback n-d. At $16.2M/IT MW (Nscale-Anthropic VR) it would be ~148 IT MW."),
 ("NQ-A", "Nebius Q2'26 deals - reading A: $22.5M per IT MW, Capacity GB300 compute capex", "Nebius", "GB300 (assumed)", 22.5, "R", 1, "IT", "E", f"={PUE}", "E",
  f"={GB3D}*{GB3K}/1000", "E", 0, "E", None, "E", 22, "S1,S2,S11,S14", "ACV = midpoint of '$20-25 million per megawatt' (R). Leased/colo capacity assumed (most late-2026 capacity is colo/build-to-suit). Reported 22 = '1 year 10 months'."),
 ("NQ-B", "Nebius Q2'26 deals - reading A': $22.5M per IT MW, IREN-disclosed GB300 capex ($29M/IT MW)", "Nebius", "GB300 (assumed)", 22.5, "R", 1, "IT", "E", f"={PUE}", "E",
  "=5800/200", "D", 0, "E", None, "E", 22, "S1,S3,S14", "Same as NQ-A but compute capex = IREN's actual Dell purchase price per IT MW (D)."),
 ("NQ-C", "Nebius Q2'26 deals - reading B: $22.5M per GROSS MW, Capacity GB300 compute capex", "Nebius", "GB300 (assumed)", 22.5, "R", 1, "gross", "E", f"={PUE}", "E",
  f"=L{{r}}*{GB3D}*{GB3K}/1000", "E", 0, "E", None, "E", 22, "S1,S11,S14", "Gross reading: ACV per IT MW = 22.5 x factor."),
]
tagcols = {}
for j, d in enumerate(D):
    r = H0 + 1 + j
    (i_, deal, prov, chip, acv, acv_t, mw, basis, mw_t, fac, fac_t, capx, capx_t, dc, dc_t, m, m_t, rep, s, note) = d
    vals = {1: i_, 2: deal, 3: prov, 4: chip, 5: acv, 6: acv_t, 7: mw, 8: basis, 9: mw_t, 10: fac, 11: fac_t,
            16: capx.format(r=r) if isinstance(capx, str) else capx, 17: capx_t, 19: dc, 20: dc_t,
            23: (OPXM(r) if m is None else m), 24: m_t, 28: rep, 30: s, 31: note}
    for c, v in vals.items():
        cell = pa.cell(r, c, v)
        if c in (5, 7, 10, 16, 19, 23) and not (isinstance(v, str) and v.startswith("=")) and v != "n-d": cell.font = BLUE
    if mw == "n-d":
        for c in (12, 13, 14, 15, 18, 21, 22, 25, 26, 27): pa.cell(r, c, "n-d")
        pa.cell(r, 29, "n-d"); continue
    pa.cell(r, 12, f'=IF(H{r}="gross",G{r}/J{r},G{r})'); pa.cell(r, 13, f'=IF(H{r}="gross",G{r},G{r}*J{r})')
    pa.cell(r, 14, f"=E{r}/L{r}"); pa.cell(r, 15, f"=E{r}/M{r}"); pa.cell(r, 18, f"=P{r}/L{r}")
    pa.cell(r, 21, f"=R{r}+S{r}"); pa.cell(r, 22, f"=U{r}*L{r}/M{r}")
    pa.cell(r, 25, f"=R{r}/(N{r}*W{r})*12"); pa.cell(r, 26, f"=U{r}/(N{r}*W{r})*12"); pa.cell(r, 27, f"=V{r}/(O{r}*W{r})*12")
    tg = [acv_t, mw_t, capx_t, dc_t, m_t]
    pa.cell(r, 29, "E" if "E" in tg else ("P" if "P" in tg else "D"))
    if i_.startswith("NQ"):
        for c in (25, 26, 27): pa.cell(r, c).fill = KEY
for r in range(H0 + 1, H0 + 1 + len(D)):
    for c in (14, 15, 18, 21, 22): pa.cell(r, c).number_format = "0.00"
    pa.cell(r, 23).number_format = "0.0%"
    for c in (25, 26, 27): pa.cell(r, c).number_format = "0.0"
    pa.cell(r, 31).alignment = WRAP; pa.cell(r, 2).alignment = WRAP
LASTD = H0 + len(D)
ROWID = {d[0]: H0 + 1 + j for j, d in enumerate(D)}
m0 = LASTD + 2
pa.cell(m0, 1, "Memo (not MW-based)").font = BOLD
pa.cell(m0 + 1, 1, "OR-hyp"); pa.cell(m0 + 1, 2, "Oracle analyst-day example: $60B / 6 yr AI contract, $6.4B cost/yr; compute 65% treated as straight-line depreciation over 6 yr -> capex = 0.65 x 6.4 x 6")
pa.cell(m0 + 1, 16, "=0.65*6.4*6"); pa.cell(m0 + 1, 17, "E"); pa.cell(m0 + 1, 5, 10); pa.cell(m0 + 1, 6, "R")
pa.cell(m0 + 1, 23, f"={ORM}"); pa.cell(m0 + 1, 24, "D"); pa.cell(m0 + 1, 26, f"=P{m0+1}/(E{m0+1}*W{m0+1})*12"); pa.cell(m0 + 1, 29, "E"); pa.cell(m0 + 1, 30, "S12")
pa.cell(m0 + 1, 26).number_format = "0.0"; pa.cell(m0 + 1, 23).number_format = "0.0%"; pa.cell(m0 + 1, 2).alignment = WRAP
pa.cell(m0 + 2, 1, "X-AB"); pa.cell(m0 + 2, 2, "Crusoe Abilene (DC developer; facility only): $15B / 1,200 MW = DC shell+power $M per MW (basis unlabeled); lease revenue to Oracle n-d -> payback n-d")
pa.cell(m0 + 2, 19, "=15000/1200"); pa.cell(m0 + 2, 20, "D"); pa.cell(m0 + 2, 26, "n-d"); pa.cell(m0 + 2, 30, "S13"); pa.cell(m0 + 2, 2).alignment = WRAP
pa.cell(m0 + 3, 1, "Peer band"); pa.cell(m0 + 3, 2, "Peer deal payback (total capex) - min / median / max over P1, N1, N1a, N2, O1 (deal-level rows; excludes CoreWeave fleet)")
band = ",".join(f"Z{ROWID[k]}" for k in ("P1", "N1", "N1a", "N2", "O1"))
pa.cell(m0 + 3, 25, f"=MIN({band})"); pa.cell(m0 + 3, 26, f"=MEDIAN({band})"); pa.cell(m0 + 3, 27, f"=MAX({band})")
pa.cell(m0 + 4, 2, "Peer deal payback (compute-only) - min / median / max over the same rows")
bandc = ",".join(f"Y{ROWID[k]}" for k in ("P1", "N1", "N1a", "N2", "O1"))
pa.cell(m0 + 4, 25, f"=MIN({bandc})"); pa.cell(m0 + 4, 26, f"=MEDIAN({bandc})"); pa.cell(m0 + 4, 27, f"=MAX({bandc})")
for rr in (m0 + 3, m0 + 4):
    for c in (25, 26, 27): pa.cell(rr, c).number_format = "0.0"; pa.cell(rr, 29, "E")
pa.freeze_panes = pa.cell(H0 + 1, 3)

# =========================== Basis_Sensitivity ===========================
bs = wb.create_sheet("Basis_Sensitivity", 2)
bs.column_dimensions["A"].width = 58
for c in range(2, 14): bs.column_dimensions[CL(c)].width = 13
bs["A1"] = "Does Nebius's ~22-month payback on $20-25M/MW reconcile only on a GROSS-MW reading?  Payback = capex per IT MW / (ACV per Nebius-MW x f - opex per IT MW) x 12, f = 1 (IT reading) or gross/IT factor (gross reading)."
bs["A1"].font = BOLD
bs["A3"] = "Inputs"; bs["A3"].font = BOLD
inp = [("ACV per MW - low / mid / high ($M/yr)", [20, 22.5, 25], "R", "S1/S2 '$20-25 million per megawatt'"),
       ("Nebius payback target (months)", [22], "R", "S1 '1 year 10 months' (fn: revenue-recognition basis, ex-prepayment, incl. related operating costs)"),
       ("Opex per IT MW-yr, leased capacity ($M)", [f"=Payback_ACV!{OPEX}"], "E", "S14"),
       ("Gross/IT factor sensitivity (PUE-type)", [1.2, 1.3, 1.4], "E", "S11 (GB300/VR base PUE 1.2; Hopper 1.3); Nscale Monarch 1.23, AIB 1.30"),
       ("Benchmark peer margin, IREN project EBITDA", [f"=Payback_ACV!{IRM}"], "R", "S4")]
for i, (lab, vs, t, s) in enumerate(inp, 4):
    bs.cell(i, 1, lab)
    for k, v in enumerate(vs): c = bs.cell(i, 2 + k, v); c.font = BLUE
    bs.cell(i, 6, t); bs.cell(i, 7, s)
ACVL, ACVM, ACVH, PBT, OPX, F12, F13, F14, MIR = "$B$4", "$C$4", "$D$4", "$B$5", "$B$6", "$B$7", "$C$7", "$D$7", "$B$8"

bs["A10"] = "Capex benchmarks ($M per IT MW, prepayments ignored)"; bs["A10"].font = BOLD
hdr_row = 10
for k, h in enumerate(["Value", "Tag", "Source"], 2): bs.cell(hdr_row, k, h).font = BOLD
BEN = [("b1  IREN actual: Dell GB300 GPUs + ancillaries $5.8B / 200 IT MW (compute only)", "=5800/200", "D", "S3/S4"),
       ("b2  IREN-implied: '>$20M/IT MW ~2-yr payback' at 85% margin = 24/12 x 20 x 85%", f"=24/12*20*{MIR}", "D", "S5/S4"),
       ("b3  Capacity GB300 compute LOW", GB300_LOW_COMPUTE, "E", "S11"),
       ("b4  Capacity GB300 compute BASE (= 491 x $81.9k)", f"=Payback_ACV!{GB3D}*Payback_ACV!{GB3K}/1000", "E", "S11"),
       ("b5  Modeling leased GB300 (PB_CapGB)", 41.3, "E", "S14"),
       ("b6  Capacity GB300 compute HIGH", GB300_HIGH_COMPUTE, "E", "S11"),
       ("b7  Capacity Vera Rubin compute BASE (= 320 x $94.4k)", f"=Payback_ACV!{VRD}*Payback_ACV!{VRK}/1000", "E", "S11 (Modeling uses 41 for Rubin: S14)"),
       ("b8  GB300 OWNED full stack base (compute + $12.5M shell)", GB300_OWNED_BASE, "E", "S11 - only relevant if Nebius owns the shell (opex then excludes rent)")]
for i, (lab, v, t, s) in enumerate(BEN, 11):
    bs.cell(i, 1, lab); c = bs.cell(i, 2, v); c.font = BLUE; c.number_format = "0.00"; bs.cell(i, 3, t); bs.cell(i, 4, s)
BROWS = list(range(11, 11 + len(BEN)))

r0 = 21
bs.cell(r0, 1, "Table A - implied payback (months) by capex benchmark x MW reading").font = BOLD
heads = [("IT @ $20", ACVL, None), ("IT @ $22.5", ACVM, None), ("IT @ $25", ACVH, None),
         ("Gross x1.2 @ $22.5", ACVM, F12), ("Gross x1.3 @ $22.5", ACVM, F13), ("Gross x1.4 @ $22.5", ACVM, F14),
         ("Gross x1.3 @ $20", ACVL, F13), ("Gross x1.3 @ $25", ACVH, F13)]
for k, (h, _, _) in enumerate(heads, 2): c = bs.cell(r0 + 1, k, h); c.font = BOLD; c.fill = HDR; c.alignment = WRAP
bs.cell(r0 + 1, 10, "Tag").font = BOLD
for j, br in enumerate(BROWS):
    r = r0 + 2 + j; bs.cell(r, 1, f"=A{br}")
    for k, (h, acv, f) in enumerate(heads, 2):
        fx = "1" if f is None else f
        c = bs.cell(r, k, f"=$B${br}/({acv}*{fx}-{OPX})*12"); c.number_format = "0.0"
    bs.cell(r, 10, "E")
TA = r0 + 2  # first row of table A
r1 = r0 + 2 + len(BROWS) + 1
bs.cell(r1, 1, "Table B - compute capex per IT MW REQUIRED to hit 22 months = 22/12 x (ACV x f - opex)").font = BOLD
for k, (h, f) in enumerate([("IT (f=1)", "1"), ("Gross x1.2", F12), ("Gross x1.3", F13), ("Gross x1.4", F14)], 2):
    c = bs.cell(r1 + 1, k, h); c.font = BOLD; c.fill = HDR
for j, (lab, acv) in enumerate([("ACV $20M/MW", ACVL), ("ACV $22.5M/MW", ACVM), ("ACV $25M/MW", ACVH)]):
    r = r1 + 2 + j; bs.cell(r, 1, lab)
    for k, f in enumerate(["1", F12, F13, F14], 2):
        c = bs.cell(r, k, f"={PBT}/12*({acv}*{f}-{OPX})"); c.number_format = "0.0"
    bs.cell(r, 7, "D/E")
TB = r1 + 2
r2 = r1 + 6
bs.cell(r2, 1, "Table C - margin-based variant (payback = capex / (ACV x f x margin)), ACV $22.5M").font = BOLD
ch = [("IT @ b1 IREN $29M", "$B$11", "1"), ("IT @ b4 Capacity $40.3M", "$B$14", "1"), ("Gross x1.3 @ b1", "$B$11", F13), ("Gross x1.3 @ b4", "$B$14", F13)]
for k, (h, _, _) in enumerate(ch, 2): c = bs.cell(r2 + 1, k, h); c.font = BOLD; c.fill = HDR; c.alignment = WRAP
MG = [("Nebius AI cloud adj. EBITDA margin Q2'26 (segment, after overhead) - D", f"=Payback_ACV!{NBM}"),
      ("CoreWeave adj. EBITDA margin Q2'26 (after DC rent) - D", f"=Payback_ACV!W{ROWID['C1']}"),
      ("Oracle AI-infra cash margin (analyst-day example) - D/E", f"=Payback_ACV!{ORM}"),
      ("Leased deal-level: 1 - opex/ACV at $22.5M IT - E", f"=1-{OPX}/{ACVM}"),
      ("IREN Microsoft project EBITDA margin (owned power) - R", f"=Payback_ACV!{IRM}")]
for j, (lab, v) in enumerate(MG):
    r = r2 + 2 + j; bs.cell(r, 1, lab); c = bs.cell(r, 6, v); c.number_format = "0.0%"
    for k, (h, cap_, f) in enumerate(ch, 2):
        c = bs.cell(r, k, f"={cap_}/({ACVM}*{f}*$F{r})*12"); c.number_format = "0.0"
bs.cell(r2 + 1, 6, "Margin").font = BOLD
r3 = r2 + 2 + len(MG) + 1
bs.cell(r3, 1, "Table D - opex sensitivity ($M per IT MW-yr), ACV $22.5M, payback months").font = BOLD
dh = [("IT @ b1", "$B$11", "1"), ("IT @ b2", "$B$12", "1"), ("IT @ b4", "$B$14", "1"), ("Gross x1.3 @ b4", "$B$14", F13), ("Gross x1.3 @ b5", "$B$15", F13)]
for k, (h, _, _) in enumerate(dh, 2): c = bs.cell(r3 + 1, k, h); c.font = BOLD; c.fill = HDR
for j, ox in enumerate([2, 3, 4.5, 6]):
    r = r3 + 2 + j; bs.cell(r, 1, f"Opex {ox} $M/IT MW-yr" + (" (Modeling base)" if ox == 4.5 else "")); bs.cell(r, 7, ox).font = BLUE
    for k, (h, cap_, f) in enumerate(dh, 2):
        c = bs.cell(r, k, f"={cap_}/({ACVM}*{f}-$G{r})*12"); c.number_format = "0.0"
bs.cell(r3 + 1, 7, "Opex").font = BOLD
r4 = r3 + 7
bs.cell(r4, 1, "Table E - ACV per IT MW-yr: Nebius readings vs peer IT-basis band").font = BOLD
for k, h in enumerate(["$M per IT MW-yr", "Tag", "Source"], 2): c = bs.cell(r4 + 1, k, h); c.font = BOLD; c.fill = HDR
EE = [("Nebius Q2 deals - IT reading (low / mid / high = 20 / 22.5 / 25)", f"={ACVM}", "R", "S1"),
      ("Nebius Q2 deals - gross reading x1.2 (mid)", f"={ACVM}*{F12}", "D", "S1 x factor"),
      ("Nebius Q2 deals - gross reading x1.3 (mid)", f"={ACVM}*{F13}", "D", ""),
      ("Nebius Q2 deals - gross reading x1.4 (mid)", f"={ACVM}*{F14}", "D", ""),
      ("Nebius Q2 deals - gross reading range top ($25M x 1.4)", f"={ACVH}*{F14}", "D", ""),
      ("IREN recent 3-yr contracts ('>$20m per MW (IT)')", 20, "R", "S5"),
      ("IREN active discussions ('~$25m per MW (IT)')", 25, "R", "S5"),
      ("Nscale - Anthropic VR (6 yr)", f"=Payback_ACV!N{ROWID['N2']}", "D", "S10"),
      ("Nebius - Microsoft (GPU-derived 203.7 IT MW)", "=17392.9/5/(100000/Payback_ACV!" + GB3D + ")", "D", "S10"),
      ("Nebius - Microsoft (300 MW read as gross, x1.3)", f"=Payback_ACV!N{ROWID['NB1']}", "E", "S10"),
      ("Oracle - OpenAI (4.5 GW read as gross /1.2)", f"=Payback_ACV!N{ROWID['O1']}", "E", "S19"),
      ("Nscale - Microsoft all SOWs", f"=Payback_ACV!N{ROWID['N1']}", "D", "S10"),
      ("IREN - Microsoft", f"=Payback_ACV!N{ROWID['P1']}", "D", "S3"),
      ("Memo: Nebius 'ACV per MW $12M 2026 base' (same slide axis as >$20M)", 12, "R", "S1 - ties to ~250 active MW at $3.0B ARR (Filings, S18) = active/IT-like axis"),
      ("Memo: Nebius '>$40M Q3 short-term capacity deals'", 40, "R", "S1")]
for j, (lab, v, t, s) in enumerate(EE):
    r = r4 + 2 + j; bs.cell(r, 1, lab); c = bs.cell(r, 2, v); c.number_format = "0.00"; bs.cell(r, 3, t); bs.cell(r, 4, s)
TE = r4 + 2
r5 = r4 + 2 + len(EE) + 1
bs.cell(r5, 1, "Table F - implied $ per GPU-hour at 100% utilization (ACV per IT MW x 1e6 / (GPUs per IT MW x 8,760))").font = BOLD
for k, h in enumerate(["GB300 IT", "GB300 gross x1.3", "Vera Rubin IT", "VR gross x1.3"], 2): c = bs.cell(r5 + 1, k, h); c.font = BOLD; c.fill = HDR
for j, (lab, acv) in enumerate([("ACV $20M/MW", ACVL), ("ACV $22.5M/MW", ACVM), ("ACV $25M/MW", ACVH)]):
    r = r5 + 2 + j; bs.cell(r, 1, lab)
    for k, (f, dens) in enumerate([("1", GB3D), (F13, GB3D), ("1", VRD), (F13, VRD)], 2):
        c = bs.cell(r, k, f"={acv}*{f}*1000000/(Payback_ACV!{dens}*8760)"); c.number_format = "0.00"
    bs.cell(r, 6, "D/E")
r6 = r5 + 6
bs.cell(r6, 1, "VERDICT METRICS (live)").font = BOLD
VM = [("IT reading @ IREN-disclosed capex b1, ACV $22.5M, opex $4.5M (months)", f"=C{TA}"),
      ("IT reading @ IREN-implied capex b2 (months)", f"=C{TA+1}"),
      ("IT reading @ Capacity GB300 base b4 (months)", f"=C{TA+3}"),
      ("IT reading @ Modeling leased GB300 b5 (months)", f"=C{TA+4}"),
      ("Gross x1.3 reading @ Capacity GB300 base b4 (months)", f"=F{TA+3}"),
      ("Gross x1.2 / x1.4 reading @ b4 (months)", f'=TEXT(E{TA+3},"0.0")&" / "&TEXT(G{TA+3},"0.0")'),
      ("Required compute capex per IT MW for 22 mo - IT reading, ACV $20 / $22.5 / $25", f'=TEXT(B{TB},"0.0")&" / "&TEXT(B{TB+1},"0.0")&" / "&TEXT(B{TB+2},"0.0")'),
      ("Required compute capex per IT MW for 22 mo - gross x1.3 reading, ACV $22.5", f"=D{TB+1}"),
      ("Peer reported analogue (explicit IT basis): IREN >$20M/IT MW -> ~24 months", 24),
      ("IT reading reconciles within +/-3 months at IREN-disclosed capex (b1 or b2)?", f'=IF(AND(ABS(C{TA}-{PBT})<=3,ABS(C{TA+1}-{PBT})<=3),"YES","NO")'),
      ("IT reading reconciles within +/-3 months at Capacity/Modeling capex (b4/b5)?", f'=IF(MIN(ABS(C{TA+3}-{PBT}),ABS(C{TA+4}-{PBT}))<=3,"YES","NO")'),
      ("Gross reading (1.2-1.4) reconciles within +/-3 months at Capacity base (b4)?", f'=IF(AND(ABS(E{TA+3}-{PBT})<=3,ABS(G{TA+3}-{PBT})<=3),"YES (whole band)",IF(OR(ABS(E{TA+3}-{PBT})<=3,ABS(G{TA+3}-{PBT})<=3),"PARTLY","NO"))')]
for j, (lab, v) in enumerate(VM):
    r = r6 + 1 + j; bs.cell(r, 1, lab).alignment = WRAP; c = bs.cell(r, 2, v); c.number_format = "0.0"; c.fill = KEY
VROW = r6 + 1
bs.cell(r6 + 1 + len(VM) + 1, 1, (
    "Verdict: NO - the 22 months does not require a gross reading. On an IT reading it reconciles with the only DISCLOSED peer evidence "
    "(IREN's actual GB300 capex $29M/IT MW, and IREN's own '>$20M per MW (IT) ~2-yr payback' on the same capex/(revenue - direct cost) "
    "definition). It needs a gross reading only if Nebius's GPU+network capex is ~$40M+/IT MW (Capacity/Modeling third-party estimate) or "
    "if the margin is the ~50% segment EBITDA margin. Confidence the IT reading is right: medium (~60-65%); basis is undisclosed (S1, S18)."))
bs.cell(r6 + 1 + len(VM) + 1, 1).alignment = WRAP
bs.row_dimensions[r6 + 1 + len(VM) + 1].height = 90

# =========================== Dilution_Data ===========================
dd = wb.create_sheet("Dilution_Data", 3)
hdr(dd, 3, ["Company", "Quarter", "Qtrs since series start", "Open MW", "Close MW", "MW tag", "Avg MW", "New MW", "Ramp ratio (new/avg)",
            "ln(avg MW)", "Revenue $M (qtr)", "Rev tag", "Rev per avg MW ($M/MW-yr)", "In fit?", "Source / caveat"],
    [11, 8, 8, 9, 9, 8, 9, 9, 10, 9, 11, 7, 12, 7, 80])
dd["A1"] = "Revenue per average active MW vs capacity. Avg MW = (open+close)/2; ramp ratio = MW added in quarter / average MW; rev/MW annualized (x4)."; dd["A1"].font = BOLD
rows = []
for i, q in enumerate(cr_q):
    rows.append(("CoreWeave", q, i + 1, cr_mw[i], cr_mw[i + 1], "E" if i < 4 else "R", cr_rev[i], "R", "Y",
                 "S8 crwv_kpis Inputs I09E/I01. 2024 MW straight-line interpolated (E); 'more than' MW wording adds +/-10-50 MW noise."))
for i, q in enumerate(ir_q):
    rows.append(("IREN", q, i + 1, ir_mw[i], ir_mw[i + 1], "E", ir_rev[i], "R", "N",
                 "iren_kpis Inputs I09E (MW = GPUs/713, E) and I01 AI Cloud revenue. Tiny MW (1-3) - shown for shape; NOT used in fit."))
rows.append(("IREN", "Q2'26", 9, "n-d", ir_mw_q226, "R", ir_rev_q226, "R", "N",
             "40 MW is the only reported AI MW (period-end, basis unstated). Rev / PERIOD-END MW (avg n-d: Mar-26 MW n-d). Not in fit."))
for i, q in enumerate(nb_q):
    rows.append(("Nebius", q, i + 1, nb_mw[i], nb_mw[i + 1], "E", nb_rev[i], "R", "N",
                 "S15 nebius_kpi_history KPI rows 28 (Capacity energized MW incl. plug, mixed IT/gross, E) and 6 (AI cloud revenue, R). Back-test only."))
DD0 = 4
for j, rw in enumerate(rows):
    r = DD0 + j
    co, q, seq, o, c_, mt, rev, rt, fit, s = rw
    for k, v in zip((1, 2, 3, 4, 5, 6, 11, 12, 14, 15), (co, q, seq, o, c_, mt, rev, rt, fit, s)):
        cell = dd.cell(r, k, v)
        if k in (4, 5, 11) and v != "n-d": cell.font = BLUE
    if o == "n-d":
        for k in (7, 8, 9, 10): dd.cell(r, k, "n-d")
        dd.cell(r, 13, f"=K{r}*4/E{r}")
    else:
        dd.cell(r, 7, f"=(D{r}+E{r})/2"); dd.cell(r, 8, f"=E{r}-D{r}"); dd.cell(r, 9, f"=H{r}/G{r}")
        dd.cell(r, 10, f"=LN(G{r})"); dd.cell(r, 13, f"=K{r}*4/G{r}")
    for k in (7, 8): dd.cell(r, k).number_format = "#,##0.0"
    dd.cell(r, 9).number_format = "0.000"; dd.cell(r, 10).number_format = "0.00"; dd.cell(r, 13).number_format = "0.00"
CR = (DD0, DD0 + 9); CR25 = (DD0 + 4, DD0 + 9)
IRR_ = (DD0 + 10, DD0 + 15); NBR = (DD0 + 17, DD0 + 24)
dd.freeze_panes = "C4"

# =========================== Dilution_Fit ===========================
df = wb.create_sheet("Dilution_Fit", 4)
df.column_dimensions["A"].width = 62
for c in range(2, 12): df.column_dimensions[CL(c)].width = 13
df["A1"] = "Dilution fits (live LINEST-equivalents: SLOPE / INTERCEPT / RSQ / STEYX). Primary: rev/avg MW = a + b x ramp ratio  ->  dilution factor = 1 - k x ramp, k = -b/a."; df["A1"].font = BOLD
hdr(df, 3, ["Fit", "Intercept a", "Slope b", "R^2", "SE of slope", "k = -b/a", "k low (b+SE)", "k high (b-SE)", "n", "Tag"])
def rng(col, a, b): return f"Dilution_Data!${col}${a}:${col}${b}"
FITS = [("F1 CoreWeave Q1'24-Q2'26: rev/MW vs ramp ratio (PRIMARY)", "I", CR),
        ("F2 CoreWeave Q1'25-Q2'26 (reported MW only): rev/MW vs ramp ratio", "I", CR25),
        ("F3 CoreWeave Q1'24-Q2'26: rev/MW vs ln(avg MW)", "J", CR),
        ("F4 CoreWeave Q1'25-Q2'26: rev/MW vs ln(avg MW) (scale effect; confounded with ramp)", "J", CR25),
        ("F5 Nebius Q1'25-Q2'26: rev/MW vs ramp ratio (info: confounded by Hopper->Blackwell repricing)", "I", (NBR[0] + 2, NBR[1])),
        ("F6 IREN Q2'24-Q3'25 (MW estimated, 1-3 MW): rev/MW vs ramp ratio (info)", "I", IRR_)]
for j, (lab, xc, (a, b)) in enumerate(FITS):
    r = 4 + j; y = rng("M", a, b); x = rng(xc, a, b)
    df.cell(r, 1, lab); df.cell(r, 2, f"=INTERCEPT({y},{x})"); df.cell(r, 3, f"=SLOPE({y},{x})"); df.cell(r, 4, f"=RSQ({y},{x})")
    df.cell(r, 5, f"=STEYX({y},{x})/SQRT(DEVSQ({x}))"); df.cell(r, 9, f"=COUNT({y})"); df.cell(r, 10, "D/E")
    if xc == "I":
        df.cell(r, 6, f"=-C{r}/B{r}"); df.cell(r, 7, f"=-(C{r}+E{r})/B{r}"); df.cell(r, 8, f"=-(C{r}-E{r})/B{r}")
    else:
        for k in (6, 7, 8): df.cell(r, k, "n/a")
    for k in range(2, 9): df.cell(r, k).number_format = "0.000"
df["A11"] = "Selected dilution coefficients (k) for Nebius_Path"; df["A11"].font = BOLD
df["A12"] = "k LOW dilution (F1 slope + 1 SE)"; df["B12"] = "=G4"
df["A13"] = "k BASE (F1)"; df["B13"] = "=F4"
df["A14"] = "k HIGH dilution (F1 slope - 1 SE)"; df["B14"] = "=H4"
df["A15"] = "Cross-check: k from F2 (2025-26 reported only)"; df["B15"] = "=F5"
for r in range(12, 16): df.cell(r, 2).number_format = "0.000"; df.cell(r, 3, "E")
df["A17"] = "Back-test: CRWV k applied to IREN and Nebius (implied 'undiluted' rev/MW = actual / (1 - k x ramp))"; df["A17"].font = BOLD
hdr(df, 18, ["Company / quarter", "Ramp ratio", "Actual rev/MW", "Factor (1-k*ramp)", "Undiluted rev/MW", "F1 fitted (CRWV only)", "Note"])
bt = list(range(CR[0], CR[1] + 1)) + list(range(IRR_[0], IRR_[1] + 1)) + list(range(NBR[0], NBR[1] + 1))
for j, src_r in enumerate(bt):
    r = 19 + j
    df.cell(r, 1, f'=Dilution_Data!A{src_r}&" "&Dilution_Data!B{src_r}'); df.cell(r, 2, f"=Dilution_Data!I{src_r}")
    df.cell(r, 3, f"=Dilution_Data!M{src_r}"); df.cell(r, 4, f"=1-$B$13*B{r}"); df.cell(r, 5, f"=C{r}/D{r}")
    if src_r <= CR[1]: df.cell(r, 6, f"=$B$4+$C$4*B{r}")
    for k in range(2, 7): df.cell(r, k).number_format = "0.00"
df.cell(19 + len(bt) + 1, 1, ("Read: CRWV rev/MW falls ~0.57 $M/MW-yr per +0.1 of ramp ratio (R^2 ~0.79, 10 qtrs). IREN's one big ramp quarter (Q4'24, ramp ~0.8) "
                              "sits on the same line (actual 5.6 vs ~9-11 in steady quarters). ln(MW) has no explanatory power on the full CRWV sample (R^2 ~0.07); "
                              "on 2025-26 alone it fits (R^2 ~0.79) but that window is also the high-ramp window - treat scale dilution as unproven. "
                              "Nebius Q4'25 (ramp 0.92) diluted LESS than CRWV k predicts (undiluted ~14 vs ~9-13 around it) - Microsoft Vineland tranches billed quickly."))
df.cell(19 + len(bt) + 1, 1).alignment = WRAP; df.row_dimensions[19 + len(bt) + 1].height = 110

# =========================== Nebius_Path ===========================
npth = wb.create_sheet("Nebius_Path", 5)
npth.column_dimensions["A"].width = 62
for c in range(2, 13): npth.column_dimensions[CL(c)].width = 11
npth["A1"] = ("Nebius quarterly revenue per average energized IT MW, Q3'26-Q4'27.  rev/MW_q = SS_q x (1 - k x ramp_q);  SS_q = (MW_Q2'26 x L0 + (avgMW_q - MW_Q2'26) x P_new) / avgMW_q;  "
              "L0 = Q2'26 AI cloud revenue x4 / avg(Q1'26, Q2'26) MW of the same case. MW = Capacity 'Energized MW low/base/high - IT-normalized' (values copied, not edited). "
              "Bands: LOW = P_new low + k high; BASE = P_new base + k base; HIGH = P_new high + k low.")
npth["A1"].font = BOLD; npth["A1"].alignment = WRAP; npth.row_dimensions[1].height = 60
pin = [("Q2'26 Nebius AI cloud revenue ($M)", f"=Dilution_Data!K{NBR[1]}", "R", "S15"),
       ("P_new LOW: incremental-MW price = '$12M 2026 base' ACV/MW", 12, "R", "S1"),
       ("P_new BASE: = Nebius-Microsoft $17.1M per IT MW (~midpoint of $12M base and $22.5M Q2 deals)", "=17392.9/5/(100000/Payback_ACV!" + GB3D + ")", "D", "S10/S11"),
       ("P_new HIGH: = midpoint of Q2-deal '$20-25M per MW' (IT reading)", 22.5, "R", "S1"),
       ("k LOW dilution", "=Dilution_Fit!B12", "E", "Dilution_Fit"),
       ("k BASE", "=Dilution_Fit!B13", "E", "Dilution_Fit"),
       ("k HIGH dilution", "=Dilution_Fit!B14", "E", "Dilution_Fit")]
for i, (lab, v, t, s_) in enumerate(pin, 3):
    npth.cell(i, 1, lab); c = npth.cell(i, 2, v); c.font = BLUE; c.number_format = "0.00"; npth.cell(i, 3, t); npth.cell(i, 4, s_)
REVQ, PL, PB, PH, KL, KB, KH = [f"$B${i}" for i in range(3, 10)]
QS = list(QCOLS.keys())  # B = Q1'26, C = Q2'26, D..I = Q3'26..Q4'27
npth.cell(11, 1, "Quarter").font = BOLD
for k, q in enumerate(QS, 2): c = npth.cell(11, k, q); c.font = BOLD; c.fill = HDR
npth.cell(11, 10, "Tag").font = BOLD; npth.cell(11, 11, "FY27 sum / source").font = BOLD
r = 12
PATHROWS = {}
for case in ("Low", "Base", "High"):
    npth.cell(r, 1, f"MW CASE: {case.upper()} (Capacity Quarterly MW row {MWCASE[case]} 'Energized MW {case.lower()} - IT-normalized')").font = BOLD
    r += 1; mwr = r
    npth.cell(r, 1, "End-of-quarter energized IT MW")
    for k, v in enumerate(mw_path[case], 2): c = npth.cell(r, k, v); c.font = BLUE; c.number_format = "#,##0.0"
    npth.cell(r, 10, "E"); npth.cell(r, 11, f"S16 row {MWCASE[case]}, cols M-T")
    r += 1; avr = r; npth.cell(r, 1, "Average MW")
    for k in range(3, 10): c = npth.cell(r, k, f"=({CL(k-1)}{mwr}+{CL(k)}{mwr})/2"); c.number_format = "#,##0.0"
    r += 1; nwr = r; npth.cell(r, 1, "New MW in quarter")
    for k in range(3, 10): c = npth.cell(r, k, f"={CL(k)}{mwr}-{CL(k-1)}{mwr}"); c.number_format = "#,##0.0"
    r += 1; rpr = r; npth.cell(r, 1, "Ramp ratio (new / avg)")
    for k in range(3, 10): c = npth.cell(r, k, f"={CL(k)}{nwr}/{CL(k)}{avr}"); c.number_format = "0.000"
    r += 1; l0r = r; npth.cell(r, 1, "L0 = Q2'26 rev x4 / avg Q2'26 MW (this case)"); c = npth.cell(r, 3, f"={REVQ}*4/C{avr}"); c.number_format = "0.00"; npth.cell(r, 10, "D/E")
    out = {}
    for band, P, K in (("LOW", PL, KH), ("BASE", PB, KB), ("HIGH", PH, KL)):
        r += 1; npth.cell(r, 1, f"Rev per avg IT MW - {band} ($M/MW-yr)")
        npth.cell(r, 3, f"=$C${l0r}")
        for k in range(4, 10):
            a_ = f"{CL(k)}{avr}"
            npth.cell(r, k, f"=(MIN({a_},$C${mwr})*$C${l0r}+MAX({a_}-$C${mwr},0)*{P})/{a_}*(1-{K}*{CL(k)}{rpr})")
        for k in range(3, 10): npth.cell(r, k).number_format = "0.00"
        npth.cell(r, 10, "E")
        if band == "BASE":
            for k in range(3, 10): npth.cell(r, k).fill = KEY
        out[band] = r
    r += 1; npth.cell(r, 1, "Implied AI cloud revenue - BASE band ($M, qtr) [memo]")
    npth.cell(r, 3, f"={REVQ}")
    for k in range(4, 10): c = npth.cell(r, k, f"={CL(k)}{out['BASE']}*{CL(k)}{avr}/4"); c.number_format = "#,##0"
    npth.cell(r, 11, f"=SUM(F{r}:I{r})"); npth.cell(r, 11).number_format = "#,##0"; npth.cell(r, 10, "E")
    out["REV"] = r
    r += 1; npth.cell(r, 1, "Implied FY27 revenue LOW band / HIGH band ($M) [memo, cols K / L]")
    npth.cell(r, 11, "=" + "+".join(f"{CL(k)}{out['LOW']}*{CL(k)}{avr}/4" for k in range(6, 10)))
    npth.cell(r, 12, "=" + "+".join(f"{CL(k)}{out['HIGH']}*{CL(k)}{avr}/4" for k in range(6, 10)))
    for k in (11, 12): npth.cell(r, k).number_format = "#,##0"
    npth.cell(r, 10, "E"); out["REVLH"] = r
    PATHROWS[case] = dict(mw=mwr, avg=avr, ramp=rpr, l0=l0r, **out)
    r += 2
npth.cell(r, 1, "Memo - YE26 / YE27 MW from other team sources (not used in path)").font = BOLD; r += 1
for case in ("Low", "Base", "High"):
    npth.cell(r, 1, f"Capacity identified-site energized {case.lower()} (mixed IT/gross basis, row {MWMIX[case]})")
    npth.cell(r, 2, mw_mix[case][0]); npth.cell(r, 3, mw_mix[case][1]); npth.cell(r, 10, "E"); r += 1
npth.cell(r, 1, "Model v0.7 Block 2 energized (build_model.py text; pre-dates Capacity 22:33 ET update): YE26 / YE27 bear-base-bull")
npth.cell(r, 2, "185/305/424"); npth.cell(r, 3, "764/1,069/1,374"); npth.cell(r, 10, "E"); r += 1
npth.cell(r, 1, f"Mgmt-guided CONNECTED MW (Capacity row {mgmt_guided['row']} plug): YE26 / YE27")
npth.cell(r, 2, mgmt_guided["YE26"]); npth.cell(r, 3, mgmt_guided["YE27"]); npth.cell(r, 10, "R/E")
npth.cell(r, 4, "Connected != energized; as a revenue denominator the Q4'26 ramp ratio would be >1 and k would cut rev/MW by >50% - not used.")
r += 2
NOTES = ["Assumptions / caveats:",
         "1) k fitted on CoreWeave (10 quarters, Dilution_Fit F1); IREN's Q4'24 ramp quarter sits on the same line; Nebius Q4'25 diluted less than k predicts (Microsoft tranches billed fast) - so k may overstate Nebius dilution.",
         "2) SS_q mixes the legacy book at L0 with incremental MW at P_new. CRWV's k also embeds some mix effect -> possible double count; the LOW band (k high + $12M) brackets it.",
         "3) P_new uses Nebius's ACV/MW disclosures on the IT reading. Under a GROSS reading P_new per IT MW would be x1.2-1.4 higher (Basis_Sensitivity Table E).",
         "4) L0 differs by case because Capacity's Q1/Q2'26 MW differ by case (IT-normalized 171/182.5 base). Revenue memo rows are sanity checks only (Modeling owns revenue)."]
for t in NOTES:
    npth.cell(r, 1, t).alignment = WRAP; npth.merge_cells(start_row=r, start_column=1, end_row=r, end_column=11); npth.row_dimensions[r].height = 30; r += 1

# =========================== README ===========================
rd = wb.create_sheet("README", 0)
rd.column_dimensions["A"].width = 150
TXT = [
 "payback_dilution.xlsx - Nebius Peer KPIs bot - built 2026-10-04 (ET) by peers/build_payback_dilution.py. Recalculated in LibreOffice: 0 error cells.",
 "",
 "QUESTION 1: Does Nebius's ~22-month payback on $20-25M/MW reconcile only on a GROSS-MW reading? (moves Modeling PT $46 -> $147)",
 "Exact Nebius wording (S1, Q2'26 letter Aug 12 2026): 'ACV per MW is stepping up(1): $12M 2026 base; >$20M Q2 deals; >$40M Q3 short-term capacity deals'; 'Payback period(2)(3) 1 year 10 months for deals closed in Q2'.",
 "  fn(1) 'ACV - Annual contract value.'  fn(2) 'Methodology is on a revenue recognition basis, excluding any prepayment.'  fn(3) 'Estimated payback period for new deals signed in Q2, based on forecast costs and contracted future capacity, including capacity not yet built.'",
 "  Body: 'yield of $20-25 million per megawatt'; 'the expected payback period for the associated capex and related operating costs for Q2 deals is 1 year and 10 months, down from our two-to-three year payback period previously.' Call (S2): '<two-year payback'. NO IT vs gross basis anywhere.",
 "Method: Payback = compute capex per IT MW / (ACV per Nebius-MW x f - opex per IT MW) x 12, f = 1 (IT) or 1.2/1.3/1.4 (gross). Prepayments excluded (fn 2). Leased capacity: shell is rent inside opex.",
 "Key point: within any one deal payback is basis-invariant (Payback_ACV cols Z = AA). The basis only matters because Nebius's ACV/MW is being compared with per-IT-MW capex benchmarks.",
 "VERDICT: NO - it does not reconcile ONLY on gross. IT reading ($22.5M per IT MW, leased, opex $4.5M/IT MW-yr) = 19.3 months at IREN's DISCLOSED GB300 capex ($29M/IT MW, Dell PO) and 22.7 months at the ~$34M/IT MW implied by IREN's own '>$20M per MW (IT), ~2 year payback' (same capex/(revenue - direct costs) definition, explicit IT basis).",
 "  The gross reading is REQUIRED only if Nebius's compute capex is ~$40M+/IT MW (Capacity base $40.3M / Modeling $41.3M: IT 26.8-27.5 mo vs gross x1.2/1.3/1.4 = 21.5/19.5/17.9 mo), or if the ~50% segment EBITDA margin is used. 22 mo on IT needs <= $33.0M/IT MW (range $28.4-37.6M at ACV $20-25M); on gross x1.3 <= $45.4M. Confidence IT reading is right: MEDIUM (~60%).",
 "  Sensitivities: compute capex (+$1M/IT MW = +0.67 mo on IT reading); opex (+$1M/IT MW-yr = +1.2-1.6 mo); ACV $20 vs $25 (IT @ $29M: 22.5 vs 17.0 mo); gross factor 1.2 vs 1.4 (@ $40.3M: 21.5 vs 17.9 mo). The capex number (third-party estimate vs IREN's disclosed purchase price) is the single swing factor.",
 "  Counter-evidence (kept): Capacity finds YE25 in-service servers+network at only $18.4M per active MW (S18) - Nebius's own MW may run close to gross, or it buys cheaper; Filings (21:39 ET) leans gross at $41M capex/IT MW.",
 "",
 "QUESTION 2: Dilution curve. Primary fit (Dilution_Fit F1): CoreWeave rev/avg MW = 10.88 - 5.68 x ramp ratio (R^2 0.79, n=10) -> factor 1 - k x ramp, k = 0.52 (0.43-0.62 at +/-1 SE). ln(MW) fit fails on full sample (R^2 0.07).",
 "  Nebius_Path applies k to Capacity's energized MW low/base/high with a legacy-book + new-price mix (P_new $12 / $17.1 / $22.5). Output: quarterly rev/avg MW Q3'26-Q4'27 by MW case and LOW/BASE/HIGH band.",
 "",
 "TABS: Payback_ACV (deal table: ACV and capex per IT and gross MW, margins, payback) | Basis_Sensitivity (IT vs gross test, grids, verdict metrics) | Dilution_Data | Dilution_Fit | Nebius_Path | Sources.",
 "TAGS: R reported; P press/secondary; D derived from R/P; E estimated; n-d not disclosed. Blue font = hard-coded input (sourced). All other numbers are live formulas.",
 "Charts: peers/charts/payback_by_basis.png, dilution_curve.png, nebius_rev_mw_path.png.",
 "Not edited: capacity/, model/, filings/, parts/ (values copied with cell citations).",
]
for i, t in enumerate(TXT, 1):
    c = rd.cell(i, 1, t); c.alignment = WRAP
    if t.startswith(("QUESTION", "VERDICT", "payback_dilution")): c.font = BOLD
wb.move_sheet("README", -wb.index(wb["README"]))
order = ["README", "Payback_ACV", "Basis_Sensitivity", "Dilution_Data", "Dilution_Fit", "Nebius_Path", "Sources"]
wb._sheets = [wb[n] for n in order]

if "--charts" not in sys.argv:
    wb.save(OUT); print("saved", OUT, {k: v for k, v in ROWID.items()}, "PATHROWS", PATHROWS, "VROW", VROW, "TA", TA)

# =========================== Charts (from recalculated workbook values) ===========================
if "--charts" in sys.argv:
    import matplotlib; matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np
    v = openpyxl.load_workbook(OUT, data_only=True)
    CH = f"{ROOT}/peers/charts"; os.makedirs(CH, exist_ok=True)
    # 1) payback by basis
    bsv = v["Basis_Sensitivity"]; pav = v["Payback_ACV"]
    labels, vals, colors = [], [], []
    for lab, cell, col in [("Nebius IT reading\n@ IREN actual $29M", "C23", "#2e7d32"), ("Nebius IT reading\n@ IREN-implied $34M", "C24", "#2e7d32"),
                           ("Nebius IT reading\n@ Capacity $40.3M", "C26", "#c62828"), ("Nebius IT reading\n@ Modeling $41.3M", "C27", "#c62828"),
                           ("Nebius gross x1.2\n@ $40.3M", "E26", "#1565c0"), ("Nebius gross x1.3\n@ $40.3M", "F26", "#1565c0"), ("Nebius gross x1.4\n@ $40.3M", "G26", "#1565c0")]:
        labels.append(lab); vals.append(bsv[cell].value); colors.append(col)
    for rid, lab in [(24, "IREN 3-yr >$20M/IT MW\n(model; reported ~24)"), (28, "Nscale-Anthropic VR\n(total capex, E)"), (29, "Oracle-OpenAI\n(compute, E)"),
                     (23, "IREN-Microsoft\n(total capex)"), (26, "Nscale-Microsoft\n(total capex, E)"), (30, "Nebius-Microsoft\n(compute, E)"), (25, "CoreWeave fleet\nQ2'26")]:
        labels.append(lab); vals.append(pav.cell(rid, 26).value); colors.append("#9e9e9e")
    fig, ax = plt.subplots(figsize=(13, 6.2))
    x = np.arange(len(vals)); ax.bar(x, vals, color=colors)
    for xi, yv in zip(x, vals): ax.text(xi, yv + 0.8, f"{yv:.1f}", ha="center", fontsize=8)
    ax.axhline(22, color="k", ls="--", lw=1); ax.axhspan(19, 25, color="#fff59d", alpha=0.35, zorder=0)
    ax.text(len(vals) - 0.6, 64, "Dashed line / yellow band:\nNebius disclosed 22 mo (+/-3)", ha="right", fontsize=9)
    ax.set_xticks(x); ax.set_xticklabels(labels, rotation=45, ha="right", fontsize=8)
    ax.set_ylabel("Payback, months (prepayments excluded)"); ax.set_ylim(0, max(vals) * 1.12)
    ax.set_title("Payback test: Nebius $22.5M/MW under IT vs gross readings vs peer deals")
    from matplotlib.patches import Patch
    ax.legend(handles=[Patch(color="#2e7d32", label="IT reading, IREN-disclosed capex"), Patch(color="#c62828", label="IT reading, third-party capex est."),
                       Patch(color="#1565c0", label="Gross reading (factor 1.2-1.4)"), Patch(color="#9e9e9e", label="Peer deals (stated basis)")],
              loc="upper left", bbox_to_anchor=(1.01, 1), fontsize=8, frameon=False)
    fig.text(0.01, -0.16, "Source: payback_dilution.xlsx Basis_Sensitivity Table A, Payback_ACV col Z. Opex $4.5M/IT MW-yr (Modeling). Tags mostly E.", fontsize=7)
    fig.savefig(f"{CH}/payback_by_basis.png", dpi=150, bbox_inches="tight"); plt.close(fig)
    # 2) dilution curve
    dd = v["Dilution_Data"]; fit = v["Dilution_Fit"]
    a_, b_ = fit["B4"].value, fit["C4"].value
    fig, axs = plt.subplots(1, 2, figsize=(14, 5.6))
    sty = {"CoreWeave": ("#1565c0", "o"), "IREN": ("#ef6c00", "s"), "Nebius": ("#2e7d32", "^")}
    for r in range(4, dd.max_row + 1):
        co = dd.cell(r, 1).value; ramp = dd.cell(r, 9).value; y = dd.cell(r, 13).value; avg = dd.cell(r, 7).value
        if co not in sty or not isinstance(y, (int, float)): continue
        c, m = sty[co]
        if isinstance(ramp, (int, float)):
            axs[0].scatter(ramp, y, color=c, marker=m, s=40, label=co)
            axs[0].annotate(dd.cell(r, 2).value, (ramp, y), fontsize=6, xytext=(3, 3), textcoords="offset points")
            axs[1].scatter(avg, y, color=c, marker=m, s=40, label=co)
            axs[1].annotate(dd.cell(r, 2).value, (avg, y), fontsize=6, xytext=(3, 3), textcoords="offset points")
    xx = np.linspace(0, 1, 50); axs[0].plot(xx, a_ + b_ * xx, color="#1565c0", lw=1.5, label=f"Fit (CRWV): {a_:.2f} {b_:+.2f} x ramp, R²={fit['D4'].value:.2f}")
    axs[0].set_xlabel("Ramp ratio = MW added in quarter / average MW"); axs[0].set_ylabel("Revenue per average MW ($M/MW-yr)")
    axs[0].set_title("Dilution vs ramp ratio (primary fit)")
    axs[1].set_xscale("log"); axs[1].set_xlabel("Average active MW (log)"); axs[1].set_title(f"Rev/MW vs scale (CRWV ln-fit R²={fit['D6'].value:.2f} full sample)")
    for ax in axs:
        h, l = ax.get_legend_handles_labels(); d = dict(zip(l, h)); ax.grid(alpha=0.3)
    h, l = axs[0].get_legend_handles_labels(); d = dict(zip(l, h))
    fig.legend(d.values(), d.keys(), loc="upper center", bbox_to_anchor=(0.5, -0.02), ncol=4, fontsize=8, frameon=False)
    fig.text(0.01, -0.07, "Sources: crwv_kpis / iren_kpis Inputs (IREN MW estimated from GPUs, E; CRWV 2024 MW interpolated, E); Nebius: filings/nebius_kpi_history.xlsx rows 6/28/58.", fontsize=7)
    fig.savefig(f"{CH}/dilution_curve.png", dpi=150, bbox_inches="tight"); plt.close(fig)
    # 3) Nebius path
    npv = v["Nebius_Path"]
    q = [npv.cell(11, k).value for k in range(3, 10)]
    fig, ax = plt.subplots(figsize=(11, 5.6))
    cc = {"Low": "#ef6c00", "Base": "#2e7d32", "High": "#1565c0"}
    for case, rows_ in PATHROWS.items():
        lo = [npv.cell(rows_["LOW"], k).value for k in range(3, 10)]; ba = [npv.cell(rows_["BASE"], k).value for k in range(3, 10)]
        hi = [npv.cell(rows_["HIGH"], k).value for k in range(3, 10)]
        ls = "-" if case == "Base" else "--"
        ax.plot(q, ba, ls, color=cc[case], lw=2 if case == "Base" else 1.2, marker="o", label=f"{case} MW case - base band")
        if case == "Base": ax.fill_between(q, lo, hi, color=cc[case], alpha=0.15, label="Base MW case - LOW/HIGH band")
    ax.axhline(22.5, color="grey", ls=":", lw=1); ax.text(0, 22.8, "Q2-deal ACV $22.5M/MW (IT reading)", fontsize=8, color="grey")
    ax.axhline(12, color="grey", ls=":", lw=1); ax.text(6, 11.2, "'$12M 2026 base' ACV/MW", fontsize=8, color="grey", ha="right")
    ax.set_ylabel("Revenue per average energized IT MW ($M/MW-yr)"); ax.set_ylim(0, 26); ax.grid(alpha=0.3)
    ax.set_title("Nebius rev/avg MW path to YE27 with peer-based (CoreWeave) dilution")
    ax.legend(loc="upper left", bbox_to_anchor=(1.01, 1), fontsize=8, frameon=False)
    fig.text(0.01, -0.04, "MW: capacity/nebius_capacity.xlsx Quarterly MW, IT-normalized energized low/base/high. k from Dilution_Fit F1. Estimated (E).", fontsize=7)
    fig.savefig(f"{CH}/nebius_rev_mw_path.png", dpi=150, bbox_inches="tight"); plt.close(fig)
    print("charts written")
