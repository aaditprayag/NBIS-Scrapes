# Standalone one-page summary table (same content as memo page 3) -> nebius_vs_peers_summary_table.xlsx
import sys; sys.path.insert(0, "/workspace/nebius/peers/memo")
from common_data import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
OUT = "/workspace/nebius/peers/memo/nebius_vs_peers_summary_table.xlsx"
wb = Workbook(); ws = wb.active; ws.title = "Summary"
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E8C")
FL = {"Better": "DCEFE0", "Worse": "F8D9D5", "In line": "EDEDED"}
thin = Side(style="thin", color="BFBFBF"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
ws["A1"] = "Nebius vs peers - one-page unit-economics summary (Thu Oct 8, 2026). Source: peers/comparison.xlsx (Comparison values/tags; Ranking flags and peer medians), rebuilt ~10:00 ET on model v0.9e locks."
ws["A1"].font = Font(bold=True, size=12, color="1F4E8C")
ws["A2"] = "Values as in the workbook (unrounded); tag column per company: R reported, P press, D derived, E estimated; n/d = not disclosed / not computable on the same definition. Flag = Nebius vs comparable-peer median (>10% gap)."
hdr = ["Group", "KPI ID", "KPI", "Unit"]
for p in PEERS: hdr += [p, f"{p} tag", f"{p} as of"]
hdr += ["Peer median", "Flag (Ranking)", "Ranking #", "Excluded from median", "Reason (Comparison)"]
ws.append([]); ws.append(hdr)
for c in ws[4]: c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = B
for lab, k, u, g in SUMMARY:
    row = [g, k, lab.replace("\u00b9", ""), u]
    for p in PEERS:
        x = raw(k, p)
        row += [x if x is not None else (ROWS[k][COL[p]] or "n/d"), stag(tag(k, p)) if x is not None else "n/d", asof(k, p) or ""]
    exc = ", ".join(p for p in PEERS if excluded(k, p))
    row += [med(k), flag(k), RANK.get(k, {}).get("rank"), exc, ROWS[k][FLAG + 1]]
    ws.append(row)
    r = ws.max_row
    for c in ws[r]: c.border = B; c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(r, len(hdr) - 3).fill = PatternFill("solid", fgColor=FL.get(flag(k), "FFFFFF"))
    for j in range(5, 5 + 3 * len(PEERS), 3):
        if isinstance(ws.cell(r, j).value, float): ws.cell(r, j).number_format = "0.00"
    ws.cell(r, 5).fill = PatternFill("solid", fgColor="E8EEF7"); ws.cell(r, 5).font = Font(bold=True)
    ws.cell(r, len(hdr) - 4).number_format = "0.00"
notes = [
 f"K02b = revenue per average BILLED IT MW (Capacity billing path; Nebius {ni('D_K02B_BASE_MW'):.1f} IT MW avg Q1-Q2'26) - NOT energized or gross MW. CoreWeave K02b = K02 (active ~ billed); other peers n/d.",
 "Nebius MW (K02/K13 etc.) is a Capacity & Power estimate (mixed IT/gross; 181.5/193 MW base, Capacity base case Quarterly MW!M130/N130).",
 "Nscale K28 18.7% = legacy $137M debt, excluded from median (marginal secured ~6.25% E). Nscale K23 pre-ramp.",
 "Periods: IREN K02 Sep-25; Oracle May-26 (FQ4 FY26), K18 Aug-26; Nscale H1'26; Lambda/Crusoe press-only, mostly n/d. K12 Nebius = FY25.",
 "CoreWeave K18 = 6.47 yrs (Q2'26) after the pack fix (was 16.67 bug).",
 f"Share count convention: N_FD_DEBT {ni('N_FD_DEBT'):.2f}M (new Feb-2030/Feb-2034 notes as debt; applies below ${FDC['PX_2030']:.2f}, incl. the ${FDC['REF_PX']:.2f} 10/7 close) vs N_FD_IFCONV {ni('N_FD_IFCONV'):.2f}M (both series converted above ${FDC['PX_2034']:.2f}; remove ALL convert debt from EV); {FDC['FD_MID']/1e6:.2f}M between ${FDC['PX_2030']:.2f} and ${FDC['PX_2034']:.2f}. Spread +{ni('D_FD_CONV_SENS'):.2f}%. Never mix share and debt sides (model v0.9e uses the debt treatment for every scenario PT below $313.46). v0.9d {ni('N_FD_V09D'):.2f}M historical only. No KPI here is per-share.",
 f"Forward bar (context, not scored): YE26 guide ${ni('D_YE26_LO_BASE'):.1f}M ($7B/{ni('N_BILL_YE26_BASE'):.1f}) / ${ni('D_YE26_HI_HIGH'):.1f}M ($9B/{ni('N_BILL_YE26_HIGH'):.1f}) per billed IT MW; Street FY27 ex-Microsoft ${ni('D_STREET_FY27_NONMSFT'):.2f}M (peer comparator; ($12,300M - ${ni('N_MSFT_FY27_REV'):,.0f}M) / ({ni('N_BILL_FY27_BASE'):.2f} - {ni('N_MSFT_FY27_MW'):.2f})); blended ${ni('D_STREET_FY27_BASE'):.2f}M memo only; high ${ni('D_STREET_FY27_HIGH'):.2f}M / ex-MSFT ${ni('D_STREET_FY27_NONMSFT_HIGH'):.2f}M. FY27 billed MW = model v0.9e engine-linked {ni('N_BILL_FY27_BASE'):.2f}/{ni('N_BILL_FY27_HIGH'):.2f} (model/nebius_model.xlsx Drivers B34/C34; Vineland {ni('N_MSFT_FY27_MW'):.2f} billed MW-yr = Microsoft ${ni('N_MSFT_FY27_REV'):,.1f}M / $14.5M). At $16M comps the Street needs 768.75 MW = {ni('D_COV16_BASE'):.2f}x base / {ni('D_COV16_HIGH'):.2f}x high.",
 "Nebius Q3'26 results: ~Nov 10 (est.) - not announced by the company.",
 BAND_NOTE + " See Contracts_$MWyr row NEW2. History: band was $9.7-16.2M (Nscale-Anthropic top) before Oct 8.",
]
ws.append([]); ws.append(["Notes"]); ws.cell(ws.max_row, 1).font = Font(bold=True)
for n_ in notes: ws.append([n_])
widths = {"A": 20, "B": 7, "C": 40, "D": 10}
for col, w in widths.items(): ws.column_dimensions[col].width = w
from openpyxl.utils import get_column_letter
for j in range(5, len(hdr) + 1): ws.column_dimensions[get_column_letter(j)].width = 10
ws.column_dimensions[get_column_letter(len(hdr))].width = 70
ws.freeze_panes = "E5"

# Contracts
wc = wb.create_sheet("Contracts_$MWyr")
wc["A1"] = "Compute contracts: implied $M per MW-yr = TCV / term / MW (comparison.xlsx Contracts_$MWyr). Facility leases listed separately below (NON-COMPUTE, shell+power)."
wc["A1"].font = Font(bold=True, color="1F4E8C")
ch = ["ID", "Provider", "Customer", "Category", "Signed", "TCV ($M)", "TCV tag", "Term (yrs)", "MW (stated)", "MW basis", "IT MW (normalized)", "IT MW method", "$M per MW-yr (stated MW)", "$M per IT MW-yr", "Source", "Notes"]
con, nd, fac = contract_rows()
def put(title, recs):
    wc.append([]); wc.append([title]); wc.cell(wc.max_row, 1).font = Font(bold=True)
    wc.append(ch)
    for c in wc[wc.max_row]: c.font = H; c.fill = HF; c.border = B; c.alignment = Alignment(wrap_text=True)
    for r in recs:
        vals = [r["id"], r["provider"], r["customer"], r["cat"], r["signed"] or "n/d", r["tcv"], r["tcvtag"], r["term"], r["mw"], r["mwbasis"], r["itmw"], r["itmeth"], r["stated"], r["it"], r["src"], r["notes"]]
        wc.append([("n/d" if v_ in (None, "") else v_) for v_ in vals])
        for c in wc[wc.max_row]:
            c.border = B
            if isinstance(c.value, float): c.number_format = "0.00"
put("A. Compute contracts with a computable $/MW-yr (plus all Nebius rows)", con)
wc.append(["Nebius-Microsoft: $14.49M per IT MW [E, house] = model Contracts!E12 (row N1 above); 11.60 / 17.08 superseded (history only)"])
put("B. Compute contracts with no computable $/MW (MW, term or value undisclosed)", nd)
put("C. NON-COMPUTE facility leases (shell + power only; not comparable to compute $/MW)", fac)
for col, w in zip("ABCDEFGHIJKLMNOP", [7, 18, 30, 20, 14, 10, 14, 8, 10, 22, 10, 22, 12, 12, 40, 50]): wc.column_dimensions[col].width = w
wb.save(OUT); print("saved", OUT)
