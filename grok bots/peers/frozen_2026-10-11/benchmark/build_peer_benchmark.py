#!/usr/bin/env python3
"""One-page peer benchmark (Sun Oct 11 finals): disclosed $ per IT MW-yr, contracted vs energized MW,
power and permit status for CoreWeave, IREN, Oracle, Nscale, Lambda, Crusoe vs Nebius.

Reuses (read-only): peers/comparison.xlsx (Contracts_$MWyr, Sites, Nebius_Inputs), peers/price_or_mw.xlsx,
peers/private/private_peers.md/.xlsx, peers/sites/appendix_* + site_verification.md, peers/{coreweave,iren,oracle} packs,
parts/contract_economics.xlsx (MidTerm_20M_Search), capacity/site_verification.md, model v0.9e via peers/nebius_locks.py.
Writes only to peers/benchmark/.
Tags: R reported/primary (filing, company release, earnings call, regulator); D derived by us from R inputs;
P press-only (or derived with any press input; footnoted); E estimate (Nebius energized MW, Capacity base case); n/d not disclosed.
"""
import sys, os, re, shutil, subprocess, datetime
sys.path.insert(0, "/workspace/nebius/peers")
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
import nebius_locks as L

D = "/workspace/nebius/peers/benchmark"
XLSX, DOCX, PDF, PNG = (f"{D}/peer_benchmark.{e}" for e in ("xlsx", "docx", "pdf", "png"))
ASOF = "Fri Oct 9, 2026 (ET)"

# ---------------------------------------------------------------- inputs (reuse, read-only)
cmp_ = openpyxl.load_workbook("/workspace/nebius/peers/comparison.xlsx", data_only=True)
CON = {r[0]: r for r in cmp_["Contracts_$MWyr"].iter_rows(min_row=4, values_only=True) if r[0]}
NI = {r[0]: r[2] for r in cmp_["Nebius_Inputs"].iter_rows(min_row=3, values_only=True) if r[0]}
ML = L.model_locks()
assert ML["VERSION"]["value"] == "v0.9e", ML["VERSION"]
LOCK = dict(MSFT_REV=ML["MSFT_REV"]["value"], VIN_MW=ML["VIN_MW"]["value"], STREET=ML["STREET_FY27"]["value"],
            FY27_BASE=ML["FY27_BASE"]["value"], FY27_HIGH=ML["FY27_HIGH"]["value"], YE26_BASE=ML["YE26_BASE"]["value"],
            YE26_HIGH=ML["YE26_HIGH"]["value"], GUIDE_LO=ML["GUIDE_LO"]["value"], GUIDE_HI=ML["GUIDE_HI"]["value"],
            HURDLE=ML["HURDLE_XMSFT"]["value"], BLENDED=ML["BLENDED"]["value"],
            EN_Q1=NI["N16"], EN_Q2=NI["N17"], CONTRACTED=NI["N19"])
# requested v0.9e locks (task brief) - hard asserts so a model move breaks the build instead of silently drifting
for k, v, tol in (("MSFT_REV", 2643.8, .05), ("VIN_MW", 182.33, .005), ("FY27_BASE", 531.84, .005), ("FY27_HIGH", 659.16, .005),
                  ("YE26_BASE", 257.7, 1e-6), ("YE26_HIGH", 284.2, 1e-6), ("HURDLE", 27.63, .005), ("BLENDED", 23.13, .005),
                  ("EN_Q1", 181.5, 1e-9), ("EN_Q2", 193, 1e-9)):
    assert abs(LOCK[k] - v) < tol, (k, LOCK[k], v)
assert abs(LOCK["MSFT_REV"] / LOCK["VIN_MW"] - 14.5) < 1e-6
assert abs(NI["D_STREET_FY27_NONMSFT"] - LOCK["HURDLE"]) < 1e-9 and abs(NI["D_STREET_FY27_BASE"] - LOCK["BLENDED"]) < 1e-9
# item 3 (Oct 9 coordinator): locked Nebius hurdles verified against model v0.9e cells (read-only)
_mv = openpyxl.load_workbook("/workspace/nebius/model/nebius_model.xlsx", data_only=True)
_mf = openpyxl.load_workbook("/workspace/nebius/model/nebius_model.xlsx")
MODEL_CELLS = [  # key, label, brief value, sheet, cell, tol, note
 ("HURDLE", "Non-MSFT hurdle (peer comparator)", 27.6, "Drivers", "B9", 0.05, "formula =(INDEX(Cons_Rev,1,2027-FYBase)-B32)/B30; = Judge_Slide!B20"),
 ("CLEAR_BASE", "Clear-spot BASE ($237.15 PT; Judge (a) base, 11% WACC)", 24.51, "Drivers", "B6", 0.005, "static full-model bisection; = Judge_Slide!B8/B14; Debate_Grid!A20 'base clear-spot (Judge line 1)'"),
 ("CLEAR_HIGH", "Clear-spot HIGH (high billed path)", 22.68, "Debate_Grid", "A18", 0.005, "Debate_Grid!C18 'high-path clear-spot' (PT $237.12, E18); Drivers!A58 text '$22.68M GB300'"),
 ("TP_BASE", "Turn-positive BASE (Judge basis; first PT > 0)", 17.05, "Debate_Grid", "A13", 0.005, "PT $0.12 (D13); Debate_Grid!B26 'BASE at $17.05M ($0 at $17.0432)'"),
 ("TP_HIGH", "Turn-positive HIGH (Judge basis; first PT > 0)", 16.70, "Debate_Grid", "A10", 0.005, "PT $0.09 (E10); Debate_Grid!B26 'HIGH first positive at $16.70M ($0 at $16.6976)'"),
 ("BLENDED", "MEMO: $23.13M = Street FY27 \u00f7 our 531.8 billed MW, all customers [D]", 23.13, "Drivers", "B10", 0.005, "formula =INDEX(Cons_Rev,1,2027-FYBase)/B34 = $12,300M / 531.84; = Judge_Slide!B16 'STREET (blended, supporting line)'"),
]
MC = {}
for k_, lab_, brief_, sh_, cl_, tol_, note_ in MODEL_CELLS:
    mv_ = _mv[sh_][cl_].value; mf_ = _mf[sh_][cl_].value
    assert isinstance(mv_, (int, float)) and abs(mv_ - brief_) < tol_, ("model cell mismatch", k_, sh_, cl_, mv_, brief_)
    MC[k_] = dict(v=mv_, ref=f"{sh_}!{cl_}", formula=mf_ if isinstance(mf_, str) and mf_.startswith("=") else "static value", note=note_, label=lab_, brief=brief_)
assert abs(MC["HURDLE"]["v"] - LOCK["HURDLE"]) < 1e-9 and abs(MC["BLENDED"]["v"] - LOCK["BLENDED"]) < 1e-9
assert "v0.9e" in str(_mv["Cover"]["A2"].value)
MSFT_E12 = _mv["Contracts"]["E12"].value
assert abs(MSFT_E12 - 17392.9 / (5 * 240)) < 1e-6 and _mf["Contracts"]["E12"].value == "=MSFT_TCV/(5*MSFT_MW)", MSFT_E12
assert _mv["Inputs"]["F215"].value == 17392.9 and _mv["Inputs"]["F216"].value == 240 and _mv["Inputs"]["F111"].value == 14.5 and _mv["Drivers"]["C11"].value == 14.5

# comp band set (common_data BAND_IDS) - must stay 9.7-16.7
BAND_IDS = ["P1", "NEW1", "P2", "O1", "NEW2", "O2"]
def _band(cid):
    r = CON[cid]; return r[22] if isinstance(r[22], (int, float)) else r[21]
BAND = {c: _band(c) for c in BAND_IDS}
assert min(BAND, key=BAND.get) == "P1" and max(BAND, key=BAND.get) == "NEW2", BAND
assert abs(BAND["P1"] - 9.7) < .05 and abs(BAND["NEW2"] - 16.67) < .01, BAND

# ---------------------------------------------------------------- sources
SRC = {
 "CW-META1": ("CoreWeave 8-K 2025-09-30 (Meta up to $14.2B through Dec-2031)", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001769628&type=8-K"),
 "CW-META2": ("CoreWeave 8-K 2026-04-09 (Meta up to $21.0B)", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001769628&type=8-K"),
 "CW-OAI": ("CoreWeave 8-Ks 2025-03-10 / 05-15 / 09-25 (OpenAI $11.9B + $4.0B + $6.5B); reused via parts/contract_economics.xlsx P6-P8", "local:/workspace/nebius/parts/contract_economics.xlsx"),
 "CW-JS": ("CoreWeave PR 2026-04-15 (Jane Street $6B; term/MW n/d)", "https://www.coreweave.com/news/jane-street-signs-6-billion-ai-cloud-agreement-with-coreweave"),
 "CW-DECK": ("CoreWeave convertible-notes investor presentation, 8-K exhibit 2026-09-17, slide 7 (~$40M/MW annualized on 3-6-month contracts)", "https://www.sec.gov/Archives/edgar/data/1769628/000176962826000429/crwv2026convertpresentat.htm"),
 "CW-Q2": ("CoreWeave Q2'26 earnings PR (8-K Ex.99.1, 2026-08-11): active power ~1.5 GW, contracted ~3.7 GW; ~4.2 GW contracted as of Aug-11", "https://www.sec.gov/Archives/edgar/data/1769628/000176962826000362/coreweave2q26earningspress.htm"),
 "CW-SITE": ("peers/sites/appendix_coreweave_denton (CORZ 8-K 2026-09-10 ERCOT status; CORZ Q2'26 deck 2026-07-28; TCEQ PBR 179143; TCEQ Title V/NSR lists 10/8; Community Impact 2026-09-24)", "https://investors.corescientific.com/sec-filings/all-sec-filings/content/0001839341-26-000023/0001839341-26-000023.pdf"),
 "IR-MSFT": ("IREN 8-K 2025-11-03 (Microsoft $9.7B, 5 yrs, 200 MW IT, Childress Horizons 1-4)", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001878848&type=8-K"),
 "IR-NV": ("IREN PR 2026-05 (NVIDIA $3.4B AI cloud contract, ~60 MW of existing Childress halls)", "https://iren.gcs-web.com/news-releases/news-release-details/iren-secures-34bn-ai-cloud-contract-nvidia"),
 "IR-FY26": ("IREN FY26 results 2026-08-27 ('Recent 3-year contracts >$20m revenue per MW (IT)'; 'Active discussions at ~$25m per MW (IT)'); 10-K FY26 (~40 MW AI operating at Jun-30)", "https://www.sec.gov/Archives/edgar/data/1878848/000187884826000051/irenreportsfy26results.htm"),
 "IR-SITE": ("peers/sites/appendix_iren_childress (IREN 8-K Ex.99.1 2026-08-13 H1 accepted; IREN PR 2024-07-24 750 MW; TCEQ Title V #4951 list data 10/8; Gov. Abbott 2026-09-21)", "https://www.sec.gov/Archives/edgar/data/1878848/000114036126032638/ef20080141_ex99-1.htm"),
 "OR-OAI": ("WSJ 2025-09-10 (OpenAI ~$300B / ~5 yrs / 4.5 GW) - press only, not in Oracle filings; reused via comparison.xlsx Contracts O2", "https://www.wsj.com/business/openai-oracle-sign-300-billion-computing-deal-among-biggest-in-history-ff27c8fe"),
 "OR-META": ("Bloomberg/Reuters 2025-09-19; CNBC 2025-10-16 (Meta ~$20B, talks) - press", "https://www.cnbc.com/"),
 "OR-TCNT": ("FT via DCD 2026-10-01 (Tencent ~100k GPUs, ~$7B / 5 yrs) - press, unconfirmed", "https://www.datacenterdynamics.com/en/news/tencent-signs-on-for-100000-gpus-via-oracle-report/"),
 "OR-CALL": ("Oracle FQ1 FY27 earnings call 2026-09-10 (850 MW delivered in quarter; Abilene 6 of 8 bldgs / 618 MW / 75%; '>10 GW secured'); peers/oracle/oracle_kpis.xlsx I09/I10/I45E", "https://www.fool.com/earnings/call-transcripts/2026/09/11/oracle-orcl-q1-2027-earnings-call-transcript/"),
 "OR-SITE": ("peers/sites/appendix_oracle_port_washington (PSC docs 593736/597276; WDNR 25-MIN-094 final 2026-06-12; TechRadar 2026-10-04; Aterio 2026-09-30; Wisconsin Watch)", "https://apps.psc.wi.gov/ERF/ERFview/viewdoc.aspx?docid=597276"),
 "NS-S1": ("Nscale Form S-1 2026-09-18 (Anthropic GSAs up to ~$44.6B, 460 MW IT Monarch; Microsoft up to $43.8B, 257.6k GPUs; 1,370 MW IT active+contracted; financing for Anthropic not committed)", "https://www.sec.gov/Archives/edgar/data/2110365/000119312526395475/ck0002110365-20260918.htm"),
 "NS-TERM": ("Bloomberg/Reuters via WV MetroNews (Anthropic six-year term); Microsoft tranche terms 5-6 yrs (5.5 midpoint, GPU Market agent) - press/assumption", "local:/workspace/nebius/parts/contract_economics.xlsx Notes_Method"),
 "NS-NARV": ("Aker-Nscale PR 2025-09-17 (Microsoft Narvik $6.2B, 5 yrs, ~52k GB300)", "https://akernscale.com/"),
 "NS-SITE": ("peers/sites/site_verification.md s.3 (Kvandal; NRK 2026-09-29 NVE) and s.7 (Monarch; WVDEP R14-0042 application 2026-04-20 / draft 2026-09-29; WV News stormwater order)", "local:/workspace/nebius/peers/sites/site_verification.md"),
 "LM-ANTH": ("WSJ / Reuters / Bloomberg 2026-08-31 (Anthropic $35B, ~350 MW, Hut 8 Beacon Point, NVIDIA holds lease); six-year term per WSJ via TheEnergyMag - press only, unconfirmed by Lambda/Anthropic", "https://www.reuters.com/technology/anthropic-signs-35-billion-cloud-deal-with-nvidia-backed-lambda-source-says-2026-08-31/"),
 "LM-MSFT": ("Lambda PR 2025-11-03 (Microsoft 'multibillion-dollar'; no TCV/MW)", "https://lambda.ai/blog/lambda-announces-multibillion-dollar-agreement-with-microsoft-to-deploy-ai-infrastructure-powered-by-tens-of-thousands-of-nvidia-gpus"),
 "LM-NV": ("DCD citing The Information 2025-09-08 (NVIDIA $1.5B / ~18k GPU 4-yr leaseback; $1.3B/10k tranche) - press", "https://www.datacenterdynamics.com/en/news/nvidia-signs-15bn-deal-to-lease-its-gpus-back-from-lambda-report/"),
 "LM-BKLG": ("TechCrunch / WSJ 2026-10-06 (investor letter: backlog $15B Jun-26 -> $50B Sep-26) - press", "https://techcrunch.com/2026/10/06/ai-computing-startup-lambda-to-raise-4b-ahead-of-planned-ipo/"),
 "LM-MW": ("peers/private/private_peers.xlsx L12/L14 (47 MW Q2'25 via The Information, P); Lambda PR 2025-10-28 (KC 24 MW, >100 MW potential)", "https://lambda.ai/blog/lambda-to-build-a-100mw-ai-factory-in-kansas-city-mo"),
 "LM-SITE": ("peers/sites/site_verification.md s.5 (MO DNR de minimis 102025-010 issued 2025-10-29; MO DNR letter 2026-04-15 rental engines; KSHB Evergy)", "https://dnr.mo.gov/air/business-industry/air-permits/airworld-data-center-kansas-city-102025-010"),
 "CR-F": ("Crusoe PR 2026-09-17 Series F ($140B+ TCV; 6 GW+ gross contracted; ~1 GW gross delivered and operational)", "https://www.crusoe.ai/resources/newsroom/crusoe-announces-series-f-funding"),
 "CR-DEALS": ("Crusoe PR 2026-03-27 (Microsoft 900 MW / 2x336 MW IT); Crusoe JV PR 2025-05-21 (Abilene 1.2 GW leased to Oracle); DCD/Bloomberg 2026-06-19 (Meta ~1.6 GW) and 2026-09-04 (Jane Street $13B/5 yrs) - press for Meta/Jane Street", "https://www.crusoe.ai/resources/newsroom/crusoe-announces-new-900-mw-ai-factory-campus-in-abilene-texas-to-support-microsoft-ai-infrastructure"),
 "CR-SITE": ("peers/sites/appendix_crusoe_abilene (Oracle FQ1 FY27 call; Mortenson; TCEQ 177263 tech review; TCEQ 182126 NORI/NAPD/PLS; NSR list 10/8; Gov. Abbott 2026-09-21; Texas Lawbook)", "https://www.tceq.texas.gov/downloads/permitting/air/bilingual/pending-permit-notices/182126-pls-english.pdf"),
 "NB-MSFT": ("Nebius 6-K 2025-09-08 (Microsoft $17.4B / up to $19.4B, 5 yrs, Vineland; '>100k GB300' is press)", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001513845&type=6-K"),
 "NB-META": ("Nebius 6-K 2025-11-12 (Meta $3B / 5 yrs) and 6-K 2026-03-16 (Meta $12B dedicated + up to $15B contingent); MW not stated", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001513845&type=6-K"),
 "NB-Q2": ("Nebius Q2'26 shareholder letter, 6-K 2026-08-12 Ex.99.2 p.3-4 (four landmark deals, 'yield of $20-25 million per megawatt' ACV; 2026 base ~$12M/MW; KPI chart 'Q3 short-term capacity deals >$40M' = management comment on Q3-to-date short-term pricing, not reported Q3 results); MW basis not defined", "https://www.sec.gov/Archives/edgar/data/1513845/000110465926094568/tm2622968d1_ex99-2.htm"),
 "NB-LOCK": ("Model v0.9e: model/nebius_model.xlsx Drivers B9/B10/B30-B34/C34, Judge_Slide B33-B35, Inputs F111 ($14.5M) via peers/nebius_locks.py; peers/comparison.xlsx Nebius_Inputs", "local:/workspace/nebius/model/nebius_model.xlsx"),
 "NB-MW": ("capacity/nebius_capacity.xlsx Quarterly MW M130/N130 (energized 181.5/193, E); Q1'26 letter p.6 (>3.5 GW contracted power); Q2'26 letter (800 MW-1 GW connected YE26 guide)", "local:/workspace/nebius/capacity/nebius_capacity.xlsx"),
 "NB-SITE": ("capacity/site_verification.md s.1 + vineland_site_appendix (NJ Governor 2026-09-22; NJDEP order; Vineland stop-construction 8/6, 8/10; Planning Board 8/17; Oct 14 agenda)", "https://www.nj.gov/governor/news/2026/20260922a.shtml"),
 "NB-CALL": ("Nebius Q2'26 earnings call 2026-08-12, CEO Arkady Volozh prepared remarks: shorter-duration capacity 'typically for up to six months' ... 'We are negotiating deals for $40 million-$50 million per megawatt range, and sometimes above ... One of such deal has been signed just recently' (Quartr transcript via StockAnalysis; local filings/altdata/transcripts/transcript_q2-2026.txt)", "https://stockanalysis.com/stocks/nbis/transcripts/655490-q2-2026/"),
 "BD-A201": ("Bitdeer AI PR 2026-09-29 (GlobeNewswire via Nasdaq): >70% off-take of ~21.7 MW A201 Johor; >$1.7B expected revenue over 5 years; GB300 NVL72; energization Q1 2027; MW not labeled IT/gross. Re-read 2026-10-09", "https://www.nasdaq.com/press-release/bitdeer-ai-further-expands-ai-cloud-services-take-commitments-more-70-its-217mw-a201"),
 "BD-MKT": ("GPU Market agent: market/contract_economics.xlsx MidTerm_20M_Search rows 24-25, 33, 36 (A201 ceiling <= $22.4M = >$1.7B/5/(0.70 x 21.7 = 15.19 MW); 'IT vs gross MW not stated'; Datatech Times 9/16 and ArusGrid call the 86.8 MW campus 'critical IT load' (secondary only); Needham 10/2 and 10/7 notes use the same $22.4M); market/hyperscaler_contract_mw.xlsx Slide_Band B2; model Drivers!B13 / Debate_Grid A17", "local:/workspace/nebius/market/contract_economics.xlsx"),
 "BD-A102": ("Bitdeer AI PR 2026-08-31 (A102 Malaysia 9.5 MW sold out, >$800M / 5 yrs; revenue from Q1 2027) via market/contract_economics.xlsx MidTerm_20M_Search row 25", "https://www.globenewswire.com/news-release/2026/08/31/3353668/0/en/"),
 "BD-A103": ("Bitdeer AI release 2026-10-08 via TMCnet (A103 Cyberjaya 60 MW IT, 8-yr data-center services agreement, energize Q1 2028, no value; secured 333.5 MW) - peer watch 10/9", "https://www.tmcnet.com/usubmit/-bitdeer-ai-secures-60mw-a103-facility-its-existing-/2026/10/08/10460312.htm"),
 "NB-MODEL": ("model/nebius_model.xlsx v0.9e (read-only): Drivers!B9 27.63 non-MSFT hurdle; Drivers!B6 = Judge_Slide!B8 24.51 base clear-spot; Debate_Grid!A18 22.68 high clear-spot (Drivers!A58); Debate_Grid!A13 17.05 / A10 16.70 turn-positive (Debate_Grid!B26); Drivers!B10 = Judge_Slide!B16 23.13 blended (memo only)", "local:/workspace/nebius/model/nebius_model.xlsx"),
 "MKT-20": ("parts/contract_economics.xlsx MidTerm_20M_Search (GPU Market agent; re-swept 2026-10-09 08:50 ET: no mid-term deal >=$20M/IT MW-yr with disclosed MW)", "local:/workspace/nebius/parts/contract_economics.xlsx"),
}

# ---------------------------------------------------------------- contracts (numeric; $/MW-yr computed by formula in xlsx)
# (id, peer, counterparty, TCV $M, TCV tag, term yrs, term tag, MW, MW basis, MW tag, result tag, src keys, note)
def c(cid, k): return CON[cid][k]
CONTRACTS = [
 ("P4", "CoreWeave", "Meta #1", c("P4",6), "R", c("P4",8), "R", None, "n/d", "n/d", "n/d", "CW-META1", "Up to; no MW or GPUs disclosed"),
 ("P5", "CoreWeave", "Meta #2", c("P5",6), "R", c("P5",8), "R", None, "n/d", "n/d", "n/d", "CW-META2", "Up to; incl. Vera Rubin; no MW"),
 ("P6-8", "CoreWeave", "OpenAI (3 orders)", c("P6",6)+c("P7",6)+c("P8",6), "R", None, "R (3.9-5.7 by order)", None, "n/d", "n/d", "n/d", "CW-OAI", "No MW disclosed"),
 ("NEW4", "CoreWeave", "Jane Street", c("NEW4",6), "R", None, "n/d", None, "n/d", "n/d", "n/d", "CW-JS", "Term and MW n/d"),
 ("P1", "IREN", "Microsoft", c("P1",6), "R", c("P1",8), "R", c("P1",9), "IT (company-stated)", "R", "D", "IR-MSFT", "Comp-band LOW"),
 ("NEW1", "IREN", "NVIDIA", c("NEW1",6), "R", c("NEW1",8), "R", c("NEW1",9), "unlabeled (~60 MW of existing halls)", "R", "D", "IR-NV", "Ramp from early 2027"),
 ("O2", "Oracle", "OpenAI (Stargate) - stated MW", c("O2",6), "P", c("O2",8), "P", c("O2",9), "unlabeled (stated)", "P", "P", "OR-OAI", "Press only; footnote 3"),
 ("O2-IT", "Oracle", "OpenAI (Stargate) - IT read", c("O2",6), "P", c("O2",8), "P", c("O2",11), "IT = 4,500 / 1.2 PUE (gross read; ~353 GPUs/MW)", "D", "P", "OR-OAI", "Same deal, IT-MW read"),
 ("NEW10", "Oracle", "Meta", c("NEW10",6), "P", None, "n/d", None, "n/d", "n/d", "n/d", "OR-META", "Talks; MW/term n/d"),
 ("TCNT", "Oracle", "Tencent (~100k GPUs)", 7000, "P", 5, "P", None, "n/d", "n/d", "n/d", "OR-TCNT", "Unconfirmed; MW n/d"),
 ("O1", "Nscale", "Anthropic (Monarch WV)", c("O1",6), "R", c("O1",8), "P", c("O1",9), "IT (S-1 facilities table)", "R", "P", "NS-S1|NS-TERM", "Up to; term press; financing not committed"),
 ("P2", "Nscale", "Microsoft (all SOWs)", c("P2",6), "R", c("P2",8), "P", c("P2",11), "IT derived = 191.6k/491 + 66k/320 GPUs per IT MW", "D", "P", "NS-S1|NS-TERM", "Up to; term midpoint of 5-6 yr tranches"),
 ("P2a", "Nscale", "Microsoft - Narvik (Aker JV)", c("P2a",6), "R", c("P2a",8), "R", c("P2a",11), "IT derived = 52k GB300 / 491", "D", "D", "NS-NARV", "Within P2"),
 ("NEW2", "Lambda", "Anthropic (Beacon Point TX)", c("NEW2",6), "P", c("NEW2",8), "P", c("NEW2",9), "unlabeled (IT vs gross unstated)", "P", "P", "LM-ANTH", "Comp-band TOP; press only; footnote 5"),
 ("P3", "Lambda", "Microsoft", None, "R ('multibillion')", None, "n/d", None, "n/d", "n/d", "n/d", "LM-MSFT", "No economics public"),
 ("NEW3", "Lambda", "NVIDIA GPU leaseback ($1.3B tranche)", c("NEW3",6), "P", c("NEW3",8), "P", None, "n/d (10k GPUs -> $3.71/GPU-hr)", "n/d", "n/d", "LM-NV", "MW n/d"),
 ("NEW9", "Crusoe", "Oracle/OpenAI (Abilene 1.2 GW lease)", None, "n/d", None, "n/d", 1200, "gross", "R", "n/d", "CR-DEALS", "Value n/d ($15B = JV project cost, not contract value)"),
 ("NEW8", "Crusoe", "Microsoft (Abilene 900 MW campus)", None, "n/d", None, "n/d", 672, "IT (2 x 336)", "R", "n/d", "CR-DEALS", "Value n/d; first bldg mid-2027"),
 ("NEW7", "Crusoe", "Meta (~1.6 GW Childress + Warrenton)", None, "n/d", None, "n/d", 1600, "unlabeled", "P", "n/d", "CR-DEALS", "Press; value n/d"),
 ("NEW6", "Crusoe", "Jane Street", c("NEW6",6), "P", c("NEW6",8), "P", None, "n/d", "n/d", "n/d", "CR-DEALS", "Press; MW n/d"),
 ("BD-A201", "Bitdeer AI", "A201 Johor off-take (unnamed customers, >70% of ~21.7 MW)", 1700, "R ('>$1.7B')", 5, "R", round(0.70 * 21.7, 2), "unlabeled in primary; 'critical IT' per secondary only", "R (floor: >70%)", "D (basis P)", "BD-A201|BD-MKT", "CEILING (MW is a floor); 100% contracted -> 15.7"),
 ("BD-A102", "Bitdeer AI", "A102 Malaysia (sold out)", 800, "R ('>$800M')", 5, "R", 9.5, "unlabeled", "R", "D", "BD-A102|BD-MKT", "Floor ('>'); revenue from Q1'27"),
 ("BD-A103", "Bitdeer AI", "A103 Cyberjaya 8-yr DC services (context)", None, "n/d", 8, "R", 60, "IT (release)", "R", "n/d", "BD-A103", "No value; energize Q1'28"),
 ("N1-H", "Nebius", "Microsoft - HOUSE (model v0.9e Contracts!E12; MW not stated by Nebius/Microsoft)", c("N1",6), "R", 5, "R", 240, "IT (model estimate, Inputs!F216)", "E", "E (house)", "NB-LOCK|NB-MODEL", "17,392.9 / (5 x 240 IT MW est.) = 14.49; agrees with the Drivers 14.5 input (Inputs!F111)"),
 ("N2", "Nebius", "Meta #1", c("N2",6), "R", c("N2",8), "R", None, "n/d (filings state no MW)", "n/d", "n/d", "NB-META", "~50 MW est. (market) not used"),
 ("N3", "Nebius", "Meta #2 (dedicated)", c("N3",6), "R", c("N3",8), "R", None, "n/d", "n/d", "n/d", "NB-META", "+ up to $15B contingent"),
]
# company KPI / non-contract $ per MW prints (value entered; >=20 flagged)
KPI_PRINTS = [
 # id, peer, label, low, high, tag, basis, src, why-high (None if <20)
 ("CW-K1", "CoreWeave", "Short-dated (3-6 month) contracts, annualized revenue per MW", 40, 40, "R", "MW basis unstated (CRWV 'active power' reads IT, E)", "CW-DECK",
  "Short term (3-6 mo) and GPU-inclusive spot pricing after a ~70% July-2026 SKU price rise; marginal deals, not the fleet ($8.24M per avg active MW Q2'26)."),
 ("CW-K2", "CoreWeave", "Denton colocation (CORZ lease ~590 MW, '$10B+' ~12 yrs) - shell+power, not compute", 1.4, 1.4, "D", "IT (CORZ leased MW)", "CW-SITE", None),
 ("IR-K1", "IREN", "Recent 3-year AI cloud contracts, revenue per MW (IT)", 20, 20, "R", "IT (company: 'per MW (IT)'); '>' = floor", "IR-FY26",
  "Short (3-yr) tenor vs the 5-yr Microsoft anchor at $9.7M, GPU-inclusive cloud revenue (~2-yr payback per IREN); portfolio KPI with no per-deal MW."),
 ("IR-K2", "IREN", "'Active discussions' pricing (not signed)", 25, 25, "R", "IT (company)", "IR-FY26",
  "Unsigned negotiating level, GPU-inclusive, short tenor; no MW or counterparty."),
 ("CR-K1", "Crusoe", "Company TCV $140B+ / 6 GW+ gross contracted (TCV per MW, NOT per year)", None, None, "D", "gross MW; term n/d", "CR-F",
  "Whole-term TCV per GROSS MW, not annual; mixes facility leases, cloud and JV capacity. Any multi-year term divides it (e.g. 5 yrs -> ~$4.7M/MW-yr, illustrative D). Not comparable to $/MW-yr."),
 ("NB-K1", "Nebius", "Q2'26 four 'landmark' deals, ACV per MW", 20, 25, "R", "MW basis not defined in the 6-K", "NB-Q2",
  "1-3-yr tenor (vs 5-yr hyperscaler anchors), GPU-inclusive full-stack cloud ACV; MW basis undefined (cannot map to billed IT MW); MW per deal n/d."),
 ("NB-K2", "Nebius", "Mgmt comment on Q3-to-date short-term pricing (Q2'26 letter KPI chart '>$40M'; CEO on Q2 call 8/12/26: negotiating $40-50M per MW, one deal signed) - NOT reported Q3 results", 40, 50, "R", "basis n/d", "NB-Q2|NB-CALL",
  "Short term (up to ~6 mo; 'typically between three and six months'), premium for time-bounded need; one signed deal, capacity online later in 2026; not underwritable as a mid-term blend."),
 ("NB-K3", "Nebius", "2026 base book, per MW (company)", 12, 12, "R", "basis n/d", "NB-Q2", None),
]

# ---------------------------------------------------------------- derived values (python mirror of xlsx formulas)
def permw(tcv, term, mw):
    return tcv / term / mw if all(isinstance(x, (int, float)) for x in (tcv, term, mw)) else None
VAL = {r[0]: permw(r[3], r[5], r[7]) for r in CONTRACTS}
for k_, ref in (("P1", 9.7), ("NEW1", 11.333), ("O2", 13.333), ("O2-IT", 16.0), ("O1", 16.159), ("P2", 13.351), ("P2a", 11.708),
                ("NEW2", 16.667), ("N1-H", 14.494)):
    assert abs(VAL[k_] - ref) < 0.01, (k_, VAL[k_], ref)
assert abs(VAL["BD-A201"] - 22.38) < 0.01 and VAL["BD-A201"] >= 20 and abs(1700 / 5 / 21.7 - 15.67) < .01  # market file: <= $22.4M ceiling / 15.7 if 100%
CR_TCV_PER_MW = 140000 / 6000
HURDLES = [  # Nebius required revenue per billed IT MW (v0.9e) - hurdles, NOT contract prints
 ("H1", "Street FY27 EX-MICROSOFT / FY27 base billed ex-Vineland (PEER COMPARATOR)", (LOCK["STREET"] - LOCK["MSFT_REV"]) / (LOCK["FY27_BASE"] - LOCK["VIN_MW"])),
 ("H2", "MEMO: $23.13M = Street FY27 \u00f7 our 531.8 billed MW, all customers [D] (= model Drivers!B10; not a lock)", LOCK["STREET"] / LOCK["FY27_BASE"]),
 ("H3", "Street FY27 ex-Microsoft / FY27 HIGH billed ex-Vineland", (LOCK["STREET"] - LOCK["MSFT_REV"]) / (LOCK["FY27_HIGH"] - LOCK["VIN_MW"])),
 ("H4", "Street FY27 blended / FY27 HIGH billed - memo only", LOCK["STREET"] / LOCK["FY27_HIGH"]),
 ("H5", "YE26 guide low $7B / YE26 base billed 257.7", LOCK["GUIDE_LO"] / LOCK["YE26_BASE"]),
 ("H6", "YE26 guide high $9B / YE26 high billed 284.2", LOCK["GUIDE_HI"] / LOCK["YE26_HIGH"]),
]
HV = {h[0]: h[2] for h in HURDLES}
assert abs(HV["H1"] - LOCK["HURDLE"]) < 1e-9 and abs(HV["H2"] - LOCK["BLENDED"]) < 1e-9

# ---------------------------------------------------------------- the one-page benchmark rows
# Inline markup: {{text}} -> bold red (print >= $20M/MW); tags in [ ].
f1 = lambda x: f"{x:.1f}"
ROWS = [
 dict(peer="CoreWeave", site="key site: Denton TX (Core Scientific)",
  usd="n/d for every named deal: Meta $14.2B/6.2y and $21.0B/6.7y, OpenAI $22.4B (3 orders), Jane Street $6B, Anthropic (no value) - no MW disclosed [R]. "
      "KPI: {{~$40M/MW-yr}} on 3-6-month deals\u00b9 [R]. Context: Denton colocation ~$1.4M/MW-yr (shell+power, not compute) [D].",
  basis="Contracts: n/a (no MW). $40M KPI: 'MW' unstated; CRWV active power reads IT [E].",
  mw="Contracted 3,700 (Jun-30) -> ~4,200 (Aug-11) [R]; active 1,500 [R] = 2.5x [D]. Denton ~260 IT active/billed [D on R].",
  power="Denton Municipal Electric PPA (~394 MW); ERCOT: 297 MW energized + 74 MW validated, both outside Batch Zero [R]. Standby gensets only [R].",
  permit="TCEQ PBR 179143 standby engines [R]; no Title V/NSR [R, absence]. City moratorium: 2nd hearing Oct 27, vote >= Dec 1; city attorney: need not stop existing projects [P].",
  src="$: 8-Ks 9/30/25, 4/9/26; PR 4/15/26; convert deck 8-K 9/17/26 sl.7. MW: Q2'26 PR 8/11/26. Pwr/permit: CORZ 8-K 9/10/26; TCEQ PBR 179143, Title V/NSR lists 10/8; Community Impact 9/24."),
 dict(peer="IREN", site="key site: Childress TX (Microsoft Horizons)",
  usd=f"Microsoft $9.7B/5y/200 IT MW = 9.7 [D]; NVIDIA $3.4B/5y/~60 MW = {VAL['NEW1']:.1f} [D] -> named range 9.7-{VAL['NEW1']:.1f}. "
      "KPI: recent 3-yr deals {{>$20M}} and discussions {{~$25M}} (unsigned)\u00b2 [R].",
  basis="Microsoft: IT (company). NVIDIA: 60 MW unlabeled. KPIs: 'per MW (IT)'.",
  mw="Contracted AI ~260 IT (200 + ~60) [D]; operating AI ~40 MW at Jun-30 [R]. Childress now ~75 active (50-100) / 50 billed IT (H1 accepted 8/13) [D on R].",
  power="AEP Texas / ERCOT; 750 MW grid, existing energized load incl. legacy mining; IREN-owned substations [R]. No on-site generation found.",
  permit="TCEQ Title V #4951 'Administratively Complete', not issued; inside the Sep 21 Texas data-center permit pause [R/P]. No NSR/PSD [R, absence].",
  src="$: 8-K 11/3/25; NVIDIA PR 5/26; FY26 results 8/27/26. MW: 10-K FY26; 8-K 8/13/26. Pwr/permit: IREN PR 7/24/24; TCEQ Title V list 10/8; Gov. Abbott 9/21."),
 dict(peer="Oracle (OCI)", site="key site: Port Washington WI (Abilene counted under Crusoe)",
  usd=f"OpenAI ~$300B/5y/4.5 GW = {VAL['O2']:.1f} stated / {VAL['O2-IT']:.1f} IT read\u00b3 [P]. Meta ~$20B, Tencent ~$7B/5y: MW n/d [P].",
  basis="4.5 GW unlabeled; ~353 GPUs/MW implies gross -> IT read = /1.2 [D].",
  mw="Contracted: n/d as filed; '>10 GW secured' (call) [R]; OpenAI 4.5 GW [P]. Installed base n/d; 850 MW delivered in FQ1 FY27 incl. Abilene 618 [R]; ~2,050 cumulative since Jun-25 [E]. Port Washington 0 [R].",
  power="Port Washington: We Energies / ATC; 0 MW connected; ATC Ozaukee upgrade built Dec-26 to Dec-27; partial power ~Oct-Dec 2027 (Aterio) [R/P]. 87 MW backup diesel [R].",
  permit="WDNR air 25-MIN-094 final 6/12/26 [R]. ATC grid case at PSC refiled; completeness decision ~Oct 19 [P].",
  src="$: WSJ 9/10/25; CNBC 10/16/25; FT via DCD 10/1/26. MW: FQ1 FY27 call 9/10/26. Pwr/permit: PSC docs 593736/597276; WDNR permit; TechRadar 10/4; Aterio 9/30."),
 dict(peer="Nscale", site="key sites: Monarch WV; Kvandal NO",
  usd=f"Anthropic up to $44.6B/6y/460 IT = {VAL['O1']:.1f} [P]\u2074; Microsoft up to $43.8B/5.5y/~596 IT = {VAL['P2']:.1f} [P]\u2074; Narvik $6.2B/5y/~106 IT = {VAL['P2a']:.1f} [D] -> range {VAL['P2a']:.1f}-{VAL['O1']:.1f}.",
  basis="IT: Monarch 460 IT per S-1; Microsoft IT MW = GPUs / 491 (GB300) or 320 (Rubin) per IT MW [D].",
  mw="Active + contracted 1,370 IT (8/31) [R]; active <= 81 (5-site sum, upper bound) [D] = ~17x. Monarch 460 contracted / 0 active [R].",
  power="Monarch: off-grid gas microgrid, 864 x 2.5 MW engines (2.16 GW nameplate), no utility interconnect [R]. Kvandal: 230 MW Statnett allocation; Vattenfall PPA from 2027 [R].",
  permit="WVDEP air R14-0042 draft 9/29, comments close Oct 30 [R]; stormwater order 7/14 [P]. Kvandal: NVE fine warning, deadline Nov 15 [P]. Anthropic financing not committed [R].",
  src="$: S-1 9/18/26; Aker-Nscale PR 9/17/25; Bloomberg/Reuters (term). MW: S-1. Pwr/permit: WVDEP R14-0042 app/draft; Nscale site page; NRK 9/29."),
 dict(peer="Lambda", site="key sites: Kansas City MO; Beacon Point TX",
  usd=f"Anthropic ~$35B/6y/~350 MW = {VAL['NEW2']:.1f} [P]\u2075 (comp-band top). Microsoft 'multibillion': no value/MW [R]. NVIDIA leaseback $1.3B/4y/10k GPUs ($3.71/GPU-hr): MW n/d [P].",
  basis="350 MW unlabeled (IT vs gross unstated).",
  mw="Company totals n/d (no S-1). Pieces: Beacon Point ~350 [P]; KC 24 (->100+) [R]. Last active 47 MW (Q2'25) [P]; KC 0-24 (base 12) [E]. Backlog $50B (Sep-26)\u2075 [P].",
  power="KC: Evergy (full supply late summer 2026) + 5 x 3 MW rental prime-power engines Jun-Oct (+ up to 6 x 3) [R]. Beacon Point: Hut 8 build, NVIDIA holds lease [P].",
  permit="KC: MO DNR de minimis 102025-010 issued 10/29/25 [R]; no go-live confirmation. Beacon Point: n/d.",
  src="$: Reuters/WSJ/Bloomberg 8/31/26; Lambda PR 11/3/25; DCD 9/8/25. MW: The Information via press; Lambda PR 10/28/25; WSJ/TechCrunch 10/6. Pwr/permit: MO DNR permit + 4/15/26 letter; KSHB."),
 dict(peer="Crusoe", site="key site: Abilene TX (Oracle/OpenAI tenant)",
  usd="n/d for every named deal: Oracle Abilene lease, Microsoft 900 MW, Meta ~1.6 GW (values n/d); Jane Street $13B/5y (MW n/d) [P]. "
      f"Company TCV/MW: {{{{~${CR_TCV_PER_MW:.1f}M per gross MW}}}}, whole term, not per year\u2076 [D].",
  basis="Gross (company states gross MW); Microsoft 672 IT.",
  mw="Contracted 6 GW+ gross / ~1 GW gross operational [R] = 6x [D]. Abilene: 618 delivered (6 of 8 bldgs, 75%) [R]; active 618 (421-618) [D on R].",
  power="ERCOT / AEP via Lancium 138 kV (200 MW) + 345 kV (1,000 MW) substations [R]; ~360 MW on-site gas turbines, backup per Crusoe [R/P].",
  permit="TCEQ 177263 issued [R]. 182126 / PSDTX1688 (Bldgs 9-10, up to ~1.3 GW gas) Technically Complete, not issued, frozen by the Sep 21 pause; TCEQ report Oct 19, ERCOT Dec 10, PUCT Dec 17 [R/P].",
  src="$: Crusoe PRs 9/17/26, 3/27/26, 5/21/25; Bloomberg via DCD 6/19, 9/4. MW: Series F PR; Oracle call 9/10. Pwr/permit: Mortenson; TCEQ 177263/182126; NSR list 10/8; Gov. Abbott 9/21."),
 dict(peer="Bitdeer AI", site="not one of the six; added Oct 9 per GPU Market lock (A201 Johor MY)",
  usd=f"A201: >$1.7B/5y / >70% of ~21.7 MW (~{0.70*21.7:.1f} MW) = {{{{<= ${VAL['BD-A201']:.1f}M ceiling}}}} [D on R; basis P]\u2079 (15.7 if 100% contracted). A102: >$800M/5y/9.5 MW = >{VAL['BD-A102']:.1f} [D]. A103 Cyberjaya 60 MW IT, 8-yr (10/8): no value [R].",
  basis="Primary releases do not label MW; 'critical IT' only in secondary sources (Datatech Times, ArusGrid) [P].",
  mw="Contracted: A201 >= 15.2 of 21.7, A102 9.5, A103 60 IT (~85) [R/D]; energized 0 at these three (A201 Q1'27, A102 Q1'27, A103 Q1'28) [R]. Company AI-cloud active MW n/d.",
  power="Energization targets only (A201 Q1'27; A103 Q1'28) [R]; utility / interconnection n/d; company 'secured 333.5 MW' [R].",
  permit="n/d (Malaysian permits not in team files).",
  src="$: Bitdeer PRs 9/29/26, 8/31/26; A103 release 10/8 (TMCnet); market/contract_economics.xlsx MidTerm_20M_Search r24-25/33/36; Needham 10/2, 10/7 (secondary)."),
 dict(peer="Nebius", site="model v0.9e locks; key site: Vineland NJ",
  usd=f"Microsoft: ${MSFT_E12:.2f}M per IT MW [E, house] (Contracts!E12)\u00b9\u2070. Meta $3B and $12B/5y: MW n/d [R]. "
      "KPI: Q2'26 deals {{$20-25M}} ACV/MW; mgmt comment on Q3-to-date short-term pricing {{>$40M}} (Q2 letter + call, 8/12/26)\u2077 [R]. "
      f"LOCKED hurdles (v0.9e, per billed IT MW): non-MSFT {{{{${MC['HURDLE']['v']:.2f}M}}}}; clear-spot {{{{${MC['CLEAR_BASE']['v']:.2f}M}}}} base / {{{{${MC['CLEAR_HIGH']['v']:.2f}M}}}} high; turn-positive ${MC['TP_BASE']['v']:.2f}M / ${MC['TP_HIGH']['v']:.2f}M\u2078 [D].\n"
      "Memo (blended, not a hurdle lock): {{$23.13M}} = Street FY27 \u00f7 our 531.8 billed MW, all customers [D].",
  basis="Microsoft MW never stated by Nebius/Microsoft (300 MW = Vineland site scalability, 2024 20-F); house uses 240 IT MW est. [E]; Q2 KPI basis undefined; hurdles = billed IT MW.",
  mw=f"Contracted > 3.5 GW (lower bound; 5 GW YE26 target) [R]; energized {LOCK['EN_Q1']:.1f} -> {LOCK['EN_Q2']:.0f} (Mar -> Jun-26) [E]; billed YE26 {LOCK['YE26_BASE']:.1f} / {LOCK['YE26_HIGH']:.1f}, FY27 avg {LOCK['FY27_BASE']:.2f} / {LOCK['FY27_HIGH']:.2f} [D]; connected guide YE26 800 MW-1 GW [R].",
  power="Vineland: 62 rented gensets (122.9 MW) + ~5 MW grid = ~107 MW IT ceiling; Bloom/LNG for the rest, no energization date [R/P/E].",
  permit="Gensets unpermitted: NJDEP order 9/22, $1.07M penalty, 45-day permit deadline ~early/mid Nov [R]. City stop-work on Bloom/LNG (8/6, 8/10); amended site plan approved 8/17; Oct 14 Planning Board: no DataOne item [R].",
  src="$: 6-Ks 9/8/25, 11/12/25, 3/16/26; Q2'26 letter + call 8/12/26; model v0.9e Drivers!B6/B9, Debate_Grid!A10/A13/A18. MW: Capacity Quarterly MW M130/N130; Q1/Q2'26 letters. Pwr/permit: NJ Gov. 9/22; capacity/site_verification s.1."),
]
FOOT = [
 "\u00b9 CoreWeave ~$40M/MW-yr (convert deck 9/17/26 sl.7): annualized revenue on 3-6-month contracts after a ~70% July SKU price rise. It is high because of the short term and GPU-inclusive spot pricing; marginal deals only (fleet $8.24M per avg active MW, Q2'26). MW basis unstated.",
 "\u00b2 IREN >$20M (recent 3-yr deals) / ~$25M (discussions, unsigned), per MW (IT): high because of the 3-yr tenor (vs the 5-yr Microsoft anchor at $9.7M) and GPU-inclusive cloud pricing (~2-yr payback). Portfolio KPI; no per-deal MW or counterparty.",
 "\u00b3 Oracle-OpenAI ~$300B / ~5 yrs / 4.5 GW is WSJ only (P); Oracle filings do not name it. IT read divides 4.5 GW by 1.2.",
 "\u2074 Nscale: TCV and MW are from the S-1 (R), but the term years are press (Anthropic 6 yrs, Bloomberg/Reuters) or a 5-6-yr tranche midpoint (Microsoft 5.5), so both prints are tagged P.",
 "\u2075 Lambda: $35B, ~350 MW (WSJ/Reuters/Bloomberg, Aug 31) and the 6-yr term (WSJ via TheEnergyMag) are press only and unconfirmed by Lambda or Anthropic; the $50B backlog is from an investor letter reported by WSJ (Oct 6). No S-1/DRS. Medium confidence.",
 f"\u2076 Crusoe ~${CR_TCV_PER_MW:.1f}M is whole-term TCV ($140B+) per GROSS contracted MW (6 GW+), not per year, and mixes facility leases, cloud and JV capacity. Not comparable to $/MW-yr (a 5-yr term would be ~${CR_TCV_PER_MW/5:.1f}M/MW-yr, illustrative).",
 "\u2077 Nebius $20-25M (four Q2'26 deals) and >$40M: high because of 1-3-yr and up-to-6-month tenors vs 5-yr anchors and GPU-inclusive full-stack ACV; MW basis undefined. Q3 is NOT reported: '>$40M' is a mgmt comment on Q3-to-date short-term pricing - Q2'26 letter KPI chart (6-K Ex.99.2, 8/12/26, sec.gov/Archives/edgar/data/1513845/000110465926094568/tm2622968d1_ex99-2.htm) and CEO Volozh, Q2'26 call 8/12/26 ('negotiating ... $40-$50 million per megawatt ... one such deal has been signed'; stockanalysis.com/stocks/nbis/transcripts/655490-q2-2026/).",
 f"\u2078 Nebius hurdles = REQUIRED revenue per billed IT MW, not contract prints (GPU-inclusive full-stack revenue). Model v0.9e cells: non-MSFT {MC['HURDLE']['v']:.2f} = Drivers!B9 = (12,300-2,643.8)/(531.84-182.33); clear-spot base {MC['CLEAR_BASE']['v']:.2f} = Drivers!B6 (= Judge_Slide!B8, PT $237.15) / high {MC['CLEAR_HIGH']['v']:.2f} = Debate_Grid!A18; turn-positive base {MC['TP_BASE']['v']:.2f} = Debate_Grid!A13 / high {MC['TP_HIGH']['v']:.2f} = Debate_Grid!A10. Separate memo line (not a hurdle): $23.13M = Street FY27 \u00f7 our 531.8 billed MW, all customers [D] = Drivers!B10 = $12,300M / 531.84 (= Judge_Slide!B16); includes the Microsoft anchor, unlike the ~$27.6M non-MSFT hurdle.",
 f"\u2079 Bitdeer A201 <= ${VAL['BD-A201']:.1f}M: a CEILING because contracted MW is a floor (>70% of ~21.7 MW; 100% = 15.7); GPU-inclusive GB300 cloud (not a powered shell); 5-yr; customers unnamed; MW basis unstated in Bitdeer's releases ('critical IT' per secondary sources only, P). Source: GPU Market agent market/contract_economics.xlsx MidTerm_20M_Search (re-read Bitdeer PR 9/29/26).",
 f"HEADLINE: Only one contract with disclosed MW reaches $20M/MW: Bitdeer, ~15 MW, a <= $22.4M ceiling, basis from secondary sources. COMP BAND unchanged at $9.7-16.7M per MW-yr (low IREN-Microsoft 9.7, IT; top Lambda-Anthropic {BAND['NEW2']:.1f}, P, Medium); Bitdeer's ceiling sits outside the band as a near-miss (model Drivers!B13). ",
 f"\u00b9\u2070 Nebius-Microsoft ${MSFT_E12:.2f}M per IT MW [E, house] = model Contracts!E12 = $17,392.9M / (5 x 240 IT MW est.); agrees with the Drivers 14.5 input (Inputs!F111). No MW stated by Nebius/Microsoft.",
 "Tags: R = reported/primary (filing, company release, earnings call, regulator record); D = derived by us from R inputs; P = press-only, or derived with any press input; E = estimate (team model/alt-data); n/d = not disclosed. Red bold = print at or above $20M per MW. "
 "Active = in service as evidenced (site appendices: Denton, Childress, Abilene, Port Washington). Nebius Q3'26 results ~Nov 10 (est.). Full sources/URLs: peers/benchmark/peer_benchmark.xlsx (Sources tab).",
]

# ---------------------------------------------------------------- XLSX
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E8C")
RED = Font(bold=True, color="C00000"); thin = Side(style="thin", color="BFBFBF"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical="top")
def hdr(ws, row, vals):
    for j, v in enumerate(vals, 1):
        cl = ws.cell(row, j, v); cl.font = H; cl.fill = HF; cl.alignment = WR; cl.border = BD
def strip(s): return s.replace("{{", "").replace("}}", "")
def tags_in(s): return ", ".join(sorted(set(t for t in re.findall(r"\[([^\]]+)\]", s))))

wb = Workbook()
ws = wb.active; ws.title = "README"
readme = [
 f"Peer benchmark - Nebius vs CoreWeave, IREN, Oracle, Nscale, Lambda, Crusoe (+ Bitdeer AI row, Oct 9). As of {ASOF}. Built by peers/benchmark/build_peer_benchmark.py (Nebius Peer KPIs).",
 "Question: disclosed $ per IT MW-yr (TCV / term / MW for each named contract), contracted vs energized MW, power and permit status; each cell sourced and tagged.",
 "Tabs: Benchmark (one row per company, value + tag + source per field) | Contracts (named contracts; $/MW-yr = live formula TCV/term/MW) | KPI_Prints (company $/MW statements) | High_Prints (every print >= $20M/MW with why) | Nebius_Locks (model v0.9e, live hurdle formulas) | Comp_Band (MIN/MAX formulas) | Sources.",
 "Tags: R reported/primary; D derived by us from R inputs; P press-only or derived with any press input (footnoted); E estimate; n/d not disclosed. Red bold = >= $20M per MW.",
 "Reused (read-only): peers/comparison.xlsx Contracts_$MWyr + Sites + Nebius_Inputs; peers/price_or_mw.xlsx; peers/private/private_peers.*; peers/sites/appendix_* + site_verification.md; peers/{coreweave,iren,oracle} packs; parts/contract_economics.xlsx MidTerm_20M_Search; capacity/site_verification.md; model v0.9e via peers/nebius_locks.py.",
 "HEADLINE: Only one contract with disclosed MW reaches $20M/MW: Bitdeer, ~15 MW, a <= $22.4M ceiling, basis from secondary sources. (Bitdeer AI row added Oct 9 per coordinator, from the GPU Market agent's locked position in market/contract_economics.xlsx MidTerm_20M_Search.)",
 "Web refresh Oct 9: Lambda-Anthropic still unconfirmed by the companies (Reuters/Bloomberg/WSJ only) -> stays P; Crusoe 6 GW+/~1 GW still the latest company figure (Series F PR 9/17); Bitdeer A201 PR 9/29 re-read (>$1.7B / 5 yrs / >70% of ~21.7 MW, MW unlabeled). Bitdeer A103 Cyberjaya 60 MW (Oct 8, 8-yr, no value) shown as context.",
 "Oct 9 corrections: (1) Bitdeer row added (red). (2) Nebius '>$40M' relabelled as a mgmt comment on Q3-to-date short-term pricing (Q2'26 letter KPI chart + CEO on Q2 call, both 8/12/26), not reported Q3 results. (3) Nebius headline = LOCKED hurdles verified to model v0.9e cells (Nebius_Locks 'Model cells' block); blended $23.13M kept as memo only (Drivers!B10).",
 "Comp band: unchanged $9.7-16.7M (asserted from comparison.xlsx Contracts_$MWyr P1/NEW1/P2/O1/NEW2/O2); Bitdeer's <= $22.4M ceiling is outside the band (near-miss, as in model Drivers!B13).",
 "Nebius Q3'26 results date: ~Nov 10 (est.), not announced.",
]
for i, t in enumerate(readme, 1): ws.cell(i, 1, t).alignment = Alignment(wrap_text=True)
ws.column_dimensions["A"].width = 180

wsB = wb.create_sheet("Benchmark")
wsB["A1"] = f"One-page peer benchmark (as of {ASOF}). Text cells carry inline [tags]; tag columns list the tags used; source columns cite documents (keys -> Sources tab)."
wsB["A1"].font = Font(bold=True, color="1F4E8C", size=12)
cols = ["Company", "Key site(s)", "Disclosed $ per IT MW-yr (named contracts) + KPI prints", "$ tags", "MW basis for $ figures", "Print >= $20M/MW?",
        "Contracted vs energized (active) MW", "MW tags", "Power status (utility/ISO, interconnection, on-site gen)", "Power tags",
        "Permit status", "Permit tags", "Sources (short)", "Source keys"]
hdr(wsB, 3, cols)
SKEYS = {"CoreWeave": "CW-META1|CW-META2|CW-OAI|CW-JS|CW-DECK|CW-Q2|CW-SITE", "IREN": "IR-MSFT|IR-NV|IR-FY26|IR-SITE",
         "Oracle (OCI)": "OR-OAI|OR-META|OR-TCNT|OR-CALL|OR-SITE", "Nscale": "NS-S1|NS-TERM|NS-NARV|NS-SITE",
         "Lambda": "LM-ANTH|LM-MSFT|LM-NV|LM-BKLG|LM-MW|LM-SITE", "Crusoe": "CR-F|CR-DEALS|CR-SITE|OR-CALL",
         "Nebius": "NB-MSFT|NB-META|NB-Q2|NB-CALL|NB-LOCK|NB-MODEL|NB-MW|NB-SITE|MKT-20", "Bitdeer AI": "BD-A201|BD-MKT|BD-A102|BD-A103"}
for i, r in enumerate(ROWS, 4):
    flag = "YES - " + "; ".join(re.findall(r"\{\{(.+?)\}\}", r["usd"])) if "{{" in r["usd"] else "no"
    vals = [r["peer"], r["site"], strip(r["usd"]), tags_in(r["usd"]), r["basis"], flag, r["mw"], tags_in(r["mw"]), r["power"], tags_in(r["power"]),
            r["permit"], tags_in(r["permit"]), r["src"], SKEYS[r["peer"]]]
    for j, v in enumerate(vals, 1):
        cl = wsB.cell(i, j, v); cl.alignment = WR; cl.border = BD
    if flag != "no": wsB.cell(i, 6).font = RED; wsB.cell(i, 3).font = Font(color="C00000")
nb = 4 + len(ROWS) + 1
wsB.cell(nb, 1, "Footnotes").font = Font(bold=True)
for k_, t in enumerate(FOOT, nb + 1): wsB.cell(k_, 1, t)
for col, w in zip("ABCDEFGHIJKLMN", [12, 22, 60, 9, 30, 22, 50, 9, 48, 9, 50, 9, 50, 30]): wsB.column_dimensions[col].width = w
wsB.freeze_panes = "C4"

wsC = wb.create_sheet("Contracts")
wsC["A1"] = "Named contracts: $M per MW-yr = TCV ($M) / term (yrs) / MW - LIVE formula; blank where any input is n/d. Reused from peers/comparison.xlsx Contracts_$MWyr (IDs kept) unless noted. Column O flags >= $20M."
wsC["A1"].font = Font(bold=True, color="1F4E8C")
hdr(wsC, 3, ["ID", "Company", "Counterparty", "TCV ($M)", "TCV tag", "Term (yrs)", "Term tag", "MW used", "MW basis", "MW tag",
             "$M per MW-yr (formula)", "Result tag", "IT-basis?", "Note", ">= $20M flag (formula)", "Source keys", "Python check"])
for i, r in enumerate(CONTRACTS, 4):
    cid, peer, cp, tcv, tt, term, tmt, mw, mwb, mwt, rt, sk, note = r
    vals = [cid, peer, cp, tcv if tcv is not None else "n/d", tt, term if term is not None else "n/d", tmt, mw if mw is not None else "n/d",
            mwb, mwt, None, rt, ("IT" if mwb.startswith("IT") else ("gross" if mwb.startswith("gross") else "unlabeled/n.a.")), note, None, sk, VAL[cid]]
    for j, v in enumerate(vals, 1):
        cl = wsC.cell(i, j, v); cl.alignment = WR; cl.border = BD
    if True:
        wsC.cell(i, 11, f'=IF(AND(ISNUMBER(D{i}),ISNUMBER(F{i}),ISNUMBER(H{i})),D{i}/F{i}/H{i},"n/d")')
    wsC.cell(i, 15, f'=IF(ISNUMBER(K{i}),IF(K{i}>=20,"FLAG >=20","<20"),"n/d")')
    wsC.cell(i, 11).number_format = "0.00"; wsC.cell(i, 17).number_format = "0.00"
last = 3 + len(CONTRACTS)
wsC.conditional_formatting.add(f"K4:K{last}", CellIsRule(operator="greaterThanOrEqual", formula=["20"], font=RED))
for col, w in zip("ABCDEFGHIJKLMNOPQ", [8, 11, 30, 10, 8, 9, 10, 9, 34, 8, 12, 8, 10, 36, 12, 18, 10]): wsC.column_dimensions[col].width = w

wsK = wb.create_sheet("KPI_Prints")
wsK["A1"] = "Company-level $ per MW statements (not single named contracts). Values as stated ('>' entered at the floor). >= $20M in red."
wsK["A1"].font = Font(bold=True, color="1F4E8C")
hdr(wsK, 3, ["ID", "Company", "Print", "Low ($M/MW)", "High ($M/MW)", "Tag", "MW basis", ">= $20M flag (formula)", "Why high (if flagged)", "Source keys"])
for i, r in enumerate(KPI_PRINTS, 4):
    pid, peer, lab, lo, hi, tg, basis, sk, why = r
    if pid == "CR-K1": lo = hi = None
    vals = [pid, peer, lab, lo, hi, tg, basis, None, why or "", sk]
    for j, v in enumerate(vals, 1):
        cl = wsK.cell(i, j, v); cl.alignment = WR; cl.border = BD
    if pid == "CR-K1":
        wsK.cell(i, 4, "=140000/6000"); wsK.cell(i, 5, "=D%d" % i)
        wsK.cell(i, 3).value = lab + "  [per MW of TCV, not per yr]"
    wsK.cell(i, 8, f'=IF(ISNUMBER(E{i}),IF(E{i}>=20,"FLAG >=20","<20"),"n/d")')
    for j in (4, 5): wsK.cell(i, j).number_format = "0.0"
lk = 3 + len(KPI_PRINTS)
wsK.conditional_formatting.add(f"D4:E{lk}", CellIsRule(operator="greaterThanOrEqual", formula=["20"], font=RED))
for col, w in zip("ABCDEFGHIJ", [8, 11, 50, 10, 10, 6, 30, 12, 70, 14]): wsK.column_dimensions[col].width = w

wsN = wb.create_sheet("Nebius_Locks")
wsN["A1"] = "Nebius model v0.9e locks (read via peers/nebius_locks.py; asserted against the task brief). Hurdle cells are live formulas."
wsN["A1"].font = Font(bold=True, color="1F4E8C")
hdr(wsN, 3, ["Key", "Item", "Value", "Unit", "Tag", "Source"])
NL = [
 ("MSFT_REV", "Microsoft (Vineland) FY27 revenue", LOCK["MSFT_REV"], "$M", "D", "model Drivers!B32"),
 ("VIN_MW", "Vineland FY27 billed MW-yr (= MS rev / $14.5M)", LOCK["VIN_MW"], "IT MW", "D", "model Drivers!B33"),
 ("STREET", "Street FY27 revenue", LOCK["STREET"], "$M", "E", "model Judge_Slide B72 / Cons_Rev"),
 ("FY27_BASE", "FY27 avg billed IT MW - base", LOCK["FY27_BASE"], "IT MW", "D", "model Drivers!B34"),
 ("FY27_HIGH", "FY27 avg billed IT MW - high", LOCK["FY27_HIGH"], "IT MW", "D", "model Drivers!C34"),
 ("YE26_BASE", "YE26 billed IT MW - base", LOCK["YE26_BASE"], "IT MW", "D", "model Judge_Slide B34"),
 ("YE26_HIGH", "YE26 billed IT MW - high", LOCK["YE26_HIGH"], "IT MW", "D", "model Judge_Slide B35"),
 ("GUIDE_LO", "YE26 ARR guide low", LOCK["GUIDE_LO"], "$M", "R", "Q2'26 letter"),
 ("GUIDE_HI", "YE26 ARR guide high", LOCK["GUIDE_HI"], "$M", "R", "Q2'26 letter"),
 ("EN_Q1", "Energized MW 31-Mar-26", LOCK["EN_Q1"], "MW", "E", "capacity Quarterly MW!M130"),
 ("EN_Q2", "Energized MW 30-Jun-26", LOCK["EN_Q2"], "MW", "E", "capacity Quarterly MW!N130"),
 ("CONTRACTED", "Contracted power (lower bound)", LOCK["CONTRACTED"], "MW", "R", "Q1'26 letter p.6"),
]
for i, r in enumerate(NL, 4):
    for j, v in enumerate(r, 1): wsN.cell(i, j, v).border = BD
row0 = 4 + len(NL) + 1
wsN.cell(row0, 1, "Hurdles ($M per billed IT MW-yr; required, not contract prints)").font = Font(bold=True)
hdr(wsN, row0 + 1, ["Key", "Item", "Value (formula)", "Unit", "Tag", ">= $20M flag", "Python check"])
HF_ = {"H1": "=(C6-C4)/(C7-C5)", "H2": "=C6/C7", "H3": "=(C6-C4)/(C8-C5)", "H4": "=C6/C8", "H5": "=C11/C9", "H6": "=C12/C10"}
for i, (hk, lab, v) in enumerate(HURDLES, row0 + 2):
    for j, x in enumerate([hk, lab, HF_[hk], "$M/IT MW-yr", "D", f'=IF(C{i}>=20,"FLAG >=20","<20")', v], 1):
        cl = wsN.cell(i, j, x); cl.border = BD
    wsN.cell(i, 3).number_format = "0.00"; wsN.cell(i, 7).number_format = "0.00"
wsN.conditional_formatting.add(f"C{row0+2}:C{row0+1+len(HURDLES)}", CellIsRule(operator="greaterThanOrEqual", formula=["20"], font=RED))
wsN.cell(row0 + 3 + len(HURDLES), 1, "House Microsoft $/IT MW-yr = 14.49 [E] (Contracts N1-H = model Contracts!E12 = 17,392.9 / (5 x 240 IT est.)); agrees with the Drivers 14.5 input behind C4/C5. Energized 181.5 -> 193 is E (Capacity base case incl. plug; mixed IT/gross).")
mr = row0 + 5 + len(HURDLES)
wsN.cell(mr, 1, "MODEL CELLS (Oct 9): locked Nebius hurdles verified against model/nebius_model.xlsx v0.9e (read-only; cached values)").font = Font(bold=True)
hdr(wsN, mr + 1, ["Key", "Item", "Model value", "Model cell", "Brief value", "Diff (formula)", "Model formula / note"])
for i, (k_, lab_, brief_, sh_, cl_, tol_, note_) in enumerate(MODEL_CELLS, mr + 2):
    d_ = MC[k_]
    for j, x in enumerate([k_, lab_, d_["v"], d_["ref"], brief_, f"=C{i}-E{i}", ("formula " + d_["formula"][1:] if d_["formula"] != "static value" else "static") + " | " + note_], 1):
        cl = wsN.cell(i, j, x); cl.border = BD; cl.alignment = WR
    wsN.cell(i, 3).number_format = "0.0000"; wsN.cell(i, 6).number_format = "0.0000"
for col, w in zip("ABCDEFG", [12, 60, 14, 16, 12, 12, 90]): wsN.column_dimensions[col].width = w

wsBand = wb.create_sheet("Comp_Band")
wsBand["A1"] = "Comp band (memo set; comparison.xlsx Contracts_$MWyr): MIN/MAX are live formulas. Unchanged $9.7-16.7M."
wsBand["A1"].font = Font(bold=True, color="1F4E8C")
hdr(wsBand, 3, ["ID", "Deal", "$M per MW-yr (IT where available)", "Tag"])
BLAB = {"P1": ("IREN-Microsoft", "D (R inputs)"), "NEW1": ("IREN-NVIDIA", "D (R inputs)"), "P2": ("Nscale-Microsoft", "P (term)"),
        "O1": ("Nscale-Anthropic", "P (term)"), "NEW2": ("Lambda-Anthropic", "P"), "O2": ("Oracle-OpenAI (IT read)", "P")}
for i, cid in enumerate(BAND_IDS, 4):
    wsBand.cell(i, 1, cid); wsBand.cell(i, 2, BLAB[cid][0]); wsBand.cell(i, 3, BAND[cid]).number_format = "0.00"; wsBand.cell(i, 4, BLAB[cid][1])
e_ = 3 + len(BAND_IDS)
wsBand.cell(e_ + 2, 2, "Band low (MIN)"); wsBand.cell(e_ + 2, 3, f"=MIN(C4:C{e_})").number_format = "0.00"
wsBand.cell(e_ + 3, 2, "Band top (MAX)"); wsBand.cell(e_ + 3, 3, f"=MAX(C4:C{e_})").number_format = "0.00"
for col, w in zip("ABCD", [8, 30, 18, 14]): wsBand.column_dimensions[col].width = w

wsH = wb.create_sheet("High_Prints")
wsH["A1"] = "Every print at or above $20M per MW, with why it is high (GPU-inclusive vs powered shell, gross vs IT, short term). Only one contract with disclosed MW reaches $20M/MW: Bitdeer, ~15 MW, a <= $22.4M ceiling, basis from secondary sources."
wsH["A1"].font = Font(bold=True, color="C00000")
hdr(wsH, 3, ["Company", "Print", "$M per MW", "Tag", "Type", "Why high", "Source keys"])
HP = [
 ("CoreWeave", "3-6-month contracts (annualized)", "~40", "R", "company KPI", KPI_PRINTS[0][8], "CW-DECK"),
 ("IREN", "Recent 3-yr contracts", ">20", "R", "company KPI", KPI_PRINTS[2][8], "IR-FY26"),
 ("IREN", "Active discussions (unsigned)", "~25", "R", "company KPI", KPI_PRINTS[3][8], "IR-FY26"),
 ("Crusoe", "TCV $140B+ / 6 GW+ gross (not annual)", f"~{CR_TCV_PER_MW:.1f}", "D", "TCV per MW (not $/MW-yr)", KPI_PRINTS[4][8], "CR-F"),
 ("Nebius", "Q2'26 four landmark deals ACV/MW", "20-25", "R", "company KPI", KPI_PRINTS[5][8], "NB-Q2"),
 ("Nebius", "Mgmt comment on Q3-to-date short-term pricing (Q2 letter + call 8/12/26; not reported Q3)", ">40 (40-50 negotiating)", "R", "mgmt comment", KPI_PRINTS[6][8], "NB-Q2|NB-CALL"),
 ("Bitdeer AI", "A201 Johor off-take, >$1.7B / 5 yrs / >70% of ~21.7 MW (~15.2 MW)", f"<= {VAL['BD-A201']:.1f}", "D (basis P)", "signed contract, disclosed MW (ceiling)", "CEILING: contracted MW is a floor (100% contracted = 15.7); GPU-inclusive GB300 cloud, not powered shell; 5-yr; unnamed customers; MW basis unstated in primary (IT per secondary only). Source: GPU Market agent MidTerm_20M_Search.", "BD-A201|BD-MKT"),
 ("Nebius", "LOCKED non-MSFT hurdle (peer comparator)", f"{MC['HURDLE']['v']:.2f}", "D", "required, not a print", f"Model {MC['HURDLE']['ref']}: (Street $12.3B - Microsoft $2,643.8M) / (531.84 - 182.33) billed IT MW. Above every disclosed contract and above mgmt's $20-25M.", "NB-MODEL"),
 ("Nebius", "LOCKED clear-spot BASE (PT = $237.15)", f"{MC['CLEAR_BASE']['v']:.2f}", "D", "required, not a print", f"Model {MC['CLEAR_BASE']['ref']} (= Judge_Slide!B8): uncontracted GB300 price per billed IT MW at which the base path's self-consistent PT equals the $237.15 close.", "NB-MODEL"),
 ("Nebius", "LOCKED clear-spot HIGH", f"{MC['CLEAR_HIGH']['v']:.2f}", "D", "required, not a print", f"Model {MC['CLEAR_HIGH']['ref']} ('high-path clear-spot'; Drivers!A58).", "NB-MODEL"),
 ("Nebius", "MEMO (separate from the non-MSFT hurdle): $23.13M = Street FY27 \u00f7 our 531.8 billed MW, all customers [D]", f"{MC['BLENDED']['v']:.2f}", "D", "memo", f"Model {MC['BLENDED']['ref']} = INDEX(Cons_Rev,1,2027-FYBase)/B34 = $12,300M / 531.84 (= Judge_Slide!B16 'supporting line'); includes the $14.5M Microsoft anchor.", "NB-MODEL"),
 ("Nebius", "YE26 guide low / YE26 base billed", f"{HV['H5']:.2f}", "D", "required, not a print", "Run-rate ARR guide ($7B) over YE26 base billed 257.7 IT MW.", "NB-LOCK"),
 ("Nebius", "YE26 guide high / YE26 high billed", f"{HV['H6']:.2f}", "D", "required, not a print", "$9B over 284.2 IT MW.", "NB-LOCK"),
 ("Nebius", "Street FY27 ex-MSFT / FY27 HIGH billed", f"{HV['H3']:.2f}", "D", "required, not a print", "High billed path (659.16) ex-Vineland.", "NB-LOCK"),
]
for i, r in enumerate(HP, 4):
    for j, v in enumerate(r, 1):
        cl = wsH.cell(i, j, v); cl.alignment = WR; cl.border = BD
    wsH.cell(i, 3).font = RED
wsH.cell(4 + len(HP) + 1, 1, "Not >= $20M: Nebius-Microsoft house 14.49 [E] (model Contracts!E12). Turn-positive (locked, < $20M): base {:.2f} (Debate_Grid!A13) / high {:.2f} (Debate_Grid!A10). Bitdeer A102 >16.8 (floor).".format(MC["TP_BASE"]["v"], MC["TP_HIGH"]["v"]))
for col, w in zip("ABCDEFG", [11, 40, 10, 6, 22, 90, 12]): wsH.column_dimensions[col].width = w

wsS = wb.create_sheet("Sources")
hdr(wsS, 1, ["Key", "Document", "URL / path", "Accessed"])
for i, (k_, (doc, url)) in enumerate(SRC.items(), 2):
    wsS.cell(i, 1, k_); wsS.cell(i, 2, doc).alignment = WR; wsS.cell(i, 3, url); wsS.cell(i, 4, "2026-10-05..10-09 (team packs); Lambda/Crusoe re-checked 2026-10-09")
for col, w in zip("ABCD", [11, 110, 90, 30]): wsS.column_dimensions[col].width = w
wb.save(XLSX)

# recalc via team script (LibreOffice); keep the recalculated copy so cached values exist; 0 errors required
r_ = subprocess.run([sys.executable, "/workspace/nebius/model/recalc.py", XLSX], capture_output=True, text=True, cwd="/workspace/nebius/model")
print(r_.stdout.strip()); assert r_.returncode == 0, r_.stderr
assert "done" in r_.stdout and not [l for l in r_.stdout.splitlines() if not l.startswith("done")], "recalc errors: " + r_.stdout
rc = "/tmp/rc/peer_benchmark.xlsx"
chk = openpyxl.load_workbook(rc, data_only=True)
for i, r in enumerate(CONTRACTS, 4):
    v = chk["Contracts"].cell(i, 11).value
    if VAL[r[0]] is None: assert v == "n/d", (r[0], v)
    else: assert abs(v - VAL[r[0]]) < 1e-9, (r[0], v, VAL[r[0]])
for i, (hk, lab, v) in enumerate(HURDLES, row0 + 2):
    assert abs(chk["Nebius_Locks"].cell(i, 3).value - v) < 1e-9, hk
assert abs(chk["Comp_Band"].cell(e_ + 2, 3).value - 9.7) < .05 and abs(chk["Comp_Band"].cell(e_ + 3, 3).value - 16.667) < .01
nform = sum(1 for wsx in openpyxl.load_workbook(XLSX) for row in wsx.iter_rows() for cl in row if isinstance(cl.value, str) and cl.value.startswith("="))
shutil.copy(rc, XLSX)
# LibreOffice drops nothing material but re-apply column widths is unnecessary; verify formulas survived
nform2 = sum(1 for wsx in openpyxl.load_workbook(XLSX) for row in wsx.iter_rows() for cl in row if isinstance(cl.value, str) and cl.value.startswith("="))
print(f"formulas: {nform} written, {nform2} after recalc copy")
assert nform2 == nform

# ---------------------------------------------------------------- DOCX (one landscape page)
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = Inches(11), Inches(8.5)
for side in ("left_margin", "right_margin"): setattr(sec, side, Inches(0.3))
sec.top_margin = Inches(0.28); sec.bottom_margin = Inches(0.25)
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(6)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
pf = st.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
NAVY = RGBColor(0x1F, 0x4E, 0x8C); REDC = RGBColor(0xC0, 0x00, 0x00)
def para(text, size=6, bold=False, color=None, after=0):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    add_runs(p, text, size, bold, color); return p
def add_runs(p, text, size=6, bold=False, color=None):
    for k_, part in enumerate(re.split(r"(\{\{.+?\}\})", text)):
        if not part: continue
        hot = part.startswith("{{")
        rr = p.add_run(part[2:-2] if hot else part)
        rr.font.size = Pt(size); rr.bold = bold or hot
        if hot: rr.font.color.rgb = REDC
        elif color: rr.font.color.rgb = color
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)
def cell_margins(tbl, m=30):
    tblPr = tbl._tbl.tblPr; mar = OxmlElement("w:tblCellMar")
    for s in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{s}"); e.set(qn("w:w"), str(m if s in ("left", "right") else 10)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)
para("Peer benchmark: $ per IT MW-yr, contracted vs energized MW, power and permits (CoreWeave, IREN, Oracle, Nscale, Lambda, Crusoe, Bitdeer vs Nebius)", 9.5, True, NAVY)
para(f"Nebius Peer KPIs | as of {ASOF} | public sources only | Nebius on model v0.9e locks | $ per MW-yr = TCV / term / MW for each named contract ($M) | comp band $9.7-16.7M unchanged | "
     "R reported/primary, D derived from R, P press (footnoted), E estimate, n/d not disclosed | {{red bold}} = print >= $20M/MW", 6.3, False, RGBColor(0x40, 0x40, 0x40), after=1)
para("Only one contract with disclosed MW reaches $20M/MW: Bitdeer, ~15 MW, a <= $22.4M ceiling, basis from secondary sources.", 7.5, True, REDC, after=2)
HDRS = ["Company", "Disclosed $M per IT MW-yr (named contracts) + company KPI prints", "MW basis of the $ figures", "Contracted vs energized (active) MW",
        "Power (utility / ISO, interconnection, on-site gen)", "Permit status", "Sources (full list: xlsx Sources tab)"]
W = [0.78, 2.45, 1.05, 1.75, 1.55, 1.55, 1.27]
tbl = doc.add_table(rows=1 + len(ROWS), cols=len(HDRS)); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl.style = "Table Grid"; tbl.autofit = False
cell_margins(tbl)
for j, h_ in enumerate(HDRS):
    cl = tbl.rows[0].cells[j]; cl.width = Inches(W[j]); shade(cl, "1F4E8C")
    p = cl.paragraphs[0]; rr = p.add_run(h_); rr.bold = True; rr.font.size = Pt(6.3); rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
for i, r in enumerate(ROWS, 1):
    cells = tbl.rows[i].cells
    for j in range(len(HDRS)): cells[j].width = Inches(W[j])
    p = cells[0].paragraphs[0]; rr = p.add_run(r["peer"]); rr.bold = True; rr.font.size = Pt(7); rr.font.color.rgb = NAVY
    p2 = cells[0].add_paragraph(); add_runs(p2, r["site"], 6, color=RGBColor(0x55, 0x55, 0x55))
    for j, key in enumerate(["usd", "basis", "mw", "power", "permit"], 1):
        add_runs(cells[j].paragraphs[0], r[key].replace(" -> ", " \u2192 "), 6.6)
    add_runs(cells[6].paragraphs[0], r["src"], 5.8, color=RGBColor(0x44, 0x44, 0x44))
    if r["peer"] == "Nebius":
        for cl in cells: shade(cl, "E8EEF7")
# repeat widths at grid level (LibreOffice honours tblGrid)
grid = tbl._tbl.tblGrid
for j, gc in enumerate(grid.findall(qn("w:gridCol"))): gc.set(qn("w:w"), str(int(W[j] * 1440)))
doc.add_paragraph().paragraph_format.space_after = Pt(1)
for t in FOOT: para(t, 6.0, color=RGBColor(0x30, 0x30, 0x30))
doc.save(DOCX)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", D, DOCX], check=True, capture_output=True, timeout=180)
pages = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", PDF], capture_output=True, text=True).stdout).group(1))
print("pdf pages:", pages); assert pages == 1, f"PDF is {pages} pages"
subprocess.run(["pdftoppm", "-png", "-r", "170", "-singlefile", PDF, PNG[:-4]], check=True)
print("wrote", XLSX, DOCX, PDF, PNG)
