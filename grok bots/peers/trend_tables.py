import openpyxl,sys
f, ids = sys.argv[1], sys.argv[2].split(',')
wb=openpyxl.load_workbook(f,data_only=True); ws=wb['KPIs']
hdr=[c.value for c in ws[3]]; fq=[c.value for c in ws[2]]
n=ws.max_column-1
rows={ws.cell(row=r,column=1).value:r for r in range(4,ws.max_row+1) if ws.cell(row=r,column=1).value}
cols=list(range(6,n+1))
print('| ID | KPI | Unit | '+' | '.join(hdr[c-1] for c in cols)+' |')
print('|'+'---|'*(3+len(cols)))
ab={'Reported':'R','Derived':'D','Estimated':'E'}
for k in ids:
    r=rows[k]; out=[]
    for c in cols:
        v=ws.cell(row=r,column=c).value; t=ws.cell(row=r+1,column=c).value
        if isinstance(v,(int,float)):
            fmt = '%.0f' if abs(v)>=100 else ('%.1f' if abs(v)>=10 else '%.2f')
            out.append((fmt%v)+('ᴱ' if t=='Estimated' else ''))
        else: out.append('n/d')
    print(f"| {k} | {ws.cell(row=r,column=2).value} | {ws.cell(row=r,column=3).value} | "+' | '.join(out)+' |')
