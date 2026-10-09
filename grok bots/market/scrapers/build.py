import sys,csv,statistics,importlib
sys.path.insert(0,'/workspace/nebius')
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os
rows=[]
for m in sys.argv[1:]: rows+=importlib.import_module(m).R
COLS=['provider','chip','config','pricing_type','term','quoted_price','quoted_unit','usd_per_gpu_hr','region','notes','source_url','retrieved_date']
with open('/workspace/nebius/price_snapshot_2026-10-04.csv','w',newline='') as f:
    w=csv.DictWriter(f,COLS); w.writeheader(); w.writerows(rows)
P='/workspace/nebius/nebius_market.xlsx'
if os.path.exists(P):
    wb=load_workbook(P)
    for n in ['Price_Snapshot_2026-10-04','Price_Summary']:
        if n in wb.sheetnames: del wb[n]
else:
    wb=Workbook(); wb.remove(wb.active)
H=Font(bold=True,color='FFFFFF'); F=PatternFill('solid',fgColor='1F3864')
def hdr(ws,row=1):
    for c in ws[row]: c.font=H; c.fill=F; c.alignment=Alignment(wrap_text=True,vertical='top')
ws=wb.create_sheet('Price_Snapshot_2026-10-04',0)
ws.append(COLS); hdr(ws)
for x in rows: ws.append([x[c] for c in COLS])
for c in ws['H'][1:]: c.number_format='$#,##0.000'
for i,wd in enumerate([14,14,38,26,24,18,26,14,30,60,44,13],1): ws.column_dimensions[get_column_letter(i)].width=wd
ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
# summary
CORE=['on-demand','reserved/committed','spot/preemptible','inference/dedicated endpoint']
CHIPS=['H100','H200','B200','B300','GB200 NVL72','GB300 NVL72','MI300X','MI325X','MI355X']
s=wb.create_sheet('Price_Summary',1)
s.append(['Price summary, $/GPU-hr, snapshot 2026-10-04 (numeric rows only; "contact sales" rows excluded; Nebius superseded pre-Oct-1 column excluded)']); s['A1'].font=Font(bold=True,size=12)
s.append(['Method: provider-level value = median of that provider\'s rows for chip x pricing type (avoids over-weighting providers with many regions/sizes). Peer median = median of non-Nebius provider-level values. Min/max = across all individual numeric rows incl. Nebius. Neocloud peers = all non-hyperscaler providers ex-Nebius; hyperscalers = AWS, Azure, Google Cloud, Oracle OCI. Nebius on-demand = column effective Oct 1, 2026. Spot "from" floors (Nebius) and AWS Capacity Blocks are included as listed; see notes.'])
s.append([])
s.append(['chip','pricing_type','n_rows','n_providers','providers','min_$/GPU-hr','median_all_rows','max_$/GPU-hr','peer_median_ex_Nebius (provider-level)','Nebius_$/GPU-hr','Nebius_vs_peer_median_%','neocloud_peer_median','Nebius_vs_neocloud_%','hyperscaler_median (AWS/Azure/GCP/OCI)','Nebius_vs_hyperscaler_%']); hdr(s,4)
out=[]
for ch in CHIPS:
  for pt in CORE:
    rr=[x for x in rows if x['chip']==ch and x['pricing_type']==pt and isinstance(x['usd_per_gpu_hr'],(int,float))]
    if not rr: continue
    prov={}
    for x in rr: prov.setdefault(x['provider'],[]).append(x['usd_per_gpu_hr'])
    pv={k:statistics.median(v) for k,v in prov.items()}
    peers=[v for k,v in pv.items() if k!='Nebius']
    pm=statistics.median(peers) if peers else None
    nb=pv.get('Nebius')
    vals=[x['usd_per_gpu_hr'] for x in rr]
    HS={'AWS','Azure','Google Cloud','Oracle OCI'}
    neo=[v for k,v in pv.items() if k!='Nebius' and k not in HS]; hs=[v for k,v in pv.items() if k in HS]
    nm=statistics.median(neo) if neo else None; hm=statistics.median(hs) if hs else None
    f=lambda a,b:(a/b-1) if (a is not None and b) else None
    s.append([ch,pt,len(rr),len(pv),', '.join(f'{k} {v:.2f}' for k,v in sorted(pv.items(),key=lambda z:z[1])),min(vals),statistics.median(vals),max(vals),pm,nb,f(nb,pm),nm,f(nb,nm),hm,f(nb,hm)])
for row in s.iter_rows(min_row=5):
    for i in (5,6,7,8,9,11,13): row[i].number_format='$#,##0.00'
    for i in (10,12,14): row[i].number_format='+0.0%;-0.0%'
for i,wd in enumerate([14,28,8,11,70,13,15,13,22,15,16,15,16,18,16],1): s.column_dimensions[get_column_letter(i)].width=wd
s.freeze_panes='C5'
wb.save(P); print(len(rows),'rows saved')
for row in s.iter_rows(min_row=5,values_only=True): print(row)
