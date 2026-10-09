import json
#!/usr/bin/env python3
"""Build /workspace/nebius/parts/demand_signals.xlsx from the CSV/JSON outputs of the demand scrapers.
Every data row carries: value, period, tag (Reported / Derived / Estimated / Observed), source_id (-> Sources sheet), note.
'Observed' = read directly off a public web page / API on the access date (third-party data, not company-reported)."""
import csv, json, os, glob, collections, datetime, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
D=os.path.dirname(os.path.abspath(__file__)); R=f'{D}/raw'; OUT='/workspace/nebius/parts/demand_signals.xlsx'
ACC='2026-10-04'
def latest(pat):
    f=sorted(glob.glob(f'{D}/{pat}')); return f[-1] if f else None
# ---------------- Sources ----------------
SRC=[]
def S(sid,desc,url,acc=ACC,kind='web'):
    SRC.append([sid,desc,url,acc,kind]); return sid
S('GH-NBIS','Nebius Greenhouse job board API (public)','https://boards-api.greenhouse.io/v1/boards/nebius/jobs?content=true')
S('GH-CRWV','CoreWeave Greenhouse job board API (public)','https://boards-api.greenhouse.io/v1/boards/coreweave/jobs?content=true')
S('ASHBY-LAMBDA','Lambda Ashby job board API (public)','https://api.ashbyhq.com/posting-api/job-board/lambda')
S('WB-JOBS','Wayback Machine captures of the same ATS API URLs (CDX list in raw/wayback/cdx_*.json)','https://web.archive.org/cdx/search/cdx?url=boards-api.greenhouse.io/v1/boards/nebius&matchType=prefix')
S('SITES','Team capacity file: Nebius announced sites (capacity/site_list.csv)','/workspace/nebius/capacity/site_list.csv','2026-10-04','team file')
S('20F-25','Nebius 20-F FY2025 (filed 2026-04-30): ~1,500 employees at 12/31/25; product-dev headcount 653->813; customer concentration FY23-25','https://www.sec.gov/Archives/edgar/data/1513845/000110465926052948/nbis-20251231x20f.htm','2026-10-04','SEC')
S('20F-24','Nebius 20-F FY2024: customer concentration FY22-24','https://www.sec.gov/Archives/edgar/data/1513845/000110465925042228/nbis-20241231x20f.htm','2026-10-04','SEC')
S('TRANCO','Tranco list API (30-day Dowdall-combined rank of CrUX, Farsight, Majestic, Cloudflare Radar, Umbrella); daily lists by date','https://tranco-list.eu/api/lists/date/{date}')
S('SW','Similarweb free data endpoint - HTTP 403 from box (blocked); not recorded','https://data.similarweb.com/api/v1/data?domain=nebius.com')
S('OR-PROV','OpenRouter providers page (RSC payload: modelCount, dailyTokens, monthlyTokens per provider)','https://openrouter.ai/providers')
S('OR-NBIS','OpenRouter Nebius Token Factory provider page (chartData: daily tokens by model, last ~90 days)','https://openrouter.ai/provider/nebius')
S('OR-API','OpenRouter public model endpoints API (per-provider pricing)','https://openrouter.ai/api/v1/models/{model}/endpoints')
S('HF-API','Hugging Face models API inference_provider filter (enum of live providers)','https://huggingface.co/api/models?inference_provider=nebius')
S('HF-DOC','Hugging Face Inference Providers doc page for Nebius','https://huggingface.co/docs/inference-providers/providers/nebius')
S('HF-ORG','Hugging Face Nebius org overview API','https://huggingface.co/api/organizations/nebius/overview')
S('NB-STORIES','nebius.com customer stories (server sitemap + each story page)','https://nebius.com/server-sitemap.xml')
S('WB-STORIES','Wayback CDX first capture per nebius.com/customer-stories/* URL','https://web.archive.org/cdx/search/cdx?url=nebius.com/customer-stories/&matchType=prefix&collapse=urlkey')
S('GITHUB','GitHub search API org:<org> (public repos, stars, forks, created)','https://api.github.com/search/repositories?q=org:nebius')
S('HN','Hacker News Algolia search API (stories/comments containing term, by created_at)','https://hn.algolia.com/api/v1/search?query=nebius')
S('SO','Stack Exchange API search/excerpts site=stackoverflow (full-text posts)','https://api.stackexchange.com/2.3/search/excerpts?site=stackoverflow&q=nebius')
S('DOCS-CL','Nebius AI Cloud docs changelog (weekly <Update> blocks, markdown)','https://docs.nebius.com/changelog.md')
S('DISCORD','No public Discord invite found on nebius.com / docs.nebius.com / tokenfactory.nebius.com (only OpenRouter\'s invite on OR pages)','https://nebius.com/')
S('Q324-PR','Nebius Q3 2024 results PR (AI Studio introduced)','https://www.sec.gov/Archives/edgar/data/1513845/000110465924112875/tm2427209d1_ex99-1.htm','2026-10-04','SEC')
S('Q424-PR','Nebius Q4 2024 results PR','https://www.sec.gov/Archives/edgar/data/1513845/000110465925015480/tm257185d1_ex99-1.htm','2026-10-04','SEC')
S('Q125-SHL','Nebius Q1 2025 shareholder letter (AI Studio >60,000 registered users; "still early from a revenue perspective")','https://www.sec.gov/Archives/edgar/data/1513845/000110465925050975/tm2515580d1_ex99-2.htm','2026-10-04','SEC')
S('Q225-SHL','Nebius Q2 2025 shareholder letter (Lightning AI 300k+ members access on-demand/preemptible)','https://www.sec.gov/Archives/edgar/data/1513845/000110465925075028/tm2522866d1_ex99-2.htm','2026-10-04','SEC')
S('Q325-SHL','Nebius Q3 2025 shareholder letter (Token Factory launch)','https://www.sec.gov/Archives/edgar/data/1513845/000110465925109806/tm2530882d1_ex99-2.htm','2026-10-04','SEC')
S('Q425-SHL','Nebius Q4 2025 shareholder letter','https://www.sec.gov/Archives/edgar/data/1513845/000110465926013946/tm266173d1_ex99-2.htm','2026-10-04','SEC')
S('Q126-SHL','Nebius Q1 2026 shareholder letter','https://www.sec.gov/Archives/edgar/data/1513845/000110465926059872/tm2614392d1_ex99-2.htm','2026-10-04','SEC')
S('Q226-SHL','Nebius Q2 2026 shareholder letter (TF workloads >3x; Microsoft all tranches delivered incl. two in past month; ACV/MW; 70% deals w/ prepayment; contract tiers)','https://www.sec.gov/Archives/edgar/data/1513845/000110465926094568/tm2622968d1_ex99-2.htm','2026-10-04','SEC')
S('Q226-FS','Nebius Q2 2026 interim FS 6-K (customer concentration Q2/1H; segment revenue; RPO)','https://www.sec.gov/Archives/edgar/data/1513845/000110465926094844/nbis-20260812xex99d2.htm','2026-10-04','SEC')
S('Q226-PR','Nebius Q2 2026 results PR (group revenue 582.3; AI cloud 574.9)','https://www.sec.gov/Archives/edgar/data/1513845/000110465926094568/tm2622968d1_ex99-1.htm','2026-10-04','SEC')
S('MSFT-6K','Microsoft agreement: TCV $17,392.9M, 5-yr (20-F Note 1 p.F-17; 6-K 2025-09-08)','https://www.sec.gov/Archives/edgar/data/1513845/000110465926052948/nbis-20251231x20f.htm','2026-10-04','SEC')
S('META-6K','Meta 2025 agreement TCV $2,880.7M 5-yr, tranches Dec-25/Feb-26 (20-F Note 1 p.F-18); Meta 2026 $12B dedicated from early 2027 (6-K 2026-03-16)','https://www.sec.gov/Archives/edgar/data/1513845/000110465926027886/tm268879d1_ex99-1.htm','2026-10-04','SEC')
S('CALL-Q324','Q3 2024 earnings call transcript (on-demand customers; AI Studio launch)','/workspace/nebius/filings/altdata/transcripts/transcript_q3-2024.txt','2026-10-04','transcript')
S('CALL-Q424','Q4 2024 earnings call transcript ("short-term on-demand contracts"; AI Studio single instances)','/workspace/nebius/filings/altdata/transcripts/transcript_q4-2024.txt','2026-10-04','transcript')
S('CALL-Q425','Q4 2025 earnings call transcript (Tavily ~700k developers; per-token monetization)','/workspace/nebius/filings/altdata/transcripts/transcript_q4-2025.txt','2026-10-04','transcript')
S('CALL-Q126','Q1 2026 earnings call transcript (Revolut, monday.com Token Factory wins; pipeline excludes hyperscaler deals)','/workspace/nebius/filings/altdata/transcripts/transcript_q1-2026.txt','2026-10-04','transcript')
S('CALL-Q226','Q2 2026 earnings call transcript (Tavily >2.5M developers; TF high-margin revenue)','/workspace/nebius/filings/altdata/transcripts/transcript_q2-2026.txt','2026-10-04','transcript')
S('CRWV-10Q','CoreWeave 10-Q Q2 2026: committed contracts = 98% of revenue (benchmark for on-demand share)','https://www.sec.gov/Archives/edgar/data/1769628/','2026-10-04','SEC')
S('FACTBASE','Filings teammate fact base (nebius_fact_base.md)','/workspace/nebius/filings/nebius_fact_base.md','2026-10-04','team file')
S('CALC','Derived/estimated in this workbook (method in note column)','demand_signals.xlsx','2026-10-04','calc')

wb=Workbook(); HDR=Font(bold=True,color='FFFFFF'); FILL=PatternFill('solid',fgColor='1F3864'); BOLD=Font(bold=True)
def sheet(name,cols,rows,widths=None,title=None,notes=None):
    ws=wb.create_sheet(name); r0=1
    if title: ws.cell(1,1,title).font=Font(bold=True,size=12); r0=2
    for n in (notes or []): ws.cell(r0,1,n).font=Font(italic=True,size=9); r0+=1
    for j,c in enumerate(cols,1):
        x=ws.cell(r0,j,c); x.font=HDR; x.fill=FILL; x.alignment=Alignment(wrap_text=True,vertical='top')
    for i,row in enumerate(rows,r0+1):
        for j,v in enumerate(row,1): ws.cell(i,j,v)
    for j in range(1,len(cols)+1):
        ws.column_dimensions[get_column_letter(j)].width=(widths[j-1] if widths and j-1<len(widths) else 16)
    ws.freeze_panes=ws.cell(r0+1,1)
    return ws
# ================= 1. JOBS =================
import importlib.util
spec=importlib.util.spec_from_file_location('cj',f'{D}/classify_jobs.py'); cj=importlib.util.module_from_spec(spec); spec.loader.exec_module(cj)
FUNCS=['Engineering / R&D / product','Sales & GTM','DC ops / facilities / construction','Hardware / supply chain / capacity','G&A','Other']
cur=list(csv.DictReader(open(latest('jobs_classified_2026-*.csv'))))
srcmap={'Nebius':'GH-NBIS','CoreWeave':'GH-CRWV','Lambda':'ASHBY-LAMBDA'}
rows=[]
for co in ['Nebius','CoreWeave','Lambda']:
    s=[x for x in cur if x['company']==co]; c=collections.Counter(x['function'] for x in s)
    rows.append([co,ACC,len(s)]+[c.get(f,0) for f in FUNCS]+[f"{100*c.get('DC ops / facilities / construction',0)/len(s):.0f}%",f"{100*c.get('Sales & GTM',0)/len(s):.0f}%",'Observed (count) / Derived (function split, rule-based title classifier)',srcmap[co],
     'Group-wide board incl. Tavily (Agentic Search), Avride/robotics, academy roles' if co=='Nebius' else ('Lambda dept field = "Data Center Business" for most roles; classified on title' if co=='Lambda' else 'Includes Weights & Biases / acquired-unit roles posted on CoreWeave board')])
# history
hist_rows=[]; hist_func=[]
hf=f'{D}/jobs_history.csv'
if os.path.exists(hf):
    H=list(csv.DictReader(open(hf)))
    snaps=collections.defaultdict(list)
    for x in H: snaps[(x['company'],x['snapshot_ts'])].append(x)
    for (co,ts),xs in sorted(snaps.items(),key=lambda k:(k[0][0],k[0][1])):
        c=collections.Counter(cj.func(x['title'],x['department']) for x in xs)
        d=datetime.date(int(ts[:4]),int(ts[4:6]),int(ts[6:8]))
        hist_rows.append([co,str(d),f'{d.year}Q{(d.month-1)//3+1}',len(xs)]+[c.get(f,0) for f in FUNCS]+['Observed (Wayback capture of public ATS API) / Derived (function split)','WB-JOBS',f'https://web.archive.org/web/{ts}id_/'+('https://boards-api.greenhouse.io/v1/boards/nebius/jobs?content=true' if co=='Nebius' else 'https://boards-api.greenhouse.io/v1/boards/coreweave/jobs?content=true' if co=='CoreWeave' else 'https://api.ashbyhq.com/posting-api/job-board/lambda?includeCompensation=true')])
# current snapshot appended to history
for r in rows: hist_rows.append([r[0],ACC,'2026Q4 (current)',r[2]]+r[3:9]+['Observed (live API)',srcmap[r[0]],'live'])
neb=[x for x in cur if x['company']=='Nebius']
cc=collections.Counter(c for x in neb for c in x['country'].split('; ') if c)
geo=[[k,v,f'{100*v/len(neb):.0f}%','Observed (multi-location postings counted once per country)','GH-NBIS'] for k,v in cc.most_common()]
key='nebius_site_(DC/hardware roles only)'
dc=[x for x in neb if x['function'] in('DC ops / facilities / construction','Hardware / supply chain / capacity')]
sc=collections.Counter(s for x in dc for s in (x[key].split('; ') if x[key] else ['Not site-specific (HQ / remote / design / supply-chain roles)']))
site_rows=[[k,v,'Derived (location+title regex -> capacity/site_list.csv)','GH-NBIS; SITES'] for k,v in sc.most_common()]
# sites with no DC postings
SITE_ALL=['Vineland, NJ (DataOne)','Kansas City / Independence, MO','Mäntsälä, Finland','Lappeenranta, Finland','Paris / Béthune, France','UK (London/Longcross/Harlow/Newport)',"Israel (Modi'in / Masmiyya / Beit Shemesh / Tel Aviv)",'Iceland (Keflavik)','Spain (Madrid)','Estonia (Tallinn)','Alabama (Birmingham)','Pennsylvania','Oklahoma','Minnesota']
for s in SITE_ALL:
    if s not in sc: site_rows.append([s,0,'Derived - no open DC/hardware posting mapped to this site','GH-NBIS; SITES'])
ws=sheet('Jobs',['company','as_of','total_open_roles']+FUNCS+['DC-ops share','Sales/GTM share','tag','source_id','note'],rows,[12,11,10,14,10,14,14,10,8,9,9,40,14,50],
 title='1. Job postings - current snapshot by function (Nebius vs CoreWeave vs Lambda)',
 notes=['Function = rule-based classifier on title (+department) in market/scrapers/demand/classify_jobs.py; counts are open requisitions on the public ATS board on the access date (not hires).',
        'Nebius board includes non-AI-cloud units (Tavily/Agentic Search 17 roles, Robotics & AI 3, Nebius Academy 6). Headcount reference: ~1,500 employees at 12/31/25 (20-F) -> 365 open roles ~24% of YE25 headcount (Derived).'])
r=ws.max_row+2; ws.cell(r,1,'Nebius open roles by country (current)').font=BOLD
for j,c in enumerate(['country','roles','share','tag','source_id'],1): x=ws.cell(r+1,j,c); x.font=HDR; x.fill=FILL
for i,g in enumerate(geo,r+2):
    for j,v in enumerate(g,1): ws.cell(i,j,v)
r=ws.max_row+2; ws.cell(r,1,f'Nebius DC-ops/facilities + hardware/supply-chain roles mapped to announced sites (n={len(dc)})').font=BOLD
for j,c in enumerate(['site (capacity/site_list.csv)','roles','tag','source_id'],1): x=ws.cell(r+1,j,c); x.font=HDR; x.fill=FILL
for i,g in enumerate(site_rows,r+2):
    for j,v in enumerate(g,1): ws.cell(i,j,v)
wsH=sheet('Jobs_History',['company','snapshot_date','quarter','total_open_roles']+FUNCS+['tag','source_id','capture_url'],hist_rows,[11,12,14,10,14,10,14,14,10,8,40,10,60],
 title='1b. Job postings over time (Wayback captures of the public ATS JSON; one row per capture)',
 notes=['Quarters with no capture are not interpolated: Nebius captures start 2024-11 (no Greenhouse API capture for Q1-Q3 2024 -> not available); CoreWeave captures start 2025-05; Lambda (Ashby) 4 captures.',
        'If a capture failed to download (Wayback rate-limit), it is absent here and listed in the README gaps.'])
# vintage of currently open roles (first_published) - proxy for hiring ramp when Wayback history is thin
r=wsH.max_row+2; wsH.cell(r,1,'1d. Currently open roles by quarter first published (ATS first_published / publishedAt) - Derived proxy for hiring ramp; survivorship-biased (filled/closed roles drop out)').font=BOLD
QS=['2024Q3','2024Q4','2025Q1','2025Q2','2025Q3','2025Q4','2026Q1','2026Q2','2026Q3','2026Q4']
for j,c in enumerate(['company']+QS+['tag','source_id'],1): x=wsH.cell(r+1,j,c); x.font=HDR; x.fill=FILL
VIN={}
for i,(co,fn,fld) in enumerate([('Nebius','nebius','first_published'),('CoreWeave','coreweave','first_published'),('Lambda','lambda','publishedAt')],r+2):
    J=json.load(open(f'{D}/raw/{fn}_jobs_2026-10-04.json'))['jobs']; c=collections.Counter()
    for x in J:
        d_=x.get(fld) or ''
        if d_: c[f'{d_[:4]}Q{(int(d_[5:7])-1)//3+1}']+=1
    VIN[co]=c
    for j,v in enumerate([co]+[c.get(q,0) for q in QS]+['Derived (count of live postings by first-publish date)',srcmap[co]],1): wsH.cell(i,j,v)
r=wsH.max_row+2; wsH.cell(r,1,'1e. Nebius currently open roles by function x quarter first published').font=BOLD
for j,c in enumerate(['function']+QS,1): x=wsH.cell(r+1,j,c); x.font=HDR; x.fill=FILL
JN=json.load(open(f'{D}/raw/nebius_jobs_2026-10-04.json'))['jobs']; fq=collections.Counter()
for x in JN:
    d_=x.get('first_published') or ''
    if d_: fq[(cj.func(x['title'],'; '.join(dd['name'] for dd in x.get('departments',[]))),f'{d_[:4]}Q{(int(d_[5:7])-1)//3+1}')]+=1
for i,f in enumerate(FUNCS,r+2):
    for j,v in enumerate([f]+[fq.get((f,q),0) for q in QS],1): wsH.cell(i,j,v)
sheet('Jobs_Detail',list(cur[0].keys()),[list(x.values()) for x in cur],[10,12,40,24,30,30,22,22,26,16,30],title='1c. Current open roles (all three boards), classified')
# ================= 2. WEB TRAFFIC =================
tr=list(csv.DictReader(open(latest('tranco_2026-*.csv'))))
doms=['nebius.com','nebius.ai','coreweave.com','lambda.ai','lambdalabs.com','crusoe.ai','together.ai','nscale.com']
bydate=collections.defaultdict(dict)
for x in tr: bydate[x['date']][x['domain']]=x['rank']
trows=[]
for d in sorted(bydate):
    dd=datetime.date.fromisoformat(d); q=f'{dd.year}Q{(dd.month-1)//3+1}'
    trows.append([d,q]+[int(bydate[d][k]) if str(bydate[d].get(k,'')).isdigit() else 'not in top 1M' for k in doms]+[bydate and 'Observed','TRANCO'])
ws=sheet('Web_Traffic',['list_date','quarter']+doms+['tag','source_id'],trows,[11,9]+[13]*8+[10,10],
 title='2. Web traffic proxy - Tranco rank (lower = more traffic) at each quarter-end',
 notes=['Tranco = 30-day combined rank of 5 independent sources incl. Cloudflare Radar; it is a RANK, not visits. Pay-level domain only: studio.nebius.com / tokenfactory.nebius.com roll up into nebius.com (no separate rank).',
        'Similarweb (data.similarweb.com) returned HTTP 403 -> not recorded. Semrush/Ahrefs public snippets and Cloudflare Radar domain pages require login/API token or JS -> not recorded. Visits: not disclosed by any source we could access.'])
r=ws.max_row+2
first=trows[0]; last=trows[-1]
ws.cell(r,1,'Derived: nebius.com rank improved from not-in-top-1M (Q1-Q2 2024; nebius.ai was the primary domain) to ~201k at YE24, ~110k at YE25 and ~56k at 2026-10-03; coreweave.com ~22k (CoreWeave ~2.5x higher-ranked); lambda.ai ~122k.').font=Font(italic=True)
# ================= 3. AI STUDIO / TOKEN FACTORY =================
prov=list(csv.DictReader(open(latest('openrouter_providers_2026-*.csv'))))
totm=sum(int(p['monthly_tokens']) for p in prov)
peers=['nebius','together','deepinfra','fireworks','baseten','coreweave','crusoe','novita','parasail','gmicloud']
orows=[]
for p in prov:
    if p['slug'] in peers:
        orows.append([p['display_name'],int(p['rank_by_monthly_tokens']),len(prov),int(p['model_count']),int(p['daily_tokens']),int(p['monthly_tokens']),float(p['share_of_listed_monthly_tokens']),'Observed (OpenRouter public page)','OR-PROV','monthly = trailing ~30d as shown on page; share = of the sum across all 92 listed providers'])
ws=sheet('AIStudio_TokenFactory',['provider','rank (monthly tokens)','providers listed','models served on OR','daily tokens','monthly tokens','share of OR monthly tokens','tag','source_id','note'],orows,[24,10,9,9,16,18,12,26,10,50],
 title='3. AI Studio / Token Factory self-serve signals',
 notes=['OpenRouter = third-party router; it captures only the serverless, OR-routed slice of Token Factory (not direct API, dedicated endpoints or enterprise contracts).'])
# daily
daily=list(csv.DictReader(open(latest('openrouter_nebius_daily_2026-*.csv'))))
mon=collections.defaultdict(lambda:[0,0])
for x in daily: mon[x['date'][:7]][0]+=int(x['total_tokens']); mon[x['date'][:7]][1]+=1
r=ws.max_row+2; ws.cell(r,1,'Nebius Token Factory tokens processed on OpenRouter, by month (from provider-page chartData; partial months flagged)').font=BOLD
for j,c in enumerate(['month','tokens','days in data','avg tokens/day','tag','source_id'],1): x=ws.cell(r+1,j,c); x.font=HDR; x.fill=FILL
i=r+2
for m,(t,n) in sorted(mon.items()):
    for j,v in enumerate([m+('' if n>=28 else f' (partial, {n}d)'),t,n,round(t/n),'Observed','OR-NBIS'],1): ws.cell(i,j,v)
    i+=1
# models + implied revenue
mods=list(csv.DictReader(open(latest('openrouter_nebius_models_2026-*.csv'))))
FALLBACK={'deepseek/deepseek-v4-flash-20260731':(0.132,0.28),'meta-llama/llama-3.3-70b-instruct':(0.52,0.72)}
r=i+1; ws.cell(r,1,'Implied OpenRouter-routed Token Factory revenue, trailing 30 days (Derived)').font=BOLD
cols=['model','Nebius tokens last 30d','Nebius price in $/M','Nebius price out $/M','price tag','rev @90/10 in/out ($)','rev @70/30 in/out ($)','providers serving model on OR']
for j,c in enumerate(cols,1): x=ws.cell(r+1,j,c); x.font=HDR; x.fill=FILL
i=r+2; lo=hi=0
for m in mods:
    t=float(m['nebius_tokens_last30d'])
    if m['nebius_price_in_$per_M']: pi,po,tag=float(m['nebius_price_in_$per_M']),float(m['nebius_price_out_$per_M']),'Observed (OR API)'
    elif m['model'] in FALLBACK: pi,po=FALLBACK[m['model']]; tag='Estimated (Nebius endpoint no longer listed; median of other providers on OR)'
    else: pi=po=None; tag='n/a (no tokens in window)'
    a=t/1e6*(0.9*pi+0.1*po) if pi is not None else 0; b=t/1e6*(0.7*pi+0.3*po) if pi is not None else 0
    lo+=a; hi+=b
    for j,v in enumerate([m['model'],t,pi,po,tag,round(a),round(b),int(m['total_provider_endpoints'])],1): ws.cell(i,j,v)
    i+=1
others=0
try: others=sum(int(x['Others']) for x in daily[-31:-1])
except Exception: pass
for j,v in enumerate(['"Others" bucket (unpriced)',others,None,None,'excluded (~0.06% of tokens)'],1): ws.cell(i,j,v)
i+=1
for j,v in enumerate(['TOTAL (30d)',None,None,None,'Derived',round(lo),round(hi)],1): ws.cell(i,j,v).font=BOLD
i+=1
for j,v in enumerate(['Annualized',None,None,None,'Derived (x12)',round(lo*12),round(hi*12)],1): ws.cell(i,j,v).font=BOLD
OR_REV=(lo,hi)
i+=2; ws.cell(i,1,'Catalog, listings and management disclosures').font=BOLD
kv=[
 ['Models served by Nebius on OpenRouter (current)',10,ACC,'Observed','OR-PROV','OR providers page modelCount; 0 free models'],
 ['OpenRouter rank / share (monthly tokens)','#57 of 92; 0.079%',ACC,'Observed / Derived','OR-PROV','470.3B of 592.2T listed monthly tokens; CoreWeave #16 (1.61%), Crusoe #58 (0.077%), Together #5 (4.39%)'],
 ['Embedding share of Nebius OR tokens (30d)','~68% (qwen3-embedding-8b, $0.01/M)',ACC,'Derived','OR-NBIS','Volume is dominated by a near-free embedding model -> revenue tiny relative to tokens'],
 ['Hugging Face Inference Providers - Nebius live?','NO - "nebius" not in API provider enum; doc page has no supported tasks',ACC,'Observed','HF-API; HF-DOC','Nebius was an HF inference provider in 2025 (doc page still exists). Live providers include Nscale (19 models), Together (18), DeepInfra (92), Novita (67). Date Nebius dropped off: not determined (Wayback of HF docs rate-limited).'],
 ['Hugging Face Nebius org: followers / members / models','311 / 33 / 10',ACC,'Observed','HF-ORG','enterprise plan; verified'],
 ['AI Studio launch','Q3 2024 (self-service inference, per-token pricing)','2024-11','Reported','Q324-PR',''],
 ['AI Studio registered users','>60,000','Q1 2025 (Mar-31-25)','Reported','Q125-SHL','"still early from a revenue perspective"; only user-count disclosure found; later quarters: not disclosed'],
 ['Token Factory launch (AI Studio successor, enterprise inference)','Released Nov 2025','Q4 2025','Reported','Q325-SHL; 20F-25','per-token monetization'],
 ['Token Factory named wins','Revolut, monday.com (Q1-26); digital-health customer (Sword Health); Revolut runs >200 H100 incl. TF','Q1-Q2 2026','Reported','CALL-Q126; NB-STORIES',''],
 ['Token Factory production inference workloads','>3x QoQ','Q2 2026','Reported','Q226-SHL','no $ or token volume disclosed'],
 ['Tavily developer community (agentic search, acquired Feb-26)','~700k -> 1M (Feb-26) -> >2.5M (Q2-26)','Q4-25..Q2-26','Reported (call)','CALL-Q425; CALL-Q226','developers, not paying customers'],
 ['AI Studio / TF revenue','not disclosed','-','-','Q125-SHL; Q226-SHL','mgmt: TF contributes "high-margin revenue" (Q2-26) but no figure'],
 ['Model catalog size over time (AI Studio/TF site)','not available - tokenfactory/studio pages are JS apps; models API requires key; Wayback catalog pages rate-limited','-','-','-','OR count history: see README gap'],
 ['Free trial (self-serve on-ramp)','Launched 2026-06-14 ($5 credit, eu-north1); SUSPENDED 2026-07-13','Q2-Q3 2026','Reported (docs changelog)','DOCS-CL','signals capacity prioritised for contracted demand'],
 ['Spot pricing for preemptible GPU VMs','Added week of 2026-09-21 (price can change every 15 min)','Q3 2026','Reported (docs changelog)','DOCS-CL','monetises idle capacity; self-serve'],
 ['Soperator (managed Slurm) GPU nodes','only with capacity block groups (reserved) from Mar-2026','Q1 2026','Reported (docs changelog)','DOCS-CL','shift of managed training toward reserved capacity'],
]
for j,c in enumerate(['item','value','period','tag','source_id','note'],1): x=ws.cell(i+1,j,c); x.font=HDR; x.fill=FILL
for k,row in enumerate(kv,i+2):
    for j,v in enumerate(row,1): ws.cell(k,j,v)
# ================= 4. CUSTOMERS =================
stories=list(csv.DictReader(open(latest('customer_stories_2026-*.csv'))))
first=json.load(open(f'{R}/wayback/stories_first_seen.json'))
CLS={'antioch':('startup','physical-AI simulation'),'chatfuel':('startup','chatbot SaaS'),'compugen':('enterprise','Nasdaq-listed biotech (Israel)'),'converge-bio':('startup','bio LLM'),'crispr-gpt-stanford':('research/OSS','Stanford/Princeton/DeepMind'),'dubformer':('startup','AI dubbing'),'helical':('startup','bio foundation models'),'higgsfield-ai':('AI lab','video/image diffusion model builder'),'jua':('AI lab','weather foundation model'),'krisp':('startup','voice AI SaaS'),'lims':('research/OSS','London Inst. for Mathematical Sciences'),'lynx-analytics':('enterprise','AI services firm'),'normacore':('startup','physical AI infra'),'positronic-robotics':('startup','robotics'),'prima-mente':('AI lab','brain foundation model'),'prime-intellect':('AI lab','distributed training/RL'),'quantori':('enterprise','IT/cheminformatics services'),'recraft':('AI lab','image model builder'),'recraft-v4':('AI lab','image model builder (V4 on HGX B200)'),'revolut':('enterprise','fintech, 70M users'),'roboforce':('startup','robotics'),'sglang':('research/OSS','inference framework'),'sievestack':('startup','molecular simulation data'),'simulacra-ai':('AI lab','quantum-chem foundation model'),'slingshot-ai':('AI lab','psychology foundation LLM ($93M raised)'),'swordhealth':('enterprise','digital health (Token Factory)'),'synthlabs':('AI lab','post-training research'),'thestage':('startup','inference optimization'),'trialhub':('startup','clinical-trial RAG'),'ultimate-bots':('startup','humanoid robot software'),'unum':('startup','open-source search'),'unum-stringzilla':('startup','open-source search (2nd story)'),'vllm':('research/OSS','inference framework'),'voxel51':('startup','CV data platform'),'wildflow':('startup','biodiversity models'),'xaid':('startup','medical imaging AI'),'yerevann':('research/OSS','non-profit ML research')}
GPU={'crispr-gpt-stanford':'single 8x H100 node','helical':'16 GPUs','krisp':'single H100 (foundation model ~2 weeks)','prima-mente':'32-GPU reserved cluster (16 nodes); scaled 1B->68B params (Q2-26 letter)','quantori':'8x H200','revolut':'>200 H100 (training + inference incl. Token Factory); 64-H100 dedicated cluster for PRAGMA','synthlabs':'starting from 128 H100','xaid':'3x H100','prime-intellect':'16 GB200 GPUs (4-node test); "hundreds to thousands" of GPUs','jua':'"thousands of the latest GPUs" (qualitative)','higgsfield-ai':'usage expanded >20x since first contract (Q2-26 letter)','lynx-analytics':'product scales 1 -> 1,000+ GPUs (capability, not a Nebius commitment)','recraft-v4':'NVIDIA HGX B200 cluster (count n/d)','voxel51':'RTX PRO 6000 Blackwell (count n/d)','roboforce':'RTX PRO 6000 + Blackwell (count n/d)','swordhealth':'dedicated Blackwell endpoints (count n/d)'}
crow=[]
for s in stories:
    sl=s['slug']; fs=first.get(sl); fsd=f'{fs[:4]}-{fs[4:6]}-{fs[6:8]}' if fs else 'not captured'
    q=f'{fs[:4]}Q{(int(fs[4:6])-1)//3+1}' if fs else ''
    cl,desc=CLS.get(sl,('unclassified',''))
    crow.append([sl,s['url'],cl,desc,fsd,q,GPU.get(sl,'not disclosed'),'not disclosed','Observed (page) / Derived (class)','NB-STORIES; WB-STORIES'])
crow.sort(key=lambda x:x[4])
ws=sheet('Customers',['customer story (slug)','url','class','what they do','first seen (Wayback)','quarter first seen','GPU count / scale disclosed','spend disclosed','tag','source_id'],crow,[22,46,13,34,13,10,48,12,22,18],
 title='4. Named customers - nebius.com customer stories (37 pages) + filings/letters',
 notes=['Class: hyperscaler / AI lab (builds own models) / enterprise (established company) / startup (AI-native app or tooling) / research/OSS. "First seen" = first Wayback capture of the URL (upper bound on publication date).'])
filed=[
 ['Microsoft','hyperscaler','TCV $17,392.9M / 5 yrs (up to ~$19.4B); ~$3.48B/yr when fully delivered (Derived = TCV/5); all tranches delivered by Aug-26 (two in Jul-Aug)','2025-09-08','Reported','MSFT-6K; Q226-SHL'],
 ['Meta (2025 agreement)','hyperscaler','TCV $2,880.7M / 5 yrs (~$576M/yr Derived); tranches Dec-25 & Feb-26','2025-11-01','Reported','META-6K'],
 ['Meta (2026 agreement)','hyperscaler','$12B dedicated (Vera Rubin) from early 2027 + up to $15B backstop of unsold capacity','2026-03-16','Reported','META-6K'],
 ['Reflection AI','AI lab','Q2-26 landmark deal (4 deals avg TCV >$1B; yield $20-25M/MW); prepayment','2026-08-12','Reported','Q226-SHL'],
 ['Cohere','AI lab','Q2-26 landmark deal; prepayment','2026-08-12','Reported','Q226-SHL'],
 ['Unnamed US AI neolab','AI lab','Q2-26 landmark deal','2026-08-12','Reported','Q226-SHL'],
 ['Unnamed large US quant trading firm','enterprise','Q2-26 landmark deal; first large financial-services win','2026-08-12','Reported','Q226-SHL'],
 ['AMI (Yann LeCun)','AI lab','compute + storage for world models','2026-08-12','Reported','Q226-SHL'],
 ['Basecamp Research','AI lab','EDEN 28B-param model trained entirely on Nebius','2026-08-12','Reported','Q226-SHL'],
 ['1X Technologies','startup','physical AI / robots on Nebius cloud','2026-05-13','Reported (call)','CALL-Q126'],
 ['Rhoda AI; Core Automation','startup / AI lab','named wins Q1-26','2026-05-13','Reported (call)','CALL-Q126'],
 ['monday.com; Revolut','enterprise','Token Factory wins Q1-26','2026-05-13','Reported (call)','CALL-Q126'],
]
r=ws.max_row+2; ws.cell(r,1,'Named in filings / letters / calls (not on customer-stories pages)').font=BOLD
for j,c in enumerate(['customer','class','disclosed spend / scale','date','tag','source_id'],1): x=ws.cell(r+1,j,c); x.font=HDR; x.fill=FILL
for k,row in enumerate(filed,r+2):
    for j,v in enumerate(row,1): ws.cell(k,j,v)
cnt=collections.Counter(x[2] for x in crow); qc=collections.Counter(x[5] for x in crow)
r=ws.max_row+2; ws.cell(r,1,'Customer-story counts (Derived)').font=BOLD
k=r+1
for c,v in cnt.most_common(): ws.cell(k,1,c); ws.cell(k,2,v); k+=1
k+=1; ws.cell(k,1,'New story URLs by quarter first seen').font=BOLD; k+=1
for q in sorted(qc): ws.cell(k,1,q); ws.cell(k,2,qc[q]); k+=1
conc=[
 ['FY2023','-','no customer >=10% (post-Toloka deconsolidation basis)','Reported','20F-25'],
 ['FY2024','A 27% / C 11% (20-F FY25 basis); A 21% / B 17% (20-F FY24 basis)','labels re-assigned between filings','Reported','20F-25; 20F-24'],
 ['Q2 2025','A 39% / B 15%','','Reported','Q226-FS'],
 ['1H 2025','A 34% / B 13%','','Reported','Q226-FS'],
 ['FY2025','A 25% / B 15%; Customer D = 83% of YE25 A/R ($597.0M) but <10% of FY25 revenue','D (20-F) = C (Q2-26 6-K): new late-2025 customer with large upfront billing','Reported','20F-25'],
 ['Q1 2026 (implied)','C ~28.9% / D ~21.0% / E <4.2%','Derived = 1H26 less Q2-26 (1H: C 26%, D 21%, E <10%; revenue 981.3 / 582.3)','Derived','Q226-FS; Q226-PR'],
 ['Q2 2026','C 24% / D 21% / E 14% (top-3 = 59%)','E new in Q2 (29% of 6/30 A/R -> prepaid); C, D present since late 2025','Reported','Q226-FS'],
 ['1H 2026','C 26% / D 21%','','Reported','Q226-FS'],
]
r=k+1; ws.cell(r,1,'Customer concentration history (% of group revenue)').font=BOLD
for j,c in enumerate(['period','10%+ customers','note','tag','source_id'],1): x=ws.cell(r+1,j,c); x.font=HDR; x.fill=FILL
for kk,row in enumerate(conc,r+2):
    for j,v in enumerate(row,1): ws.cell(kk,j,v)
# ================= 5. DEVELOPER =================
gh=list(csv.DictReader(open(latest('github_2026-*.csv'))))
drows=[]
for org in ['nebius','coreweave','lambdal']:
    s=[x for x in gh if x['org']==org]
    drows.append([org,len(s),sum(int(x['stars']) for x in s),sum(int(x['forks']) for x in s),max(s,key=lambda x:int(x['stars']))['repo'] if s else '',sum(1 for x in s if x['pushed']>='2026-09-01'),'Observed','GITHUB','lambdal = Lambda Labs legacy org (3 public repos); Lambda also publishes under other orgs - not counted' if org=='lambdal' else ''])
ws=sheet('Developer',['GitHub org','public repos','stars','forks','top repo','repos pushed since 2026-09-01','tag','source_id','note'],drows,[12,9,8,8,24,12,10,10,50],
 title='5. Developer presence',
 notes=['Contributors over time: not collected (GitHub unauthenticated core API rate-limited from box; search API used for repo-level data). Stars over time: not available without per-star history; repo creation by quarter shown instead.'])
qrep=collections.Counter(); 
for x in gh:
    if x['org']=='nebius':
        d=x['created']; qrep[f'{d[:4]}Q{(int(d[5:7])-1)//3+1}']+=1
r=ws.max_row+2; ws.cell(r,1,'nebius org: public repos created per quarter (cumulative) - Derived from repo created_at').font=BOLD
for j,c in enumerate(['quarter','new repos','cumulative'],1): x=ws.cell(r+1,j,c); x.font=HDR; x.fill=FILL
k=r+2; cum=0
allq=[f'{y}Q{q}' for y in (2023,2024,2025,2026) for q in (1,2,3,4)]
for q in allq:
    if q>'2026Q4': break
    cum+=qrep.get(q,0)
    if q>='2024Q1' or qrep.get(q): ws.cell(k,1,q); ws.cell(k,2,qrep.get(q,0)); ws.cell(k,3,cum); k+=1
men=list(csv.DictReader(open(latest('mentions_2026-*.csv'))))
mq=collections.defaultdict(dict)
for x in men: mq[(x['site'],x['entity'],x['type'])][x['quarter']]=x['count']
quarters=sorted({x['quarter'] for x in men})
r=k+1; ws.cell(r,1,'Hacker News / Stack Overflow mentions by quarter (2026Q4 = Oct 1-4 only)').font=BOLD
for j,c in enumerate(['site','entity','type']+quarters+['tag','source_id'],1): x=ws.cell(r+1,j,c); x.font=HDR; x.fill=FILL
k=r+2
for key_,v in mq.items():
    for j,val in enumerate(list(key_)+[int(v.get(q,0)) for q in quarters]+['Observed',key_[0] if key_[0]=='HN' else 'SO'],1): ws.cell(k,j,val)
    k+=1
cl=list(csv.DictReader(open(latest('docs_changelog_2026-*.csv'))))
cq=collections.defaultdict(lambda:[0,0,0])
for x in cl: cq[x['quarter']][0]+=1; cq[x['quarter']][1]+=int(x['entries']); cq[x['quarter']][2]+=int(x['added'])
r=k+1; ws.cell(r,1,'Docs changelog cadence (docs.nebius.com/changelog; weekly since ~Sep-2025; earlier weeks not on page)').font=BOLD
for j,c in enumerate(['quarter','weekly updates','entries','of which "Added"','tag','source_id'],1): x=ws.cell(r+1,j,c); x.font=HDR; x.fill=FILL
k=r+2
for q in sorted(cq):
    for j,v in enumerate([q]+cq[q]+['Observed','DOCS-CL'],1): ws.cell(k,j,v)
    k+=1
k+=1
for row in [['Discord / community size','not found - no public Nebius Discord invite on nebius.com, docs or Token Factory site','DISCORD'],
            ['Hugging Face org followers','311 (members 33, models 10, datasets 13)','HF-ORG'],
            ['Stack Overflow','1 post all-time matching "nebius" (0 in Q1-24..Q4-26 window); CoreWeave 0 -> SO not a useful signal for GPU clouds','SO']]:
    for j,v in enumerate(row,1): ws.cell(k,j,v)
    k+=1
# ================= 6. MIX ESTIMATE =================
REV=582.3; AIC=574.9; OTHER=REV-AIC
legs=[
 # leg, low, base, high, method, evidence, weakest?
 ['(a) Hyperscaler take-or-pay (Microsoft + Meta-2025)',0.40,0.45,0.46,
  'Customer C (24%) + D (21%) of Q2-26 group revenue = 45%. Identification (Estimated): C and D are the only 10%+ customers that were <10% in FY25 and present since late 2025, matching Microsoft (tranche 1 Nov-25) and Meta-2025 (tranches Dec-25/Feb-26); C held 83% of YE25 A/R ($597M) = upfront hyperscaler billing. Cross-check: Meta-2025 fully delivered = $2,880.7M/5/4 = $144M/qtr = 24.7% of Q2 revenue (fits C or D). High 46% = rounding of 24%+21%. Low 40% allows part of C/D revenue to be non-take-or-pay add-ons or labels to be mis-identified.',
  'Q226-FS; 20F-25; MSFT-6K; META-6K','Moderate - the % is reported; identity of C/D is inferred (customers unnamed).'],
 ['(b) Reserved / committed contracts (AI labs, AI natives, enterprises; 3-month to 3-yr)',None,None,None,
  'Plug = 100% - (a) - (c) - non-AI-cloud segments (TripleTen/Avride net = (582.3-574.9)/582.3 = 1.3%). Supporting evidence: Customer E (14%, new in Q2, prepaid) is a non-hyperscaler contract; mgmt calls 1-3-yr mid-term contracts "our core business"; ~70% of Q2 deals included prepayment; four >$1B-TCV deals in Q2 (Reflection, Cohere, neolab, quant firm).',
  'Q226-SHL; Q226-FS','Inherits all error from (a) and (c).'],
 ['(c) On-demand GPU + Token Factory / AI Studio (self-serve, usage-priced)',0.013,0.04,0.09,
  'Two sub-legs. (c1) On-demand/preemptible GPU: low 1% / base 3% / high 6% of revenue. Anchor: CoreWeave committed = 98% of Q2-26 revenue (on-demand 2%, Reported). Nebius historically more on-demand ("a lot of on-demand customers" Q3-24; "short-term on-demand contracts" Q4-24) but now sold out, free trial suspended Jul-26, Soperator GPU nodes reserved-only since Mar-26 -> base slightly above CoreWeave. (c2) Token Factory: low 0.3% / base 1% / high 3%. Anchor: OpenRouter-routed TF revenue only ~$<A>-<B>k per 30 days (Derived, ~$<C>-<D>M annualized); OR is one channel - if OR = 10% of serverless TF traffic, serverless ~ $<E>-<F>M/yr; plus dedicated enterprise endpoints (Revolut >200 H100 incl. TF ~ $4-5M/yr at $2-2.5/GPU-hr list, Estimated; monday.com, Sword Health). Mgmt: TF "high-margin" contributor, workloads >3x QoQ, but no $ disclosed.',
  'CRWV-10Q; CALL-Q324; CALL-Q424; DOCS-CL; OR-NBIS; OR-API; NB-STORIES; Q226-SHL','WEAKEST - no company disclosure of either sub-leg; OR captures an unknown fraction of TF; on-demand share inferred from a peer benchmark and qualitative signals.'],
]
lo_,hi_=OR_REV
for k_,v_ in zip('ABCDEF',(f'{lo_/1e3:.0f}',f'{hi_/1e3:.0f}',f'{lo_*12/1e6:.1f}',f'{hi_*12/1e6:.1f}',f'{lo_*12/1e6*10:.0f}',f'{hi_*12/1e6*10:.0f}')): legs[2][4]=legs[2][4].replace(f'<{k_}>',v_)
oth=OTHER/REV
def plug(a,c): return 1-a-c-oth
scen={'Low self-serve':(legs[0][3],legs[2][1]),'Base':(legs[0][2],legs[2][2]),'High self-serve':(legs[0][1],legs[2][3])}
mrows=[]
for nm,(a,c) in scen.items():
    b=plug(a,c)
    mrows.append([nm,a,b,c,oth,a*REV,b*REV,c*REV,OTHER,a+b+c+oth,'Estimated','CALC'])
ws=sheet('Mix_Estimate',['scenario','(a) hyperscaler %','(b) reserved/commit %','(c) on-demand + TF %','other segments %','(a) $M','(b) $M','(c) $M','other $M','check (=100%)','tag','source_id'],mrows,[18,12,12,12,10,9,9,9,9,9,10,8],
 title='6. Self-serve vs contract mix - Q2 2026 group revenue $582.3M (ESTIMATED)',
 notes=['Basis: Q2-26 group revenue $582.3M (Reported, Q226-PR); AI cloud $574.9M; concentration % are of group revenue (Q226-FS). Scenarios pair the leg ranges so they sum to 100%; (b) is the plug.'])
for row in ws.iter_rows(min_row=4,max_row=ws.max_row):
    for c in row[1:5]+row[9:10]: c.number_format='0.0%'
    for c in row[5:9]: c.number_format='#,##0'
r=ws.max_row+2
for j,c in enumerate(['leg','low','base','high','method','evidence (source_id)','confidence'],1): x=ws.cell(r,j,c); x.font=HDR; x.fill=FILL
for k,row in enumerate(legs,r+1):
    for j,v in enumerate(row,1):
        x=ws.cell(k,j,v); x.alignment=Alignment(wrap_text=True,vertical='top')
        if j in(2,3,4) and isinstance(v,float): x.number_format='0.0%'
    if row[1] is None:
        ws.cell(k,2,plug(legs[0][3],legs[2][3])).number_format='0.0%'; ws.cell(k,3,plug(legs[0][2],legs[2][2])).number_format='0.0%'; ws.cell(k,4,plug(legs[0][1],legs[2][1])).number_format='0.0%'
ws.column_dimensions['E'].width=90; ws.column_dimensions['F'].width=28; ws.column_dimensions['G'].width=40
r=ws.max_row+2; ws.cell(r,1,'Forward cross-check: exit-2026 run-rate').font=BOLD
fw=[
 ['Microsoft annual revenue when fully delivered','~$3.48B','$17,392.9M TCV / 5 yrs','Derived','MSFT-6K'],
 ['Meta-2025 annual revenue','~$0.58B','$2,880.7M / 5 yrs','Derived','META-6K'],
 ['Hyperscaler run-rate total','~$4.06B','sum','Derived','CALC'],
 ['FY26 exit ARR guidance','$7-9B','AI cloud last-month revenue x12','Reported','Q126-SHL'],
 ['Hyperscaler share of exit-2026 ARR','~45-58%','4.06 / 9 to 4.06 / 7','Derived','CALC'],
 ['Implied Microsoft Q2-26 revenue vs full run-rate','~$122-140M vs ~$870M/qtr','if Microsoft = C or D; only ~14-16% of full run-rate despite 7 of 9 tranches by 6/30 -> large step-up in 2H26 (two tranches delivered Jul-Aug)','Estimated','Q226-FS; Q226-SHL'],
 ['Meta-2026 ($12B, from early 2027)','~$2.4B/yr','$12B / 5 yrs; not in 2026','Derived','META-6K'],
]
for j,c in enumerate(['item','value','method','tag','source_id'],1): x=ws.cell(r+1,j,c); x.font=HDR; x.fill=FILL
for k,row in enumerate(fw,r+2):
    for j,v in enumerate(row,1): ws.cell(k,j,v)
# ================= Sources =================
sheet('Sources',['source_id','description','url / path','accessed','type'],SRC,[14,90,80,11,10],title='Sources (every row elsewhere cites a source_id here)')
# Summary first
ws=wb['Sheet']; ws.title='README'
lines=['Nebius demand signals - built '+datetime.datetime.now().strftime('%Y-%m-%d %H:%M ET')+' by market/scrapers/demand/build_demand_xlsx.py',
 'Tags: Reported = company filing/letter/call; Observed = read off a public third-party page/API on the access date; Derived = arithmetic on Reported/Observed (method shown); Estimated = judgment/assumption (method shown).',
 '"not disclosed" / "not available" is written wherever we could not find a value - no gaps are filled silently.',
 'Sheets: Jobs, Jobs_History, Jobs_Detail, Web_Traffic, AIStudio_TokenFactory, Customers, Developer, Mix_Estimate, Sources.',
 'Re-run: python3 jobs_snapshot.py; classify_jobs.py jobs_<date>.csv; wayback_jobs.py; tranco_history.py; openrouter_stats.py; customer_stories.py; github_stats.py; mentions.py; docs_changelog.py; build_demand_xlsx.py; demand_charts.py']
for i,l in enumerate(lines,1): ws.cell(i,1,l)
ws.column_dimensions['A'].width=160
wb.save(OUT); print('saved',OUT,'OR rev 30d',OR_REV)
