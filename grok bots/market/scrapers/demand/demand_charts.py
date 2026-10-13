#!/usr/bin/env python3
"""Deck-ready PNG charts for the Nebius demand-signals section -> /workspace/nebius/market/charts/demand_*.png
Reads the CSV/JSON outputs of the demand scrapers (same folder)."""
import textwrap, csv, glob, os, collections, datetime, json
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt, matplotlib.dates as mdates
D=os.path.dirname(os.path.abspath(__file__)); C='/workspace/nebius/market/charts'; os.makedirs(C,exist_ok=True)
def latest(p): return sorted(glob.glob(f'{D}/{p}'))[-1]
plt.rcParams.update({'font.size':10,'axes.spines.top':False,'axes.spines.right':False,'figure.dpi':150})
NB='#1F3864'; COL={'Engineering / R&D / product':'#1F3864','Sales & GTM':'#2E86C1','DC ops / facilities / construction':'#E67E22','Hardware / supply chain / capacity':'#F5B041','G&A':'#95A5A6','Other':'#D5D8DC'}
FUN=list(COL)
import importlib.util
spec=importlib.util.spec_from_file_location('cj',f'{D}/classify_jobs.py'); cj=importlib.util.module_from_spec(spec); spec.loader.exec_module(cj)
SRC='Source: '
# 1 jobs over time by function (Nebius) + peers totals
cur=list(csv.DictReader(open(latest('jobs_classified_2026-*.csv'))))
pts=collections.defaultdict(dict)  # company -> date -> Counter
if os.path.exists(f'{D}/jobs_history.csv'):
    for x in csv.DictReader(open(f'{D}/jobs_history.csv')):
        d=datetime.date(int(x['snapshot_ts'][:4]),int(x['snapshot_ts'][4:6]),int(x['snapshot_ts'][6:8]))
        pts[x['company']].setdefault(d,collections.Counter())[cj.func(x['title'],x['department'])]+=1
today=datetime.date(2026,10,4)
for x in cur: pts[x['company']].setdefault(today,collections.Counter())[x['function']]+=1
fig,(a1,a2)=plt.subplots(1,2,figsize=(13,5),gridspec_kw={'width_ratios':[1.6,1]})
ds=sorted(pts['Nebius']); bottom=[0]*len(ds); labels=[d.strftime('%b-%y') for d in ds]
for f in FUN:
    v=[pts['Nebius'][d].get(f,0) for d in ds]
    if sum(v)==0: continue
    a1.bar(labels,v,bottom=bottom,color=COL[f],label=f); bottom=[b+x for b,x in zip(bottom,v)]
for i,b in enumerate(bottom): a1.text(i,b+3,str(b),ha='center',fontsize=8)
a1.set_title('Nebius open roles by function (public Greenhouse board; Wayback captures + live)',fontsize=10,loc='left'); a1.legend(fontsize=7,frameon=False,loc='upper left',bbox_to_anchor=(0,1)); a1.tick_params(axis='x',rotation=45,labelsize=8)
for co,c in [('Nebius',NB),('CoreWeave','#7D3C98'),('Lambda','#27AE60')]:
    dd=sorted(pts[co]); a2.plot(dd,[sum(pts[co][d].values()) for d in dd],marker='o',color=c,label=co)
a2.set_title('Total open roles: Nebius vs CoreWeave vs Lambda',fontsize=10,loc='left'); a2.legend(frameon=False); a2.xaxis.set_major_formatter(mdates.DateFormatter('%b-%y')); a2.tick_params(axis='x',rotation=45,labelsize=8)
fig.text(0.01,0.005,'\n'.join(textwrap.wrap(SRC+'boards-api.greenhouse.io (nebius, coreweave), api.ashbyhq.com (lambda); history = web.archive.org captures; function = rule-based title classifier (classify_jobs.py). Accessed 2026-10-04.',170)),fontsize=7,color='grey')
plt.tight_layout(rect=(0,0.06,1,1)); plt.savefig(f'{C}/demand_job_postings_by_function.png'); plt.close()
# 2 OpenRouter
prov=list(csv.DictReader(open(latest('openrouter_providers_2026-*.csv'))))
daily=list(csv.DictReader(open(latest('openrouter_nebius_daily_2026-*.csv'))))
fig,(a1,a2)=plt.subplots(1,2,figsize=(13,4.8))
dd=[datetime.date.fromisoformat(x['date']) for x in daily][:-1]; tot=[int(x['total_tokens'])/1e9 for x in daily][:-1]
emb=[int(x.get('qwen/qwen3-embedding-8b',0))/1e9 for x in daily][:-1]
a1.fill_between(dd,0,emb,color='#AED6F1',label='qwen3-embedding-8b ($0.01/M)'); a1.fill_between(dd,emb,tot,color=NB,label='all other models')
a1.set_ylabel('tokens per day (billions)'); a1.set_title('Nebius Token Factory tokens routed via OpenRouter (daily)',fontsize=10,loc='left'); a1.legend(frameon=False,fontsize=8); a1.xaxis.set_major_formatter(mdates.DateFormatter('%d-%b'))
peers=['together','deepinfra','coreweave','fireworks','baseten','nebius','crusoe']
P=[p for p in prov if p['slug'] in peers]; P.sort(key=lambda p:float(p['share_of_listed_monthly_tokens']))
a2.barh([f"{p['display_name']} (#{p['rank_by_monthly_tokens']})" for p in P],[100*float(p['share_of_listed_monthly_tokens']) for p in P],color=[NB if p['slug']=='nebius' else '#BFC9CA' for p in P])
for i,p in enumerate(P): a2.text(100*float(p['share_of_listed_monthly_tokens'])+0.05,i,f"{100*float(p['share_of_listed_monthly_tokens']):.2f}%",va='center',fontsize=8)
a2.set_xlabel('% of OpenRouter monthly tokens (92 providers)'); a2.set_title('OpenRouter token share, trailing month (rank in brackets)',fontsize=10,loc='left')
fig.text(0.01,0.005,'\n'.join(textwrap.wrap(SRC+'openrouter.ai/provider/nebius (chartData) and openrouter.ai/providers (monthlyTokens), accessed 2026-10-04. OR covers only the OR-routed serverless slice of Token Factory.',170)),fontsize=7,color='grey')
plt.tight_layout(rect=(0,0.07,1,1)); plt.savefig(f'{C}/demand_openrouter_token_share.png'); plt.close()
# 3 Tranco
tr=list(csv.DictReader(open(latest('tranco_2026-*.csv'))))
fig,ax=plt.subplots(figsize=(9,4.8))
for dom,c in [('nebius.com',NB),('coreweave.com','#7D3C98'),('lambda.ai','#27AE60'),('lambdalabs.com','#82E0AA'),('together.ai','#E67E22'),('crusoe.ai','#B03A2E'),('nebius.ai','#85C1E9')]:
    s=[(datetime.date.fromisoformat(x['date']),int(x['rank'])) for x in tr if x['domain']==dom and x['rank'].isdigit()]
    if s: ax.plot([a for a,_ in s],[b for _,b in s],marker='o',ms=3,color=c,label=dom,lw=2.2 if dom=='nebius.com' else 1.2)
ax.set_yscale('log'); ax.set_yticks([20000,50000,100000,200000,500000,1000000]); ax.yaxis.set_minor_formatter(matplotlib.ticker.NullFormatter()); ax.invert_yaxis(); ax.set_ylabel('Tranco rank (log; up = more traffic)'); ax.legend(frameon=False,fontsize=8,ncol=2)
ax.set_title('Web traffic proxy: Tranco rank at quarter-ends, Q1-24 to Oct-26',fontsize=10,loc='left')
ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v,_: '1M' if v>=1e6 else f'{int(v/1000)}k'))
fig.text(0.01,0.005,'\n'.join(textwrap.wrap(SRC+'tranco-list.eu daily lists (30-day combined CrUX/Farsight/Majestic/Cloudflare Radar/Umbrella). Rank, not visits. Missing points = outside top 1M.',170)),fontsize=7,color='grey')
plt.tight_layout(rect=(0,0.07,1,1)); plt.savefig(f'{C}/demand_web_rank_tranco.png'); plt.close()
# 4 Mix estimate
from openpyxl import load_workbook
ws=load_workbook('/workspace/nebius/parts/demand_signals.xlsx',data_only=True)['Mix_Estimate']
rows=[r for r in ws.iter_rows(min_row=4,max_row=6,values_only=True)]
fig,ax=plt.subplots(figsize=(9,4.5)); names=[r[0] for r in rows]; left=[0]*3
for idx,lab,c in [(1,'(a) Hyperscaler take-or-pay (MSFT, Meta)','#1F3864'),(2,'(b) Reserved / committed contracts','#2E86C1'),(3,'(c) On-demand GPU + Token Factory','#E67E22'),(4,'Other segments (TripleTen, Avride)','#BFC9CA')]:
    v=[100*r[idx] for r in rows]; ax.barh(names,v,left=left,color=c,label=lab)
    for i,(l,x) in enumerate(zip(left,v)):
        if x>2.5: ax.text(l+x/2,i,f'{x:.0f}%',ha='center',va='center',color='white',fontsize=9)
    left=[l+x for l,x in zip(left,v)]
ax.set_xlim(0,100); ax.legend(frameon=False,fontsize=8,loc='upper center',bbox_to_anchor=(0.5,-0.12),ncol=2); ax.invert_yaxis()
ax.set_title('ESTIMATED revenue mix, Q2-26 ($582.3M group revenue)',fontsize=10,loc='left')
fig.text(0.01,0.005,'\n'.join(textwrap.wrap('Estimate. (a) = Customers C+D (24%+21%, Q2-26 6-K) identified as Microsoft/Meta by timing & A/R; (c) from CoreWeave 98% committed benchmark + OpenRouter-implied Token Factory revenue; (b) = plug.',170)),fontsize=7,color='grey')
plt.tight_layout(rect=(0,0.07,1,1)); plt.savefig(f'{C}/demand_mix_estimate.png'); plt.close()
# 5 customer stories by quarter & class
wsC=load_workbook('/workspace/nebius/parts/demand_signals.xlsx',data_only=True)['Customers']
st=[r for r in wsC.iter_rows(min_row=4,values_only=True) if r[0] and r[5] and str(r[5]).startswith('20')][:37]
qs=sorted({r[5] for r in st}); classes=['AI lab','startup','enterprise','research/OSS']; cc={'AI lab':NB,'startup':'#2E86C1','enterprise':'#E67E22','research/OSS':'#BFC9CA'}
fig,ax=plt.subplots(figsize=(9,4.5)); bottom=[0]*len(qs)
for cl in classes:
    v=[sum(1 for r in st if r[5]==q and r[2]==cl) for q in qs]; ax.bar(qs,v,bottom=bottom,color=cc[cl],label=cl); bottom=[b+x for b,x in zip(bottom,v)]
ax.set_ylabel('new customer-story pages'); ax.legend(frameon=False,fontsize=8); ax.set_title('nebius.com customer stories by quarter first seen (Wayback) and customer class',fontsize=10,loc='left')
fig.text(0.01,0.005,'\n'.join(textwrap.wrap(SRC+'nebius.com/server-sitemap.xml (37 story pages) and web.archive.org first capture per URL; classification by analyst. Hyperscalers (Microsoft, Meta) appear only in filings, not stories.',170)),fontsize=7,color='grey')
plt.tight_layout(rect=(0,0.07,1,1)); plt.savefig(f'{C}/demand_customer_stories.png'); plt.close()
print('charts written')
# 6 vintage of open roles
QS=['2024Q3','2024Q4','2025Q1','2025Q2','2025Q3','2025Q4','2026Q1','2026Q2','2026Q3','2026Q4*']
fig,ax=plt.subplots(figsize=(10,4.6)); w=0.27
for k,(co,fn,fld,c) in enumerate([('Nebius','nebius','first_published',NB),('CoreWeave','coreweave','first_published','#7D3C98'),('Lambda','lambda','publishedAt','#27AE60')]):
    J=json.load(open(f'{D}/raw/{fn}_jobs_2026-10-04.json'))['jobs']; cnt=collections.Counter()
    for x in J:
        d_=x.get(fld) or ''
        if d_: cnt[f'{d_[:4]}Q{(int(d_[5:7])-1)//3+1}']+=1
    v=[cnt.get(q.rstrip('*'),0) for q in QS]; ax.bar([i+(k-1)*w for i in range(len(QS))],v,w,color=c,label=f'{co} ({len(J)} open)')
ax.set_xticks(range(len(QS))); ax.set_xticklabels(QS); ax.legend(frameon=False); ax.set_ylabel('open roles (live, 2026-10-04)')
ax.set_title('Hiring ramp proxy: currently open roles by quarter first posted (Nebius board: 54 roles Nov-24 -> 365 Oct-26)',fontsize=10,loc='left')
fig.text(0.01,0.005,'\n'.join(textwrap.wrap('Source: Greenhouse first_published / Ashby publishedAt on live boards, accessed 2026-10-04; Nov-24 count from Wayback capture of the Greenhouse API. Survivorship-biased (filled roles drop out). *Q4-26 = 4 days.',170)),fontsize=7,color='grey')
plt.tight_layout(rect=(0,0.07,1,1)); plt.savefig(f'{C}/demand_job_postings_vintage.png'); plt.close()
print('vintage chart')
