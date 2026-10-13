"""Build /workspace/nebius/parts/time_to_deploy.xlsx.  usage: python build_workbook.py [--partial]"""
import sys, datetime as dt, statistics as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import analyze as A
from data_curated import MILESTONES, AVAIL, REF, ANN, GENS
OUT='/workspace/nebius/parts/time_to_deploy.xlsx'
ALT={'H200':('2024-07-01','TrendForce: H200 volume shipments Q3 2024'),
     'GB200 NVL72':('2025-05-28','NVIDIA Q1 FY26 call: GB200 NVL racks generally available'),
     'B300 / GB300 NVL72':('2025-07-03','First GB300 NVL72 deployment (CoreWeave/Dell) as market GA proxy')}
PARTIAL='--partial' in sys.argv
rows=A.load_obs(); S=A.series(rows)
wb=Workbook(); wb.remove(wb.active)
HF=Font(bold=True,color='FFFFFF'); HFILL=PatternFill('solid',fgColor='1F3864'); NEB=PatternFill('solid',fgColor='FFF2CC')
SRC=set()
def sheet(name,header,data,widths=None,note=None):
    ws=wb.create_sheet(name); r0=1
    if note: ws.cell(1,1,note).font=Font(italic=True,size=9); r0=2
    for j,h in enumerate(header,1):
        c=ws.cell(r0,j,h); c.font=HF; c.fill=HFILL; c.alignment=Alignment(wrap_text=True,vertical='top')
    for i,row in enumerate(data,r0+1):
        for j,v in enumerate(row,1):
            if isinstance(v,float): v=round(v,4)
            c=ws.cell(i,j,v)
            if isinstance(v,dt.date): c.number_format='yyyy-mm-dd'
            if isinstance(v,str) and v.startswith('http'): SRC.add(v)
        if 'Nebius' in [str(x) for x in row[:2]]:
            for j in range(1,len(header)+1): ws.cell(i,j).fill=NEB
    ws.freeze_panes=ws.cell(r0+1,1)
    for j,h in enumerate(header,1):
        ws.column_dimensions[get_column_letter(j)].width=(widths or {}).get(h,min(max(12,len(h)+2),40))
    return ws
D=A.D
# ---- Chip_Milestones
sheet('Chip_Milestones',['Chip / generation','Milestone','Date','Tag','Definition / method','Source URL','Used as lag reference?'],
      [(m[0],m[1],D(m[2]),m[3],m[4],m[5],'') for m in MILESTONES]+
      [(g,'REFERENCE "NVIDIA GA" used for lags',D(REF[g]),'Reported' if g=='H100' else 'Estimated','See milestone rows above for this chip; alternate definitions in Lag_Ranking',
        next((m[5] for m in MILESTONES if m[0].split()[0] in g and m[2]==REF[g]),''),'YES') for g in GENS],
      {'Definition / method':70,'Source URL':60,'Milestone':34},
      note='NVIDIA milestones; multiple definitions kept where they differ. Dates = public announcement/earnings-call statements. Tag: Reported / Derived / Estimated.')
# ---- Provider_Availability
def first_wb(prov,gen):
    chips=[c for c,g in A.CHIP2GEN.items() if g==gen]; best=None
    for (p,c),ser in S.items():
        if p!=prov or c not in chips: continue
        f=ser[0]; fp=next((o for o in ser if o[1] is not None),None)
        if best is None or f[0]<best[0][0]: best=(f,fp,c)
    return best
av=[]
for a in AVAIL:
    w=first_wb(a[0],a[1])
    av.append((a[0],a[1],D(a[2]),a[3],D(a[4]),a[5],a[6],D(a[7]),a[8],a[9],a[10],
               w[2] if w else '',w[0][0] if w else None,w[0][2] if w else '',
               w[1][0] if w and w[1] else None,w[1][1] if w and w[1] else None,w[1][2] if w and w[1] else '',a[11]))
sheet('Provider_Availability',['Provider','Generation','(a) First announcement date','(a) Source','(b) Preview / early access date','(b) Note','(b) Source',
      '(c) GA / first public price date (used for lag)','(c) Tag','(c) Definition','(c) Source',
      'Wayback: chip key','Wayback: first capture listing chip (upper bound)','Wayback: archive URL','Wayback: first capture with $ price','Wayback: first $/GPU-hr','Wayback: price archive URL','Notes'],
      av,{'(a) Source':45,'(b) Note':40,'(b) Source':45,'(c) Definition':45,'(c) Source':50,'Wayback: archive URL':50,'Wayback: price archive URL':50,'Notes':60},
      note='Dates are earliest public evidence found. Wayback dates are upper bounds (capture cadence ~monthly). Blank = not found in public sources (gap), not necessarily unavailable.')
if PARTIAL:
    wb.save(OUT); print('partial saved'); sys.exit()
# ---- Lag_Ranking
L=A.lag_table(); lr=[]
for r in L:
    alt=ALT.get(r['gen']); altlag=(r['ga']-D(alt[0])).days if alt and r['ga'] else None
    math=f"= {r['ga']} - {r['ref']}" if r['ga'] else f"censored: not GA as of {A.TODAY}; >= {A.TODAY} - {r['ref']} = {r['censored_min']}d"
    lr.append((r['gen'],r['provider'],r['ref'],r['ga'],r['lag'],'Derived',math,r['censored_min'],r['rank'],r['n_ranked'],
               r['lag_from_ann'],r['peer_median_ex_nebius'],r['nebius_minus_peer_median'],r['fastest'],
               alt[0] if alt else '',altlag,r['ga_tag'],r['ga_def'],r['ga_src']))
sheet('Lag_Ranking',['Generation','Provider','NVIDIA reference GA','Provider GA','Lag (days) NVIDIA GA -> provider GA','Tag','Math','Censored: min lag if not GA (days)','Rank (1=fastest)','# ranked',
      'Days from NVIDIA announcement','Peer median lag ex-Nebius (days)','Nebius minus peer median (days, +=slower)','Fastest provider',
      'Alt NVIDIA GA definition','Lag vs alt definition (days)','Provider GA tag','Provider GA definition','Provider GA source'],lr,
      {'Math':42,'Provider GA definition':45,'Provider GA source':55},
      note='Lag = provider GA (or first public price) minus NVIDIA reference GA (Chip_Milestones). Negative = before NVIDIA reference date. Ranking unaffected by choice of reference (same ref for all providers in a generation).')
# summary block
ws=wb['Lag_Ranking']; r0=ws.max_row+3
ws.cell(r0,1,'SUMMARY: Nebius vs peers').font=Font(bold=True)
hdr=['Generation','Nebius lag (d)','Nebius rank','# ranked','Peer median ex-Nebius (d)','Nebius - median (d)','Fastest','Fastest lag (d)','Nebius - fastest (d)','Nebius - CoreWeave (d)']
for j,h in enumerate(hdr,1): c=ws.cell(r0+1,j,h); c.font=HF; c.fill=HFILL
SUMM=[]
for i,g in enumerate(GENS):
    rs=[r for r in L if r['gen']==g]; n=next(r for r in rs if r['provider']=='Nebius'); cw=next(r for r in rs if r['provider']=='CoreWeave')
    f=min([r for r in rs if r['lag'] is not None],key=lambda r:r['lag'])
    nl=n['lag'] if n['lag'] is not None else f">={n['censored_min']} (not GA)"
    row=[g,nl,n['rank'] or 'not GA',n['n_ranked'],n['peer_median_ex_nebius'],n['nebius_minus_peer_median'],f['provider'],f['lag'],
         (n['lag']-f['lag']) if n['lag'] is not None else f">={n['censored_min']-f['lag']}",(n['lag']-cw['lag']) if n['lag'] is not None and cw['lag'] is not None else (f">={n['censored_min']-cw['lag']}" if cw['lag'] is not None else '')]
    SUMM.append(row)
    for j,v in enumerate(row,1): ws.cell(r0+2+i,j,v)
# ---- Launch_Premium_Decay
pdx=A.premium_decay(S); pr=[]
for r in pdx:
    excl=r['provider'] in ('Together AI',) ; aws=r['provider']=='AWS'
    conf='low' if excl else ('medium' if aws or not r['is_launch'] else 'high')
    def chg(m): return r[f'chg{m}']
    pr.append((r['provider'],r['chip'],'EC2 Capacity Blocks (reserved, per accelerator)' if aws else 'On-demand list',
               r['provider_ga'],r['launch_obs'],r['first_price_gap'],'Yes' if r['is_launch'] else 'No (first public price found later than GA+60d)',
               r['launch_price'],'Reported',r['launch_url'],
               r['p3'],chg(3),r['p6'],chg(6),r['p12'],chg(12),'Reported (as-of step function of Wayback captures)',
               r['pred'],r['pred_price'],r['prem_pred'],f"= {r['launch_price']}/{r['pred_price']} - 1" if r['pred_price'] else '',
               r['peer_med'],'; '.join(f'{p} ${v:.2f}' for p,v in r['peers']),r['prem_peer'],f"= {r['launch_price']}/{r['peer_med']} - 1" if r['peer_med'] else '',
               r['fit12'],r['n12'],r['fitall'],r['nall'],r['last_obs'],
               'Estimated: OLS of ln(price) on month index, step function sampled monthly from launch; %/mo = exp(slope)-1',conf,
               'Together: multiple tables on page, noisy parse; availability only' if excl else ('Capacity Blocks are reservation prices, not on-demand; not comparable with peers' if aws else '')))
sheet('Launch_Premium_Decay',['Provider','Chip','Price basis','Provider GA (Provider_Availability)','First observed $ price date','Days GA -> first price','Is launch price?',
      'Launch $/GPU-hr','Tag','Launch archive URL','+3m $/GPU-hr','+3m chg %','+6m $/GPU-hr','+6m chg %','+12m $/GPU-hr','+12m chg %','Tag (+m prices)',
      'Predecessor chip (same provider)','Predecessor $ same date','Premium vs predecessor %','Math (Derived)',
      'Peer median same chip same date $','Peers used','Premium vs peer median %','Math (Derived)',
      'Decay %/mo first 12m','# monthly pts','Decay %/mo full history','# monthly pts (full)','Last capture','Decay method','Confidence','Notes'],pr,
      {'Launch archive URL':50,'Peers used':40,'Decay method':45,'Notes':45,'Is launch price?':22},
      note='$/GPU-hr from Wayback captures of public pricing pages (see Sources). Negative decay = price falling. Blank = no capture after that horizon. CoreWeave 2025+ prices are 8-GPU instance price/8 (bundled CPU/RAM). Peer set (on-demand): Nebius, CoreWeave, Lambda, Crusoe, Together.')
# ---- Pricing_Page_Changes
ev=A.change_log(rows)
sheet('Pricing_Page_Changes',['Provider','Chip','Change type','Old $/GPU-hr','New $/GPU-hr','% change','Previous capture (bracket start)','Capture showing change (bracket end)','Change window (days)','Prev archive URL','Archive URL','Page dir','Tag','Note'],
      [(e['provider'],e['chip'],e['type'],e['old'],e['new'],(e['new']/e['old']-1)*100 if e['old'] and e['new'] else None,
        e['prev_ts'][:8] and D(f"{e['prev_ts'][:4]}-{e['prev_ts'][4:6]}-{e['prev_ts'][6:8]}") if e['prev_ts'] else None,
        D(f"{e['ts'][:4]}-{e['ts'][4:6]}-{e['ts'][6:8]}") if e['ts'] else None,
        (D(f"{e['ts'][:4]}-{e['ts'][4:6]}-{e['ts'][6:8]}")-D(f"{e['prev_ts'][:4]}-{e['prev_ts'][4:6]}-{e['prev_ts'][6:8]}")).days if e['prev_ts'] and e['ts'] else None,
        e['prev_url'],e['url'],e['page'],'Reported (Wayback); % Derived',e['note']) for e in ev],
      {'Prev archive URL':50,'Archive URL':50,'Note':60,'Change type':26},
      note='Every detected chip add/remove and price change on Nebius, CoreWeave, Lambda pricing pages from digest-collapsed Wayback CDX captures. Change happened between bracket start and end. "Contact us" chips shown with blank price.')
# ---- Sources
from collections import OrderedDict
srcs=OrderedDict()
for m in MILESTONES: srcs.setdefault(m[5],'NVIDIA milestone: '+m[0]+' '+m[1])
for a in AVAIL:
    for u in (a[3],a[6],a[10]):
        if u and u.startswith('http'): srcs.setdefault(u,f'{a[0]} {a[1]} availability')
for u in sorted(SRC):
    srcs.setdefault(u,'Wayback capture (id_ raw)' if 'web.archive.org' in u else 'cited')
srcs.setdefault('http://web.archive.org/cdx/search/cdx','Wayback CDX API (digest-collapsed listings)')
sheet('Sources',['URL','Used for'],list(srcs.items()),{'URL':100,'Used for':50})
wb.save(OUT); print('saved',OUT)
import json; json.dump(SUMM,open(A.BASE+'/out/lag_summary.json','w'),default=str,indent=1)
