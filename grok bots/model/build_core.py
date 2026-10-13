# Core helpers for the Nebius driver model (openpyxl, live formulas)
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

NQ = 28              # Q1'24 .. Q4'30
NHIST = 10           # Q1'24 .. Q2'26
FIRST_QCOL = 6       # F  (E = opening / index 0)
YEARS = list(range(2024, 2031))
FY_COL0 = 35         # AI
TERM_COL = 42        # AP

def qc(n):  # quarter n (0..28) -> column letter
    return CL(5 + n)
def fyc(y):
    return CL(FY_COL0 + YEARS.index(y))
QL = []
QEND = []
for y in range(2024, 2031):
    for q, (m, d) in enumerate([(3, 31), (6, 30), (9, 30), (12, 31)], start=1):
        QL.append(f"Q{q}'{str(y)[2:]}")
        QEND.append(dt.date(y, m, d))

BLUE = Font(name="Calibri", size=10, color="0000FF")
BLACK = Font(name="Calibri", size=10, color="000000")
BOLD = Font(name="Calibri", size=10, bold=True)
TITLE = Font(name="Calibri", size=14, bold=True, color="1F3864")
HDR = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
HDRFILL = PatternFill("solid", fgColor="1F3864")
SECFILL = PatternFill("solid", fgColor="D9E1F2")
F_PLACE = PatternFill("solid", fgColor="FCE4D6")   # placeholder
F_EST = PatternFill("solid", fgColor="FFF2CC")     # estimated / consensus
F_HISTCOL = PatternFill("solid", fgColor="F2F2F2")
ITAL = Font(name="Calibri", size=9, italic=True, color="595959")
thin = Side(style="thin", color="BFBFBF")

STATUS_FILL = {"Reported": None, "Derived": None, "Estimated": F_EST, "Consensus": F_EST,
               "Placeholder": F_PLACE, "Not disclosed": F_PLACE, "Input": None}

def setw(ws, widths):
    for k, v in widths.items():
        ws.column_dimensions[k].width = v

def inp(cell, value, status="Placeholder", fmt=None):
    cell.value = value
    cell.font = BLUE
    f = STATUS_FILL.get(status)
    if f: cell.fill = f
    if fmt: cell.number_format = fmt

def fml(cell, formula, fmt=None, bold=False):
    cell.value = formula
    cell.font = BOLD if bold else BLACK
    if fmt: cell.number_format = fmt

def name(wb, nm, sheet, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{sheet}'!{ref}")

class TS:
    """Time-series sheet: A label, B unit, C owner, D tag/source, E opening(idx0), F..AG quarters, AI..AO FY, AP terminal"""
    def __init__(self, wb, title, subtitle, fy=True):
        self.ws = wb.create_sheet(title)
        self.title = title
        self.r = 9
        self.fy = fy
        ws = self.ws
        ws["A1"] = title; ws["A1"].font = TITLE
        ws["A2"] = subtitle; ws["A2"].font = ITAL
        for c, t in zip("ABCDE", ["Line item", "Unit", "Owner", "Tag / source / method", "Open"]):
            ws[f"{c}3"] = t; ws[f"{c}3"].font = HDR; ws[f"{c}3"].fill = HDRFILL
        for rr, lab in [(4, "Period end"), (5, "Index"), (6, "A = actual / F = forecast"), (7, "Fiscal year")]:
            ws[f"A{rr}"] = lab; ws[f"A{rr}"].font = ITAL
        ws["E5"] = 0
        for n in range(1, NQ + 1):
            c = qc(n)
            ws[f"{c}3"] = f"=Timeline!{c}3"; ws[f"{c}3"].font = HDR; ws[f"{c}3"].fill = HDRFILL
            ws[f"{c}4"] = f"=Timeline!{c}4"; ws[f"{c}4"].number_format = "mmm-yy"; ws[f"{c}4"].font = ITAL
            ws[f"{c}5"] = f"=Timeline!{c}5"; ws[f"{c}5"].font = ITAL
            ws[f"{c}6"] = f"=Timeline!{c}6"; ws[f"{c}6"].font = ITAL
            ws[f"{c}7"] = f"=Timeline!{c}7"; ws[f"{c}7"].font = ITAL
        if fy:
            for y in YEARS:
                c = fyc(y)
                ws[f"{c}3"] = f"FY{y}"; ws[f"{c}3"].font = HDR; ws[f"{c}3"].fill = HDRFILL
                ws[f"{c}7"] = y; ws[f"{c}7"].font = ITAL
        setw(ws, {"A": 46, "B": 10, "C": 13, "D": 40, "E": 9})
        for n in range(1, NQ + 1):
            ws.column_dimensions[qc(n)].width = 10
        ws.column_dimensions["AH"].width = 2
        for y in YEARS:
            ws.column_dimensions[fyc(y)].width = 11
        ws.freeze_panes = "F8"

    def section(self, text):
        ws = self.ws
        self.r += 1
        ws[f"A{self.r}"] = text; ws[f"A{self.r}"].font = BOLD
        for c in range(1, FY_COL0 + len(YEARS)):
            ws.cell(self.r, c).fill = SECFILL
        self.r += 1

    def row(self, label, unit="", owner="", note="", fn=None, kind="flow", fmt="#,##0.0;(#,##0.0);-",
            e=None, values=None, status=None, bold=False, hist_fill=True, fy_fn=None):
        """fn(n) -> formula string or number for quarter n (1..28). values: dict n->(value,status) for inputs."""
        ws = self.ws; r = self.r
        ws[f"A{r}"] = label; ws[f"A{r}"].font = BOLD if bold else BLACK
        ws[f"B{r}"] = unit; ws[f"B{r}"].font = ITAL
        ws[f"C{r}"] = owner; ws[f"C{r}"].font = ITAL
        if isinstance(note, str) and note.startswith("="): note = "calc: " + note[1:].strip()
        ws[f"D{r}"] = note; ws[f"D{r}"].font = ITAL
        if e is not None:
            if isinstance(e, tuple):
                inp(ws[f"E{r}"], e[0], e[1], fmt)
            else:
                fml(ws[f"E{r}"], e, fmt)
        for n in range(1, NQ + 1):
            cell = ws[f"{qc(n)}{r}"]
            if values is not None and n in values:
                v, st = values[n]
                inp(cell, v, st, fmt)
            elif fn is not None:
                v = fn(n)
                if v is None: continue
                if isinstance(v, tuple):
                    inp(cell, v[0], v[1], fmt)
                else:
                    fml(cell, v, fmt, bold)
            if n <= NHIST and hist_fill and cell.fill.fgColor.rgb in (None, "00000000"):
                cell.fill = F_HISTCOL
        if self.fy:
            for y in YEARS:
                c = fyc(y); cell = ws[f"{c}{r}"]
                rng = f"$F{r}:$AG{r}"
                if fy_fn is not None:
                    v = fy_fn(y, c)
                    if v is not None: fml(cell, v, fmt, bold)
                elif kind == "flow":
                    fml(cell, f"=SUMIF($F$7:$AG$7,{c}$7,{rng})", fmt, bold)
                elif kind == "stock":
                    fml(cell, f"=INDEX({rng},1,MATCH({c}$7,$F$7:$AG$7,1))", fmt, bold)
                elif kind == "avg":
                    fml(cell, f"=AVERAGEIF($F$7:$AG$7,{c}$7,{rng})", fmt, bold)
        self.r += 1
        return r

    def blank(self, k=1):
        self.r += k
