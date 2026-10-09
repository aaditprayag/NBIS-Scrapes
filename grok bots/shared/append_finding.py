#!/usr/bin/env python3
# usage: append_finding.py bot finding sowhat source confidence
import sys, os, fcntl, datetime
from openpyxl import Workbook, load_workbook
p='/workspace/nebius/shared/findings_log.xlsx'
with open(p+'.lock','w') as lk:
    fcntl.flock(lk, fcntl.LOCK_EX)
    if os.path.exists(p): wb=load_workbook(p); ws=wb.active
    else:
        wb=Workbook(); ws=wb.active; ws.title='Findings'
        ws.append(['timestamp','bot','finding','so-what','source','confidence'])
    ws.append([datetime.datetime.now().strftime('%Y-%m-%d %H:%M ET')]+sys.argv[1:6])
    wb.save(p)
